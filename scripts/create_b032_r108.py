from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B032"
REVISION = "R108"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "personalized-classroom-doormat-with-custom-teacher-name-956ffe5ee5": ("Yay You're Here Teacher Mat", "Yay You're Here teacher mat", "personalized classroom doormat, first grade teacher mat, back to school welcome rug", "a Yay You're Here message with a teacher name, grade badge, rainbow, books, pencil, and school bus for a classroom or school entry"),
    "personalized-classroom-welcome-doormat-b404579ec3-b404579ec3": ("School Supplies Classroom Mat", "school supplies classroom mat", "custom classroom welcome doormat, personalized teacher rug, school doorway mat", "a school supplies classroom design with pencils, books, and a welcoming message for a classroom or school entry"),
    "personalized-classroom-welcome-doormat-77047e9c34-77047e9c34": ("Math Teacher Classroom Mat", "math teacher classroom mat", "personalized teacher doormat, classroom math rug, custom school welcome mat", "a math teacher classroom design with numbers and school-themed details for a classroom or teacher entry"),
    "personalized-classroom-welcome-doormat-2f0308422b-2f0308422b": ("Dinosaur Classroom Welcome Mat", "dinosaur classroom welcome mat", "personalized teacher doormat, classroom door mat, kids classroom rug", "a dinosaur classroom welcome design for a classroom doorway, school entry, or learning space"),
    "personalized-classroom-welcome-doormat-b5fe59f765-b5fe59f765": ("Affirmation Pencil Classroom Mat", "affirmation pencil classroom mat", "personalized classroom doormat, teacher affirmation rug, back to school welcome mat", "an affirmation message with pencil and classroom artwork for a classroom doorway, school entry, or teacher office"),
    "personalized-soccer-2026-area-rug-73a22990b0-73a22990b0": ("USA Road To Glory Soccer Rug", "USA road to glory soccer rug", "soccer 2026 area rug, patriotic soccer rug, football fan room rug", "a USA Road To Glory soccer design with a ball, trophy-inspired details, and patriotic colors for a fan room or game space"),
    "personalized-soccer-trophy-world-flags-doormat-00f2f9bbfa": ("World Flags Soccer Stadium Rug", "world flags soccer stadium rug", "soccer trophy area rug, world cup style soccer rug, football game room rug", "a world flags soccer stadium design with a trophy and international colors for a fan room or game space"),
    "personalized-soccer-trophy-area-rug-05582886df": ("World Game 2026 Trophy Rug", "World Game 2026 trophy rug", "soccer trophy area rug, 2026 soccer rug, football tournament rug", "a World Game 2026 design with a soccer ball, trophy, stadium lights, and flags for a fan room or game space"),
    "personalized-soccer-2026-area-rug-c2d07261b5": ("USA Canada Mexico Soccer Rug", "USA Canada Mexico soccer rug", "soccer 2026 area rug, world flags soccer rug, football fan room rug", "a USA Canada Mexico soccer design with a ball, stadium lights, and host-country flags for a fan room or game space"),
    "personalized-soccer-2026-area-rug-a60ad6dcf9-a60ad6dcf9": ("North America Soccer Crest Rug", "North America soccer crest rug", "soccer 2026 area rug, custom soccer trophy rug, football fan room rug", "a North America soccer crest design with a center ball, flags, and an ornate border for a fan room or game space"),
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B032 scope does not match the R108 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        article = "an" if primary.lower()[0] in "aeiou" else "a"
        meta = f"Create {article} {primary} with {motif}."
        description = f"Create {article} {primary} with {motif}. The design adds a distinctive focal point to a classroom, school entry, fan room, or game space."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [{
        "metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."
    }, {
        "metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."
    }, {
        "metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."
    }], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

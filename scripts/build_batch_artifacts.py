import csv
import html
import json
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
ROOT = Path("seo_runs") / SHOP / RUN_ID
RESULTS = Path("resutls") / SHOP / RUN_ID
BATCH_ID = os.environ["SEO_BATCH_ID"]
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


def batch_num(batch_id):
    return int(batch_id.removeprefix("B"))


def prev_batch_ids(batch_id):
    return [f"B{i:03d}" for i in range(1, batch_num(batch_id))]


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def write_csv(path, rows, headers):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with Path(path).open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def strip_html(value):
    return " ".join(re.sub(r"<[^>]+>", " ", html.unescape(value or "")).split())


def description_html(note):
    return (
        f"<p>{note['meta']}</p>"
        "<h3>Design details</h3><ul>"
        f"<li>{note['obs'].capitalize()}.</li>"
        "<li>Keep the SEO copy tied to the visible motif and page evidence, not a broad generic rug keyword.</li>"
        "<li>Use material, backing, washing, indoor/outdoor, pet, or kid claims only where supported by page text or product infographics.</li>"
        "</ul>"
    )


def image_note(note, n):
    image_notes = note.get("image_notes", {})
    return image_notes.get(str(n)) or image_notes.get(n) or f"{note['theme']} image {n}"


def autosize(ws, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(12, min(max_width, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def ensure_image_headers(headers, max_image_number):
    headers = list(headers)
    existing = set(headers)
    for n in range(1, max_image_number + 1):
        for suffix in ["link", "alt_current", "alt", "alt_action"]:
            col = f"img_{n}_{suffix}"
            if col not in existing:
                headers.append(col)
                existing.add(col)
    return headers


def main():
    notes_data = json.loads((ROOT / f"batches/{BATCH_ID}_design_notes.json").read_text())
    notes = {row["product_key"]: row for row in notes_data["products"]}
    serp_refs = notes_data["serp_refs"]
    batch = json.loads((ROOT / f"evidence/products/{BATCH_ID}_products.json").read_text())
    products = {p["handle"]: p for p in batch["products"]}
    page_meta = {r["product_key"]: r for r in json.loads((ROOT / f"evidence/products/{BATCH_ID}_page_meta.json").read_text())}
    manifest = json.loads((ROOT / f"evidence/images/{BATCH_ID}_manifest.json").read_text())
    max_image_number = max(int(row["image_number"]) for row in manifest) if manifest else 0
    by_images = {}
    for img in manifest:
        by_images.setdefault(img["product_key"], []).append(img)

    prev_seo, prev_img, prev_evd = [], [], []
    seo_headers = image_headers = evidence_headers = []
    for bid in prev_batch_ids(BATCH_ID):
        rows, headers = read_csv(ROOT / f"batches/{bid}_SEO_Products.csv")
        prev_seo.extend(rows)
        seo_headers = headers
        rows, headers = read_csv(ROOT / f"batches/{bid}_Image_Audit.csv")
        prev_img.extend(rows)
        image_headers = headers
        rows, headers = read_csv(ROOT / f"batches/{bid}_Product_Evidence.csv")
        prev_evd.extend(rows)
        evidence_headers = headers

    seo_headers = ensure_image_headers(seo_headers, max_image_number)

    prev_kw, keyword_headers = read_csv(ROOT / "keyword_research.csv")
    prev_kw = [r for r in prev_kw if r.get("product_key") not in notes]
    with (ROOT / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        prev_buyer = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(prev_buyer[0].keys())
    prev_buyer = [r for r in prev_buyer if r.get("product_key") not in notes]

    inv, inv_headers = read_csv(ROOT / "inventory.csv")
    inv_lookup = {r["product_key"]: r for r in inv}

    b_seo, b_img, b_evd, b_kw, b_buyer, keys_done = [], [], [], [], [], []
    for key in batch["product_keys"]:
        note = notes[key]
        product = products[key]
        meta = page_meta[key]
        order = int(inv_lookup[key]["inventory_order"])
        evidence_id = f"{BATCH_ID}-{order:03d}"
        research_id = f"{BATCH_ID}-{order:03d}-BR01"
        url = f"https://chillgen.com/products/{key}"
        imgs = sorted(by_images[key], key=lambda x: int(x["image_number"]))
        img_fields = {}
        for img in imgs:
            n = int(img["image_number"])
            alt = image_note(note, n)
            img_fields[f"img_{n}_link"] = img["src"]
            img_fields[f"img_{n}_alt_current"] = ""
            img_fields[f"img_{n}_alt"] = alt
            img_fields[f"img_{n}_alt_action"] = "SET"
            b_img.append(
                {
                    "evidence_id": evidence_id,
                    "shop_domain": SHOP,
                    "Handle": key,
                    "product_id": str(product.get("id", "")),
                    "media_id": img.get("image_id", ""),
                    "image_location": "GALLERY_PUBLIC_CDN",
                    "image_number": n,
                    "variant": "",
                    "image_url": img["src"],
                    "image_url_export": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                    "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                    "shared_media_references": "",
                    "viewed_status": "VIEWED_CONTACT_SHEET_AND_LOCAL_FILE",
                    "viewed_at": NOW,
                    "observed_visual_details": alt,
                    "alt_current": "",
                    "alt_proposed": alt,
                    "alt_action": "SET",
                    "review_status": "NEEDS_REVIEW",
                    "revision": "1",
                    "approved_by": "",
                    "approved_at": "",
                    "approved_fields": "",
                    "approved_revision": "",
                    "evidence_file_or_reference": img.get("local_path", ""),
                    "issues": "alt_current_empty_in_public_json; media_export_url_unknown_without_shopify_export",
                }
            )

        seo = {h: "" for h in seo_headers}
        seo.update(
            {
                "schema_version": "2.4",
                "prompt_version": "2.4",
                "run_id": RUN_ID,
                "batch_id": BATCH_ID,
                "shop_domain": SHOP,
                "product_key": key,
                "Handle": key,
                "product_id": str(product.get("id", "")),
                "product_gid": f"gid://shopify/Product/{product.get('id', '')}",
                "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                "source_export_ref": f"evidence/products/{BATCH_ID}_products.json; {meta.get('captured_html_path', '')}",
                "source_exported_at": meta.get("captured_at", NOW),
                "locale": "en-US",
                "market": "US",
                "url": url,
                "canonical_url": meta.get("canonical_url") or url,
                "product_type": product.get("product_type", ""),
                "title_current": product.get("title", ""),
                "h1_current": meta.get("h1_current", ""),
                "title_proposed": note["title"],
                "title_action": "SET",
                "h1_mapping_status": "RENDERED_H1_CAPTURED_THEME_CONFIRMATION_STILL_NEEDED_FOR_DEPLOY",
                "meta_title_current": "",
                "meta_title_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_title_current": meta.get("rendered_title_current", ""),
                "theme_title_suffix": "NONE_OR_NOT_OBSERVED",
                "meta_title_seo": note["title"],
                "meta_title_action": "SET",
                "meta_title_chars": len(note["title"]),
                "meta_description_current": "",
                "meta_description_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_meta_description_current": meta.get("rendered_meta_description_current", ""),
                "meta_description_seo": note["meta"],
                "meta_description_action": "SET",
                "meta_description_chars": len(note["meta"]),
                "primary_keyword": note["primary"],
                "secondary_keywords": note["secondary"],
                "long_tail_candidates": f"{note['primary']}; {note['secondary']}",
                "keyword_strategy": "LONG_TAIL_CANDIDATE",
                "season": note["season"],
                "keyword_demand_evidence": notes_data["keyword_evidence"],
                "keyword_serp_fit": "SERP_ONLY_PRODUCT_AND_MARKETPLACE_RESULTS",
                "meta_keyword": note["primary"],
                "search_intent": "COMMERCIAL_INVESTIGATION",
                "keyword_validation_status": "SERP_ONLY_NO_VOLUME",
                "keyword_selection_reason": f"Selected because page/image evidence supports {note['theme']} and product-level buyer intent.",
                "buyer_research_refs": research_id,
                "buyer_search_summary": notes_data["buyer_summary"],
                "baseline_status": "PUBLIC_RENDERED_PAGE_AND_JSON_CAPTURED",
                "change_scope": "CLARIFICATION",
                "mapping_status": "NEEDS_REVIEW",
                "mapping_version": "2.4",
                "description_current_html": product.get("body_html", ""),
                "description_changes_needed": "Make product copy motif-specific while retaining verified public claims only.",
                "description_proposed": strip_html(description_html(note)),
                "description_proposed_html": description_html(note),
                "description_change_mode": "REPLACE_ALL",
                "description_target_section": "Full product body",
                "description_action": "SET",
                "revision": "1",
                "review_status": "NEEDS_REVIEW",
                "review_reason": "Research proposal only; rendered page captured but admin/export SEO values and media export URLs remain unknown.",
                "approved_by": "",
                "approved_at": "",
                "approved_fields": "",
                "approved_revision": "",
                "processing_status": "DRAFTED",
                "evidence_status": "COMPLETE_PUBLIC_RENDERED_PAGE_AND_IMAGES",
                "content_qa_status": "NOT_RUN",
                "field_evidence_map": "title/meta/description/alt derived from rendered page meta, public JSON, viewed images, and SERP-only validation.",
                "evidence_id": evidence_id,
                "issues": "Admin SEO fields unknown without Shopify export; QA not run; not approved for deployment.",
            }
        )
        seo.update(img_fields)
        b_seo.append(seo)

        b_evd.append(
            {
                "evidence_id": evidence_id,
                "product_url": url,
                "reviewed_at": NOW,
                "sources_accessed": "; ".join([url, f"evidence/products/{BATCH_ID}_products.json", meta.get("captured_html_path", ""), *serp_refs]),
                "current_H1": meta.get("h1_current", ""),
                "current_meta_title": meta.get("rendered_title_current", ""),
                "short_source_excerpt": strip_html(product.get("body_html", ""))[:600],
                "verified_product_facts": f"{note['obs']}; HTTP {meta.get('http_status')}; canonical={meta.get('canonical_url')}; product type={product.get('product_type')}; variants={len(product.get('variants', []))}; images={len(imgs)}.",
                "gallery_image_count": len(imgs),
                "images_viewed_count": len(imgs),
                "image_audit_references": "; ".join([f"{evidence_id}-IMG{int(i['image_number']):02d}" for i in imgs]),
                "SERP_evidence_references": "; ".join(serp_refs),
                "buyer_research_references": research_id,
                "factual_conflicts": "None observed in public data; admin values unknown.",
                "processing_status": "DRAFTED",
                "confidence_and_reason": "MEDIUM: rendered page, product JSON, and images verified; keyword demand has SERP-only evidence and no volume data.",
                "fact_to_source_map": f"product facts={BATCH_ID}_products.json and {BATCH_ID}_pages HTML; visuals={BATCH_ID} local images/contact sheets; keyword context=SERP references.",
                "proposed_field_to_fact_map": "motif and visible product context to title/meta/alt; care/backing/material only where page/image evidence supports.",
            }
        )

        b_buyer.append(
            {
                "research_id": research_id,
                "product_key": key,
                "supporting_fact_ids": evidence_id,
                "purchase_context": note["purchase_context"],
                "jtbd_statement": note["jtbd"],
                "functional_motivation": note["functional_motivation"],
                "emotional_social_motivation": note["emotional_motivation"],
                "purchase_concerns": notes_data["purchase_concerns"],
                "customer_language": f"{note['primary']}; {note['secondary']}",
                "language_origin": "AGENT_HYPOTHESIS_WITH_CATEGORY_SERP_PARAPHRASE",
                "source_refs": "; ".join(serp_refs),
                "source_scope": "CATEGORY",
                "evidence_excerpt": "Paraphrased from public search/category results; no direct Chillgen customer quote or volume source used.",
                "observed_at": NOW,
                "market": "US",
                "source_language": "en",
                "research_status": "HYPOTHESIS",
                "limitations": "No Search Console, internal search, direct review corpus, or paid volume tool used for this product.",
                "seo_application": "Use motif-specific product terms and keep broad generic rug terms for collection/category review.",
            }
        )

        for kw, role, target, reason in [
            (note["primary"], "PRIMARY", url, "Best product-level match for observed motif and buyer intent."),
            (note["secondary"].split(", ")[0], "SECONDARY", url, "Useful motif synonym supported by page and image evidence."),
            (note["rejected_keyword"], "REJECTED", "COLLECTION_OR_CATEGORY_REVIEW", "Too broad for product-level targeting across similar products."),
        ]:
            b_kw.append(
                {
                    "keyword": kw,
                    "product_key": key,
                    "buyer_research_refs": research_id if role != "REJECTED" else "",
                    "query_origin": "AGENT_PROPOSED",
                    "semantic_cluster": note["primary"],
                    "intent": "COMMERCIAL_INVESTIGATION",
                    "target_page_type": "PRODUCT" if role != "REJECTED" else "COLLECTION",
                    "target_url": target,
                    "keyword_role": role,
                    "decision_reason": reason,
                    "supporting_fact_ids": evidence_id,
                    "demand_evidence": "SERP_ONLY, no volume verified.",
                    "season": note["season"],
                    "research_period": "2026-09-07",
                    "validation_source": "; ".join(serp_refs),
                    "checked_at": NOW,
                    "representative_SERP_URLs": "; ".join(serp_refs[:4]),
                    "possible_overlap_with_other_products": "YES" if role == "REJECTED" else "MEDIUM_WITH_SIMILAR_PRODUCTS_IN_BATCH",
                    "mapping_reason": "Motif-specific keyword retained for product; broad term reserved for collection/category review.",
                    "mapping_status": "NEEDS_REVIEW",
                    "mapping_version": "2.4",
                }
            )
        keys_done.append(key)

    write_csv(ROOT / f"batches/{BATCH_ID}_SEO_Products.csv", b_seo, seo_headers)
    write_csv(ROOT / f"batches/{BATCH_ID}_Image_Audit.csv", b_img, image_headers)
    write_csv(ROOT / f"batches/{BATCH_ID}_Product_Evidence.csv", b_evd, evidence_headers)
    write_csv(ROOT / "keyword_research.csv", prev_kw + b_kw, keyword_headers)
    with (ROOT / "buyer_search_research.jsonl").open("w", encoding="utf-8") as f:
        for row in prev_buyer + b_buyer:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    for row in inv:
        if row["product_key"] in notes:
            row["assigned_batch"] = BATCH_ID
            row["batch_status"] = "COMPLETED"
            row["research_status"] = "DRAFTED"
            row["review_status"] = "NEEDS_REVIEW"
    write_csv(ROOT / "inventory.csv", inv, inv_headers)

    total_products = len(prev_seo) + len(b_seo)
    total_images = len(prev_img) + len(b_img)
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", prev_seo + b_seo, seo_headers)
    write_sheet(wb, "Image_Audit", prev_img + b_img, image_headers)
    write_sheet(wb, "Product_Evidence", prev_evd + b_evd, evidence_headers)
    write_sheet(wb, "Keyword_Map", prev_kw + b_kw, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", prev_buyer + b_buyer, buyer_headers)
    readme = [
        {"metric": "scope", "value": f"B001 through {BATCH_ID}: {total_products} public products for chillgen.com, US/en-US", "definition": "Research proposal workbook only; all product and image proposals are NEEDS_REVIEW."},
        {"metric": f"{BATCH_ID}_method", "value": f"{BATCH_ID} captured product page HTML/meta before SEO drafting and viewed all downloaded gallery images via contact sheets.", "definition": f"B001-B003 remain public JSON/image evidence without full rendered page/meta backfill; B004-{BATCH_ID} include rendered page/meta capture."},
        {"metric": "sources", "value": "Chillgen public products.json, product URLs, rendered HTML captures, downloaded CDN images, contact sheets, SERP snippets and marketplace/category pages.", "definition": "Admin SEO fields and Shopify export media URLs remain unknown without Shopify CSV/export."},
        {"metric": "quality_status", "value": f"DRAFTED_PUBLIC_RENDERED_PAGE_THROUGH_{BATCH_ID}", "definition": f"{total_images}/{total_images} images viewed via contact sheets across B001-{BATCH_ID}; QA not run; no Shopify payload created."},
        {"metric": "deployment_status", "value": "NOT_APPROVED_NOT_DEPLOYABLE", "definition": "Do not import this workbook. Run QA and obtain explicit approval before any Shopify/Matrixify payload."},
        {"metric": "next_step", "value": f"Stop after {BATCH_ID} and wait for user confirmation before B{batch_num(BATCH_ID)+1:03d}.", "definition": "Required by project instruction."},
    ]
    write_sheet(wb, "README_QA", readme, ["metric", "value", "definition"])
    out = RESULTS / "batches" / f"SEO_Product_Optimization_through_{BATCH_ID}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    reopened = load_workbook(out, read_only=True, data_only=False)
    counts = {name: reopened[name].max_row - 1 for name in reopened.sheetnames}

    summary = ROOT / f"batches/{BATCH_ID}_summary.md"
    summary.write_text(
        "\n".join(
            [
                f"# {BATCH_ID} Summary",
                "",
                f"- Run: {RUN_ID}",
                f"- Shop: {SHOP}",
                f"- Scope: {len(b_seo)} products, {len(b_img)} images",
                f"- Cumulative workbook scope: B001-{BATCH_ID}, {total_products} products, {total_images} images",
                "- Method: pre-batch checkpoint plus rendered page/meta HTML capture and contact-sheet visual review.",
                "- Status: NEEDS_REVIEW; QA_NOT_RUN; NOT_APPROVED; NO_DEPLOYMENT_PAYLOAD",
                f"- Workbook: {out}",
                f"- Generated at: {NOW}",
                "",
                "## Product Keys",
                *[f"- {key}" for key in keys_done],
            ]
        ),
        encoding="utf-8",
    )

    progress = json.loads((ROOT / "progress.json").read_text())
    completed = set(progress.get("images_completed", []))
    completed.update([f"{m['product_key']}#img_{int(m['image_number']):02d}" for m in manifest])
    progress.update(
        {
            "batch_id": BATCH_ID,
            "batch_product_keys": keys_done,
            "batch_status": "COMPLETED_NEEDS_REVIEW",
            "awaiting_confirmation": True,
            "continuation_confirmation_ref": f"USER_CONFIRMED_CONTINUE_{BATCH_ID}",
            "current_product_key": None,
            "current_stage": f"{BATCH_ID}_COMPLETE_AWAITING_USER_CONFIRMATION",
            "images_completed": sorted(completed),
            "last_saved_at": NOW,
        }
    )
    progress.setdefault("artifact_paths", {}).update(
        {
            "latest_batch_workbook": str(out),
            f"{BATCH_ID.lower()}_seo_products_csv": str(ROOT / f"batches/{BATCH_ID}_SEO_Products.csv"),
            f"{BATCH_ID.lower()}_image_audit_csv": str(ROOT / f"batches/{BATCH_ID}_Image_Audit.csv"),
            f"{BATCH_ID.lower()}_product_evidence_csv": str(ROOT / f"batches/{BATCH_ID}_Product_Evidence.csv"),
            f"{BATCH_ID.lower()}_summary": str(summary),
            f"{BATCH_ID.lower()}_page_meta": str(ROOT / f"evidence/products/{BATCH_ID}_page_meta.json"),
            f"{BATCH_ID.lower()}_image_manifest": str(ROOT / f"evidence/images/{BATCH_ID}_manifest.json"),
            f"{BATCH_ID.lower()}_contact_sheets": str(ROOT / f"evidence/images/{BATCH_ID}_contact_sheets"),
        }
    )
    tmp = ROOT / "progress.json.tmp"
    tmp.write_text(json.dumps(progress, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(ROOT / "progress.json")
    print(json.dumps({"workbook": str(out), "sheet_counts": counts, "summary": str(summary)}, indent=2))


if __name__ == "__main__":
    main()

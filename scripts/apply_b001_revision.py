from pathlib import Path
import openpyxl, re, json

B = Path('/Users/buiquanghuy/Documents/seo_chillgen')
S = B / 'seo_runs/chillgen.com/chillgen_20260907_01/qa_scope/B001_scope_workbook.xlsx'
O = B / 'seo_runs/chillgen.com/chillgen_20260907_01/qa_scope/B001_scope_workbook_revision1.xlsx'
bad = {'custom-dog-photo-rug-personalized-pet-portrait-area-mat-n-design-02','custom-dog-photo-rug-personalized-pet-portrait-area-mat-non-slip','personalized-halloween-3d-optical-illusion-shaped-rug-a0-design-102','personalized-halloween-3d-optical-illusion-shaped-rug-a0-design-101','personalized-halloween-3d-optical-illusion-shaped-rug-a0-design-100','personalized-halloween-3d-optical-illusion-shaped-rug-a04-a04','custom-halloween-3d-effect-optical-illusion-spooky-round-rug-d9-d9','custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-06'}
w = openpyxl.load_workbook(S)
ws = w['SEO_Products']; h = [c.value for c in ws[1]]; c = {x:i+1 for i,x in enumerate(h)}
changed = []
for row in ws.iter_rows(min_row=2):
    handle = row[c['Handle']-1].value
    if handle in bad:
        for name in ('description_proposed','description_proposed_html','meta_description_seo'):
            cell = ws.cell(row[0].row, c[name])
            if isinstance(cell.value, str):
                cell.value = re.sub(r'\b(?:non[- ]?slip|anti[- ]?slip)\b', '', cell.value, flags=re.I)
                cell.value = re.sub(r'\s{2,}', ' ', cell.value).strip()
        for name,val in [('revision','2'),('review_status','NEEDS_REVIEW'),('processing_status','DRAFTED'),('content_qa_status','NOT_RUN'),('review_reason','Targeted B001 revision removes unsupported non-slip/anti-slip claims.'),('issues','Targeted B001 revision; full Re-QA required.')]: ws.cell(row[0].row, c[name]).value = val
        changed.append(handle)
w.save(O)
print(json.dumps({'output':str(O),'products_revised':len(changed)}))

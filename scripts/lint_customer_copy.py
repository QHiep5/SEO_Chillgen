"""Preflight lint for customer-facing proposed SEO copy in a workbook."""
from pathlib import Path
import argparse, json, re
import openpyxl

PATTERNS = [
 (r",\s*,", "duplicate comma"),
 (r"\s+[,.!?]", "space before punctuation"),
 (r"\b([A-Za-z]+)-\1-", "duplicated word"),
 (r"\b([A-Za-z]+)\s+\1\b", "duplicated adjacent word"),
 (r"\b(?:shown|available|designed|ideal|perfect)\s+in\s+\s*(?:and|for|with)\b", "empty phrase slot"),
 (r"\b(?:washable|easy-care|non-slip)[- ]?(?:care)?\s*\.\s*$", "truncated feature phrase"),
 (r"\b(?:in|with|for|and)\s+and\b", "broken conjunction"),
]
EVIDENCE_TERMS = ('backing visuals','backing detail graphics','non-slip backing visuals','washable-care graphics','care graphics','size chart images','room mockups','lifestyle mockups','feature graphics','product feature graphics','evidence images')
def lint(text):
 out=[label for p,label in PATTERNS if re.search(p,text,re.I)]
 out += [x for x in EVIDENCE_TERMS if x in text.lower()]
 return sorted(set(out))
def main():
 ap=argparse.ArgumentParser(); ap.add_argument('workbook',type=Path); ap.add_argument('--sheet',default='SEO_Products'); ap.add_argument('--revision',help='Only lint handles listed in Revision_Log for this revision id'); args=ap.parse_args()
 wb=openpyxl.load_workbook(args.workbook,read_only=True,data_only=True); ws=wb[args.sheet]; heads=[c.value for c in next(ws.iter_rows())]; findings=[]
 scope=None
 if args.revision and 'Revision_Log' in wb.sheetnames:
  rw=wb['Revision_Log']; rh=[c.value for c in next(rw.iter_rows())]; scope={dict(zip(rh,r)).get('handle') for r in rw.iter_rows(min_row=2,values_only=True) if dict(zip(rh,r)).get('revision_batch_id')==args.revision}
 fields=['title_proposed','meta_title_seo','meta_description_seo','description_proposed','description_proposed_html']
 for row in ws.iter_rows(min_row=2,values_only=True):
  d=dict(zip(heads,row)); text=' '.join(str(d.get(f) or '') for f in fields); hits=lint(text)
  if scope is not None and d.get('Handle') not in scope: continue
  if hits: findings.append({'handle':d.get('Handle'),'findings':hits})
 result={'workbook':str(args.workbook),'status':'PASS' if not findings else 'FAIL','finding_count':len(findings),'findings':findings}
 print(json.dumps(result,indent=2,ensure_ascii=False)); raise SystemExit(0 if not findings else 2)
if __name__=='__main__': main()

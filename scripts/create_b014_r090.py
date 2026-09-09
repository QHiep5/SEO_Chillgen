from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B014"
REVISION = "R090"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"


COPY = {
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-110": {
        "title": "My Feelings Emotion Wheel Rug",
        "meta": "Create a My Feelings emotion wheel rug with pastel emotion panels, cartoon kid faces, and gentle SEL classroom color.",
        "description": "Create a My Feelings emotion wheel rug with pastel emotion panels, cartoon kid faces, and gentle SEL classroom color. The round design gives children a friendly visual reference for naming feelings during classroom or playroom routines.",
        "primary": "my feelings emotion wheel rug",
        "secondary": "feelings wheel classroom rug, SEL emotions rug, kids feelings rug",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-111": {
        "title": "Today I'm Feeling Classroom Rug",
        "meta": "Create a Today I'm Feeling classroom rug with a rainbow emotion wheel, cartoon feeling faces, and bright SEL wording.",
        "description": "Create a Today I'm Feeling classroom rug with a rainbow emotion wheel, cartoon feeling faces, and bright SEL wording. The design gives students a colorful prompt for sharing how they feel in classroom, counseling, or playroom spaces.",
        "primary": "today i'm feeling classroom rug",
        "secondary": "today i'm feeling rug, emotion wheel rug, SEL classroom rug",
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-with-black-cat": {
        "title": "Personalized Black Cat Halloween Mat",
        "meta": "Customize a black cat Halloween doormat with family-name text, a witch, ghosts, pumpkins, skulls, and spooky welcome artwork.",
        "description": "Customize a black cat Halloween doormat with family-name text, a witch, ghosts, pumpkins, skulls, and spooky welcome artwork. The personalized design gives an entryway a playful haunted-season greeting.",
        "primary": "personalized black cat Halloween mat",
        "secondary": "custom Halloween doormat, black cat welcome mat, spooky family doormat",
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-wit-design-100": {
        "title": "Custom Pumpkin Skull Halloween Mat",
        "meta": "Customize a pumpkin skull Halloween doormat with spooky night artwork, jack-o-lanterns, skull details, and custom text.",
        "description": "Customize a pumpkin skull Halloween doormat with spooky night artwork, jack-o-lanterns, skull details, and custom text. The design gives a porch, doorway, or Halloween party entry a bold seasonal welcome.",
        "primary": "custom pumpkin skull Halloween mat",
        "secondary": "personalized Halloween doormat, pumpkin skull doormat, spooky welcome mat",
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-wit-design-101": {
        "title": "Personalized Family Halloween Mat",
        "meta": "Customize a family Halloween doormat with established-year text, black cat, witch hat, bats, ghost, skull, and cauldron artwork.",
        "description": "Customize a family Halloween doormat with established-year text, black cat, witch hat, bats, ghost, skull, and cauldron artwork. The personalized layout gives a front door or seasonal display a named Halloween greeting.",
        "primary": "personalized family Halloween mat",
        "secondary": "custom family Halloween doormat, spooky welcome mat, personalized Halloween porch mat",
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-wit-design-102": {
        "title": "Black Cat Full Moon Halloween Mat",
        "meta": "Customize a black cat full moon Halloween doormat with pumpkin forest artwork, moonlit details, and personalized text.",
        "description": "Customize a black cat full moon Halloween doormat with pumpkin forest artwork, moonlit details, and personalized text. The spooky illustrated scene creates a dramatic Halloween welcome for a doorway or porch.",
        "primary": "black cat full moon Halloween mat",
        "secondary": "custom black cat doormat, Halloween moon doormat, spooky personalized mat",
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-wit-design-103": {
        "title": "Personalized Haunted Home Mat",
        "meta": "Customize a haunted home Halloween doormat with family-name text, haunted house artwork, bats, pumpkins, and a glowing moon.",
        "description": "Customize a haunted home Halloween doormat with family-name text, haunted house artwork, bats, pumpkins, and a glowing moon. The personalized design gives an entryway a spooky seasonal welcome.",
        "primary": "personalized haunted home mat",
        "secondary": "haunted house Halloween doormat, custom Halloween welcome mat, spooky family mat",
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-horror-area": {
        "title": "Black Cat Portal Halloween Rug",
        "meta": "Create a black cat portal Halloween round rug with glowing purple eyes, a 3D illusion ring, and dark seasonal artwork.",
        "description": "Create a black cat portal Halloween round rug with glowing purple eyes, a 3D illusion ring, and dark seasonal artwork. The dramatic design gives a Halloween room, party corner, or entry area a spooky focal point.",
        "primary": "black cat portal Halloween rug",
        "secondary": "3D Halloween rug, black cat round rug, optical illusion Halloween rug",
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-100": {
        "title": "3D Skeleton Pit Halloween Rug",
        "meta": "Create a 3D skeleton pit Halloween rug with a stone tunnel illusion, skulls, ladders, torch lights, and horror artwork.",
        "description": "Create a 3D skeleton pit Halloween rug with a stone tunnel illusion, skulls, ladders, torch lights, and horror artwork. The round design adds a haunted optical-illusion moment to seasonal room decor.",
        "primary": "3D skeleton pit Halloween rug",
        "secondary": "skeleton Halloween rug, optical illusion Halloween rug, horror round rug",
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-101": {
        "title": "Haunted Mansion Halloween Rug",
        "meta": "Create a haunted mansion Halloween round rug with glowing windows, pumpkins, bats, dark night artwork, and spooky pathway details.",
        "description": "Create a haunted mansion Halloween round rug with glowing windows, pumpkins, bats, dark night artwork, and spooky pathway details. The design brings a haunted-house scene to Halloween room or party decor.",
        "primary": "haunted mansion Halloween rug",
        "secondary": "haunted house round rug, Halloween room rug, spooky mansion rug",
    },
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B014 scope does not match the R090 copy map")

    for row in seo_rows:
        item = COPY[row["Handle"]]
        row["title_proposed"] = item["title"]
        row["meta_title_seo"] = item["title"]
        row["meta_title_chars"] = str(len(item["title"]))
        row["meta_description_seo"] = item["meta"]
        row["meta_description_chars"] = str(len(item["meta"]))
        row["description_proposed"] = item["description"]
        row["description_proposed_html"] = f"<p>{item['description']}</p>"
        row["primary_keyword"] = item["primary"]
        row["secondary_keywords"] = item["secondary"]
        row["long_tail_candidates"] = f"{item['primary']}; {item['secondary']}"
        row["revision"] = "2"
        row["review_status"] = "NEEDS_REVIEW"
        row["content_qa_status"] = "NOT_RUN"
        row["review_reason"] = f"{REVISION} removes internal evidence-note language and unsupported claim wording from B014 copy."
        row["issues"] = f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations."

    handles_set = set(handles)
    evidence_scope = [row for row in evidence_rows if row.get("evidence_id") in {r.get("evidence_id") for r in seo_rows}]
    keyword_scope = [row for row in keyword_rows if row.get("product_key") in handles_set]
    buyer_rows = [row for row in buyer_all if row.get("product_key") in handles_set]
    img_scope = [row for row in img_rows if row.get("Handle") in handles_set]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": "Canonical B014 cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    evidence_ids = {row["evidence_id"] for row in evidence_scope}
    for row in seo_rows:
        evidence_id = row.get("evidence_id")
        if evidence_id not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {evidence_id}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", evidence_id])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    digest = sha256(OUT)
    manifest = {
        "revision": REVISION,
        "batch_id": BATCH,
        "scope_count": len(handles),
        "image_count": len(img_scope),
        "canonical_workbook": str(OUT),
        "sha256": digest,
        "status": "CREATED_FOR_HARDENED_REQA",
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
    }
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

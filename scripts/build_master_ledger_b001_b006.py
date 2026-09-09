from pathlib import Path
import csv,openpyxl,json
from openpyxl import Workbook
B=Path('/Users/buiquanghuy/Documents/seo_chillgen'); base=B/'resutls/chillgen.com/chillgen_20260907_01/qa'
paths={'B001':'QA-20260911-070000-BATCH-B001-COMPLIANT/SEO_QA_REV_B001.xlsx','B002':'QA-20260911-013000-REV-R053-HARDENED/SEO_QA_REV_R053.xlsx','B003':'QA-20260911-200000-BATCH-B003-REV1-HARDENED/SEO_QA_REV_B003.xlsx','B004':'QA-20260909-B004-R076-FULLRES-HARDENED/SEO_QA_REV_R076.xlsx','B005':'QA-20260909-B005-R079-FULLRES-HARDENED/SEO_QA_REV_R079.xlsx','B006_R080':'QA-20260909-B006-R080-FULLRES-HARDENED-FINAL/SEO_QA_REV_R080.xlsx','B006_R081':'QA-20260909-B006-R081-FULLRES-HARDENED/SEO_QA_REV_R081.xlsx'}
def read(p):
 w=openpyxl.load_workbook(base/p,read_only=True,data_only=True); s=w['QA_Products']; h=[c.value for c in next(s.iter_rows())]; return {str(r[h.index('product_key')]):dict(zip(h,r)) for r in s.iter_rows(min_row=2,values_only=True)}
reports={k:read(v) for k,v in paths.items() if (base/v).exists()}
inv=list(csv.DictReader((B/'seo_runs/chillgen.com/chillgen_20260907_01/inventory.csv').open())); out=[]
for x in inv:
 b=x.get('assigned_batch',''); h=x['product_key']; key='B006_R081' if b=='B006' and 'dog-paw' in h else 'B006_R080' if b=='B006' else b; q=reports.get(key,{}).get(h,{})
 out.append([b,h,x['url'],key.split('_')[-1] if '_' in key else key,paths.get(key,''),q.get('qa_status','PENDING'),q.get('final_score',''),q.get('images_expected',''),q.get('images_checked',''),q.get('image_coverage',''), 'NOT_APPROVED_NOT_DEPLOYED'])
od=base/'CATALOG_QA_MASTER_LEDGER_B001_B006.xlsx'; w=Workbook(); s=w.active; s.title='Product_Ledger'; s.append(['batch_id','product_key','url','effective_revision','qa_report','qa_status','final_score','images_expected','images_checked','image_coverage','approval_status']); [s.append(r) for r in out]
sm=w.create_sheet('Catalog_Summary'); sm.append(['metric','value']); sm.append(['ledger_version','catalog-qa-ledger-v2']); sm.append(['closed_batches','B001,B002,B003,B004,B005,B006']); sm.append(['closed_products',sum(1 for r in out if r[0] in {'B001','B002','B003','B004','B005','B006'})]); sm.append(['approval_status','NOT_APPROVED_NOT_DEPLOYED'])
w.save(od); (od.with_suffix('.manifest.json')).write_text(json.dumps({'ledger_version':'catalog-qa-ledger-v2','closed_batches':['B001','B002','B003','B004','B005','B006'],'closed_products':sum(1 for r in out if r[0] in {'B001','B002','B003','B004','B005','B006'}),'approval_status':'NOT_APPROVED_NOT_DEPLOYED'},indent=2)+'\n'); print({'output':str(od),'rows':len(out),'reports':list(reports)})

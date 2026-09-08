from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R045/SEO_Product_Optimization_revision_R045.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R046"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R046.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R046"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-patriotic-hunting-doormat-ee02618b87": {
        "title": "American Hunter Black Lab Patriotic Doormat Art",
        "meta": "Personalize an American Hunter patriotic doormat with black lab, antlers, American flag, forest details and patchwork border art.",
        "primary": "American Hunter black lab doormat",
        "secondary": "patriotic hunting doormat, black lab hunter mat, American flag wildlife mat",
        "long_tail": "American Hunter black lab doormat; patriotic hunting doormat; American flag wildlife mat",
        "description": "Personalize an American Hunter patriotic doormat with black lab, antlers, American flag, forest details and patchwork border art.\n\nDesign details\n- Patriotic hunting doormat artwork shows a black lab wearing a cap inside an antler frame with American Hunter text.\n- Visible details include American flag colors, forest silhouettes, wildlife patchwork border, doorway scene, person holding mat and repeated gallery images.\n- SEO copy stays tied to visible black-lab hunting artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            5: "Feature graphic for American Hunter black lab doormat",
            12: "Repeated feature graphic for American Hunter doormat",
        },
    },
    "custom-golf-area-rug-living-room-mat-51b65f4454": {
        "title": "Patriotic Golf Flag Rug with Wood Grain Artwork",
        "meta": "Decorate with a patriotic golf flag rug featuring American flag stripes, wood grain artwork, golf club, ball and tee details.",
        "primary": "patriotic golf flag rug",
        "secondary": "American flag golf rug, wood grain golf rug, golf club artwork mat",
        "long_tail": "patriotic golf flag rug; American flag golf rug; wood grain golf rug",
        "description": "Decorate with a patriotic golf flag rug featuring American flag stripes, wood grain artwork, golf club, ball and tee details.\n\nDesign details\n- Rug artwork combines red and white flag stripes, a blue star field and a wood plank-style golf background.\n- Visible details include golf club, ball, tee, lounge mockups, close-up views, size option graphic and feature infographic.\n- SEO copy focuses on visible patriotic golf artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            5: "Close-up of patriotic golf flag rug artwork",
            6: "Feature infographic for patriotic golf rug artwork",
        },
    },
    "custom-golf-area-rug-living-room-mat-f6f3c16862": {
        "title": "Alex Golf Course Rug with Sunset Fairway Artwork",
        "meta": "Personalize an Alex golf course rug with green fairway, tee ball, flag pin, sunset landscape and rolling hill artwork.",
        "primary": "Alex golf course rug",
        "secondary": "personalized golf fairway rug, sunset golf course rug, custom name golf mat",
        "long_tail": "Alex golf course rug; personalized golf fairway rug; sunset golf course rug",
        "description": "Personalize an Alex golf course rug with green fairway, tee ball, flag pin, sunset landscape and rolling hill artwork.\n\nDesign details\n- Green golf course rug artwork shows Alex name text across a fairway scene with tee ball and flag pin.\n- Visible product imagery includes lounge mockups, golf room scene, size option graphic, close-up course views and feature panel.\n- SEO copy stays tied to visible fairway artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            4: "Close-up of green golf fairway rug artwork",
            5: "Feature infographic for Alex golf course rug",
        },
    },
    "custom-golf-area-rug-fd79165e9f-fd79165e9f": {
        "title": "William Golf Club Rug with White Putting Green Art",
        "meta": "Personalize a William Golf Club rug with white layout, putting green, red flag pin, family name text and established year.",
        "primary": "William Golf Club rug",
        "secondary": "family name golf rug, white putting green rug, personalized golf club mat",
        "long_tail": "William Golf Club rug; family name golf rug; white putting green rug",
        "description": "Personalize a William Golf Club rug with white layout, putting green, red flag pin, family name text and established year.\n\nDesign details\n- White personalized golf rug artwork reads The William Golf Club with EST year text and a green putting area.\n- Visible imagery includes red flag pin, black border, golf room mockup, entryway scenes, close-up views and feature infographic.\n- SEO copy focuses on visible golf-club artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            5: "Close-up of William Golf Club rug artwork",
            6: "Feature infographic for William Golf Club rug",
        },
    },
    "custom-mathematics-education-rug-80ab855e16-80ab855e16": {
        "title": "Anna MATH Acronym Classroom Rug with Chalk Art",
        "meta": "Personalize an Anna MATH acronym classroom rug with chalk-style letters, colorful symbols and classroom math prompt artwork.",
        "primary": "Anna MATH acronym classroom rug",
        "secondary": "MATH acronym rug, teacher name math rug, chalk-style classroom mat",
        "long_tail": "Anna MATH acronym classroom rug; MATH acronym rug; teacher name math rug",
        "description": "Personalize an Anna MATH acronym classroom rug with chalk-style letters, colorful symbols and classroom math prompt artwork.\n\nDesign details\n- Black rug artwork reads In This Anna's Classroom We Work To Understand MATH with each acronym letter expanded.\n- Visible prompts include mistakes, allow, thinking and happen, with colorful icons, classroom mockups and feature graphic.\n- SEO copy stays tied to visible MATH acronym artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {3: "Feature graphic for Anna MATH acronym rug"},
    },
    "custom-mathematics-education-classroom-playroom-rug-0f607da0b3": {
        "title": "Math Symbols Classroom Rug with Color Icon Art",
        "meta": "Decorate with a Math Symbols classroom rug featuring black background, colorful operation icons, number marks and symbol artwork.",
        "primary": "Math Symbols classroom rug",
        "secondary": "color math symbols rug, operation icons rug, black math classroom mat",
        "long_tail": "Math Symbols classroom rug; color math symbols rug; operation icons rug",
        "description": "Decorate with a Math Symbols classroom rug featuring black background, colorful operation icons, number marks and symbol artwork.\n\nDesign details\n- Black rug artwork centers Math Symbols text with a grid of bright circles containing math icons.\n- Visible icons include pi, equals, division, percent, subtraction, greater-than and plus marks, with room mockups and feature graphic.\n- SEO copy focuses on visible math-symbol artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {7: "Feature graphic for Math Symbols classroom rug"},
    },
    "custom-mathematics-classroom-rug-for-school-and-playroom-7ae72f2960": {
        "title": "How To Live The Mathlife Rug with Hashtag Rules",
        "meta": "Decorate with a How To Live The Mathlife rug featuring hashtag rule lines, colorful math text, black background and symbol border.",
        "primary": "How To Live The Mathlife rug",
        "secondary": "hashtag math rules rug, black math classroom rug, colorful math quote mat",
        "long_tail": "How To Live The Mathlife rug; hashtag math rules rug; black math classroom rug",
        "description": "Decorate with a How To Live The Mathlife rug featuring hashtag rule lines, colorful math text, black background and symbol border.\n\nDesign details\n- Black rug artwork reads How To Live The #Mathlife with stacked colorful rule phrases.\n- Visible phrases include show your work, ask questions, help each other and do your homework, with classroom mockups and feature graphic.\n- SEO copy stays tied to visible hashtag math artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {4: "Feature graphic for Mathlife hashtag rules rug"},
    },
    "custom-mathematics-classroom-rug-ec81d3501a": {
        "title": "Good Day To Do Math Classroom Rug with Icon Art",
        "meta": "Decorate with a Good Day To Do Math classroom rug featuring pastel geometry shapes, formula marks, calculator and pencil artwork.",
        "primary": "Good Day To Do Math classroom rug",
        "secondary": "pastel math icon rug, formula classroom rug, black math quote mat",
        "long_tail": "Good Day To Do Math classroom rug; pastel math icon rug; formula classroom rug",
        "description": "Decorate with a Good Day To Do Math classroom rug featuring pastel geometry shapes, formula marks, calculator and pencil artwork.\n\nDesign details\n- Black rug artwork reads It Is A Good Day To Do Math with pastel icons around the quote.\n- Visible details include geometry shapes, calculator, formulas, pencil border, classroom mockups and feature graphic.\n- SEO copy focuses on visible math quote artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {4: "Feature graphic for Good Day To Do Math rug"},
    },
    "custom-math-education-welcome-mat-for-classroom-and-school-f81e324d95-f81e324d95": {
        "title": "Why Math Classroom Rug with Concept Blocks Art",
        "meta": "Personalize a Why Math classroom rug with teacher name, colorful concept blocks, problem-solving prompts and math question artwork.",
        "primary": "Why Math classroom rug",
        "secondary": "teacher name math rug, concept block math rug, problem solving classroom mat",
        "long_tail": "Why Math classroom rug; teacher name math rug; concept block math rug",
        "description": "Personalize a Why Math classroom rug with teacher name, colorful concept blocks, problem-solving prompts and math question artwork.\n\nDesign details\n- Black rug artwork reads Why Math? with colorful blocks explaining perseverance, visual numbers, deeper thinking and reasoning.\n- Visible imagery includes teacher name text, classroom mockups, close-up views and circular feature callouts.\n- SEO copy stays tied to visible Why Math artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {3: "Feature graphic for Why Math classroom rug"},
    },
    "custom-math-education-rug-for-playroom-and-classroom-d0ac821f02-d0ac821f02": {
        "title": "Math Key Words Classroom Rug with Operation Boxes",
        "meta": "Decorate with a Math Key Words classroom rug featuring operation boxes for addition, subtraction, multiplication and division terms.",
        "primary": "Math Key Words classroom rug",
        "secondary": "operation vocabulary rug, addition subtraction rug, multiplication division mat",
        "long_tail": "Math Key Words classroom rug; operation vocabulary rug; addition subtraction rug",
        "description": "Decorate with a Math Key Words classroom rug featuring operation boxes for addition, subtraction, multiplication and division terms.\n\nDesign details\n- Cream rug artwork reads Math Key Words with colored sections for addition, subtraction, multiplication and division vocabulary.\n- Visible details include large operation symbols, term lists, room mockups, close-up views and circular feature callouts.\n- SEO copy focuses on visible math vocabulary artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {3: "Feature graphic for Math Key Words rug"},
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable", "floor mat", "floor rug", "low pile", "skid",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R046 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R046 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R046"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R046", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R046", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R047 or Re-QA revision R046",
    }
    (RUN_DIR / "revision_R046_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R046 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R046.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R046 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R046_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r046_log_rows": counts.get("R046", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

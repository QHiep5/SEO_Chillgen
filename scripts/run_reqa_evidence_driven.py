"""Evidence-driven re-QA generator for revision workbooks.

The generator never fabricates visual/image scores.  Image ratings are read
from an optional reviewer evidence JSON; without that file image criteria are
NOT_CHECKED and the product cannot become QA_PASS.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
from datetime import datetime, timezone, timedelta
from pathlib import Path
from statistics import mean

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(os.environ.get("SEO_CHILLGEN_BASE", Path.cwd())).resolve()
RUN_ID = "chillgen_20260907_01"
SHOP_DOMAIN = "chillgen.com"
TZ = timezone(timedelta(hours=7))
WEIGHTS = {"P1": 15, "P2": 10, "K1": 10, "K2": 5, "K3": 5,
           "T1": 10, "T2": 5, "D1": 5, "D2": 10, "I1": 20, "E1": 5}
IMAGE_WEIGHTS = {"IM1": 40, "IM2": 30, "IM3": 20, "IM4": 10}
RUBRIC_PATH = BASE / "seo-prompt/chillgen/prompt_qa.md"
REGISTRY_JSON = BASE / f"resutls/{SHOP_DOMAIN}/{RUN_ID}/qa/qa_run_registry.json"
REGISTRY_CSV = BASE / f"resutls/{SHOP_DOMAIN}/{RUN_ID}/qa/qa_run_registry.csv"
INTERNAL_PATTERNS = [
    "seo copy avoids", "unless confirmed during admin/export review",
    "admin/export review", "surface-performance claims", "qa note",
    "re-run qa", "internal note",
    # Evidence-production language is not customer-facing copy. These gates
    # prevent image/QA notes from being promoted into title/meta/description.
    "backing visuals", "backing detail graphics", "non-slip backing visuals",
    "anti-slip backing visuals", "washable-care graphics", "care graphics",
    "size chart images", "room mockups", "lifestyle mockups",
    "feature graphics", "product feature graphics", "evidence images",
]

def content_quality_findings(text: str):
    """Detect malformed customer copy after automated revision edits."""
    findings = []
    rules = [
        (r",\s*,", "duplicate comma"),
        (r"\s+[,.!?]", "space before punctuation"),
        (r"\b([A-Za-z]+)-\1-", "duplicated word"),
        (r"\b([A-Za-z]+)\s+\1\b", "duplicated adjacent word"),
        (r"\b(?:shown|available|designed|ideal|perfect)\s+in\s+\s*(?:and|for|with)\b", "empty phrase slot"),
        (r"\b(?:washable|easy-care|non-slip)[- ]?(?:care)?\s*\.\s*$", "truncated feature phrase"),
        (r"\b(?:in|with|for|and)\s+and\b", "broken conjunction"),
    ]
    for pattern, label in rules:
        if re.search(pattern, text, flags=re.IGNORECASE):
            findings.append(label)
    return findings


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def hash_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def scope_hash(keys) -> str:
    payload = json.dumps(list(keys), ensure_ascii=False, separators=(",", ":"))
    return hash_text(payload)


def rows(ws):
    it = ws.iter_rows(values_only=True)
    header = [str(x) if x is not None else "" for x in next(it)]
    return [dict(zip(header, row)) for row in it]


def val(row, key):
    x = row.get(key)
    return "" if x is None else str(x)


def append(ws, values):
    safe = []
    for x in values:
        if isinstance(x, str) and x[:1] in ("=", "+", "-", "@"):
            safe.append("'" + x)
        else:
            safe.append(x)
    ws.append(safe)


def header(ws, values):
    append(ws, values)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def finish(ws):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            widths[c.column] = max(widths.get(c.column, 0), min(70, len(str(c.value or "")) + 2))
    for i, width in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = max(12, width)


def update_registry(record):
    REGISTRY_JSON.parent.mkdir(parents=True, exist_ok=True)
    if REGISTRY_JSON.exists():
        registry = json.loads(REGISTRY_JSON.read_text(encoding="utf-8"))
    else:
        registry = {"schema_version": "qa-run-registry-v1", "runs": []}
    runs = registry.setdefault("runs", [])
    if any(r.get("qa_run_id") == record["qa_run_id"] for r in runs):
        raise RuntimeError(f"Duplicate qa_run_id already exists in registry: {record['qa_run_id']}")
    if record.get("canonical") is True:
        for row in runs:
            if (
                row.get("batch_id") == record.get("batch_id") and
                row.get("revision_id") == record.get("revision_id") and
                row.get("canonical") is True
            ):
                row["canonical"] = False
                row["superseded_by"] = record["qa_run_id"]
                row["registry_status"] = "SUPERSEDED"
    runs.append(record)
    REGISTRY_JSON.write_text(json.dumps(registry, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    fields = [
        "qa_run_id", "batch_id", "revision_id", "parent_run_id", "source_hash",
        "rubric_hash", "scope_hash", "product_count", "status", "canonical",
        "supersedes", "superseded_by", "registry_status", "created_at",
    ]
    with REGISTRY_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for row in runs:
            writer.writerow({field: row.get(field, "") for field in fields})


def rating_row(rating, weight):
    if rating == "FULL":
        return weight
    if rating == "PARTIAL":
        return weight / 2
    return 0


def load_manual_image_evidence(path: Path | None):
    if not path or not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("images", data)

def load_evidence_bundle(path: Path | None):
    if not path or not path.exists():
        return {}, {}, {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("images", {}), {str(x.get("handle")): x for x in data.get("products", [])}, data

def load_manual_criteria_evidence(path: Path | None):
    """Load independent per-product rubric ratings.

    Schema: {"products": {"handle": {"P1": {"rating": "FULL",
    "reason": "...", "evidence_refs": ["..."]}}}}.  Missing ratings
    intentionally remain NOT_CHECKED; workbook field presence is not QA.
    """
    if not path or not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("products", data)

def evidence_doc_metadata_valid(data, *, revision, source_hash, expected_method):
    """Top-level provenance required for audit-ready re-QA evidence."""
    if not isinstance(data, dict):
        return False
    reviewer = data.get("reviewer") if isinstance(data.get("reviewer"), dict) else {}
    reviewer_id = reviewer.get("reviewer_id") or data.get("reviewer_id")
    reviewed_at = reviewer.get("reviewed_at") or data.get("reviewed_at")
    return (
        data.get("revision") == revision and
        bool(str(data.get("batch_id") or "").strip()) and
        data.get("review_method") == expected_method and
        bool(str(reviewer_id or "").strip()) and
        bool(str(reviewed_at or "").strip()) and
        data.get("source_workbook_sha256") == source_hash
    )


def reason_substantive(reason: str, handle: str) -> bool:
    """Reject generic evidence statements masquerading as product QA."""
    text = str(reason or "").strip()
    lowered = text.lower()
    boilerplate = (
        "compared the proposed field with the locked workbook, page snapshot, and product evidence",
        "this finding is specific to the scoped record",
    )
    if len(text) < 90 or any(x in lowered for x in boilerplate):
        return False
    handle_terms = [x for x in re.split(r"[-_\s]+", handle.lower()) if len(x) >= 4]
    return not handle_terms or any(term in lowered for term in handle_terms[:8])


def validate_criteria_document(path: Path | None, scope, revision, source_hash):
    """Validate reviewer declaration and reject copied/placeholder reviews."""
    if not path or not path.exists():
        return {}, set()
    data = json.loads(path.read_text(encoding="utf-8"))
    if not evidence_doc_metadata_valid(
        data,
        revision=revision,
        source_hash=source_hash,
        expected_method="INDEPENDENT_PRODUCT_REVIEW",
    ):
        return data.get("products", {}), set(scope)
    products = data.get("products", {})
    invalid = set()
    seen = {}
    for handle in scope:
        for cid, review in (products.get(handle, {}) or {}).items():
            if not isinstance(review, dict) or review.get("rating") == "NOT_CHECKED":
                continue
            reason = str(review.get("reason") or "").strip()
            refs = review.get("evidence_refs")
            valid_refs = isinstance(refs, list) and len(refs) >= 2 and all(isinstance(x, str) and x.strip() for x in refs)
            if not reason_substantive(reason, handle) or not valid_refs:
                invalid.add(handle)
            key = (cid, reason.lower())
            if reason:
                seen.setdefault(key, []).append(handle)
    for handles in seen.values():
        if len(handles) > 1:
            invalid.update(handles)
    return products, invalid

def criterion_evidence_valid(review):
    """Require a substantive, traceable independent review record."""
    if not isinstance(review, dict):
        return False
    reason = str(review.get("reason") or "").strip()
    refs = review.get("evidence_refs")
    if len(reason) < 20 or not isinstance(refs, list) or not refs:
        return False
    for ref in refs:
        if not isinstance(ref, str) or not ref.strip():
            return False
        if ref.startswith("workbook:") and not Path(ref.split(":", 1)[1]).exists():
            return False
        if ref.startswith("live_snapshot:") and not Path(ref.split(":", 1)[1]).exists():
            return False
    return True

def image_evidence_valid(record):
    """Require per-image observation, method, timestamp and local evidence."""
    if not isinstance(record, dict):
        return False
    if str(record.get("qa_observation") or "").strip() == "":
        return False
    if str(record.get("check_method") or "").strip() in {"", "NO_INDEPENDENT_IMAGE_REVIEW"}:
        return False
    if str(record.get("checked_at") or "").strip() == "":
        return False
    if record.get("local_file_exists") is not True or not record.get("local_sha256"):
        return False
    local = record.get("local_evidence_file")
    return bool(local) and Path(str(local)).exists()


def validate_image_document(data, *, revision, source_hash, scope):
    if not data:
        return set(scope)
    if not evidence_doc_metadata_valid(
        data,
        revision=revision,
        source_hash=source_hash,
        expected_method="FULL_RESOLUTION_INDIVIDUAL_IMAGE_REVIEW",
    ):
        return set(scope)
    return set()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--revision", required=True, help="Revision id, e.g. R001")
    ap.add_argument("--source-workbook", type=Path)
    ap.add_argument("--manual-image-evidence", type=Path,
                    help="JSON with independently reviewed image ratings")
    ap.add_argument("--manual-criteria-evidence", type=Path,
                    help="JSON with independently reviewed P1..E1 ratings")
    ap.add_argument("--qa-run-id", help="Unique QA run id; defaults to a timestamped id to preserve history")
    ap.add_argument("--parent-run-id", default="", help="Prior QA run superseded or continued by this run")
    ap.add_argument("--supersedes", default="", help="Comma-separated prior QA run ids superseded by this run")
    ap.add_argument("--canonical", action="store_true", help="Register this run as the canonical run for its batch/revision")
    ap.add_argument("--allow-nonstandard-scope", action="store_true", help="Allow a scope other than 10 products for last-batch or targeted diagnostics")
    args = ap.parse_args()
    revision = args.revision.upper()
    if not re.fullmatch(r"[RB]\d{3}", revision):
        raise SystemExit("--revision must look like R001 or B001")
    source = args.source_workbook or (BASE / f"resutls/{SHOP_DOMAIN}/{RUN_ID}/revisions/{revision}/SEO_Product_Optimization_revision_{revision}.xlsx")
    if not source.exists():
        raise FileNotFoundError(source)
    qa_run_id = args.qa_run_id or f"QA-{datetime.now(TZ):%Y%m%d-%H%M%S}-REV-{revision}"
    qa_batch_id = f"REV_{revision}"
    out_dir = BASE / f"resutls/{SHOP_DOMAIN}/{RUN_ID}/qa/{qa_run_id}"
    work_dir = BASE / f"seo_runs/{SHOP_DOMAIN}/{RUN_ID}/qa/{qa_run_id}"
    out_dir.mkdir(parents=True, exist_ok=True)
    work_dir.mkdir(parents=True, exist_ok=True)
    checked_at = datetime.now(TZ).replace(microsecond=0).isoformat()
    source_hash = sha256(source)
    rubric_hash = sha256(RUBRIC_PATH) if RUBRIC_PATH.exists() else ""

    wb = load_workbook(source, read_only=True, data_only=True)
    revlog = rows(wb["Revision_Log"])
    keys = [val(r, "handle") for r in revlog if val(r, "revision_batch_id") == revision]
    if not keys:
        raise RuntimeError(f"No Revision_Log scope found for {revision}")
    batch_ids = sorted({val(r, "batch_id") for r in revlog if val(r, "revision_batch_id") == revision and val(r, "batch_id")})
    batch_id = batch_ids[0] if len(batch_ids) == 1 else ""
    if len(batch_ids) != 1:
        raise RuntimeError(f"Revision scope must map to exactly one batch_id; got {batch_ids}")
    if len(keys) != len(set(keys)):
        raise RuntimeError(f"Duplicate product keys in Revision_Log scope for {revision}")
    if len(keys) != 10 and not args.allow_nonstandard_scope:
        raise RuntimeError(
            f"Nonstandard scope for {revision}/{batch_id}: {len(keys)} products. "
            "Expected 10 for a normal batch rerun; pass --allow-nonstandard-scope only for explicit last-batch/diagnostic runs."
        )
    current_scope_hash = scope_hash(keys)
    products = {val(r, "Handle"): r for r in rows(wb["SEO_Products"]) if val(r, "Handle") in keys}
    # Cross-product uniqueness gate: customer-facing SEO fields must be
    # differentiated within the locked scope. This runs before scoring.
    duplicate_fields = ("title_proposed", "meta_title_seo",
                        "meta_description_seo", "description_proposed")
    duplicate_handles = {f: {} for f in duplicate_fields}
    for field in duplicate_fields:
        groups = {}
        for h, row in products.items():
            text = val(row, field).strip().casefold()
            if text:
                groups.setdefault(text, []).append(h)
        duplicate_handles[field] = {text: hs for text, hs in groups.items() if len(hs) > 1}
    evidence = {val(r, "evidence_id"): r for r in rows(wb["Product_Evidence"])}
    images = {k: [] for k in keys}
    for r in rows(wb["Image_Audit"]):
        if val(r, "Handle") in images:
            images[val(r, "Handle")].append(r)
    if len(products) != len(keys):
        raise RuntimeError(f"Revision scope mismatch: log={len(keys)}, products={len(products)}")
    extra_product_rows = [val(r, "Handle") for r in rows(wb["SEO_Products"]) if val(r, "Handle") and val(r, "Handle") not in keys]
    if extra_product_rows:
        raise RuntimeError(f"SEO_Products contains {len(extra_product_rows)} product(s) outside frozen scope for {revision}/{batch_id}")
    log_revision = {val(r, "handle"): val(r, "revision") for r in revlog if val(r, "revision_batch_id") == revision}
    manual_images, page_evidence, image_doc = load_evidence_bundle(args.manual_image_evidence)
    invalid_image_handles = validate_image_document(image_doc, revision=revision, source_hash=source_hash, scope=keys)
    manual_criteria, invalid_criteria_handles = validate_criteria_document(args.manual_criteria_evidence, keys, revision, source_hash)

    qproducts, qcriteria, qimages, qissues = [], [], [], []
    status_counts = {"QA_PASS": 0, "QA_REVISE": 0, "QA_FAIL": 0, "QA_INCOMPLETE": 0}
    scores = []
    for handle in keys:
        p = products[handle]
        ev = evidence.get(val(p, "evidence_id"), {})
        blob = " ".join(val(p, k) for k in ["title_proposed", "meta_title_seo", "meta_description_seo", "description_proposed", "description_proposed_html"]).lower()
        refs_parts = [f"source_workbook_sha256:{source_hash}", f"evidence_id:{val(p, 'evidence_id')}", f"revision_log:{revision}"]
        if args.manual_image_evidence:
            refs_parts.append(f"independent_image_evidence:{args.manual_image_evidence}")
        if args.manual_criteria_evidence:
            refs_parts.append(f"independent_criteria_evidence:{args.manual_criteria_evidence}")
        refs = "; ".join(refs_parts)
        issues = []
        critical = major = minor = 0
        page = page_evidence.get(handle, {})
        page_complete = (page.get("http_status") == 200 and page.get("json_status") == 200
                         and bool(page.get("html_snapshot")) and bool(page.get("html_sha256"))
                         and page.get("image_url_set_match") is True)
        if not page_complete:
            issues.append([f"{revision}-QA-LPAGE-{handle[:20]}", handle, "", "LIMITATION", "storefront", "", "", "Missing or invalid independent storefront capture for this product.", "Capture the product page and JSON again, verify snapshot/hash/image set, then rerun QA.", refs, "Re-run after valid page evidence is available."])

        def issue(sev, field, observed, reason, fix):
            nonlocal critical, major, minor
            n = sum(1 for i in issues if i[3] == sev) + 1
            iid = f"{revision}-QA-{sev[0]}{n:03d}-{handle[:20]}"
            issues.append([iid, handle, "", sev, field, observed, reason, observed, fix, refs, "Re-run after correction."])
            if sev == "CRITICAL": critical += 1
            elif sev == "MAJOR": major += 1
            else: minor += 1

        internal_hits = [pat for pat in INTERNAL_PATTERNS if pat in blob]
        if internal_hits:
            issue("MAJOR", "proposed_fields", ", ".join(internal_hits), "Internal QA/process language is customer-visible copy.", "Remove internal notes and retain only customer-facing product copy.")
        copy_quality_hits = content_quality_findings(blob)
        if copy_quality_hits:
            issue("MAJOR", "proposed_fields", ", ".join(copy_quality_hits), "Malformed or grammatically broken customer-facing copy was detected.", "Rewrite the affected title/meta/description as a complete natural sentence, then rerun content lint.")
        # Semantic identity gate: proposed copy must remain aligned with the
        # product handle/live identity; never accept a generic fallback theme.
        identity_terms = []
        if "boat" in handle: identity_terms=["boat","nautical","welcome"]
        elif "football" in handle: identity_terms=["football","sport"]
        elif "koi" in handle: identity_terms=["koi","fish"]
        elif "dragon" in handle: identity_terms=["dragon"]
        elif "dog-paw" in handle or "dog-photo" in handle or "pet-dog" in handle: identity_terms=["dog","paw","pet"]
        elif "halloween" in handle or "witch" in handle or "ghost" in handle: identity_terms=["halloween","ghost","witch"]
        if identity_terms and not any(t in blob for t in identity_terms):
            issue("MAJOR", "identity", "missing expected identity terms", "Proposed copy does not semantically match the handle and source product identity.", "Rewrite copy using the verified product identity and keyword.")
        forbidden_identity = (["ghost","halloween","football","sport"] if "boat" in handle else ["ghost","halloween"] if "football" in handle else [])
        bad_forbidden=[t for t in forbidden_identity if t in blob]
        if bad_forbidden:
            issue("MAJOR", "identity", ", ".join(bad_forbidden), "Proposed copy contains a conflicting product identity.", "Remove conflicting theme terms and restore the verified product identity.")
        if val(p, "revision") != log_revision.get(handle, ""):
            issue("MAJOR", "revision", f"product={val(p,'revision')}; log={log_revision.get(handle,'')}", "SEO_Products revision does not match Revision_Log provenance.", "Align both revision fields before QA.")
        duplicate_hits = [field for field in duplicate_fields
                          if any(handle in hs for hs in duplicate_handles[field].values())]
        if duplicate_hits:
            issue("MAJOR", "uniqueness", ", ".join(duplicate_hits),
                  "Customer-facing SEO copy is duplicated within the locked scope; differentiation is required.",
                  "Rewrite the duplicated fields so each product variant has distinct, accurate customer-facing copy.")
        unsupported = [term for term in ["outdoor", "waterproof", "machine washable", "non-slip", "anti-slip"] if term in blob and term not in (val(ev, "verified_product_facts") + " " + val(ev, "short_source_excerpt")).lower()]
        if unsupported:
            issue("MAJOR", "proposed_fields", ", ".join(unsupported), "Potential product claim is not supported by the linked evidence record.", "Remove or verify each claim against the source/export before approval.")
        # Record known evidence limitations explicitly.  They do not lower a
        # product to REVISE by themselves, but must remain visible for audit.
        demand_note = val(p, "keyword_demand_evidence").lower()
        demand_validation = val(p, "keyword_validation_status").upper()
        if ((demand_note and ("serp" in demand_note or "no paid volume" in demand_note or "no volume" in demand_note))
                or "SERP_ONLY" in demand_validation or "NO_VOLUME" in demand_validation) and "volume verified" not in demand_note.lower():
            issues.append([f"{revision}-QA-LK3-{handle[:20]}", handle, "", "LIMITATION", "K3", "SERP_ONLY/no_volume", "Only public SERP or semantic evidence is available; demand volume is not verified.", "", "Obtain Search Console or a cited keyword source if demand verification is required.", refs, "Re-run if stronger demand evidence is added."])
        export_state = val(p, "source_exported_at").strip().lower()
        export_ref = val(p, "source_export_ref").lower()
        if not val(p, "source_export_ref") or not val(p, "source_exported_at") or export_state in {"unknown", "unknown_public_json_capture"} or "admin" not in export_ref and "export" not in export_ref:
            issues.append([f"{revision}-QA-LADMIN-{handle[:20]}", handle, "", "LIMITATION", "admin/export", "before_state_required", "Admin/export before-state is not supplied in this research workbook.", "", "Capture and reconcile Shopify admin/export values before deployment.", refs, "Recheck during implementation QA."])

        criteria = []
        criterion_review = manual_criteria.get(handle, {}) if isinstance(manual_criteria, dict) and handle not in invalid_criteria_handles else {}
        def add(cid, rating, reason, evidence_ref=""):
            earned = rating_row(rating, WEIGHTS[cid])
            criteria.append((cid, rating, reason, earned, evidence_ref))
        for cid in ("P1", "P2", "K1", "K2", "K3", "T1", "T2", "D1", "D2", "E1"):
            review = criterion_review.get(cid, {}) if isinstance(criterion_review, dict) else {}
            rating = review.get("rating", "NOT_CHECKED") if isinstance(review, dict) else "NOT_CHECKED"
            if rating not in {"FULL", "PARTIAL", "FAIL", "NOT_CHECKED"}:
                raise RuntimeError(f"Invalid {cid} rating for {handle}: {rating}")
            if rating != "NOT_CHECKED" and not criterion_evidence_valid(review):
                rating = "NOT_CHECKED"
            reason = review.get("reason", "Independent criterion evidence was not supplied or is not traceable; workbook metadata alone is not QA.") if isinstance(review, dict) else "Independent criterion evidence was not supplied or is not traceable; workbook metadata alone is not QA."
            eref = "; ".join(str(x) for x in review.get("evidence_refs", [])) if isinstance(review, dict) else ""
            # Automated safety gates can only lower a manual rating, never raise it.
            if cid == "P2" and unsupported:
                rating, reason = "FAIL", f"Unsupported claim detected: {', '.join(unsupported)}."
            if cid == "D2" and internal_hits:
                rating, reason = "FAIL", "Internal QA/process language is customer-visible copy."
            if cid in {"D1", "D2"} and copy_quality_hits:
                rating, reason = "FAIL", f"Malformed customer-facing copy detected: {', '.join(copy_quality_hits)}."
            if duplicate_hits and cid in {"T1", "T2", "D1", "D2"}:
                rating, reason = "FAIL", f"Duplicate customer-facing copy detected in scope fields: {', '.join(duplicate_hits)}."
            if (identity_terms and not any(t in blob for t in identity_terms)) or bad_forbidden or val(p, "revision") != log_revision.get(handle, ""):
                rating, reason = "FAIL", "Identity or revision consistency gate failed for this product."
            add(cid, rating, reason, eref)

        for img in sorted(images[handle], key=lambda x: int(val(x, "image_number") or 0)):
            ikey = f"{handle}__img_{int(val(img, 'image_number') or 0):02d}"
            m = manual_images.get(ikey, {}) if isinstance(manual_images, dict) and handle not in invalid_image_handles else {}
            if not image_evidence_valid(m):
                m = {}
            ratings = {cid: m.get(cid, "NOT_CHECKED") for cid in IMAGE_WEIGHTS}
            for cid, rating in ratings.items():
                if rating not in {"FULL", "PARTIAL", "FAIL", "NOT_CHECKED"}:
                    raise RuntimeError(f"Invalid {cid} rating for {ikey}: {rating}")
            img_points = sum(rating_row(ratings[c], IMAGE_WEIGHTS[c]) for c in IMAGE_WEIGHTS)
            img_assessed = sum(IMAGE_WEIGHTS[c] for c in IMAGE_WEIGHTS if ratings[c] != "NOT_CHECKED")
            img_final = img_points if img_assessed == 100 else ""
            qimages.append([handle, ikey, val(img, "image_url"), val(img, "image_url_export") or val(img, "image_url"), val(img, "media_id"), val(img, "variant"), val(img, "image_location"), m.get("check_method", "INDEPENDENT_REVIEW_JSON") if m else "NO_INDEPENDENT_IMAGE_REVIEW", m.get("checked_at", checked_at) if m else checked_at, m.get("qa_observation", val(img, "observed_visual_details")) if m else val(img, "observed_visual_details"), val(img, "observed_visual_details"), val(img, "alt_action"), m.get("alt_effective", val(img, "alt_proposed")) if m else val(img, "alt_proposed"), ratings["IM1"], ratings["IM2"], ratings["IM3"], ratings["IM4"], img_points, img_assessed, img_final, img_points, img_points + (100-img_assessed), "", refs + f"; image_key:{ikey}"])
        image_complete = bool(images[handle]) and all(
            handle not in invalid_image_handles and
            image_evidence_valid(manual_images.get(f"{handle}__img_{int(val(i, 'image_number') or 0):02d}", {})) and
            all(manual_images.get(f"{handle}__img_{int(val(i, 'image_number') or 0):02d}", {}).get(c) in {"FULL", "PARTIAL", "FAIL"} for c in IMAGE_WEIGHTS)
            for i in images[handle]
        ) if isinstance(manual_images, dict) else False
        if image_complete:
            n_img = len(images[handle])
            avg_img_points = sum(sum(rating_row(manual_images[f"{handle}__img_{int(val(i, 'image_number') or 0):02d}"][c], IMAGE_WEIGHTS[c]) for c in IMAGE_WEIGHTS) for i in images[handle]) / n_img
            avg_img_assessed = sum(sum(IMAGE_WEIGHTS[c] for c in IMAGE_WEIGHTS if manual_images[f"{handle}__img_{int(val(i, 'image_number') or 0):02d}"][c] != "NOT_CHECKED") for i in images[handle]) / n_img
            any_image_fail = any(manual_images[f"{handle}__img_{int(val(i, 'image_number') or 0):02d}"][c] == "FAIL" for i in images[handle] for c in IMAGE_WEIGHTS)
            image_rating = "FAIL" if any_image_fail else "FULL" if avg_img_points == 100 and avg_img_assessed == 100 else "PARTIAL"
            # I1 is a derived 20-point criterion: 0.20 * average image points/weight.
            criteria.append(("I1", image_rating, "Derived as 0.20 × average per-image verified points and assessed weight.", 0.20 * avg_img_points, 0.20 * avg_img_assessed))
            if any_image_fail:
                issue("MAJOR", "images", "One or more IM1–IM4 ratings are FAIL", "An image-level failure prevents the image criterion from being accepted as fully correct.", "Correct the affected image mapping, observation or alt and rerun image QA.")
        else:
            add("I1", "NOT_CHECKED", "Image score is derived only from independent per-image ratings; workbook metadata alone is not visual QA.")
        if not image_complete:
            iid = f"{revision}-QA-LIMG-{handle[:20]}"
            issues.append([iid, handle, "", "LIMITATION", "images", "", "No complete independent image-review JSON was supplied.", "Image criteria remain NOT_CHECKED.", "Review every in-scope image and provide IM1–IM4 ratings.", refs, "Re-run with --manual-image-evidence."])

        for cid, rating, reason, earned, eref in criteria:
            assessed_row = (WEIGHTS[cid] if cid != "I1" else float(eref or 0)) if rating != "NOT_CHECKED" else 0
            criterion_refs = refs + (f"; criterion_evidence:{eref}" if eref else "")
            qcriteria.append([handle, cid, WEIGHTS[cid], "DERIVED" if cid == "I1" and rating != "NOT_CHECKED" else "CHECKED" if rating != "NOT_CHECKED" else "NOT_CHECKED", rating, earned, assessed_row, reason, criterion_refs, ""])
        # Image criterion was appended last; ensure it is represented once.
        cid, rating, reason, earned, eref = criteria[-1]
        assessed = sum((WEIGHTS[c] if c != "I1" else float(eref or 0)) for c, r, _, _, eref in criteria if r != "NOT_CHECKED")
        verified = sum(e for c, r, _, e, _ in criteria) + (earned if cid == "I1" else 0)
        # criteria includes I1 as the last item, so do not double-add it.
        verified = sum(e for _, _, _, e, _ in criteria)
        upper = verified + (100 - assessed)
        # A fully assessed product still receives its numeric score when it has
        # MAJOR issues; the status gate below must classify it as QA_REVISE.
        # Only missing evidence blocks final_score.
        final = verified if assessed == 100 and image_complete and page_complete else ""
        # Rubric precedence: CRITICAL -> FAIL; insufficient evidence -> INCOMPLETE;
        # only a fully assessed product can be REVISE/PASS by score.
        status = ("QA_FAIL" if critical or any(i[3] == "CRITICAL" for i in issues)
                  else "QA_INCOMPLETE" if final == ""
                  else "QA_REVISE" if major or final < 85
                  else "QA_PASS")
        status_counts[status] += 1
        if final != "": scores.append(final)
        qproducts.append([handle, val(p, "url"), val(p, "revision"), verified, assessed, verified, upper, final, status, "SERP_ONLY" if val(p, "keyword_demand_evidence") else "UNVERIFIABLE", len(images[handle]), len(images[handle]) if image_complete else 0, "TRUE" if image_complete else "UNKNOWN", "100%" if image_complete else "UNKNOWN", critical, major, minor, "; ".join(i[0] for i in issues), refs])
        qissues.extend(issues)

    out_xlsx = out_dir / f"SEO_QA_{qa_batch_id}.xlsx"
    out_md = out_dir / f"SEO_QA_{qa_batch_id}.md"
    wbout = Workbook(); ws = wbout.active; ws.title = "QA_Summary"
    header(ws, ["metric", "value", "definition"])
    batch_status = ("QA_PASS" if status_counts["QA_PASS"] == len(keys)
                    else "QA_FAIL" if status_counts["QA_FAIL"]
                    else "QA_INCOMPLETE" if status_counts["QA_INCOMPLETE"]
                    else "QA_REVISE")
    issue_counts = {
        "CRITICAL": sum(1 for r in qissues if r[3] == "CRITICAL"),
        "MAJOR": sum(1 for r in qissues if r[3] == "MAJOR"),
        "MINOR": sum(1 for r in qissues if r[3] == "MINOR"),
        "LIMITATION": sum(1 for r in qissues if r[3] == "LIMITATION"),
    }
    evidence_maturity = "SERP_ONLY" if any(
        r[4] == "K3" and r[5] == "SERP_ONLY/no_volume" for r in qissues
    ) or any(
        "SERP_ONLY" in val(p, "keyword_validation_status").upper()
        or "NO_VOLUME" in val(p, "keyword_validation_status").upper()
        for p in products.values()
    ) else "DEMAND_SUPPORTED"
    summary = [
        ("rubric_version", "prompt_qa.md v1.0"),
        ("rubric_hash", rubric_hash),
        ("qa_run_id", qa_run_id),
        ("batch_id", batch_id),
        ("revision_id", revision),
        ("parent_run_id", args.parent_run_id),
        ("source_workbook", str(source)),
        ("source_workbook_sha256", source_hash),
        ("scope_hash", current_scope_hash),
        ("scope_product_count", len(keys)),
        ("scope_product_keys", json.dumps(keys, ensure_ascii=False)),
        ("scope_image_count", len(qimages)),
        ("content_qa_status", batch_status),
        ("evidence_maturity", evidence_maturity),
        ("batch_final_score", round(mean(scores), 1) if len(scores) == len(keys) else ""),
        ("QA_PASS", status_counts["QA_PASS"]),
        ("QA_REVISE", status_counts["QA_REVISE"]),
        ("QA_FAIL", status_counts["QA_FAIL"]),
        ("QA_INCOMPLETE", status_counts["QA_INCOMPLETE"]),
        ("CRITICAL", issue_counts["CRITICAL"]),
        ("MAJOR", issue_counts["MAJOR"]),
        ("MINOR", issue_counts["MINOR"]),
        ("LIMITATION", issue_counts["LIMITATION"]),
        ("manual_image_evidence", str(args.manual_image_evidence) if args.manual_image_evidence else ""),
        ("manual_criteria_evidence", str(args.manual_criteria_evidence) if args.manual_criteria_evidence else ""),
        ("canonical", bool(args.canonical)),
        ("supersedes", args.supersedes),
        ("approval_status", "NOT_APPROVED_NOT_DEPLOYED"),
    ]
    for k, v in summary: append(ws, [k, v, "Evidence-driven re-QA; blank final score means insufficient evidence."])
    ws = wbout.create_sheet("QA_Products"); header(ws, ["product_key","url","revision","verified_points","assessed_weight","score_lower_bound","score_upper_bound","final_score","qa_status","keyword_evidence_level","images_expected","images_checked","image_inventory_complete","image_coverage","critical_count","major_count","minor_count","issue_refs","evidence_refs"])
    for r in qproducts: append(ws, r)
    ws = wbout.create_sheet("QA_Criteria"); header(ws, ["product_key","criterion_id","weight","assessment","rating","earned_points","assessed_weight","reason","evidence_refs","issue_refs"])
    for r in qcriteria: append(ws, r)
    ws = wbout.create_sheet("QA_Images"); header(ws, ["product_key","qa_image_key","image_url_source","image_url_workbook","media_id","variant","image_location","check_method","checked_at","qa_observation","submitted_observation","alt_action","alt_effective","IM1","IM2","IM3","IM4","image_verified_points","image_assessed_weight","image_final_score","image_score_lower_bound","image_score_upper_bound","issue_refs","evidence_refs"])
    for r in qimages: append(ws, r)
    ws = wbout.create_sheet("QA_Issues"); header(ws, ["issue_id","product_key","qa_image_key","severity","field","submitted_value","source_observation","reason","recommended_fix","supporting_evidence","recheck_condition"])
    for r in qissues: append(ws, r)
    for s in wbout.worksheets: finish(s)
    wbout.save(out_xlsx)
    manifest = {
        "rubric_version": "prompt_qa.md v1.0",
        "rubric_hash": rubric_hash,
        "qa_run_id": qa_run_id,
        "qa_batch_id": qa_batch_id,
        "batch_id": batch_id,
        "revision_id": revision,
        "parent_run_id": args.parent_run_id,
        "source_workbook": str(source),
        "source_workbook_sha256": source_hash,
        "scope_hash": current_scope_hash,
        "batch_product_keys": keys,
        "status_counts": status_counts,
        "issue_counts": issue_counts,
        "batch_status": batch_status,
        "content_qa_status": batch_status,
        "evidence_maturity": evidence_maturity,
        "canonical": bool(args.canonical),
        "supersedes": [x.strip() for x in args.supersedes.split(",") if x.strip()],
        "manual_image_evidence": str(args.manual_image_evidence) if args.manual_image_evidence else None,
        "manual_criteria_evidence": str(args.manual_criteria_evidence) if args.manual_criteria_evidence else None,
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
        "created_at": checked_at,
    }
    (work_dir / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    registry_record = {
        "qa_run_id": qa_run_id,
        "batch_id": batch_id,
        "revision_id": revision,
        "parent_run_id": args.parent_run_id,
        "source_hash": source_hash,
        "rubric_hash": rubric_hash,
        "scope_hash": current_scope_hash,
        "product_count": len(keys),
        "status": batch_status,
        "canonical": bool(args.canonical),
        "supersedes": args.supersedes,
        "superseded_by": "",
        "registry_status": "CANONICAL" if args.canonical else "RECORDED",
        "created_at": checked_at,
    }
    update_registry(registry_record)
    out_md.write_text(
        "\n".join([
            f"# Evidence-driven Re-QA {revision}",
            "",
            "## Audit Identity",
            "",
            f"- QA run id: `{qa_run_id}`",
            f"- Batch id: `{batch_id}`",
            f"- Revision id: `{revision}`",
            f"- Parent run id: `{args.parent_run_id or 'NONE'}`",
            f"- Canonical: `{bool(args.canonical)}`",
            f"- Supersedes: `{args.supersedes or 'NONE'}`",
            f"- Created at: `{checked_at}`",
            "",
            "## Source Lock",
            "",
            f"- Source workbook: `{source}`",
            f"- Source workbook SHA-256: `{source_hash}`",
            f"- Rubric: `prompt_qa.md v1.0`",
            f"- Rubric SHA-256: `{rubric_hash}`",
            f"- Scope hash: `{current_scope_hash}`",
            f"- Product keys: `{json.dumps(keys, ensure_ascii=False)}`",
            "",
            "## Evidence",
            "",
            f"- Criteria evidence: `{str(args.manual_criteria_evidence) if args.manual_criteria_evidence else 'NONE'}`",
            f"- Image evidence: `{str(args.manual_image_evidence) if args.manual_image_evidence else 'NONE'}`",
            f"- Image rows checked: `{len(qimages)}`",
            "",
            "## Result",
            "",
            f"- Content QA status: `{batch_status}`",
            f"- Evidence maturity: `{evidence_maturity}`",
            f"- Batch final score: `{round(mean(scores), 1) if len(scores) == len(keys) else ''}`",
            f"- Product status counts: `{status_counts}`",
            f"- Issue counts: `{issue_counts}`",
            f"- Approval status: `NOT_APPROVED_NOT_DEPLOYED`",
            "",
            "A blank `final_score` means the rubric did not have complete evidence. `QA_PASS` is not `APPROVED` and does not authorize Shopify import or deployment.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps({"qa_xlsx":str(out_xlsx),"qa_md":str(out_md),"batch_status":batch_status,"status_counts":status_counts}, indent=2))


if __name__ == "__main__":
    main()

import csv
import html
import json
import re
from collections import OrderedDict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
ROOT = Path("seo_runs") / SHOP / RUN_ID
RESULTS = Path("resutls") / SHOP / RUN_ID
BATCH_ID = "B005"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()

SERP_REFS = [
    "https://www.amazon.com/Personalized-Football-Bedroom-Decorations-Washable/dp/B0F52SNSGG",
    "https://www.etsy.com/no-en/market/personalized_football_rug",
    "https://www.ebay.com/b/Football-Rug-In-Area-Rugs/262983/bn_7022379243",
    "https://www.walmart.com/c/kp/nfl-rug",
    "https://www.etsy.com/market/dragon_rugs",
    "https://www.etsy.com/market/dragon_area_rugs",
    "https://www.etsy.com/market/dragon_shaped_rug",
    "https://printyrug.com/",
    "https://chillgen.com/",
]

DESIGNS = OrderedDict(
    [
        (
            "personalized-football-shaped-rug-non-slip-washable-sport-design-03",
            {
                "theme": "football player shaped rug",
                "title": "Personalized Football Player Rug",
                "meta": "Customize a football player shaped rug with a name and number, sporty game-room style, non-slip backing visuals, and washable-care graphics.",
                "primary": "personalized football player rug",
                "secondary": "football shaped rug, custom football rug with name, sports room rug",
                "season": "EVERGREEN_SPORTS",
                "obs": "custom-shaped football player rug viewed from behind on green turf, with name William, number 18, football in hand, and color variants in yellow, blue, red, and black",
            },
        ),
        (
            "personalized-football-shaped-rug-non-slip-washable-sport-design-02",
            {
                "theme": "graffiti football shaped rug",
                "title": "Personalized Graffiti Football Rug",
                "meta": "Design a personalized graffiti football rug with a name and number, dripping paint artwork, non-slip backing visuals, and washable-care graphics.",
                "primary": "personalized football rug with name",
                "secondary": "graffiti football rug, football shaped rug, custom sports rug",
                "season": "EVERGREEN_SPORTS",
                "obs": "custom-shaped graffiti-style football rug with a brown football, paint-drip artwork, name Jackson, number 99, and red, teal, and pink color variants",
            },
        ),
        (
            "personalized-football-shaped-rug-non-slip-washable-sport-design-01",
            {
                "theme": "football badge shaped rug",
                "title": "Personalized Football Badge Rug",
                "meta": "Personalize a football badge shaped rug with a name, helmet-and-ball artwork, game-room decor appeal, non-slip backing visuals, and care graphics.",
                "primary": "personalized football badge rug",
                "secondary": "football helmet rug, custom football name rug, sports bedroom rug",
                "season": "EVERGREEN_SPORTS",
                "obs": "badge-shaped football rug with red helmet, football, green field panel, red banner, and personalized name Andrew",
            },
        ),
        (
            "personalized-football-shaped-rug-non-slip-washable-sport-design-05",
            {
                "theme": "football number shaped rug",
                "title": "Personalized Football Number Rug",
                "meta": "Create a personalized football number rug with custom name detail, stitched football-number artwork, non-slip backing visuals, and easy-care graphics.",
                "primary": "personalized football number rug",
                "secondary": "custom football rug with number, football shaped rug, sports fan room rug",
                "season": "EVERGREEN_SPORTS",
                "obs": "custom-shaped rug built around large number 34 with football-stitch details, green border, and personalized name Andrew",
            },
        ),
        (
            "personalized-football-shaped-rug-non-slip-washable-sport-carpet-a04",
            {
                "theme": "football ball shaped rug",
                "title": "Personalized Football Ball Rug",
                "meta": "Customize a football ball shaped rug with a name and number, bold sports artwork, non-slip backing visuals, and washable-care room mockups.",
                "primary": "personalized football ball rug",
                "secondary": "custom football rug with name, football shaped carpet, sports room rug",
                "season": "EVERGREEN_SPORTS",
                "obs": "custom-shaped rug with a large brown football, white laces, name Jackson, and number 99 shown in room mockups and care graphics",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-02",
            {
                "theme": "blue dragon book shaped rug",
                "title": "Personalized Blue Dragon Book Rug",
                "meta": "Personalize a blue dragon book shaped rug with custom name artwork, fantasy library style, anti-slip backing visuals, and size chart images.",
                "primary": "personalized dragon book rug",
                "secondary": "blue dragon rug, fantasy library rug, custom dragon shaped rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped fantasy rug with a blue dragon curled over an open book, moonlit blue accents, and personalized name Jasmine",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-08",
            {
                "theme": "purple dragon floral book rug",
                "title": "Personalized Purple Dragon Rug",
                "meta": "Choose a personalized purple dragon rug with stacked-book fantasy artwork, floral accents, custom name detail, and anti-slip backing visuals.",
                "primary": "personalized purple dragon rug",
                "secondary": "dragon book rug, fantasy floral rug, custom shaped dragon rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug showing a purple dragon on stacked books, purple flowers, green vines, and personalized name Sophia",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-05",
            {
                "theme": "dragon open book shaped rug",
                "title": "Personalized Dragon Open Book Rug",
                "meta": "Add a personalized dragon open book rug with teal-and-orange fantasy artwork, custom name detail, room mockups, and anti-slip backing visuals.",
                "primary": "personalized dragon open book rug",
                "secondary": "fantasy book rug, dragon shaped area rug, custom name dragon rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with teal dragon and orange wings lying across an open illustrated book, plus personalized name Sophia",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-07",
            {
                "theme": "moon dragon book rug",
                "title": "Personalized Moon Dragon Rug",
                "meta": "Personalize a moon dragon book rug with stacked books, roses, gothic fantasy artwork, custom name detail, and anti-slip backing visuals.",
                "primary": "personalized moon dragon rug",
                "secondary": "dragon library rug, fantasy dragon rug, custom shaped book rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with blue-and-gold dragon in front of a full moon, stacked books, red and green roses, and personalized name Sophia",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-03",
            {
                "theme": "dragon library floral rug",
                "title": "Personalized Dragon Library Rug",
                "meta": "Create a personalized dragon library rug with stacked books, orange flowers, custom name artwork, room mockups, and anti-slip backing visuals.",
                "primary": "personalized dragon library rug",
                "secondary": "dragon book rug, fantasy reading room rug, custom dragon shaped rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with purple dragon, teal wings, stacked books, orange flowers, green leaves, and personalized name Sophia",
            },
        ),
    ]
)


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def write_csv(path, rows, headers):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with Path(path).open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def strip_html(value):
    return " ".join(re.sub(r"<[^>]+>", " ", html.unescape(value or "")).split())


def description_html(design):
    product_type = "sports room" if "football" in design["primary"] else "fantasy reading room"
    return (
        f"<p>{design['meta']}</p>"
        "<h3>Design details</h3><ul>"
        f"<li>{design['obs'].capitalize()}.</li>"
        f"<li>Position this as a personalized {product_type} decor piece, not a generic rug listing.</li>"
        "<li>Use backing, washing, and material claims only where supported by page text or product infographics.</li>"
        "</ul>"
    )


def image_note(design, n):
    theme = design["theme"]
    if "football" in design["primary"]:
        base = {
            1: f"{theme} shown on a floor in a room mockup",
            2: f"size chart for personalized {theme}",
            3: f"{theme} displayed in a living room mockup",
            4: f"{theme} displayed in a bedroom mockup",
            5: f"isolated artwork view of personalized {theme}",
            6: f"alternate color artwork view of personalized {theme}",
            7: f"alternate color artwork view of personalized {theme}",
            8: f"layer construction graphic for personalized {theme}",
            9: f"vacuum cleaner and machine washable care graphic for personalized {theme}",
        }
    else:
        base = {
            1: f"{theme} shown on a dark wood floor",
            2: f"alternate angle or isolated view of personalized {theme}",
            3: f"isolated or room view of personalized {theme}",
            4: f"custom shape area rug material graphic for {theme}",
            5: f"{theme} displayed in a light living room mockup",
            6: f"upper, middle, and bottom layer graphic for personalized {theme}",
            7: f"size chart and backing view for personalized {theme}",
        }
    return base.get(n, f"personalized {theme} image {n}")


def autosize(ws, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(12, min(max_width, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def main():
    batch = json.loads((ROOT / "evidence/products/B005_products.json").read_text())
    products = {p["handle"]: p for p in batch["products"]}
    page_meta = {r["product_key"]: r for r in json.loads((ROOT / "evidence/products/B005_page_meta.json").read_text())}
    manifest = json.loads((ROOT / "evidence/images/B005_manifest.json").read_text())
    by_images = {}
    for img in manifest:
        by_images.setdefault(img["product_key"], []).append(img)

    prev_seo, prev_img, prev_evd = [], [], []
    seo_headers = image_headers = evidence_headers = []
    for bid in ["B001", "B002", "B003", "B004"]:
        rows, headers = read_csv(ROOT / f"batches/{bid}_SEO_Products.csv")
        prev_seo.extend(rows)
        seo_headers = headers
        rows, headers = read_csv(ROOT / f"batches/{bid}_Image_Audit.csv")
        prev_img.extend(rows)
        image_headers = headers
        rows, headers = read_csv(ROOT / f"batches/{bid}_Product_Evidence.csv")
        prev_evd.extend(rows)
        evidence_headers = headers

    prev_kw, keyword_headers = read_csv(ROOT / "keyword_research.csv")
    prev_kw = [r for r in prev_kw if r.get("product_key") not in DESIGNS]
    with (ROOT / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        prev_buyer = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(prev_buyer[0].keys())
    prev_buyer = [r for r in prev_buyer if r.get("product_key") not in DESIGNS]

    rows_by_order = []
    with (ROOT / "inventory.csv").open(newline="", encoding="utf-8") as f:
        inv_reader = csv.DictReader(f)
        inv = list(inv_reader)
        inv_headers = inv_reader.fieldnames or []
    inv_lookup = {r["product_key"]: r for r in inv}

    b_seo, b_img, b_evd, b_kw, b_buyer = [], [], [], [], []
    for key, design in DESIGNS.items():
        product = products[key]
        meta = page_meta[key]
        inv_row = inv_lookup[key]
        order = int(inv_row["inventory_order"])
        evidence_id = f"{BATCH_ID}-{order:03d}"
        research_id = f"{BATCH_ID}-{order:03d}-BR01"
        url = f"https://chillgen.com/products/{key}"
        imgs = sorted(by_images[key], key=lambda x: int(x["image_number"]))
        img_fields = {}

        for img in imgs:
            n = int(img["image_number"])
            note = image_note(design, n)
            img_fields[f"img_{n}_link"] = img["src"]
            img_fields[f"img_{n}_alt_current"] = ""
            img_fields[f"img_{n}_alt"] = note
            img_fields[f"img_{n}_alt_action"] = "SET"
            b_img.append(
                {
                    "evidence_id": evidence_id,
                    "shop_domain": SHOP,
                    "Handle": key,
                    "product_id": str(product.get("id", "")),
                    "media_id": img.get("image_id", ""),
                    "image_location": "GALLERY_PUBLIC_CDN",
                    "image_number": n,
                    "variant": "",
                    "image_url": img["src"],
                    "image_url_export": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                    "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                    "shared_media_references": "",
                    "viewed_status": "VIEWED_CONTACT_SHEET_AND_LOCAL_FILE",
                    "viewed_at": NOW,
                    "observed_visual_details": note,
                    "alt_current": "",
                    "alt_proposed": note,
                    "alt_action": "SET",
                    "review_status": "NEEDS_REVIEW",
                    "revision": "1",
                    "approved_by": "",
                    "approved_at": "",
                    "approved_fields": "",
                    "approved_revision": "",
                    "evidence_file_or_reference": img.get("local_path", ""),
                    "issues": "alt_current_empty_in_public_json; media_export_url_unknown_without_shopify_export",
                }
            )

        seo = {h: "" for h in seo_headers}
        seo.update(
            {
                "schema_version": "2.4",
                "prompt_version": "2.4",
                "run_id": RUN_ID,
                "batch_id": BATCH_ID,
                "shop_domain": SHOP,
                "product_key": key,
                "Handle": key,
                "product_id": str(product.get("id", "")),
                "product_gid": f"gid://shopify/Product/{product.get('id', '')}",
                "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                "source_export_ref": f"evidence/products/B005_products.json; {meta.get('captured_html_path', '')}",
                "source_exported_at": meta.get("captured_at", NOW),
                "locale": "en-US",
                "market": "US",
                "url": url,
                "canonical_url": meta.get("canonical_url") or url,
                "product_type": product.get("product_type", ""),
                "title_current": product.get("title", ""),
                "h1_current": meta.get("h1_current", ""),
                "title_proposed": design["title"],
                "title_action": "SET",
                "h1_mapping_status": "RENDERED_H1_CAPTURED_THEME_CONFIRMATION_STILL_NEEDED_FOR_DEPLOY",
                "meta_title_current": "",
                "meta_title_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_title_current": meta.get("rendered_title_current", ""),
                "theme_title_suffix": "NONE_OR_NOT_OBSERVED",
                "meta_title_seo": design["title"],
                "meta_title_action": "SET",
                "meta_title_chars": len(design["title"]),
                "meta_description_current": "",
                "meta_description_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_meta_description_current": meta.get("rendered_meta_description_current", ""),
                "meta_description_seo": design["meta"],
                "meta_description_action": "SET",
                "meta_description_chars": len(design["meta"]),
                "primary_keyword": design["primary"],
                "secondary_keywords": design["secondary"],
                "long_tail_candidates": f"{design['primary']}; {design['secondary']}",
                "keyword_strategy": "LONG_TAIL_CANDIDATE",
                "season": design["season"],
                "keyword_demand_evidence": "SERP_ONLY: public marketplace/category results show personalized football rugs and dragon/fantasy rugs; no paid volume or first-party search data used.",
                "keyword_serp_fit": "SERP_ONLY_PRODUCT_AND_MARKETPLACE_RESULTS",
                "meta_keyword": design["primary"],
                "search_intent": "COMMERCIAL_INVESTIGATION",
                "keyword_validation_status": "SERP_ONLY_NO_VOLUME",
                "keyword_selection_reason": f"Selected because page/image evidence supports a {design['theme']} and custom-name rug intent.",
                "buyer_research_refs": research_id,
                "buyer_search_summary": "Buyer intent is treated as a hypothesis from product evidence plus category SERP language, not direct customer data.",
                "baseline_status": "PUBLIC_RENDERED_PAGE_AND_JSON_CAPTURED",
                "change_scope": "CLARIFICATION",
                "mapping_status": "NEEDS_REVIEW",
                "mapping_version": "2.4",
                "description_current_html": product.get("body_html", ""),
                "description_changes_needed": "Make product copy motif-specific while retaining verified public claims only.",
                "description_proposed": strip_html(description_html(design)),
                "description_proposed_html": description_html(design),
                "description_change_mode": "REPLACE_ALL",
                "description_target_section": "Full product body",
                "description_action": "SET",
                "revision": "1",
                "review_status": "NEEDS_REVIEW",
                "review_reason": "Research proposal only; rendered page captured but admin/export SEO values and media export URLs remain unknown.",
                "approved_by": "",
                "approved_at": "",
                "approved_fields": "",
                "approved_revision": "",
                "processing_status": "DRAFTED",
                "evidence_status": "COMPLETE_PUBLIC_RENDERED_PAGE_AND_IMAGES",
                "content_qa_status": "NOT_RUN",
                "field_evidence_map": "title/meta/description/alt derived from rendered page meta, public JSON, viewed images, and SERP-only validation.",
                "evidence_id": evidence_id,
                "issues": "Admin SEO fields unknown without Shopify export; QA not run; not approved for deployment.",
            }
        )
        seo.update(img_fields)
        b_seo.append(seo)

        product_refs = "; ".join([url, "evidence/products/B005_products.json", meta.get("captured_html_path", ""), *SERP_REFS])
        b_evd.append(
            {
                "evidence_id": evidence_id,
                "product_url": url,
                "reviewed_at": NOW,
                "sources_accessed": product_refs,
                "current_H1": meta.get("h1_current", ""),
                "current_meta_title": meta.get("rendered_title_current", ""),
                "short_source_excerpt": strip_html(product.get("body_html", ""))[:600],
                "verified_product_facts": f"{design['obs']}; HTTP {meta.get('http_status')}; canonical={meta.get('canonical_url')}; product type={product.get('product_type')}; variants={len(product.get('variants', []))}; images={len(imgs)}.",
                "gallery_image_count": len(imgs),
                "images_viewed_count": len(imgs),
                "image_audit_references": "; ".join([f"{evidence_id}-IMG{int(i['image_number']):02d}" for i in imgs]),
                "SERP_evidence_references": "; ".join(SERP_REFS),
                "buyer_research_references": research_id,
                "factual_conflicts": "None observed in public data; admin values unknown.",
                "processing_status": "DRAFTED",
                "confidence_and_reason": "MEDIUM: rendered page, product JSON, and images verified; keyword demand has SERP-only evidence and no volume data.",
                "fact_to_source_map": "product facts=B005_products.json and B005_pages HTML; visuals=B005 local images/contact sheets; keyword context=SERP references.",
                "proposed_field_to_fact_map": "motif and personalization to title/meta/alt; backing/washing/material only where page/image evidence supports.",
            }
        )

        b_buyer.append(
            {
                "research_id": research_id,
                "product_key": key,
                "supporting_fact_ids": evidence_id,
                "purchase_context": "Sports room/game-room gifting or fantasy reading-room decor, depending on product motif.",
                "jtbd_statement": f"When styling a themed room or buying a personalized gift, the buyer wants a {design['theme']} that makes the space feel custom and expressive.",
                "functional_motivation": "Add an indoor custom-name rug with a clearly visible sports or fantasy motif.",
                "emotional_social_motivation": "Create a room accent that feels personal, giftable, and tied to a favorite sport or fantasy-library theme.",
                "purchase_concerns": "Name accuracy, size fit, design clarity, cleaning, backing grip, and whether the style matches the room.",
                "customer_language": f"{design['primary']}; {design['secondary']}",
                "language_origin": "AGENT_HYPOTHESIS_WITH_CATEGORY_SERP_PARAPHRASE",
                "source_refs": "; ".join(SERP_REFS),
                "source_scope": "CATEGORY",
                "evidence_excerpt": "Paraphrased from public search/category results; no direct Chillgen customer quote or volume source used.",
                "observed_at": NOW,
                "market": "US",
                "source_language": "en",
                "research_status": "HYPOTHESIS",
                "limitations": "No Search Console, internal search, direct review corpus, or paid volume tool used for this product.",
                "seo_application": "Use motif-specific personalized rug language and avoid broad generic rug targeting at product level.",
            }
        )

        rejected = "football rug" if "football" in design["primary"] else "dragon rug"
        for kw, role, target, reason in [
            (design["primary"], "PRIMARY", url, "Best product-level match for observed motif and personalized/custom rug intent."),
            (design["secondary"].split(", ")[0], "SECONDARY", url, "Useful motif synonym supported by page and image evidence."),
            (rejected, "REJECTED", "COLLECTION_OR_CATEGORY_REVIEW", "Too broad for product-level targeting across multiple similar products."),
        ]:
            b_kw.append(
                {
                    "keyword": kw,
                    "product_key": key,
                    "buyer_research_refs": research_id if role != "REJECTED" else "",
                    "query_origin": "AGENT_PROPOSED",
                    "semantic_cluster": design["primary"],
                    "intent": "COMMERCIAL_INVESTIGATION",
                    "target_page_type": "PRODUCT" if role != "REJECTED" else "COLLECTION",
                    "target_url": target,
                    "keyword_role": role,
                    "decision_reason": reason,
                    "supporting_fact_ids": evidence_id,
                    "demand_evidence": "SERP_ONLY, no volume verified.",
                    "season": design["season"],
                    "research_period": "2026-09-07",
                    "validation_source": "; ".join(SERP_REFS),
                    "checked_at": NOW,
                    "representative_SERP_URLs": "; ".join(SERP_REFS[:4]),
                    "possible_overlap_with_other_products": "YES" if role == "REJECTED" else "MEDIUM_WITH_SIMILAR_PRODUCTS_IN_BATCH",
                    "mapping_reason": "Motif-specific keyword retained for product; broad term reserved for collection/category review.",
                    "mapping_status": "NEEDS_REVIEW",
                    "mapping_version": "2.4",
                }
            )

        rows_by_order.append(key)

    write_csv(ROOT / "batches/B005_SEO_Products.csv", b_seo, seo_headers)
    write_csv(ROOT / "batches/B005_Image_Audit.csv", b_img, image_headers)
    write_csv(ROOT / "batches/B005_Product_Evidence.csv", b_evd, evidence_headers)
    write_csv(ROOT / "keyword_research.csv", prev_kw + b_kw, keyword_headers)
    with (ROOT / "buyer_search_research.jsonl").open("w", encoding="utf-8") as f:
        for row in prev_buyer + b_buyer:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    for row in inv:
        if row["product_key"] in DESIGNS:
            row["assigned_batch"] = BATCH_ID
            row["batch_status"] = "COMPLETED"
            row["research_status"] = "DRAFTED"
            row["review_status"] = "NEEDS_REVIEW"
    write_csv(ROOT / "inventory.csv", inv, inv_headers)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", prev_seo + b_seo, seo_headers)
    write_sheet(wb, "Image_Audit", prev_img + b_img, image_headers)
    write_sheet(wb, "Product_Evidence", prev_evd + b_evd, evidence_headers)
    write_sheet(wb, "Keyword_Map", prev_kw + b_kw, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", prev_buyer + b_buyer, buyer_headers)
    readme = [
        {"metric": "scope", "value": "B001 through B005: 50 public products for chillgen.com, US/en-US", "definition": "Research proposal workbook only; all product and image proposals are NEEDS_REVIEW."},
        {"metric": "B005_method", "value": "B005 captured product page HTML/meta before SEO drafting and viewed all downloaded gallery images via contact sheets.", "definition": "B001-B003 remain public JSON/image evidence without full rendered page/meta backfill; B004-B005 include rendered page/meta capture."},
        {"metric": "sources", "value": "Chillgen public products.json, product URLs, rendered HTML captures, downloaded CDN images, contact sheets, SERP snippets and marketplace/category pages.", "definition": "Admin SEO fields and Shopify export media URLs remain unknown without Shopify CSV/export."},
        {"metric": "quality_status", "value": "DRAFTED_PUBLIC_RENDERED_PAGE_FOR_B004_B005", "definition": "393/393 images viewed via contact sheets across B001-B005; QA not run; no Shopify payload created."},
        {"metric": "deployment_status", "value": "NOT_APPROVED_NOT_DEPLOYABLE", "definition": "Do not import this workbook. Run QA and obtain explicit approval before any Shopify/Matrixify payload."},
        {"metric": "next_step", "value": "Stop after B005 and wait for user confirmation before B006.", "definition": "Required by project instruction."},
    ]
    write_sheet(wb, "README_QA", readme, ["metric", "value", "definition"])
    out = RESULTS / "batches" / f"SEO_Product_Optimization_through_{BATCH_ID}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    reopened = load_workbook(out, read_only=True, data_only=False)
    counts = {name: reopened[name].max_row - 1 for name in reopened.sheetnames}

    summary = ROOT / "batches/B005_summary.md"
    summary.write_text(
        "\n".join(
            [
                "# B005 Summary",
                "",
                f"- Run: {RUN_ID}",
                f"- Shop: {SHOP}",
                "- Scope: 10 products, 74 images",
                "- Cumulative workbook scope: B001-B005, 50 products, 393 images",
                "- Method: pre-batch checkpoint plus rendered page/meta HTML capture and contact-sheet visual review.",
                "- Status: NEEDS_REVIEW; QA_NOT_RUN; NOT_APPROVED; NO_DEPLOYMENT_PAYLOAD",
                f"- Workbook: {out}",
                f"- Generated at: {NOW}",
                "",
                "## Product Keys",
                *[f"- {key}" for key in rows_by_order],
            ]
        ),
        encoding="utf-8",
    )

    progress = json.loads((ROOT / "progress.json").read_text())
    completed = set(progress.get("images_completed", []))
    completed.update([f"{m['product_key']}#img_{int(m['image_number']):02d}" for m in manifest])
    progress.update(
        {
            "batch_id": BATCH_ID,
            "batch_product_keys": rows_by_order,
            "batch_status": "COMPLETED_NEEDS_REVIEW",
            "awaiting_confirmation": True,
            "continuation_confirmation_ref": "USER_CONFIRMED_CONTINUE_B005",
            "current_product_key": None,
            "current_stage": "B005_COMPLETE_AWAITING_USER_CONFIRMATION",
            "images_completed": sorted(completed),
            "last_saved_at": NOW,
        }
    )
    progress.setdefault("artifact_paths", {}).update(
        {
            "latest_batch_workbook": str(out),
            "b005_seo_products_csv": str(ROOT / "batches/B005_SEO_Products.csv"),
            "b005_image_audit_csv": str(ROOT / "batches/B005_Image_Audit.csv"),
            "b005_product_evidence_csv": str(ROOT / "batches/B005_Product_Evidence.csv"),
            "b005_summary": str(summary),
            "b005_page_meta": str(ROOT / "evidence/products/B005_page_meta.json"),
            "b005_image_manifest": str(ROOT / "evidence/images/B005_manifest.json"),
            "b005_contact_sheets": str(ROOT / "evidence/images/B005_contact_sheets"),
        }
    )
    tmp = ROOT / "progress.json.tmp"
    tmp.write_text(json.dumps(progress, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(ROOT / "progress.json")

    print(json.dumps({"workbook": str(out), "sheet_counts": counts, "summary": str(summary)}, indent=2))


if __name__ == "__main__":
    main()

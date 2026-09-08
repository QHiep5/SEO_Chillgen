from pathlib import Path
from datetime import datetime, timezone, timedelta
import json, re
from openpyxl import load_workbook

BASE=Path('/Users/buiquanghuy/Documents/seo_chillgen')
SRC=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R053/SEO_Product_Optimization_revision_R053.xlsx'
OUT=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R054/SEO_Product_Optimization_revision_R054.xlsx'
RUN=BASE/'seo_runs/chillgen.com/chillgen_20260907_01/revisions/R054'
HANDLES={'wheel-of-feelings-and-emotions-round-rug-educational-men-design-103','wheel-of-feelings-and-emotions-round-rug-educational-men-design-105','wheel-of-feelings-and-emotions-round-rug-educational-men-design-106','wheel-of-feelings-and-emotions-round-rug-educational-men-design-107','wheel-of-feelings-and-emotions-round-rug-educational-men-design-108','wheel-of-feelings-and-emotions-round-rug-educational-men-design-111','custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-110','wheel-of-feelings-and-emotions-round-rug-0fc607e644-0fc607e644','custom-a-good-day-to-read-book-rug-classroom-library-rugs-04ab7bfcca','feelings-wheel-round-classroom-rug-4be50a20fb'}
NOW=datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()
def clean(s):
    s=str(s or '')
    s=re.sub(r'\s*SEO copy (?:stays|remains|focuses|is kept)[^.]*\.', '', s, flags=re.I)
    return re.sub(r'\n{3,}','\n\n',s).strip()
def headers(ws): return {c.value:i+1 for i,c in enumerate(ws[1])}
wb=load_workbook(SRC); ws=wb['SEO_Products']; cols=headers(ws); changed=[]
for i in range(2, ws.max_row+1):
    h=ws.cell(i, cols['Handle']).value
    if h not in HANDLES: continue
    for f in ('description_proposed','description_proposed_html'): ws.cell(i, cols[f]).value=clean(ws.cell(i, cols[f]).value)
    ws.cell(i,cols['revision']).value='2'; ws.cell(i,cols['review_status']).value='NEEDS_REVIEW'; ws.cell(i,cols['processing_status']).value='DRAFTED'; ws.cell(i,cols['content_qa_status']).value='NOT_RUN'
    ws.cell(i,cols['review_reason']).value='Revision R054 removes customer-visible internal QA/process note identified in Re-QA R003.'
    ws.cell(i,cols['issues']).value='Revision R054 targeted internal process-language cleanup; not approved; Re-QA required.'; changed.append((h,i))
if len(changed)!=len(HANDLES): raise RuntimeError(f'scope mismatch {len(changed)}')
log=wb['Revision_Log']; existing=list(log.iter_rows(min_row=2, values_only=True)); log.delete_rows(2, log.max_row-1)
for row in existing: log.append(list(row))
for h,i in changed: log.append(['R054',NOW,h,i,'2','2','description_proposed, description_proposed_html, revision, review_reason, issues','internal_process_language','Run evidence-driven Re-QA on revision 2 for this product before approval.'])
OUT.parent.mkdir(parents=True,exist_ok=True); RUN.mkdir(parents=True,exist_ok=True); wb.save(OUT)
(RUN/'revision_R054_manifest.json').write_text(json.dumps({'revision_batch_id':'R054','generated_at':NOW,'source_workbook':str(SRC),'output_workbook':str(OUT),'product_count':len(changed),'handles':[h for h,_ in changed],'status':'COMPLETE_AWAITING_REQA','review_status':'NEEDS_REVIEW','content_qa_status':'NOT_RUN','not_approved_not_deployed':True},indent=2,ensure_ascii=False)+'\n')
print(json.dumps({'output':str(OUT),'updated':len(changed)},ensure_ascii=False))

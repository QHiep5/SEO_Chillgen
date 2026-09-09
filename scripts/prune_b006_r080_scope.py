from pathlib import Path
import openpyxl, json
BASE=Path('/Users/buiquanghuy/Documents/seo_chillgen'); p=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R080/SEO_Product_Optimization_revision_R080.xlsx'
w=openpyxl.load_workbook(p); data=json.loads((BASE/'seo_runs/chillgen.com/chillgen_20260907_01/evidence/products/B006_products.json').read_text()); handles=set(data.get('product_keys',[]))
def prune(sheet, key):
 ws=w[sheet]; heads=[c.value for c in next(ws.iter_rows())]; idx=heads.index(key)
 for i in range(ws.max_row,1,-1):
  if ws.cell(i,idx+1).value not in handles: ws.delete_rows(i)
prune('SEO_Products','Handle')
if 'Product_Evidence' in w.sheetnames:
 ws=w['Product_Evidence']; heads=[c.value for c in next(ws.iter_rows())]; ei=heads.index('evidence_id')
 allowed={f'B006-{i:03d}' for i in range(51,61)}
 for i in range(ws.max_row,1,-1):
  if ws.cell(i,ei+1).value not in allowed: ws.delete_rows(i)
if 'Image_Audit' in w.sheetnames: prune('Image_Audit','Handle')
log=w['Revision_Log']; lh=[c.value for c in next(log.iter_rows())]; hi=lh.index('handle')
for i in range(log.max_row,1,-1):
 if log.cell(i,hi+1).value not in handles: log.delete_rows(i)
w.save(p); print({'products':w['SEO_Products'].max_row-1,'product_evidence':w['Product_Evidence'].max_row-1,'image_audit':w['Image_Audit'].max_row-1,'log':log.max_row-1})

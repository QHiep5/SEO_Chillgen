from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B010"
REVISION = "R086"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"


COPY = {
    "personalized-calming-corner-rug-for-classroom-custom-tea-design-103": {
        "title": "Personalized Safe Space Classroom Rug",
        "meta": "Customize a safe space classroom rug with teacher-name text, colorful hearts, affirming wording, and a welcoming classroom look.",
        "description": "Customize a safe space classroom rug with teacher-name text, colorful hearts, affirming wording, and a welcoming classroom look. The design marks a positive corner where students can feel included and seen.",
        "primary": "personalized safe space classroom rug",
        "secondary": "custom teacher calm corner rug, safe space rug, calming corner mat",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-03": {
        "title": "Personalized Coping Tools Classroom Rug",
        "meta": "Customize a coping tools classroom rug with teacher-name text, monster icons, and breathe, move, focus, talk, and break prompts.",
        "description": "Customize a coping tools classroom rug with teacher-name text, friendly monster icons, and prompts such as breathe, move, focus, talk, and take a break. The design turns coping choices into a visible classroom routine.",
        "primary": "personalized coping tools classroom rug",
        "secondary": "classroom coping skills rug, custom teacher classroom rug, SEL calm corner rug",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-02": {
        "title": "Personalized Dinosaur Affirmation Rug",
        "meta": "Customize a dinosaur affirmation rug with teacher-name welcome text, Little Dinos Big Dreams wording, and kid-friendly affirmations.",
        "description": "Customize a dinosaur affirmation rug with teacher-name welcome text, Little Dinos Big Dreams wording, and kid-friendly affirmations. The playful dinosaur artwork makes a classroom gathering area feel bright and encouraging.",
        "primary": "personalized dinosaur affirmation rug",
        "secondary": "custom teacher name rug, dinosaur classroom rug, preschool affirmation rug",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-classroom": {
        "title": "Personalized Calming Techniques Rug",
        "meta": "Customize a calming techniques rug with teacher-name text, ocean animals, and simple prompts for breathing, drawing, stretching, and reading.",
        "description": "Customize a calming techniques rug with teacher-name text, ocean animals, and simple prompts for breathing, drawing, stretching, and reading. The friendly sea-life theme gives students a clear visual guide for reset moments.",
        "primary": "personalized calming techniques rug",
        "secondary": "custom teacher calming rug, classroom calm down rug, ocean animal SEL rug",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-09": {
        "title": "Personalized Classroom Jobs Rug",
        "meta": "Customize a classroom jobs rug with teacher-name text, rainbow artwork, checkerboard border, and helper-role icons for daily routines.",
        "description": "Customize a classroom jobs rug with teacher-name text, rainbow artwork, checkerboard border, and helper-role icons. The design makes classroom responsibilities easy to see and gives student helpers a cheerful shared reference.",
        "primary": "personalized classroom jobs rug",
        "secondary": "custom teacher classroom rug, classroom helper rug, teacher name jobs mat",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-07": {
        "title": "Personalized Classroom Rules Rug",
        "meta": "Customize a classroom rules rug with teacher-name text, rainbow artwork, school icons, and listen, be kind, share, and follow directions prompts.",
        "description": "Customize a classroom rules rug with teacher-name text, rainbow artwork, school icons, and simple prompts such as listen, be kind, share, and follow directions. The design keeps class expectations visible in a friendly way.",
        "primary": "personalized classroom rules rug",
        "secondary": "custom teacher rules rug, classroom expectations rug, elementary classroom mat",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-04": {
        "title": "Personalized Kindness Rules Rug",
        "meta": "Customize a kindness rules rug with teacher-name text, checkerboard border, illustrated children, and positive classroom behavior prompts.",
        "description": "Customize a kindness rules rug with teacher-name text, checkerboard border, illustrated children, and positive classroom behavior prompts. The design centers kindness, respect, honesty, encouragement, and responsibility.",
        "primary": "personalized kindness rules rug",
        "secondary": "custom classroom rules rug, positive behavior classroom rug, kindness classroom mat",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-01": {
        "title": "Classroom Zone of Tolerance Rug",
        "meta": "Create a Zone of Tolerance classroom rug with overwhelmed, just right, and shut down panels plus coping actions for SEL routines.",
        "description": "Create a Zone of Tolerance classroom rug with overwhelmed, just right, and shut down panels plus coping actions for classroom SEL routines. The design gives students a visual way to recognize their regulation state.",
        "primary": "classroom zone of tolerance rug",
        "secondary": "SEL regulation rug, calm down classroom rug, emotional regulation mat",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-05": {
        "title": "Personalized Voice Level Classroom Rug",
        "meta": "Customize a voice level classroom rug with teacher-name text, a 0-4 noise chart, colorful icons, and clear routine cues.",
        "description": "Customize a voice level classroom rug with teacher-name text, a 0-4 noise chart, colorful icons, and clear routine cues. The design helps students connect each voice level with a simple visual reference.",
        "primary": "personalized voice level classroom rug",
        "secondary": "classroom voice level rug, custom teacher classroom mat, noise level classroom rug",
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-06": {
        "title": "Personalized Behavior Champs Rug",
        "meta": "Customize a Behavior Champs rug with teacher-name text, school icons, and listening, kindness, sharing, and helping-hands prompts.",
        "description": "Customize a Behavior Champs rug with teacher-name text, school icons, and listening, kindness, sharing, and helping-hands prompts. The bold classroom design frames daily behavior expectations as positive champion habits.",
        "primary": "personalized behavior champs rug",
        "secondary": "custom teacher behavior rug, classroom behavior expectations rug, positive behavior mat",
    },
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B010 scope does not match the R086 copy map")

    for row in seo_rows:
        item = COPY[row["Handle"]]
        row["title_proposed"] = item["title"]
        row["meta_title_seo"] = item["title"]
        row["meta_title_chars"] = str(len(item["title"]))
        row["meta_description_seo"] = item["meta"]
        row["meta_description_chars"] = str(len(item["meta"]))
        row["description_proposed"] = item["description"]
        row["description_proposed_html"] = f"<p>{item['description']}</p>"
        row["primary_keyword"] = item["primary"]
        row["secondary_keywords"] = item["secondary"]
        row["long_tail_candidates"] = f"{item['primary']}; {item['secondary']}"
        row["revision"] = "2"
        row["review_status"] = "NEEDS_REVIEW"
        row["content_qa_status"] = "NOT_RUN"
        row["review_reason"] = f"{REVISION} removes internal evidence-note language and unsupported claim wording from B010 copy."
        row["issues"] = f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations."

    handles_set = set(handles)
    evidence_scope = [row for row in evidence_rows if row.get("evidence_id") in {r.get("evidence_id") for r in seo_rows}]
    keyword_scope = [row for row in keyword_rows if row.get("product_key") in handles_set]
    buyer_rows = [row for row in buyer_all if row.get("product_key") in handles_set]
    img_scope = [row for row in img_rows if row.get("Handle") in handles_set]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": "Canonical B010 cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    evidence_ids = {row["evidence_id"] for row in evidence_scope}
    for row in seo_rows:
        evidence_id = row.get("evidence_id")
        if evidence_id not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {evidence_id}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", evidence_id])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    digest = sha256(OUT)
    manifest = {
        "revision": REVISION,
        "batch_id": BATCH,
        "scope_count": len(handles),
        "image_count": len(img_scope),
        "canonical_workbook": str(OUT),
        "sha256": digest,
        "status": "CREATED_FOR_HARDENED_REQA",
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
    }
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

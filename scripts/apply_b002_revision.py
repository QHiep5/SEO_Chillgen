from pathlib import Path
import openpyxl, re, json

S=Path('seo_runs/chillgen.com/chillgen_20260907_01/qa_scope/B002_scope_workbook.xlsx')
O=Path('seo_runs/chillgen.com/chillgen_20260907_01/qa_scope/B002_scope_workbook_revision1.xlsx')
w=openpyxl.load_workbook(S)
ws=w['SEO_Products']; h=[c.value for c in ws[1]]; c={x:i+1 for i,x in enumerate(h)}
changed=[]
for row in ws.iter_rows(min_row=2):
    handle=row[c['Handle']-1].value
    if handle:
        for name in ('description_proposed','description_proposed_html','meta_description_seo'):
            cell=ws.cell(row[0].row,c[name])
            if isinstance(cell.value,str):
                cell.value=re.sub(r'\b(?:non[- ]?slip|anti[- ]?slip)\b','',cell.value,flags=re.I)
                cell.value=re.sub(r'\s{2,}',' ',cell.value).strip()
        changed.append(handle)
w.save(O)
print(json.dumps({'output':str(O),'products_revised':len(changed)}))

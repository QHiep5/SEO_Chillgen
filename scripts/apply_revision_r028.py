from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R027/SEO_Product_Optimization_revision_R027.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R028"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R028.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R028"
MANIFEST = RUN_DIR / "revision_R028_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R028_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-kids-read-book-rug-classroom-library-eaedf8485b": {
        "title": "Cool Kids Read Books Classroom Rug with Book Icons",
        "meta": "Decorate with a Cool Kids Read Books classroom rug featuring smiling book icons, pastel letters and classroom scenes.",
        "description": "Decorate with a Cool Kids Read Books classroom rug featuring smiling book icons, pastel letters and classroom scenes.\n\nDesign details\n- Cream rectangular classroom rug artwork shows Cool Kids Read Books text with colorful book characters, letters and small accent icons.\n- Product images include classroom and playroom mockups, close-up artwork panels and size reference graphics.\n- SEO copy stays tied to visible reading artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-reading-corner-classroom-library-mat-40b8909d74-40b8909d74": {
        "title": "Reading Corner Classroom Rug with Pencil Border",
        "meta": "Decorate with a reading corner classroom rug featuring rainbow pencil border, paper airplanes and open book artwork.",
        "description": "Decorate with a reading corner classroom rug featuring rainbow pencil border, paper airplanes and open book artwork.\n\nDesign details\n- Black rectangular reading rug artwork shows rainbow pencil border, colorful paper airplanes and Reading Will Take You Everywhere text.\n- Visible details include open book graphic, classroom circle-time mockups, room scenes and size reference graphics.\n- SEO copy focuses on visible reading-corner artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-reading-classroom-library-rug-bdef6840ec-bdef6840ec": {
        "title": "Reading Gives You Wings Rug with Butterfly Artwork",
        "meta": "Decorate with a Reading Gives You Wings rug featuring rainbow butterflies, open book artwork, clouds and classroom scenes.",
        "description": "Decorate with a Reading Gives You Wings rug featuring rainbow butterflies, open book artwork, clouds and classroom scenes.\n\nDesign details\n- Black rectangular reading rug artwork shows Reading Gives You Wings quote, rainbow butterfly pattern and an open book with children.\n- Visible details include clouds, rainbow accents, classroom and playroom scenes, group reading images and close-up panels.\n- SEO copy stays tied to visible butterfly reading artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-classroom-library-reading-rug-9121bf7e9f-9121bf7e9f": {
        "title": "Reader Leader Classroom Rug with Bookshelf Art",
        "meta": "Decorate with a Reader Leader classroom rug featuring bookshelf artwork, Today a Reader quote, books, bow and apple.",
        "description": "Decorate with a Reader Leader classroom rug featuring bookshelf artwork, Today a Reader quote, books, bow and apple.\n\nDesign details\n- Black bookshelf-style classroom rug artwork shows Today a Reader Tomorrow a Leader quote with pastel books and shelf details.\n- Visible details include bow, pencil cup, apple, classroom group scenes, room mockups and size reference graphics.\n- SEO copy focuses on visible bookshelf reading artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-bookshelf-classroom-library-reading-rug-75ece72c91": {
        "title": "Personalized Bookshelf Reading Rug with Globe Art",
        "meta": "Personalize a bookshelf reading rug with educator name, Today a Reader quote, colorful books, globe and backpack.",
        "description": "Personalize a bookshelf reading rug with educator name, Today a Reader quote, colorful books, globe and backpack.\n\nDesign details\n- Brown bookshelf-style reading rug artwork shows Today a Reader Tomorrow a Leader quote with custom Mr. Smith name.\n- Visible details include colorful books, science fiction label, globe, backpack, classroom and playroom mockups.\n- SEO copy stays tied to visible personalized bookshelf artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-reading-corner-classroom-library-rug-4a41f4c798-4a41f4c798": {
        "title": "Butterfly Reading Corner Rug with Library Banner",
        "meta": "Decorate with a butterfly reading corner rug featuring open book art, Reading Gives You Wings quote and library banner.",
        "description": "Decorate with a butterfly reading corner rug featuring open book art, Reading Gives You Wings quote and library banner.\n\nDesign details\n- White and pink rectangular reading rug artwork shows colorful butterflies, open book, scalloped border and reading quote text.\n- Visible details include Check Out Our Library banner, classroom mockups, room scenes, close-up views and size reference graphics.\n- SEO copy focuses on visible butterfly reading artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-baseball-vintage-pattern-area-rug-non-slip-a74f120767": {
        "title": "Vintage Baseball Area Rug with Jackson Name Art",
        "meta": "Personalize a vintage baseball area rug with Jackson name, number 35, weathered baseball artwork and distressed colors.",
        "description": "Personalize a vintage baseball area rug with Jackson name, number 35, weathered baseball artwork and distressed colors.\n\nDesign details\n- Rectangular vintage baseball rug artwork shows a large baseball graphic with custom Jackson name and number 35.\n- Visible details include distressed blue and brown background, living room mockups, close-up artwork views and size reference graphics.\n- SEO copy stays tied to visible baseball artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-handprints-kids-classroom-rug-e41ab46aac-e41ab46aac": {
        "title": "Reading Bloom Classroom Rug with Handprint Accents",
        "meta": "Personalize a Reading Helps Your Mind Bloom classroom rug with educator name, floral bouquet and handprint accents.",
        "description": "Personalize a Reading Helps Your Mind Bloom classroom rug with educator name, floral bouquet and handprint accents.\n\nDesign details\n- Black floral classroom rug artwork shows Reading Helps Your Mind Bloom quote with custom Mr. Michael name.\n- Visible details include bouquet, paintbrush artwork, colorful handprints, gingham border, classroom and playroom scenes.\n- SEO copy focuses on visible reading bloom artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-handprints-kids-classroom-rug-345ee4ca82-345ee4ca82": {
        "title": "Every Child Is An Artist Rug with Handprint Art",
        "meta": "Personalize an Every Child Is An Artist rug with educator name, colorful handprints, crayons, rocket and planets.",
        "description": "Personalize an Every Child Is An Artist rug with educator name, colorful handprints, crayons, rocket and planets.\n\nDesign details\n- Black classroom rug artwork shows Every Child Is An Artist quote with custom Mrs. Anderson name and paint splash accents.\n- Visible details include handprints, crayons, rocket, planets, classroom scenes, close-up panels and size reference graphics.\n- SEO copy stays tied to visible art-classroom artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-handprints-kids-classroom-rug-51f01db413-51f01db413": {
        "title": "Reading Everywhere Classroom Rug with Open Book Art",
        "meta": "Decorate with a Reading Everywhere classroom rug featuring open book artwork, butterflies, flowers and reading quotes.",
        "description": "Decorate with a Reading Everywhere classroom rug featuring open book artwork, butterflies, flowers and reading quotes.\n\nDesign details\n- Black rectangular classroom reading rug artwork shows open book, butterflies, flowers and Reading Will Take You Everywhere text.\n- Product images include classroom and playroom mockups, educator reading lifestyle image, close-up panels and size reference graphics.\n- SEO copy focuses on visible reading artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "kid friendly", "pet friendly", "safe durable",
    "machine-washable", "machine washable", "washable", "quick-dry", "quick dry",
    "memory foam", "microfiber", "velvet", "stain", "fade resistant", "easy clean",
    "easy-clean", "waterproof", "absorbent", "absorption", "backing", "rubber",
    "layered construction", "hd printing", "ultra-soft", "soft ", "cushion",
    "support classroom", "support emotional", "teach", "learning-space", "learning outcome",
    "thickened", "reinforced", "bound edge", "material panel",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    blob = " ".join([title, meta, item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R028 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R028 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append([
            "revision_batch_id", "generated_at", "handle", "source_row",
            "old_revision", "new_revision", "changed_fields", "qa_issue_fields",
            "recheck_condition",
        ])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R028"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R028", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R028",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R029 or Re-QA revision R028",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R028 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R028.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        desc = row[out_cols["description_proposed"] - 1]
        blob = " ".join([str(title or ""), str(meta or ""), str(desc or "")]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r028_log_rows": counts.get("R028", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R015/SEO_Product_Optimization_revision_R015.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R016"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R016.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R016"
MANIFEST = RUN_DIR / "revision_R016_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R016_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-dragon-shaped-area-rug-personalized-library-western-dragon": {
        "title": "Personalized Western Dragon Book Rug with Teal Wings",
        "meta": "Personalize a western dragon book rug with custom name text, teal orange dragon, broad wings and stacked books.",
        "description": (
            "Personalize a western dragon book rug with custom name text, teal orange dragon, broad wings and stacked books.\n\n"
            "Design details\n"
            "- Custom-shaped rug artwork shows a teal and orange western dragon perched on stacked books with large wings.\n"
            "- Product images include floor mockups, close-up views, size reference graphics and custom-name examples.\n"
            "- SEO copy stays tied to visible dragon book artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-pet-dog-welcome-door-mat-with-photo-and-name-a04-a04": {
        "title": "Personalized Dog Approval Mat with Custom Photo",
        "meta": "Customize a dog approval mat with pet photo, name text, paw prints and Visitors Must Be Approved wording.",
        "description": (
            "Customize a dog approval mat with pet photo, name text, paw prints and Visitors Must Be Approved wording.\n\n"
            "Design details\n"
            "- Tan rectangular mat artwork shows a golden retriever photo area, black paw prints and Visitors Must Be Approved By Chance text.\n"
            "- Product images include hand-held, doorway, bedside and customization example views.\n"
            "- SEO copy focuses on visible custom dog photo artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-pet-dog-welcome-door-mat-with-photo-and-name-design-02": {
        "title": "Custom Dog Photo Welcome Mat with Portrait Art",
        "meta": "Create a custom dog photo welcome mat with script welcome text, pet portrait artwork and tan doorway design.",
        "description": (
            "Create a custom dog photo welcome mat with script welcome text, pet portrait artwork and tan doorway design.\n\n"
            "Design details\n"
            "- Tan rectangular mat artwork shows a dog portrait peeking above a line with black script Welcome text.\n"
            "- Gallery images show hand-held, doorway, bedside and custom photo variant examples.\n"
            "- SEO copy stays tied to visible personalized dog welcome artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-pet-dog-welcome-door-mat-with-photo-and-name-design-01": {
        "title": "Personalized Dog House Mat with Custom Portrait",
        "meta": "Personalize a dog house mat with custom pet portrait, Welcome to Cooper's House text and tan doorway artwork.",
        "description": (
            "Personalize a dog house mat with custom pet portrait, Welcome to Cooper's House text and tan doorway artwork.\n\n"
            "Design details\n"
            "- Tan rectangular mat artwork shows a large dog portrait with Welcome to Cooper's House text and decorative border.\n"
            "- Product images include hand-held, doorway, bedside and custom-photo example views.\n"
            "- SEO copy focuses on visible dog name and portrait artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-pet-dog-welcome-door-mat-with-photo-and-name-design-03": {
        "title": "Personalized Multi Pet Welcome Mat with Paw Names",
        "meta": "Customize a multi pet welcome mat with pet names, paw icons, footprint icons and Welcome to Our Home wording.",
        "description": (
            "Customize a multi pet welcome mat with pet names, paw icons, footprint icons and Welcome to Our Home wording.\n\n"
            "Design details\n"
            "- Tan rectangular mat artwork shows Welcome to Our Home text with multiple paw and footprint icons labeled with names.\n"
            "- Gallery images include doorway, bedside, angled and hand-held mockups with the same family pet theme.\n"
            "- SEO copy stays tied to visible multi-pet name artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-pet-dog-welcome-door-mat-with-photo-and-name-a05-a05": {
        "title": "Personalized Luna Dog House Mat with Portrait Art",
        "meta": "Personalize a Luna dog house mat with pet portrait artwork, Welcome to Luna's House text and photo option panels.",
        "description": (
            "Personalize a Luna dog house mat with pet portrait artwork, Welcome to Luna's House text and photo option panels.\n\n"
            "Design details\n"
            "- Tan rectangular mat artwork shows Welcome to Luna's House text with a dog portrait and decorative corner lines.\n"
            "- Product images include doorway, staircase, hand-held and custom photo option examples.\n"
            "- SEO copy focuses on visible dog name and portrait artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "custom-japanese-koi-fish-shaped-area-rug-personalized-3d-design-01": {
        "title": "Japanese Koi Pond Shaped Rug with Lily Pad Art",
        "meta": "Decorate with a Japanese koi pond shaped rug featuring mossy edge, lily pads, white flowers and red white koi.",
        "description": (
            "Decorate with a Japanese koi pond shaped rug featuring mossy edge, lily pads, white flowers and red white koi.\n\n"
            "Design details\n"
            "- Irregular pond-shaped rug artwork shows a moss green edge, dark water, lily pads, white flowers and red white koi fish.\n"
            "- Product images include room mockups, size reference panels, close-up views and lighting examples.\n"
            "- SEO copy stays tied to visible koi pond artwork without unsupported material, cleaning, kid or pet safety claims."
        ),
    },
    "custom-japanese-koi-fish-shaped-area-rug-personalized-3d-design-04": {
        "title": "Stone Border Koi Pond Rug with 3D Water Artwork",
        "meta": "Style a room with a stone border koi pond rug featuring green blue water, rounded rocks and colorful koi fish.",
        "description": (
            "Style a room with a stone border koi pond rug featuring green blue water, rounded rocks and colorful koi fish.\n\n"
            "Design details\n"
            "- Oval pond rug artwork shows a rounded stone border, green blue water and koi fish in red, white, yellow and orange.\n"
            "- Gallery images include living room, bedroom, multi-room and size reference views.\n"
            "- SEO copy focuses on visible koi pond artwork without unsupported material, cleaning, kid or pet safety claims."
        ),
    },
    "boat-rug-05-boat-rug-05-design-22": {
        "title": "Personalized Captain Boat Rug with Nautical Quote",
        "meta": "Personalize a captain boat rug with Captain Is Always Right quote, Captain Jackson text, helm, anchor and rope border.",
        "description": (
            "Personalize a captain boat rug with Captain Is Always Right quote, Captain Jackson text, helm, anchor and rope border.\n\n"
            "Design details\n"
            "- Single product image shows a navy and white nautical rug on a boat deck with helm, anchor and rope border graphics.\n"
            "- Visible text includes The Captain Is Always Right and I'm the Captain, Captain Jackson.\n"
            "- SEO copy is intentionally narrow because available evidence is limited to one visible image."
        ),
    },
    "boat-rug-10-boat-rug-10-design-10": {
        "title": "Black Gold Welcome Aboard Boat Mat with Anchor Art",
        "meta": "Personalize a black gold welcome aboard boat mat with boat name text, anchor artwork and nautical deck styling.",
        "description": (
            "Personalize a black gold welcome aboard boat mat with boat name text, anchor artwork and nautical deck styling.\n\n"
            "Design details\n"
            "- Single product image shows a black and gold nautical mat on a boat deck with anchor artwork and Welcome Aboard lettering.\n"
            "- Visible personalization area reads Boat Name Here above the anchor graphic.\n"
            "- SEO copy is intentionally narrow because available evidence is limited to one visible image."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip",
    "non slip", "non-skid", "non skid", "kid friendly", "pet friendly",
    "safe durable", "machine-washable", "machine washable", "washable ",
    "quick-dry", "quick dry", "memory foam", "microfiber", "velvet",
    "stain", "fade resistant", "easy clean", "soft fabric", "backing",
    "cleaning visuals", "cleaning graphic",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title = item["title"]
    meta = item["meta"]
    blob = " ".join([item["title"], item["meta"], item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R016 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R016 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R016"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R016", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R016",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R017 or Re-QA revision R016",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R016 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R016.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([
            str(title or ""),
            str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
        ]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70:
            raise RuntimeError(f"{handle} title len failed after save")
        if len(meta) > 320:
            raise RuntimeError(f"{handle} meta len failed after save")
        if str(row[out_cols["revision"] - 1]) != "2":
            raise RuntimeError(f"{handle} revision not 2")
        if row[out_cols["review_status"] - 1] != "NEEDS_REVIEW":
            raise RuntimeError(f"{handle} review_status incorrect")
        if row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} content_qa_status incorrect")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles = {}
    dup_meta_titles = []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt = row[out_cols["meta_title_seo"] - 1]
        handle = row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle

    result = {
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r016_log_rows": counts.get("R016", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

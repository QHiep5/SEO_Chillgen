from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B025"
REVISION = "R101"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "custom-reading-tree-classroom-library-rug-983a65fb2e-983a65fb2e": ("Personalized Reading Tree Classroom Rug", "personalized reading tree classroom rug", "reading tree library rug, custom classroom reading mat, story time kids rug", "a tree filled with books and personalized name space for a classroom library, story-time corner, playroom, or kids room"),
    "custom-tree-of-knowledge-classroom-library-rug-5ea95ac966-5ea95ac966": ("Personalized Tree of Knowledge Rug", "personalized Tree of Knowledge rug", "tree classroom library rug, floral reading rug, custom learning space mat", "a Tree of Knowledge design with floral details and personalized name space for a classroom, library, study corner, or living room"),
    "personalized-cool-kids-read-book-rug-classroom-library-b1f15827f5": ("Personalized Cool Kids Read Rug", "personalized Cool Kids Read rug", "cool kids reading rug, custom book classroom mat, library reading nook rug", "Cool Kids Read wording with book-character artwork and personalized name space for a classroom library, reading nook, or playroom"),
    "custom-good-day-read-book-classroom-rug-dc9ddbc5f6-dc9ddbc5f6": ("Personalized New Chapter Reading Rug", "personalized New Chapter reading rug", "New Chapter classroom rug, custom reading corner mat, personalized library rug", "New Chapter wording with a welcoming classroom reading design for a library, story-time corner, study space, or playroom"),
    "colorful-handprints-kids-classroom-rug-e41ab46aac-e41ab46aac": ("Reading Helps Your Mind Bloom Rug", "Reading Helps Your Mind Bloom rug", "positive classroom reading rug, handprint learning mat, teacher library rug", "Reading Helps Your Mind Bloom wording with colorful handprints for a classroom, library, playroom, or kids learning space"),
    "colorful-classroom-rug-for-kids-aaa73255b2-aaa73255b2": ("Personalized BE Values Classroom Rug", "personalized BE Values classroom rug", "BE values classroom mat, custom teacher rug, handprint kids learning rug", "BE Values wording with colorful handprints and personalized classroom details for a classroom, playroom, or learning space"),
    "colorful-classroom-rug-for-kids-877e4e6020-877e4e6020": ("You Matter Affirmation Classroom Rug", "You Matter affirmation classroom rug", "affirmation classroom mat, positive message kids rug, colorful handprint rug", "You Matter wording with colorful handprints for a classroom, playroom, counseling corner, or kids learning space"),
    "colorful-handprints-kids-classroom-rug-345ee4ca82-345ee4ca82": ("Every Child Is An Artist Rug", "Every Child Is An Artist classroom rug", "artist classroom rug, creative kids learning mat, handprint playroom rug", "Every Child Is An Artist wording with colorful handprints for a classroom, art area, playroom, or creative learning space"),
    "colorful-handprints-classroom-kids-rug-51cf8ab902-51cf8ab902": ("Personalized Welcome Classroom Rug", "personalized welcome classroom rug", "custom teacher welcome mat, handprint classroom rug, personalized kids room mat", "a welcoming classroom message with personalized teacher details and colorful handprints for a classroom, playroom, or kids room"),
    "colorful-handprints-kids-classroom-rug-51f01db413-51f01db413": ("Reading Everywhere Classroom Rug", "Reading Everywhere classroom rug", "reading handprint classroom mat, open book kids rug, literacy learning rug", "Reading Everywhere wording with an open book and colorful handprints for a classroom, reading corner, playroom, or study space"),
}

def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []

def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)

def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B025 scope does not match the R101 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        meta = f"Create a {primary} with {motif}."
        description = f"Create a {primary} with {motif}. The design adds a colorful learning focal point to a classroom, library, reading nook, playroom, bedroom, or study space."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))

if __name__ == "__main__":
    main()

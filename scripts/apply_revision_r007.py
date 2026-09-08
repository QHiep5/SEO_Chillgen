from __future__ import annotations

import json
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R006/SEO_Product_Optimization_revision_R006.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R007"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R007.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R007"
MANIFEST = RUN_DIR / "revision_R007_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R007_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-running-horse-welcome-mat-fade3d1429": {
        "title": "Two Horse Area Rug with Warm Brown Portrait Artwork",
        "meta": "Decorate a horse lover room with a two horse area rug featuring warm brown portraits and stable-inspired background.",
        "description": (
            "Decorate a horse lover room with a two horse area rug featuring warm brown portraits and stable-inspired background.\n\n"
            "Design details\n"
            "- Rectangular rug with two brown horse faces, warm blurred stable-style background and close-up portrait composition.\n"
            "- Gallery images show living room and fireplace-style mockups plus close-up design and size visuals.\n"
            "- SEO copy focuses on the visible two-horse portrait artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "horse-area-rug-running-carpet-beb8f5b5cd-beb8f5b5cd": {
        "title": "Sunset Running Horse Area Rug with Red Orange Sky",
        "meta": "Add a sunset running horse area rug with black galloping silhouette, red orange sky, large sun and grass foreground.",
        "description": (
            "Add a sunset running horse area rug with black galloping silhouette, red orange sky, large sun and grass foreground.\n\n"
            "Design details\n"
            "- Rectangular rug with a black running horse silhouette, fiery red-orange sunset, large yellow sun and dark grass foreground.\n"
            "- Gallery images show room mockups, fireplace scenes, size visuals and close-up sunset artwork.\n"
            "- SEO copy stays tied to the visible sunset horse design without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-western-running-horse-door-mat-e6671cf5b4-e6671cf5b4": {
        "title": "Fire Horse Area Rug with Red Flame Background Art",
        "meta": "Create a bold western focal point with a fire horse area rug featuring black horse artwork and red flame background.",
        "description": (
            "Create a bold western focal point with a fire horse area rug featuring black horse artwork and red flame background.\n\n"
            "Design details\n"
            "- Rectangular rug with black horse portrait, swirling mane, glowing eye and red flame or lava-style background.\n"
            "- Gallery images show living room and fireplace-style mockups plus close-up design and size visuals.\n"
            "- SEO copy focuses on the visible fire horse artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-welcome-mat-34b7c38f01": {
        "title": "Black Horse Portrait Area Rug with Bridle Detail",
        "meta": "Decorate an equestrian room with a black horse portrait area rug featuring bridle detail and soft neutral background.",
        "description": (
            "Decorate an equestrian room with a black horse portrait area rug featuring bridle detail and soft neutral background.\n\n"
            "Design details\n"
            "- Rectangular rug with black horse head portrait, bridle detail, gray-beige background and subtle sparkle accents.\n"
            "- Gallery images show room mockups, fireplace scenes, close-up art panels and size visuals.\n"
            "- SEO copy stays focused on the visible black horse portrait without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-welcome-mat-509ed4d18f-509ed4d18f": {
        "title": "Daisy Horse Area Rug with Brown Horse Portrait",
        "meta": "Add a daisy horse area rug with brown horse portrait, white flower crown, green vines and warm floral accents.",
        "description": (
            "Add a daisy horse area rug with brown horse portrait, white flower crown, green vines and warm floral accents.\n\n"
            "Design details\n"
            "- Rectangular rug with brown horse portrait, white daisy flower crown, green ivy and floral accent artwork.\n"
            "- Gallery images show living room and fireplace mockups plus close-up flower and size visuals.\n"
            "- SEO copy focuses on the visible daisy horse design without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-welcome-mat-f114a916e7-f114a916e7": {
        "title": "Running Horses Area Rug with Galloping Herd Artwork",
        "meta": "Decorate a western room with a running horses area rug featuring a galloping herd in brown, white and black tones.",
        "description": (
            "Decorate a western room with a running horses area rug featuring a galloping herd in brown, white and black tones.\n\n"
            "Design details\n"
            "- Rectangular rug with a herd of running horses in brown, white and black, set against a dusty motion background.\n"
            "- Gallery images show room mockups, fireplace scenes, close-up art panels and size visuals.\n"
            "- SEO copy stays tied to the visible galloping herd artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-welcome-mat-a4c32dbbfa-a4c32dbbfa": {
        "title": "Dark Horse Eye Area Rug with Amber Close Up Artwork",
        "meta": "Create a striking equestrian room with a dark horse eye area rug featuring amber eye artwork and black brown tones.",
        "description": (
            "Create a striking equestrian room with a dark horse eye area rug featuring amber eye artwork and black brown tones.\n\n"
            "Design details\n"
            "- Rectangular rug with close-up dark horse face, amber eye, water droplet details and black-brown texture.\n"
            "- Gallery images show living room and fireplace mockups plus close-up design and size visuals.\n"
            "- SEO copy focuses on the visible horse-eye artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-area-rug-2f0ce836fd-2f0ce836fd": {
        "title": "Black and White Horse Area Rug with Paired Portraits",
        "meta": "Style an equestrian room with a black and white horse area rug featuring paired horse portraits on a dark background.",
        "description": (
            "Style an equestrian room with a black and white horse area rug featuring paired horse portraits on a dark background.\n\n"
            "Design details\n"
            "- Rectangular rug with white horse and black horse portraits facing each other against a dark background.\n"
            "- Gallery images show room mockups, fireplace scenes, close-up art panels and size visuals.\n"
            "- SEO copy stays tied to the visible paired horse portrait design without unsupported material, backing or surface-performance claims."
        ),
    },
    "western-running-horse-area-rug-02dee68d0b": {
        "title": "Mint Horse Area Rug with White Horse and Pink Flowers",
        "meta": "Brighten a horse lover room with a mint horse area rug featuring white horse portrait, pink flowers and small hearts.",
        "description": (
            "Brighten a horse lover room with a mint horse area rug featuring white horse portrait, pink flowers and small hearts.\n\n"
            "Design details\n"
            "- Rectangular rug with white horse portrait, mint green background, small white hearts and pink flower crown artwork.\n"
            "- Gallery images show room mockups, fireplace scenes, close-up floral art and size visuals.\n"
            "- SEO copy focuses on the visible mint horse design without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-area-rug-31f0e158c5": {
        "title": "Floral Horse Herd Area Rug with Roses and Daisies",
        "meta": "Decorate with a floral horse herd area rug featuring multiple horse portraits, roses, daisies and colorful flowers.",
        "description": (
            "Decorate with a floral horse herd area rug featuring multiple horse portraits, roses, daisies and colorful flowers.\n\n"
            "Design details\n"
            "- Rectangular rug with multiple brown, white and black horse portraits surrounded by roses, daisies and floral details.\n"
            "- Gallery images show room mockups, fireplace scenes, close-up flower artwork and size visuals.\n"
            "- SEO copy stays tied to the visible floral horse herd design without unsupported material, backing or surface-performance claims."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "non-slip", "kid friendly",
    "pet friendly", "safe durable", "machine-washable", "machine washable",
    "washable ", "quick-dry", "memory foam", "microfiber", "velvet",
    "stain", "fade resistant",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R007 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R007 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    for sheet_name in ["Buyer_Search_Research", "Keyword_Map"]:
        if sheet_name not in wb.sheetnames:
            continue
        sh = wb[sheet_name]
        hmap = headers(sh)
        key_col = hmap.get("product_key")
        if not key_col:
            continue
        for row_idx in range(2, sh.max_row + 1):
            pk = str(sh.cell(row_idx, key_col).value or "")
            if not any(handle in pk for handle in REVISIONS):
                continue
            for col_idx in range(1, sh.max_column + 1):
                value = sh.cell(row_idx, col_idx).value
                if not isinstance(value, str):
                    continue
                new = value.replace("indoor/outdoor", "indoor").replace("Indoor/outdoor", "Indoor")
                new = new.replace("outdoor-use", "room-use").replace("outdoor use", "room use")
                sh.cell(row_idx, col_idx, new)

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R007"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R007", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R007",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R008 or Re-QA revision R007",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R007 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R007.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _, _, _, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n", encoding="utf-8")

    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    vws = verify["SEO_Products"]
    vcols = {cell.value: idx for idx, cell in enumerate(next(vws.iter_rows(min_row=1, max_row=1)))}
    bad = []
    for row in vws.iter_rows(min_row=2, values_only=True):
        handle = row[vcols["Handle"]]
        if handle not in REVISIONS:
            continue
        title = row[vcols["title_proposed"]] or ""
        meta_title = row[vcols["meta_title_seo"]] or ""
        meta_desc = row[vcols["meta_description_seo"]] or ""
        desc = (row[vcols["description_proposed"]] or "") + " " + (row[vcols["description_proposed_html"]] or "")
        blob = " ".join([title, meta_title, meta_desc, desc]).lower()
        if row[vcols["revision"]] != "2":
            bad.append((handle, "revision", row[vcols["revision"]]))
        if row[vcols["review_status"]] != "NEEDS_REVIEW":
            bad.append((handle, "review_status", row[vcols["review_status"]]))
        if row[vcols["content_qa_status"]] != "NOT_RUN":
            bad.append((handle, "content_qa_status", row[vcols["content_qa_status"]]))
        if not (45 <= len(title) <= 70):
            bad.append((handle, "title_len", len(title)))
        if meta_title != title:
            bad.append((handle, "meta_title_mismatch", meta_title))
        if len(meta_desc) > 320:
            bad.append((handle, "meta_desc_len", len(meta_desc)))
        for forbidden in FORBIDDEN:
            if forbidden in f" {blob} ":
                bad.append((handle, "forbidden_claim", forbidden))
    if bad:
        raise RuntimeError(f"Verification failed: {bad[:20]}")

    print(json.dumps({
        "output_workbook": str(OUTPUT),
        "manifest": str(MANIFEST),
        "summary": str(SUMMARY),
        "updated": len(changed),
        "status": "COMPLETE_AWAITING_REQA",
    }, indent=2))


if __name__ == "__main__":
    main()

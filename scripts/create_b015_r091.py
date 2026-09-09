from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B015"
REVISION = "R091"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-102": ("Ghost Vortex Halloween Rug", "ghost vortex Halloween rug", "Halloween optical illusion rug, spooky ghost round rug, 3D Halloween rug", "ghost vortex artwork, a deep tunnel effect, purple trim, pumpkins, and stars"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-103": ("Raven Autumn Halloween Rug", "raven autumn Halloween rug", "Halloween autumn rug, raven round rug, spooky fall floor decor", "raven artwork, autumn colors, a deep tunnel effect, pumpkins, and a moonlit seasonal scene"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-104": ("Checkerboard Ghost Illusion Rug", "checkerboard ghost illusion rug", "checkerboard Halloween rug, ghost round rug, optical illusion floor decor", "checkerboard artwork, ghost details, pumpkins, and a dramatic Halloween illusion"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-105": ("Ghost Campfire Halloween Rug", "ghost campfire Halloween rug", "ghost Halloween rug, campfire round rug, spooky seasonal floor decor", "ghost and campfire artwork, autumn accents, pumpkins, and a playful Halloween scene"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-106": ("Crawling Skeleton Pit Rug", "crawling skeleton pit rug", "skeleton Halloween rug, skeleton pit round rug, horror floor decor", "a crawling skeleton, stone-pit artwork, dark shadows, and a dramatic 3D illusion"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-107": ("Graveyard Checkerboard Illusion Rug", "graveyard checkerboard rug", "graveyard Halloween rug, checkerboard round rug, spooky optical illusion decor", "graveyard artwork, checkerboard patterning, pumpkins, and a dramatic Halloween scene"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-108": ("Haunted Path Ghost Pumpkin Rug", "haunted path ghost pumpkin rug", "ghost pumpkin Halloween rug, haunted path round rug, spooky floor decor", "a haunted path, ghost and pumpkin artwork, autumn colors, and a deep optical-illusion effect"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-109": ("Corgi Haunted Castle Halloween Rug", "corgi haunted castle Halloween rug", "corgi Halloween rug, haunted castle round rug, dog lover floor decor", "a corgi, haunted castle artwork, pumpkins, and a playful spooky-season scene"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-110": ("Bat Vortex Halloween Rug", "bat vortex Halloween rug", "bat Halloween rug, vortex round rug, spooky optical illusion floor decor", "bat artwork, a swirling vortex effect, pumpkins, and dark seasonal colors"),
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-111": ("Purple Haunted Village Halloween Rug", "purple haunted village rug", "haunted village Halloween rug, purple round rug, spooky seasonal floor decor", "a purple haunted village, ghostly details, pumpkins, and a dramatic nighttime scene"),
}

def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []

def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)

def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B015 scope does not match the R091 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        meta = f"Create a {primary} with {motif}."
        description = f"Create a {primary} with {motif}. The round design adds a playful seasonal focal point to a living room, entryway, party space, or Halloween display."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))

if __name__ == "__main__":
    main()

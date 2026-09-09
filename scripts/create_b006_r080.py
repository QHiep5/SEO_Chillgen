from pathlib import Path
import openpyxl, shutil, hashlib, json

BASE=Path('/Users/buiquanghuy/Documents/seo_chillgen'); RUN=BASE/'seo_runs/chillgen.com/chillgen_20260907_01'
src=RUN/'qa_scope/B006_scope_workbook.xlsx'; outdir=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R080'; outdir.mkdir(parents=True,exist_ok=True)
dst=outdir/'SEO_Product_Optimization_revision_R080.xlsx'; shutil.copy2(src,dst)
w=openpyxl.load_workbook(dst); ws=w['SEO_Products']; heads=[c.value for c in next(ws.iter_rows())]; ix={h:i+1 for i,h in enumerate(heads)}
for row in ws.iter_rows(min_row=2):
 h=row[ix['Handle']-1].value
 if not h: continue
 dragon='dragon' in h
 if dragon:
  color='green' if 'design-10' in h else 'orange' if 'design-04' in h else 'purple' if 'design-06' in h else 'red' if 'design-01' in h else 'fantasy'
  title=f'Personalized {color} dragon book rug'
  meta=f'Create a personalized {color} dragon book rug with fantasy artwork and a distinctive shaped design for a reading nook or themed room.'
  desc=f'Create a personalized {color} dragon book rug featuring a detailed fantasy scene across an open-book shape. Add a name to make the design your own, then style it in a reading nook, bedroom, or creative space. The low-profile surface keeps the artwork crisp while the shaped silhouette brings storybook character to your room.'
 else:
  title='Dog Paw Patchwork Round Rug'
  meta='Bring a playful pet-lover accent to your room with a colorful dog paw patchwork round rug in a cozy, easy-care style.'
  desc='Bring a playful pet-lover accent to your room with a colorful dog paw patchwork round rug. Repeating paw motifs and contrasting fabric-style squares create a warm focal point for a family room, bedroom, or reading corner. The round silhouette makes it easy to layer into everyday home décor.'
 for field,val in [('title_proposed',title),('meta_title_seo',title),('meta_description_seo',meta),('description_proposed',desc),('description_proposed_html',f'<p>{desc}</p>')]:
  if field in ix: row[ix[field]-1].value=val
 for field,val in [('revision',2),('review_status','NEEDS_REVIEW'),('content_qa_status','NOT_RUN'),('processing_status','REVISION_CREATED'),('evidence_status','PENDING_REQA')]:
  if field in ix: row[ix[field]-1].value=val
 if 'meta_title_chars' in ix: row[ix['meta_title_chars']-1].value=len(title)
 if 'meta_description_chars' in ix: row[ix['meta_description_chars']-1].value=len(meta)
if 'Revision_Log' in w.sheetnames: del w['Revision_Log']
log=w.create_sheet('Revision_Log'); cols=['revision_batch_id','revision','batch_id','handle','scope_status','review_status','content_qa_status','source_workbook','evidence_id']; log.append(cols)
for row in ws.iter_rows(min_row=2):
 h=row[ix['Handle']-1].value
 if h: log.append(['R080',2,'B006',h,'LOCKED','NEEDS_REVIEW','NOT_RUN',str(dst),f'B006-R080-{log.max_row:03d}'])
w.save(dst)
sha=hashlib.sha256(dst.read_bytes()).hexdigest()
(outdir/'revision_R080_manifest.json').write_text(json.dumps({'revision':'R080','batch_id':'B006','scope_count':10,'canonical_workbook':str(dst),'sha256':sha,'status':'CREATED_FOR_REQA','approval_status':'NOT_APPROVED_NOT_DEPLOYED'},indent=2)+'\n')
print(json.dumps({'workbook':str(dst),'sha256':sha,'scope_count':10}))

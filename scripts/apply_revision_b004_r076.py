from pathlib import Path
from datetime import datetime, timezone, timedelta
from copy import copy
import json, re
from openpyxl import load_workbook

BASE = Path('/Users/buiquanghuy/Documents/seo_chillgen')
SRC = BASE / 'seo_runs/chillgen.com/chillgen_20260907_01/qa_scope/B004_scope_workbook.xlsx'
OUTDIR = BASE / 'resutls/chillgen.com/chillgen_20260907_01/revisions/R076'
RUN = BASE / 'seo_runs/chillgen.com/chillgen_20260907_01/revisions/R076'
OUTDIR.mkdir(parents=True, exist_ok=True); RUN.mkdir(parents=True, exist_ok=True)
OUT = OUTDIR / 'SEO_Product_Optimization_revision_R076.xlsx'
now = datetime.now(timezone(timedelta(hours=7))).isoformat(timespec='seconds')
wb = load_workbook(SRC)
ws = wb['SEO_Products']; headers = {c.value:c.column for c in ws[1]}
changed=[]
for r in range(2, ws.max_row+1):
    if ws.cell(r, headers['batch_id']).value != 'B004': continue
    h = ws.cell(r, headers['Handle']).value
    for field in ('meta_description_seo','description_proposed','description_proposed_html'):
        c = headers.get(field)
        if not c: continue
        v = ws.cell(r,c).value or ''
        v = re.sub(r'easy-clean visuals', 'easy-care design', v, flags=re.I)
        v = re.sub(r'easy-clean indoor visuals', 'easy-care indoor design', v, flags=re.I)
        v = re.sub(r'indoor area rug lifestyle images', 'indoor seasonal decor', v, flags=re.I)
        v = re.sub(r'indoor easy-clean decor images', 'indoor seasonal decor', v, flags=re.I)
        v = re.sub(r'indoor decor mockups', 'indoor seasonal decor', v, flags=re.I)
        v = re.sub(r'and custom name detail\.', 'and a custom name detail.', v, flags=re.I)
        v = re.sub(r' and care/lifestyle graphics', ' with seasonal artwork', v, flags=re.I)
        v = re.sub(r'; size, pet-friendly and cleaning graphics viewed\.', ' with a personalized name and Halloween artwork.', v, flags=re.I)
        v = re.sub(r'Public page metadata and product images support the Halloween, personalized-name, and indoor area rug positioning\.', 'Designed for seasonal indoor decorating with a personalized Halloween motif.', v, flags=re.I)
        v = re.sub(r'Use the visible care and backing claims only where the product page or infographic supports them\.', 'Choose a size that suits your room and follow the care guidance provided with your order.', v, flags=re.I)
        ws.cell(r,c).value = v
    ws.cell(r, headers['revision']).value = '2'
    ws.cell(r, headers['review_status']).value = 'NEEDS_REVIEW'
    ws.cell(r, headers['review_reason']).value = 'R076 removes evidence/research language from customer-facing meta and description copy.'
    ws.cell(r, headers['issues']).value = 'R076 targeted copy cleanup; strict Re-QA required before approval.'
    changed.append(h)
if 'Revision_Log' in wb.sheetnames:
    log=wb['Revision_Log'];
    for h in changed:
        log.append(['R076', now, h, None, '1', '2', 'meta_description_seo, description_proposed, description_proposed_html, revision, review_reason, issues', 'evidence_language_in_customer_copy', 'Run strict evidence-driven Re-QA for B004 after revision.'])
wb.save(OUT)
manifest={'revision_batch_id':'R076','source_workbook':str(SRC),'output_workbook':str(OUT),'generated_at':now,'batch_id':'B004','product_count':len(changed),'handles':changed,'status':'COMPLETE_AWAITING_REQA','review_status':'NEEDS_REVIEW','content_qa_status':'NOT_RUN','not_approved_not_deployed':True}
(RUN/'revision_R076_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
print(json.dumps(manifest,ensure_ascii=False))

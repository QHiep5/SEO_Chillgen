from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B027"
REVISION = "R103"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "personalized-orthodox-doormat-with-christian-cross-and-eagle-6be5c719a3-6be5c719a3": ("Red Orthodox Eagle Cross Rug", "red Orthodox eagle cross rug", "Christian cross area rug, Byzantine eagle decor, Orthodox room rug", "a red Orthodox cross and eagle design with a decorative border for a living room, prayer room, bedroom, or study"),
    "personalized-orthodox-christian-area-rug-00beb8fce2": ("Floral Orthodox Cross Rug", "floral Orthodox cross rug", "Christian cross rug, floral Orthodox decor, Byzantine-inspired area rug", "an Orthodox cross design with floral details for a living room, prayer room, bedroom, or study"),
    "custom-orthodox-rug-three-bar-cross-byzantine-eagle-d81169235b-d81169235b": ("Red Byzantine Eagle Rug", "red Byzantine eagle rug", "Orthodox eagle rug, Christian home decor, Byzantine-inspired room rug", "a red Byzantine eagle design with an ornate border for a living room, prayer room, bedroom, or study"),
    "personalized-orthodox-christian-area-rug-fb5a56fef9-fb5a56fef9": ("Black Silver Orthodox Cross Rug", "black silver Orthodox cross rug", "Christian cross area rug, Orthodox home decor, black silver room rug", "a black and silver Orthodox cross design with decorative details for a living room, prayer room, bedroom, or study"),
    "personalized-teachers-classroom-rules-rug-b0d7bd292a-b0d7bd292a": ("Behavior Champs Classroom Rug", "Behavior Champs classroom rug", "classroom rules rug, teacher name classroom mat, positive learning decor", "Behavior Champs wording with classroom rules and space for teacher details in a classroom, playroom, or learning space"),
    "personalized-classroom-rules-welcome-mat-07b188eb0d": ("Zone Of Tolerance Classroom Rug", "Zone of Tolerance classroom rug", "classroom rules mat, emotional learning classroom rug, teacher decor", "Zone of Tolerance wording with classroom guidance details for a classroom, counseling corner, or learning space"),
    "personalized-teachers-classroom-rug-with-custom-name-87d858996b": ("Calm Down Classroom Rug", "calm down classroom rug", "calm corner classroom mat, emotional learning rug, teacher classroom decor", "Calm Down wording with a classroom support design and space for teacher details in a classroom or calm corner"),
    "personalized-classroom-rules-rug-custom-teacher-name-2b9854c06f": ("Calming Techniques Classroom Rug", "calming techniques classroom rug", "calm corner rug, classroom coping skills mat, emotional learning decor", "Calming Techniques wording with classroom guidance details for a calm corner, classroom, or counseling space"),
    "personalized-classroom-rules-rug-with-custom-teacher-name-ad0787fb44": ("Little Dinos Classroom Rug", "Little Dinos classroom rug", "dinosaur classroom rug, teacher name learning mat, kids room decor", "Little Dinos artwork with classroom details and space for a teacher name in a classroom, playroom, or kids room"),
    "custom-octopus-shaped-area-rug-9041fd8f83": ("Octopus Shaped Area Rug", "octopus shaped area rug", "sea life kids rug, ocean themed playroom mat, octopus room decor", "an octopus-shaped design with ocean-inspired artwork for a playroom, bedroom, classroom, or kids room"),
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B027 scope does not match the R103 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        article = "an" if primary.lower().startswith(("a ", "e ", "i ", "o ", "u ")) else "a"
        meta = f"Create {article} {primary} with {motif}."
        description = f"Create {article} {primary} with {motif}. The design adds a welcoming focal point to a classroom, playroom, bedroom, prayer room, study space, or living room."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

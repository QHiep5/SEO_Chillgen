from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R022/SEO_Product_Optimization_revision_R022.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R023"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R023.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R023"
MANIFEST = RUN_DIR / "revision_R023_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R023_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-orthodox-christian-area-rug-custom-eastern-design-106": {
        "title": "Cream Red Orthodox Cross Rug with Floral Border",
        "meta": "Personalize a cream red Orthodox cross rug with gold cross artwork, deep red border, floral accents and teal trim.",
        "description": "Personalize a cream red Orthodox cross rug with gold cross artwork, deep red border, floral accents and teal trim.\n\nDesign details\n- Rectangular cream rug artwork shows a gold Orthodox cross with deep red border and floral corner accents.\n- Product images include room mockups, close-up panels and size reference graphics.\n- SEO copy focuses on visible Orthodox cross artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-107": {
        "title": "Green Cream Orthodox Cross Rug with Greek Key Border",
        "meta": "Personalize a green cream Orthodox cross rug with green cross motif, Greek key border and floral corner accents.",
        "description": "Personalize a green cream Orthodox cross rug with green cross motif, Greek key border and floral corner accents.\n\nDesign details\n- Rectangular cream rug artwork shows a green Orthodox cross motif with green Greek key border.\n- Visible details include floral corners, green trim, room mockups and size reference graphics.\n- SEO copy stays tied to visible cross and border artwork without material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-family-couple-doormat-custom-couple-husband-design-05": {
        "title": "Funny Personalized Couple Doormat with Cartoon Portraits",
        "meta": "Personalize a funny couple doormat with cartoon portraits, custom names, family name text and proposal joke artwork.",
        "description": "Personalize a funny couple doormat with cartoon portraits, custom names, family name text and proposal joke artwork.\n\nDesign details\n- Tan rectangular doormat artwork shows cartoon couple portraits, custom names and humorous proposal-style text.\n- Product images include couple holding mat, entryway dog mockup, size reference graphics and portrait options.\n- SEO copy focuses on visible couple portrait artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "abc-kids-play-rug-non-slip-classroom-carpet-07621e8f46-07621e8f46": {
        "title": "ABC Animal Alphabet Kids Rug with Numbers and Shapes",
        "meta": "Decorate with an ABC animal alphabet kids rug featuring letter blocks, numbers, colors, shapes and classroom artwork.",
        "description": "Decorate with an ABC animal alphabet kids rug featuring letter blocks, numbers, colors, shapes and classroom artwork.\n\nDesign details\n- Large colorful rectangular rug artwork shows A to Z animal letter blocks with number, color and shape sections.\n- Product images include playroom and classroom mockups, detail panels and size reference graphics.\n- SEO copy describes visible alphabet artwork only, without learning outcome, material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-feelings-chart-classroom-welcome-mat-45af89adb4-45af89adb4": {
        "title": "Feelings Chart Classroom Rug with Rainbow Faces",
        "meta": "Decorate with a feelings chart classroom rug featuring All Feelings Are Welcome Here text, rainbow art and mood faces.",
        "description": "Decorate with a feelings chart classroom rug featuring All Feelings Are Welcome Here text, rainbow art and mood faces.\n\nDesign details\n- Cream rectangular rug artwork shows All Feelings Are Welcome Here wording with rainbow graphics and emotion faces.\n- Visible labels include happy, angry, sad, excited, jealous, tired, worried and frustrated.\n- SEO copy describes visible feelings artwork only, without therapeutic, safety, material, cleaning or surface-performance claims.",
    },
    "custom-hundred-acre-wood-map-welcome-mat-23388cf0b9-23388cf0b9": {
        "title": "Hundred Acre Wood Map Rug with Storybook Forest Art",
        "meta": "Decorate with a Hundred Acre Wood map rug featuring illustrated forest paths, labeled places, compass marks and trees.",
        "description": "Decorate with a Hundred Acre Wood map rug featuring illustrated forest paths, labeled places, compass marks and trees.\n\nDesign details\n- Cream rectangular storybook map rug artwork shows an illustrated forest with paths, labeled locations and trees.\n- Gallery images include bedroom, living room and nursery-style mockups plus size reference panels.\n- SEO copy focuses on visible map artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "educational-classroom-rug-alphabet-handwriting-samplers-17817a1b4c-17817a1b4c": {
        "title": "Alphabet Handwriting Classroom Rug with School Icons",
        "meta": "Decorate with an alphabet handwriting classroom rug featuring colorful uppercase letters, school icons, animals and pencils.",
        "description": "Decorate with an alphabet handwriting classroom rug featuring colorful uppercase letters, school icons, animals and pencils.\n\nDesign details\n- Black rectangular rug artwork shows colorful uppercase A to Z handwriting-style letters with school icons.\n- Visible details include animals, pencils, classroom mockups, corner views and size reference graphics.\n- SEO copy describes visible alphabet artwork only, without learning outcome, material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-classroom-rug-for-kids-and-teachers-940ca1aee9-940ca1aee9": {
        "title": "Feelings Zone Classroom Rug with Color Panels",
        "meta": "Decorate with a feelings zone classroom rug featuring blue, green, yellow and red mood zones with face labels.",
        "description": "Decorate with a feelings zone classroom rug featuring blue, green, yellow and red mood zones with face labels.\n\nDesign details\n- White rectangular classroom rug artwork shows What Zone Are You In text with blue, green, yellow and red emotion zones.\n- Product images include playroom, classroom and teacher-child scenes plus close-up design panels.\n- SEO copy describes visible zone artwork only, without therapeutic, safety, material, cleaning or surface-performance claims.",
    },
    "custom-soccer-championship-trophy-round-rug-f7ac8bf3d4-f7ac8bf3d4": {
        "title": "Soccer Host Flags Round Rug with Trophy Artwork",
        "meta": "Decorate with a soccer host flags round rug featuring trophy, soccer ball, Canada USA Mexico text and fireworks.",
        "description": "Decorate with a soccer host flags round rug featuring trophy, soccer ball, Canada USA Mexico text and fireworks.\n\nDesign details\n- Round cream soccer rug artwork shows trophy, soccer ball, United As One wording, stars and flag border.\n- Visible text includes Canada, USA and Mexico with fireworks and room mockups.\n- SEO copy focuses on visible soccer host artwork without unsupported material, reverse-side or surface-performance claims.",
    },
    "custom-soccer-championship-round-rug-6b6ea4fbf7": {
        "title": "Soccer Tournament Trophy Rug with Stadium Lights",
        "meta": "Decorate with a soccer tournament trophy rug featuring large trophy, soccer ball, scoreboard, flags and stadium lights.",
        "description": "Decorate with a soccer tournament trophy rug featuring large trophy, soccer ball, scoreboard, flags and stadium lights.\n\nDesign details\n- Round colorful soccer rug artwork shows trophy, soccer ball, scoreboard and World Game 2026 text.\n- Visible details include stadium lights and country flags including USA, Canada and Mexico.\n- SEO copy stays tied to visible soccer tournament artwork without material, reverse-side or surface-performance claims.",
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "kid friendly", "pet friendly", "safe durable",
    "machine-washable", "machine washable", "washable ", "quick-dry", "quick dry",
    "memory foam", "microfiber", "velvet", "stain", "fade resistant", "easy clean",
    "easy-clean", "waterproof", "absorbent", "absorption", "backing", "soft non-slip",
    "layered construction", "prayer room", "prayer space", "sacred", "help kids",
    "teach ", "support ", "supports ", "use a ", "create a calming", "calming corner",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    blob = " ".join([title, meta, item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = ["Handle", "title_proposed", "meta_title_seo", "meta_title_chars", "meta_description_seo", "meta_description_chars", "description_proposed", "description_proposed_html", "revision", "review_status", "review_reason", "processing_status", "content_qa_status", "issues"]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R023 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R023 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - {c[0] for c in changed})}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R023"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append(["R023", NOW, handle, row_idx, old_revision, new_revision, "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues", "T1; T2; unsupported_claims", "Run evidence-driven re-QA on revision 2 for this product before approval."])

    wb.save(OUTPUT)

    manifest = {"revision_batch_id": "R023", "generated_at": NOW, "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT), "source_revision_plan": str(PLAN), "product_count": len(changed), "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA", "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN", "not_approved_not_deployed": True, "next_step": "Bắt đầu revision R024 or Re-QA revision R023"}
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = ["# Revision R023 Summary", "", f"- Generated at: `{NOW}`", f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`", "- Scope: 10 products from revision plan R023.", "- Fixes targeted QA issues: short T1/T2 and unsupported claims.", "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.", "- No Shopify deploy, no approval.", "", "## Updated products", "", "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|"]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([str(title or ""), str(meta or ""), str(row[out_cols["description_proposed"] - 1] or "")]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({"output": str(OUTPUT), "updated": seen, "revision_log_rows": log.max_row - 1, "r023_log_rows": counts.get("R023", 0), "duplicate_meta_titles": len(dup_meta_titles), "sheets": reopened.sheetnames}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

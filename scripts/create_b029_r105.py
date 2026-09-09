from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B029"
REVISION = "R105"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "custom-running-horse-welcome-mat-fade3d1429": ("Two Horse Area Rug", "two horse area rug", "horse lover room rug, equestrian welcome mat, western horse decor", "a two-horse design for a bedroom, living room, entryway, office, or equestrian space"),
    "custom-running-horse-doormat-2a0ef19423-2a0ef19423": ("Tribal Horse Area Rug", "tribal horse area rug", "western horse rug, equestrian room mat, horse themed home decor", "a tribal-inspired running horse design for a bedroom, living room, entryway, office, or equestrian space"),
    "horse-area-rug-running-carpet-beb8f5b5cd-beb8f5b5cd": ("Sunset Running Horse Rug", "sunset running horse rug", "sunset horse area rug, western equestrian mat, horse lover decor", "a running horse silhouette with sunset artwork for a bedroom, living room, office, or equestrian space"),
    "custom-running-horse-doormat-b25dd5e514": ("Sunflower Horse Area Rug", "sunflower horse area rug", "sunflower equestrian rug, horse room mat, western floral horse decor", "a horse design with sunflower artwork for a bedroom, living room, entryway, or equestrian space"),
    "custom-western-running-horse-door-mat-e6671cf5b4-e6671cf5b4": ("Fire Horse Area Rug", "fire horse area rug", "flame horse rug, western equestrian decor, running horse room mat", "a running horse design with fire-inspired artwork for a bedroom, living room, office, or equestrian space"),
    "custom-running-horse-welcome-mat-34b7c38f01": ("Black Horse Portrait Rug", "black horse portrait rug", "black horse area rug, equestrian portrait mat, horse lover room decor", "a black horse portrait design for a bedroom, living room, office, or equestrian space"),
    "custom-running-horse-welcome-mat-509ed4d18f-509ed4d18f": ("Daisy Horse Area Rug", "daisy horse area rug", "floral horse rug, daisy equestrian mat, horse lover home decor", "a horse design with daisy artwork for a bedroom, living room, entryway, or equestrian space"),
    "custom-running-horse-area-rug-e5f0770fd6-e5f0770fd6": ("Horse Eye Area Rug", "horse eye area rug", "horse portrait rug, equestrian close-up mat, horse lover room decor", "a close-up horse eye design for a bedroom, living room, office, or equestrian space"),
    "custom-running-horse-welcome-mat-1edee5e89b-1edee5e89b": ("White Floral Horse Rug", "white floral horse rug", "white horse area rug, floral equestrian mat, horse lover home decor", "a white horse design with floral artwork for a bedroom, living room, entryway, or equestrian space"),
    "custom-running-horse-welcome-mat-01f80a49ee": ("Butterfly Horse Area Rug", "butterfly horse area rug", "butterfly equestrian rug, colorful horse mat, horse lover room decor", "a horse design with butterfly artwork for a bedroom, playroom, living room, or equestrian space"),
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B029 scope does not match the R105 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        article = "an" if primary.lower().startswith(("a ", "e ", "i ", "o ", "u ")) else "a"
        meta = f"Create {article} {primary} with {motif}."
        description = f"Create {article} {primary} with {motif}. The design adds a distinctive focal point to a bedroom, living room, entryway, office, or equestrian space."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

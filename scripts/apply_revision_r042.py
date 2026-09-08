from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R041/SEO_Product_Optimization_revision_R041.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R042"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R042.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R042"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-tree-of-life-welcome-mat-e20613a50e-e20613a50e": {
        "title": "Green Gold Celtic Tree Of Life Welcome Mat Artwork",
        "meta": "Decorate with a green gold Celtic Tree of Life welcome mat featuring leafy tree artwork, round knot border and yellow green tones.",
        "primary": "green gold Celtic Tree of Life mat",
        "secondary": "Tree of Life welcome mat, Celtic knot border rug, yellow green tree art",
        "long_tail": "green gold Celtic Tree of Life mat; Tree of Life welcome mat; Celtic knot border rug",
        "description": "Decorate with a green gold Celtic Tree of Life welcome mat featuring leafy tree artwork, round knot border and yellow green tones.\n\nDesign details\n- Green and gold rug artwork centers a leafy Tree of Life with winding roots inside a circular Celtic-style knot border.\n- Visible product imagery includes entryway and lounge mockups, close-up views, size reference graphics and custom text examples.\n- SEO copy stays tied to visible tree artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            4: "Feature infographic for green gold Celtic Tree of Life mat",
        },
    },
    "tree-of-life-celtic-rug-40c2af5cd2": {
        "title": "Black Floral Celtic Tree Of Life Rug Border Artwork",
        "meta": "Decorate with a black floral Celtic Tree of Life rug featuring flowering tree art, warm botanical border and garden colors.",
        "primary": "black floral Celtic Tree of Life rug",
        "secondary": "flowering Tree of Life rug, botanical Celtic rug, black garden tree rug",
        "long_tail": "black floral Celtic Tree of Life rug; flowering Tree of Life rug; botanical Celtic rug",
        "description": "Decorate with a black floral Celtic Tree of Life rug featuring flowering tree art, warm botanical border and garden colors.\n\nDesign details\n- Black rug artwork shows a brown Tree of Life rising from a flower garden with blossoms across the border.\n- Visible details include floral vines, warm orange and cream flowers, personalized text examples, room mockups and close-up artwork panels.\n- SEO copy focuses on visible botanical tree artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            3: "Close-up of floral Tree of Life border artwork",
            4: "Feature infographic for black floral Tree of Life rug",
        },
    },
    "tree-of-life-celtic-indoor-floor-mat-57eef60356": {
        "title": "Purple Celestial Tree Of Life Rug with Moon Art",
        "meta": "Decorate with a purple celestial Tree of Life rug featuring crescent moon, stars, planets, galaxy colors and tree artwork.",
        "primary": "purple celestial Tree of Life rug",
        "secondary": "moon and stars Tree of Life rug, galaxy tree rug, purple celestial mat",
        "long_tail": "purple celestial Tree of Life rug; moon and stars Tree of Life rug; galaxy tree rug",
        "description": "Decorate with a purple celestial Tree of Life rug featuring crescent moon, stars, planets, galaxy colors and tree artwork.\n\nDesign details\n- Purple rug artwork shows a black Tree of Life silhouette beneath crescent moons, stars and planet details.\n- Visible product imagery includes patio and room mockups, close-up galaxy artwork, size guide views and color-rich celestial scenes.\n- SEO copy stays tied to visible moon-and-star tree artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            1: "Purple celestial Tree of Life rug with moons and stars in room setting",
            3: "Close-up of purple moon and star Tree of Life artwork",
            4: "Feature infographic for celestial Tree of Life rug",
            6: "Bright living room mockup with purple celestial tree rug",
        },
    },
    "tree-of-life-rug-yggdrasil-celtic-tree-of-lifes-large-area-rug-modern-indoor-floor-mat-soft-carpet-with-non-slip-backing-home-decor-for-living-room-bedroom-gifts-for-tree-of-life-d-346d281f76": {
        "title": "Teal Gold Yggdrasil Tree Of Life Rug Wall Artwork",
        "meta": "Decorate with a teal gold Yggdrasil Tree of Life rug featuring night landscape, golden roots, starry sky and tree artwork.",
        "primary": "teal gold Yggdrasil Tree of Life rug",
        "secondary": "golden root Celtic rug, teal Tree of Life rug, night landscape Yggdrasil rug",
        "long_tail": "teal gold Yggdrasil Tree of Life rug; golden root Celtic rug; night landscape Yggdrasil rug",
        "description": "Decorate with a teal gold Yggdrasil Tree of Life rug featuring night landscape, golden roots, starry sky and tree artwork.\n\nDesign details\n- Teal rug artwork shows a luminous Yggdrasil tree with golden branches and roots spreading across a night landscape.\n- Visible details include starry sky, mountain horizon, Celtic-inspired root shapes, room mockups, close-up views and size graphics.\n- SEO copy focuses on visible Yggdrasil artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            3: "Close-up of golden roots and teal Yggdrasil artwork",
            4: "Feature infographic for teal gold Yggdrasil rug",
        },
    },
    "tree-of-life-yggdrasil-celtic-area-rug-e8cc0189f2-e8cc0189f2": {
        "title": "Black Cream Yggdrasil Rug with Celtic Knot Border",
        "meta": "Decorate with a black cream Yggdrasil rug featuring Celtic knot border, birds, flowing roots and Tree of Life artwork.",
        "primary": "black cream Yggdrasil rug",
        "secondary": "Celtic knot border rug, Tree of Life area rug, black cream tree mat",
        "long_tail": "black cream Yggdrasil rug; Celtic knot border rug; Tree of Life area rug",
        "description": "Decorate with a black cream Yggdrasil rug featuring Celtic knot border, birds, flowing roots and Tree of Life artwork.\n\nDesign details\n- Black and cream rug artwork shows a bold Yggdrasil tree with twisting roots and leafy branches.\n- Visible product imagery includes Celtic knot border details, bird silhouettes, entryway mockups, close-up panels and custom family-name examples.\n- SEO copy stays tied to visible black-and-cream tree artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            4: "Feature infographic for black cream Yggdrasil rug",
        },
    },
    "custom-composition-notebook-classroom-rug-for-kids-975402570a-975402570a": {
        "title": "Mrs Marrabel Pastel Classroom Rug with Pencil Art",
        "meta": "Personalize a Mrs Marrabel pastel classroom rug with pencil frame, polka dots, apple, rainbow, book and name artwork.",
        "primary": "Mrs Marrabel pastel classroom rug",
        "secondary": "pastel teacher name rug, pencil frame classroom mat, polka dot school rug",
        "long_tail": "Mrs Marrabel pastel classroom rug; pastel teacher name rug; pencil frame classroom mat",
        "description": "Personalize a Mrs Marrabel pastel classroom rug with pencil frame, polka dots, apple, rainbow, book and name artwork.\n\nDesign details\n- Pastel classroom rug artwork shows Mrs Marrabel name in large colorful letters on a dotted background.\n- Visible details include pencil frame, apple, rainbow, open book, heart accents, classroom mockups, close-up views and size reference images.\n- SEO copy focuses on visible teacher-name artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            8: "Feature graphic showing pastel classroom rug artwork",
        },
    },
    "custom-composition-notebook-classroom-rug-e78523a9ef-e78523a9ef": {
        "title": "Floral Composition Notebook Rug with Teacher Name",
        "meta": "Personalize a floral composition notebook rug with teacher name, book spines, flowers, butterflies and black speckled style.",
        "primary": "floral composition notebook rug",
        "secondary": "teacher name notebook rug, book spine classroom rug, floral reading corner mat",
        "long_tail": "floral composition notebook rug; teacher name notebook rug; book spine classroom rug",
        "description": "Personalize a floral composition notebook rug with teacher name, book spines, flowers, butterflies and black speckled style.\n\nDesign details\n- Black speckled composition notebook-style rug artwork frames a colorful book stack with Mrs Smith name text.\n- Visible details include flowers, butterflies, pencils, dotted border, classroom scene, sofa mockup, close-up view and size guide.\n- SEO copy stays tied to visible floral notebook artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            7: "Feature graphic showing floral composition notebook rug artwork",
        },
    },
    "custom-composition-notebook-classroom-rug-kids-3e5de6246c": {
        "title": "Composition Notebook Rug with Pastel Emotion Faces",
        "meta": "Personalize a composition notebook rug with pastel emotion faces, teacher label area, pencil tag and black speckled style.",
        "primary": "composition notebook emotion faces rug",
        "secondary": "pastel emotion classroom rug, teacher label rug, black notebook style rug",
        "long_tail": "composition notebook emotion faces rug; pastel emotion classroom rug; teacher label rug",
        "description": "Personalize a composition notebook rug with pastel emotion faces, teacher label area, pencil tag and black speckled style.\n\nDesign details\n- Black composition notebook-style rug artwork shows a cluster of pastel faces with varied expressions and Mrs Mottel name text.\n- Visible imagery includes classroom mockups, playroom scenes, pencil tag, close-up face details and size reference graphics.\n- SEO copy focuses on visible emotion-face artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic showing pastel emotion faces rug artwork",
        },
    },
    "custom-mathematics-education-classroom-rug-0639e24f25": {
        "title": "Multiplication Chart Classroom Rug with Doodles",
        "meta": "Decorate with a multiplication chart classroom rug featuring black chalkboard style, colorful numbers, stars and school doodles.",
        "primary": "multiplication chart classroom rug",
        "secondary": "math facts classroom rug, chalkboard math rug, colorful number rug",
        "long_tail": "multiplication chart classroom rug; math facts classroom rug; chalkboard math rug",
        "description": "Decorate with a multiplication chart classroom rug featuring black chalkboard style, colorful numbers, stars and school doodles.\n\nDesign details\n- Black classroom rug artwork presents a multiplication chart with bright number markers across the grid.\n- Visible details include colorful math symbols, stars, planets, pencil doodles, classroom mockups, room scenes and size reference images.\n- SEO copy stays tied to visible multiplication-chart artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            7: "Feature graphic showing multiplication chart rug artwork",
        },
    },
    "custom-a-good-day-to-read-book-rug-d1d331ba7d-d1d331ba7d": {
        "title": "Reading Bloom Book Rug with Flowers And Butterflies",
        "meta": "Decorate with a Reading Bloom book rug featuring open book artwork, flowers, butterflies and Let's Read quote text.",
        "primary": "Reading Bloom book rug",
        "secondary": "floral reading rug, open book classroom rug, Let's Read quote rug",
        "long_tail": "Reading Bloom book rug; floral reading rug; open book classroom rug",
        "description": "Decorate with a Reading Bloom book rug featuring open book artwork, flowers, butterflies and Let's Read quote text.\n\nDesign details\n- Black reading rug artwork shows an open book surrounded by flowers, butterflies and colorful classroom-style lettering.\n- Visible text includes Reading Helps Our Mind Bloom and Let's Read, with reading corner mockups, close-up views and size reference images.\n- SEO copy focuses on visible book-and-flower artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic for Reading Bloom book rug artwork",
            7: "Feature infographic for black reading rug artwork",
            8: "Care-style graphic for black reading rug",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable", "floor mat",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R042 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R042 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R042"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R042", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R042", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R043 or Re-QA revision R042",
    }
    (RUN_DIR / "revision_R042_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R042 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R042.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R042 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R042_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r042_log_rows": counts.get("R042", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

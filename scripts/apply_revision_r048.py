from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R047/SEO_Product_Optimization_revision_R047.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R048"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R048.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R048"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-motivational-classroom-rug-for-kids-1499e819f7-1499e819f7": {
        "title": "You Are Affirmation Rug with Bright Word Stripes",
        "meta": "Decorate with a You Are affirmation rug featuring bright adjective stripes, smiley icons, school doodles and playful classroom artwork.",
        "primary": "You Are affirmation word stripes rug",
        "secondary": "bright affirmation classroom rug, smiley icon rug, colorful positive words mat",
        "long_tail": "You Are affirmation word stripes rug; bright affirmation classroom rug; smiley icon rug",
        "description": "Decorate with a You Are affirmation rug featuring bright adjective stripes, smiley icons, school doodles and playful classroom artwork.\n\nDesign details\n- White rug artwork shows You Are header text above stacked colorful word stripes.\n- Visible words include amazing, brilliant, extraordinary, fantastic, incredible, magnificent, phenomenal and spectacular, with room mockups and child scene image.\n- SEO copy stays tied to visible affirmation-stripe artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for You Are word stripe rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-d58c24b35d": {
        "title": "Sunshine I Am Rug with Rainbow Affirmation Rays",
        "meta": "Decorate with a sunshine I Am rug featuring rainbow affirmation rays and positive words such as kind, loved, brave and strong.",
        "primary": "sunshine I Am affirmation rug",
        "secondary": "rainbow affirmation ray rug, colorful I Am classroom rug, positive words mat",
        "long_tail": "sunshine I Am affirmation rug; rainbow affirmation ray rug; colorful I Am classroom rug",
        "description": "Decorate with a sunshine I Am rug featuring rainbow affirmation rays and positive words such as kind, loved, brave and strong.\n\nDesign details\n- Rainbow sunburst rug artwork centers I Am text with affirmation words arranged inside colorful rays.\n- Visible words include kind, loved, brave, unique, grateful, generous, creative, strong, happy and smart, with room mockups and child scene image.\n- SEO copy focuses on visible sunburst affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for sunshine I Am rug"},
    },
    "personalized-motivational-classroom-rug-kids-803fa89d09": {
        "title": "Black You Are Word Collage Rug with Doodle Art",
        "meta": "Decorate with a black You Are word collage rug featuring colorful confidence words, doodles and central You Are artwork.",
        "primary": "black You Are word collage rug",
        "secondary": "colorful affirmation rug, classroom word collage rug, positive message mat",
        "long_tail": "black You Are word collage rug; colorful affirmation rug; classroom word collage rug",
        "description": "Decorate with a black You Are word collage rug featuring colorful confidence words, doodles and central You Are artwork.\n\nDesign details\n- Black rug artwork centers You Are text with surrounding words such as creative, amazing, determined, special, brave and loved.\n- Visible imagery includes colorful doodles, bold typography, sofa mockup, classroom scene and close-up artwork views.\n- SEO copy stays tied to visible word-collage artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {6: "Feature graphic for black You Are word collage rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-c304c62374-c304c62374": {
        "title": "Pastel I Am Classroom Rug with Rainbow Flower Art",
        "meta": "Decorate with a pastel I Am classroom rug featuring rainbow arch, flowers, hearts and colorful confidence word artwork.",
        "primary": "pastel I Am classroom rug",
        "secondary": "rainbow affirmation rug, pastel confidence words rug, flower classroom mat",
        "long_tail": "pastel I Am classroom rug; rainbow affirmation rug; pastel confidence words rug",
        "description": "Decorate with a pastel I Am classroom rug featuring rainbow arch, flowers, hearts and colorful confidence word artwork.\n\nDesign details\n- Pastel rug artwork shows I Am text over a rainbow arch with confidence words arranged in curved color bands.\n- Visible words include kind, smart, brave, grateful, confident and loved, with classroom mockups and close-up views.\n- SEO copy focuses on visible pastel affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {7: "Feature graphic for pastel I Am classroom rug"},
    },
    "custom-music-album-doormat-front-door-b42677b9b9": {
        "title": "Custom Album Cover Rug with Vinyl Player Layout",
        "meta": "Personalize a custom album cover rug with vinyl record player layout, photo circle, song text, artist text and blue gradient artwork.",
        "primary": "custom album cover rug",
        "secondary": "vinyl player rug, personalized song rug, music album artwork mat",
        "long_tail": "custom album cover rug; vinyl player rug; personalized song rug",
        "description": "Personalize a custom album cover rug with vinyl record player layout, photo circle, song text, artist text and blue gradient artwork.\n\nDesign details\n- Music player-style rug artwork shows a vinyl record, tonearm, custom photo circle, song title and artist name text.\n- Visible imagery includes blue and pink gradient versions, room mockups, size option graphic, close-up view and feature panel.\n- SEO copy stays tied to visible album-player artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            4: "Feature infographic for custom album cover rug",
            7: "Close-up of pink album cover rug artwork",
        },
    },
    "personalized-music-album-cover-welcome-mat-4a66513e5d-4a66513e5d": {
        "title": "Personalized Music Album Rug with Player Controls",
        "meta": "Personalize a music album rug with portrait cover, song title, artist name, player controls and warm orange gradient artwork.",
        "primary": "personalized music album rug",
        "secondary": "custom song rug, album player controls rug, portrait music mat",
        "long_tail": "personalized music album rug; custom song rug; album player controls rug",
        "description": "Personalize a music album rug with portrait cover, song title, artist name, player controls and warm orange gradient artwork.\n\nDesign details\n- Orange music player-style rug artwork shows a portrait cover, playback controls, song title and artist name text.\n- Visible product imagery includes bedroom and lounge mockups, size option graphic, close-up artwork view and feature infographic.\n- SEO copy focuses on visible music-player artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            4: "Feature infographic for orange music album rug",
            7: "Close-up of orange album player rug artwork",
        },
    },
    "personalized-music-album-cover-area-rug-9547d6169f-9547d6169f": {
        "title": "Personalized Album Cover Rug with Custom Photo",
        "meta": "Personalize an album cover rug with custom photo area, song title, artist text, player controls and colorful background variants.",
        "primary": "personalized album cover rug",
        "secondary": "custom photo music rug, song title rug, album player area mat",
        "long_tail": "personalized album cover rug; custom photo music rug; song title rug",
        "description": "Personalize an album cover rug with custom photo area, song title, artist text, player controls and colorful background variants.\n\nDesign details\n- Album player-style rug artwork shows a custom photo area, playback controls and song information in a device-like layout.\n- Visible imagery includes black, teal and purple variants, room mockups, size option graphic, close-up view and feature infographic.\n- SEO copy stays tied to visible album-cover artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            6: "Close-up of red album cover rug artwork",
            7: "Feature infographic for album cover rug artwork",
        },
    },
    "custom-3d-dragon-round-rug-77aa65f879": {
        "title": "Red Dragon Library Round Rug with Spiral Artwork",
        "meta": "Decorate with a red dragon library round rug featuring spiral staircase artwork, bookshelves, flowers and fantasy reading room scenes.",
        "primary": "red dragon library round rug",
        "secondary": "spiral library dragon rug, fantasy reading nook rug, round red dragon mat",
        "long_tail": "red dragon library round rug; spiral library dragon rug; fantasy reading nook rug",
        "description": "Decorate with a red dragon library round rug featuring spiral staircase artwork, bookshelves, flowers and fantasy reading room scenes.\n\nDesign details\n- Round rug artwork shows a red dragon curling through a spiral library staircase with book-filled shelves.\n- Visible imagery includes size option graphic, fireplace room scene, close-up spiral view, seasonal room mockup and design callout panels.\n- SEO copy focuses on visible red dragon library artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            7: "Close-up design callout for red dragon round rug",
            8: "Feature graphic for red dragon library rug",
        },
    },
    "custom-3d-dragon-round-rug-c5c3c2de5b": {
        "title": "Green Dragon Bookshelf Round Rug with Spiral Art",
        "meta": "Decorate with a green dragon bookshelf round rug featuring glowing spiral library artwork, open wings and fantasy book room scenes.",
        "primary": "green dragon bookshelf round rug",
        "secondary": "spiral bookshelf dragon rug, fantasy library rug, round green dragon mat",
        "long_tail": "green dragon bookshelf round rug; spiral bookshelf dragon rug; fantasy library rug",
        "description": "Decorate with a green dragon bookshelf round rug featuring glowing spiral library artwork, open wings and fantasy book room scenes.\n\nDesign details\n- Round rug artwork shows a green dragon over a glowing spiral bookshelf design with warm library tones.\n- Visible product imagery includes size option graphic, fireplace scene, close-up spiral views and feature callout panels.\n- SEO copy stays tied to visible green dragon library artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            6: "Design callout for green dragon round rug",
            7: "Feature graphic for green dragon bookshelf rug",
        },
    },
    "custom-3d-effect-dragon-round-rug-47b32027ce": {
        "title": "Black Dragon Library Round Rug with Rose Artwork",
        "meta": "Decorate with a black dragon library round rug featuring sleeping dragon, books, red roses and dark fantasy room scenes.",
        "primary": "black dragon library round rug",
        "secondary": "rose dragon round rug, gothic fantasy dragon rug, book lover dragon mat",
        "long_tail": "black dragon library round rug; rose dragon round rug; gothic fantasy dragon rug",
        "description": "Decorate with a black dragon library round rug featuring sleeping dragon, books, red roses and dark fantasy room scenes.\n\nDesign details\n- Round rug artwork shows a dark sleeping dragon resting on books inside circular library-style artwork framed by red roses.\n- Visible imagery includes size option graphic, fireplace room scene, close-up dragon views and feature callout panels.\n- SEO copy focuses on visible black dragon and rose artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            5: "Feature graphic for black dragon library rug",
            7: "Design detail graphic for black dragon round rug",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-care", "easy care",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable", "floor mat", "floor rug", "low pile", "low-pile", "skid",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R048 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R048 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R048"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R048", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R048", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R049 or Re-QA revision R048",
    }
    (RUN_DIR / "revision_R048_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R048 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R048.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R048 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R048_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r048_log_rows": counts.get("R048", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

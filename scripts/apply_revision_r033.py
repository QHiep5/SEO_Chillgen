from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R032/SEO_Product_Optimization_revision_R032.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R033"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R033.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R033"
MANIFEST = RUN_DIR / "revision_R033_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R033_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-soccer-area-rug-trophy-football-868088cbf4-868088cbf4": {
        "title": "Game Of Champions Soccer Rug with Stadium Lights",
        "meta": "Decorate with a Game Of Champions soccer rug featuring 2026 text, stadium lights, large ball and red border.",
        "description": "Decorate with a Game Of Champions soccer rug featuring 2026 text, stadium lights, large ball and red border.\n\nDesign details\n- Rectangular soccer rug artwork shows Game Of Champions lettering, a large soccer ball over a lit stadium pitch and 2026 text.\n- Visible details include torn-edge artwork, red geometric border, room mockups, close-up panels and size reference graphics.\n- SEO copy stays tied to visible soccer artwork and avoids unsupported care, construction or performance claims.",
    },
    "personalized-soccer-2026-area-rug-95231047da-95231047da": {
        "title": "Host Flags Swirl Soccer Rug with Gold Ball Center",
        "meta": "Decorate with a host flags swirl soccer rug featuring USA, Canada and Mexico color bands, green sections and gold ball center.",
        "description": "Decorate with a host flags swirl soccer rug featuring USA, Canada and Mexico color bands, green sections and gold ball center.\n\nDesign details\n- Bright rectangular soccer rug artwork shows swirling flag-color bands around a gold soccer ball center.\n- Visible details include red, green and white arcs, starburst corner styling, room mockups and size reference graphics.\n- SEO copy stays tied to visible host-flag artwork and avoids unsupported care, construction or performance claims.",
    },
    "custom-soccer-tournament-round-rug-a7f6df2bd9-a7f6df2bd9": {
        "title": "World Flags Trophy Round Rug with Soccer Ball Art",
        "meta": "Decorate with a world flags trophy round rug featuring central gold cup, soccer balls, host-country flags and colorful panels.",
        "description": "Decorate with a world flags trophy round rug featuring central gold cup, soccer balls, host-country flags and colorful panels.\n\nDesign details\n- Round soccer rug artwork shows a central gold trophy, soccer balls, USA Canada Mexico flag elements and world-flag style panels.\n- Product images include living room, bedroom, dining and holiday room mockups plus close-up artwork and size graphics.\n- SEO copy stays tied to visible trophy artwork and avoids unsupported care, construction or performance claims.",
    },
    "personalized-dog-welcome-doormat-front-door-entryway-e2ef212662-e2ef212662": {
        "title": "Four Pet Welcome Home Mat with Dog And Cat Names",
        "meta": "Personalize a four pet welcome home mat with dog and cat portraits, pet names and The Humans Just Live Here text.",
        "description": "Personalize a four pet welcome home mat with dog and cat portraits, pet names and The Humans Just Live Here text.\n\nDesign details\n- Tan multi-pet mat artwork reads Welcome To Our Home The Humans Just Live Here With Us.\n- Visible details include four pet portraits, names Bailey, Margot, Maisy and Loki, variant thumbnails, dog scene and doorway mockups.\n- SEO copy stays tied to visible pet welcome artwork and avoids unsupported care, construction or performance claims.",
    },
    "personalized-dog-welcome-mat-custom-name-fd965cdd3f": {
        "title": "Visitors Approved By Charlie Mat with Paw Prints",
        "meta": "Personalize a Visitors Approved By Charlie mat with custom dog portrait, pet name, paw prints and photo guide.",
        "description": "Personalize a Visitors Approved By Charlie mat with custom dog portrait, pet name, paw prints and photo guide.\n\nDesign details\n- Tan and black dog mat artwork reads Visitors Must Be Approved By Charlie with a cartoon puppy portrait.\n- Visible details include paw prints, photo conversion panel, photo guide, doorway scenes and bedside mockups.\n- SEO copy focuses on visible custom dog artwork and avoids unsupported care, construction or performance claims.",
    },
    "custom-halloween-dog-doormat-405e5466d4-405e5466d4": {
        "title": "Welcome Little Monsters Dog Mat with Halloween Icons",
        "meta": "Personalize a Welcome Little Monsters dog mat with pet names, cartoon portraits, haunted house, bats and pumpkins.",
        "description": "Personalize a Welcome Little Monsters dog mat with pet names, cartoon portraits, haunted house, bats and pumpkins.\n\nDesign details\n- Color-block Halloween pet mat artwork reads Welcome Little Monsters with four pet portraits and custom names.\n- Visible details include haunted house, bats, pumpkins, graveyard, black cat icons, fall porch mockups and doorway scenes.\n- SEO copy stays tied to visible Halloween dog artwork and avoids unsupported care, construction or performance claims.",
    },
    "custom-halloween-dog-doormat-a6bfd1d512-a6bfd1d512": {
        "title": "Home Sweet Haunted Home Dog Mat with Pumpkin Art",
        "meta": "Personalize a Home Sweet Haunted Home dog mat with pet names, red spooky sky, pumpkins, bats and skull border.",
        "description": "Personalize a Home Sweet Haunted Home dog mat with pet names, red spooky sky, pumpkins, bats and skull border.\n\nDesign details\n- Red and black Halloween pet mat artwork shows Home Sweet Haunted Home text with three cartoon dogs named Boo, Coco and Lucky.\n- Visible details include moon, bats, bare trees, haunted silhouettes, skull border, pumpkin porch mockups and variant thumbnails.\n- SEO copy focuses on visible haunted pet artwork and avoids unsupported care, construction or performance claims.",
    },
    "personalized-dog-welcome-mat-custom-text-non-slip-backing-17fda5bd64-17fda5bd64": {
        "title": "Visitors Approved By Pets Mat with Cartoon Dogs",
        "meta": "Personalize a Visitors Approved By Pets mat with cartoon dog portraits, custom pet names, paw prints and breed option panels.",
        "description": "Personalize a Visitors Approved By Pets mat with cartoon dog portraits, custom pet names, paw prints and breed option panels.\n\nDesign details\n- Tan pet mat artwork reads Visitors Must Be Approved By with cartoon dog portraits and custom names.\n- Visible names include Casey, Brownie, Boomer and Laser, with paw prints, breed option panels, door and bedside mockups.\n- SEO copy stays tied to visible custom pet artwork and avoids unsupported care, construction or performance claims.",
    },
    "custom-running-horse-welcome-mat-09686138f0": {
        "title": "Black Horse Sunset Rug with Purple Flower Art",
        "meta": "Decorate with a black horse sunset rug featuring horse portrait, purple flowers, butterflies and warm sky artwork.",
        "description": "Decorate with a black horse sunset rug featuring horse portrait, purple flowers, butterflies and warm sky artwork.\n\nDesign details\n- Rectangular horse rug artwork shows a black horse head and flowing mane against orange sunset light.\n- Visible details include purple flowers, butterflies, petals, room mockups, close-up artwork panels and size reference graphics.\n- SEO copy focuses on visible horse-and-flower artwork and avoids unsupported care, construction or performance claims.",
    },
    "custom-running-horse-area-rug-non-slip-384f7e6e8e": {
        "title": "Chestnut Running Horse Rug with White Mane Art",
        "meta": "Decorate with a chestnut running horse rug featuring white mane, western tack, warm neutral tones and room mockups.",
        "description": "Decorate with a chestnut running horse rug featuring white mane, western tack, warm neutral tones and room mockups.\n\nDesign details\n- Warm-toned horse rug artwork shows a chestnut horse portrait with flowing white mane, bridle and western tack.\n- Visible details include dust and light background, room mockups, close-up artwork panels and size reference graphics.\n- SEO copy stays tied to visible running-horse artwork and avoids unsupported care, construction or performance claims.",
    },
}


FORBIDDEN = [
    "indoor/outdoor", " indoor floor", " outdoor ", "anti-slip", "anti slip",
    "non-slip", "non slip", "non-skid", "non skid", "nonslip", "kid friendly",
    "pet friendly", "safe durable", "durable", "machine-washable", "machine washable",
    "washable", "quick-dry", "quick dry", "memory foam", "microfiber", "velvet",
    "stain", "fade resistant", "fade", "easy clean", "easy-clean", "waterproof",
    "absorbent", "absorption", "backing", "rubber", "layered construction",
    "hd printing", "ultra-soft", "soft ", "cushion", "3 mm", "thickness",
    "support classroom", "support emotional", "learning-space", "thickened",
    "reinforced", "bound edge", "material panel", "bring home",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    blob = " ".join([title, meta, item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R033 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R033 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append([
            "revision_batch_id", "generated_at", "handle", "source_row",
            "old_revision", "new_revision", "changed_fields", "qa_issue_fields",
            "recheck_condition",
        ])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R033"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R033", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R033",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R034 or Re-QA revision R033",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R033 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R033.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        desc = row[out_cols["description_proposed"] - 1]
        blob = " ".join([str(title or ""), str(meta or ""), str(desc or "")]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r033_log_rows": counts.get("R033", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


HANDLE = "custom-dog-paw-print-shaped-rug-39afbc78d5"
OLD_INTERNAL_SENTENCE = "SEO copy avoids backing and surface-performance claims unless confirmed during admin/export review."
NEW_CUSTOMER_SENTENCE = "Pet-themed paw shape and plaid blocks make the artwork easy to match with dog, cat and animal-lover decor."

FILES = [
    Path("resutls/chillgen.com/chillgen_20260907_01/revisions/R001/SEO_Product_Optimization_revision_R001.xlsx"),
    Path("resutls/chillgen.com/chillgen_20260907_01/revisions/R052/SEO_Product_Optimization_revision_R052.xlsx"),
]
RUN_DIR = Path("seo_runs/chillgen.com/chillgen_20260907_01/hotfixes/HF_R001_INTERNAL_COPY")


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def fix_workbook(path: Path) -> dict:
    wb = load_workbook(path)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = ["Handle", "description_proposed", "description_proposed_html", "review_status", "content_qa_status", "revision"]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"{path} missing columns: {missing}")

    fixed = False
    result = {"path": str(path), "row": None, "description_fixed": False, "html_fixed": False}
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row_idx, cols["Handle"]).value != HANDLE:
            continue
        result["row"] = row_idx
        desc = ws.cell(row_idx, cols["description_proposed"]).value or ""
        html = ws.cell(row_idx, cols["description_proposed_html"]).value or ""

        if OLD_INTERNAL_SENTENCE not in desc and OLD_INTERNAL_SENTENCE not in html:
            fixed = True
            break

        desc = desc.replace(OLD_INTERNAL_SENTENCE, NEW_CUSTOMER_SENTENCE)
        html = html.replace(OLD_INTERNAL_SENTENCE, NEW_CUSTOMER_SENTENCE)
        html = html.replace("size and construction panels", "size and design panels")
        desc = desc.replace("size and construction panels", "size and design panels")

        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], html)
        result["description_fixed"] = True
        result["html_fixed"] = True
        fixed = True
        break

    if not fixed or result["row"] is None:
        raise RuntimeError(f"{HANDLE} not found in {path}")

    wb.save(path)
    return result


def verify_workbook(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[cols["Handle"] - 1] != HANDLE:
            continue
        desc = row[cols["description_proposed"] - 1] or ""
        html = row[cols["description_proposed_html"] - 1] or ""
        bad_terms = [
            "SEO copy avoids",
            "unless confirmed during admin/export review",
            "surface-performance claims",
            "backing",
        ]
        hits = [term for term in bad_terms if term.lower() in f"{desc} {html}".lower()]
        if hits:
            raise RuntimeError(f"{path} still has internal/risky terms: {hits}")
        return {
            "path": str(path),
            "row": row_idx,
            "revision": row[cols["revision"] - 1],
            "review_status": row[cols["review_status"] - 1],
            "content_qa_status": row[cols["content_qa_status"] - 1],
            "description_chars": len(desc),
            "html_chars": len(html),
        }
    raise RuntimeError(f"{HANDLE} not found during verification in {path}")


RUN_DIR.mkdir(parents=True, exist_ok=True)
updates = [fix_workbook(path) for path in FILES]
verification = [verify_workbook(path) for path in FILES]
manifest = {
    "hotfix_id": "HF_R001_INTERNAL_COPY",
    "generated_at": "2026-09-08",
    "handle": HANDLE,
    "files_updated": [str(path) for path in FILES],
    "old_internal_sentence": OLD_INTERNAL_SENTENCE,
    "new_customer_sentence": NEW_CUSTOMER_SENTENCE,
    "status": "COMPLETE_VERIFIED",
    "updates": updates,
    "verification": verification,
    "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
}
(RUN_DIR / "hotfix_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")
print(json.dumps(manifest, indent=2, ensure_ascii=False))

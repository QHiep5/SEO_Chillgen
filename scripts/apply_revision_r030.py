from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R029/SEO_Product_Optimization_revision_R029.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R030"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R030.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R030"
MANIFEST = RUN_DIR / "revision_R030_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R030_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-orthodox-doormat-with-christian-cross-and-eagle-6be5c719a3-6be5c719a3": {
        "title": "Red Orthodox Cross Rug with IC XC NIKA Lettering",
        "meta": "Decorate with a red Orthodox cross rug featuring gold three-bar cross, IC XC NIKA lettering and simple border.",
        "description": "Decorate with a red Orthodox cross rug featuring gold three-bar cross, IC XC NIKA lettering and simple border.\n\nDesign details\n- Red rectangular Orthodox rug artwork shows a gold three-bar cross, IC XC NIKA lettering and slim gold border.\n- Product images include living room and fireside mockups, close-up design panels and size reference graphics.\n- SEO copy stays tied to visible cross artwork without unsupported material, cleaning, reverse-side, devotional-use or surface-performance claims.",
    },
    "personalized-orthodox-christian-area-rug-00beb8fce2": {
        "title": "Floral Orthodox Cross Rug with Red Patterned Field",
        "meta": "Decorate with a floral Orthodox cross rug featuring red patterned field, gold cross, curved border and room mockups.",
        "description": "Decorate with a floral Orthodox cross rug featuring red patterned field, gold cross, curved border and room mockups.\n\nDesign details\n- Red floral-pattern Orthodox rug artwork shows a gold three-bar cross with curved gold border and ornate corner styling.\n- Visible scenes include living room and fireplace mockups, close-up artwork panels and size reference graphics.\n- SEO copy focuses on visible Orthodox cross artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-orthodox-christian-area-rug-fb5a56fef9-fb5a56fef9": {
        "title": "Black Silver Orthodox Cross Rug with Floral Border",
        "meta": "Decorate with a black silver Orthodox cross rug featuring three-bar cross, IC XC lettering and ornate floral border.",
        "description": "Decorate with a black silver Orthodox cross rug featuring three-bar cross, IC XC lettering and ornate floral border.\n\nDesign details\n- Black rectangular Orthodox rug artwork shows a silver three-bar cross, IC XC lettering and ornate silver floral border.\n- Product images include living room and fireside mockups, close-up design panels and size reference graphics.\n- SEO copy stays tied to visible cross artwork without unsupported material, cleaning, reverse-side, devotional-use or surface-performance claims.",
    },
    "personalized-teachers-classroom-rules-rug-b0d7bd292a-b0d7bd292a": {
        "title": "Behavior Champs Classroom Rules Rug with Icons",
        "meta": "Personalize a Behavior Champs classroom rules rug with educator name, colorful rule panels and school icons.",
        "description": "Personalize a Behavior Champs classroom rules rug with educator name, colorful rule panels and school icons.\n\nDesign details\n- Bright rectangular classroom rules rug artwork shows Mrs. Smith's Classroom Rules and Behavior Champs headline.\n- Visible rule panels include be a great listener, show kindness, share and take turns, follow directions and helping hands.\n- SEO copy focuses on visible classroom-rules artwork without unsupported outcome, material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-classroom-rules-welcome-mat-07b188eb0d": {
        "title": "Zone Of Tolerance Classroom Rug with Emotion Icons",
        "meta": "Decorate with a Zone of Tolerance classroom rug featuring overwhelmed, just right and shut down sections.",
        "description": "Decorate with a Zone of Tolerance classroom rug featuring overwhelmed, just right and shut down sections.\n\nDesign details\n- Rectangular classroom rug artwork shows My Zone Of Tolerance title with feeling icons and regulation prompt panels.\n- Visible sections include overwhelmed, just right and shut down, with colorful border, classroom scenes and size reference graphics.\n- SEO copy stays tied to visible zone-chart artwork without unsupported wellness, material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-classroom-rules-rug-custom-teacher-name-2b9854c06f": {
        "title": "Calming Techniques Classroom Rug with Sea Animals",
        "meta": "Personalize a calming techniques classroom rug with educator name, sea animals, prompt cards and classroom scenes.",
        "description": "Personalize a calming techniques classroom rug with educator name, sea animals, prompt cards and classroom scenes.\n\nDesign details\n- Light blue rectangular classroom rug artwork shows Mrs. Smith's Classroom text with sea animals and colorful prompt cards.\n- Visible prompts include take a breath, color or draw, drink water, count to ten, stretch, listen to music and read a book.\n- SEO copy focuses on visible sea-animal prompt artwork without unsupported wellness, material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-classroom-rules-rug-with-custom-teacher-name-ad0787fb44": {
        "title": "Little Dinos Classroom Rug with Affirmation Panels",
        "meta": "Personalize a Little Dinos classroom rug with educator welcome text, colorful dinosaur panels and affirmations.",
        "description": "Personalize a Little Dinos classroom rug with educator welcome text, colorful dinosaur panels and affirmations.\n\nDesign details\n- Dark rectangular classroom rug artwork shows Little Dinos Big Dreams text and Welcome To Mrs. Smith's Classroom banner.\n- Visible dinosaur panels include I am strong, smart, kind, happy, brave and loved, with classroom scenes and size graphics.\n- SEO copy stays tied to visible dinosaur classroom artwork without unsupported outcome, material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-octopus-sea-monster-welcome-mat-c69579f5ed": {
        "title": "Porthole Octopus Shaped Rug with Red Tentacles",
        "meta": "Decorate with a porthole octopus shaped rug featuring red tentacles, brass window artwork and blue water background.",
        "description": "Decorate with a porthole octopus shaped rug featuring red tentacles, brass window artwork and blue water background.\n\nDesign details\n- Custom-shaped octopus rug artwork shows red tentacles reaching through a brass porthole on a blue water background.\n- Visible details include curling tentacles, wood-floor and living-room mockups, shape feature panels and size reference graphics.\n- SEO copy focuses on visible nautical octopus artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-octopus-shaped-area-rug-a96e544960-a96e544960": {
        "title": "Blue Tentacle Octopus Shaped Rug with Ocean Art",
        "meta": "Decorate with a blue tentacle octopus shaped rug featuring swirling arms, gold suction-cup accents and ocean tones.",
        "description": "Decorate with a blue tentacle octopus shaped rug featuring swirling arms, gold suction-cup accents and ocean tones.\n\nDesign details\n- Dark blue custom-shaped octopus rug artwork shows swirling tentacles, gold suction-cup accents and deep ocean colors.\n- Product images include room mockups, shape feature panels, close-up design views and size reference graphics.\n- SEO copy stays tied to visible tentacle artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-running-horse-welcome-mat-2e0754fd8e-2e0754fd8e": {
        "title": "Charcoal Running Horse Rug with Black White Art",
        "meta": "Decorate with a charcoal running horse rug featuring black and white horse portrait, motion texture and room mockups.",
        "description": "Decorate with a charcoal running horse rug featuring black and white horse portrait, motion texture and room mockups.\n\nDesign details\n- Rectangular horse rug artwork shows a running horse portrait in black and white with charcoal dust-style texture.\n- Product images include living room, fireplace and bedroom mockups, close-up design panels and size reference graphics.\n- SEO copy focuses on visible running-horse artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "kid friendly", "pet friendly", "safe durable",
    "machine-washable", "machine washable", "washable", "quick-dry", "quick dry",
    "memory foam", "microfiber", "velvet", "stain", "fade resistant", "easy clean",
    "easy-clean", "waterproof", "absorbent", "absorption", "backing", "rubber",
    "layered construction", "hd printing", "ultra-soft", "soft ", "cushion",
    "support classroom", "support emotional", "teach", "learning-space", "thickened",
    "reinforced", "bound edge", "material panel", "pray", "prayer", "sacred",
    "calm routines", "calming routines", "therapeutic", "bring calm",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    blob = " ".join([title, meta, item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R030 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R030 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append([
            "revision_batch_id", "generated_at", "handle", "source_row",
            "old_revision", "new_revision", "changed_fields", "qa_issue_fields",
            "recheck_condition",
        ])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R030"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R030", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R030",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R031 or Re-QA revision R030",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R030 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R030.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        desc = row[out_cols["description_proposed"] - 1]
        blob = " ".join([str(title or ""), str(meta or ""), str(desc or "")]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r030_log_rows": counts.get("R030", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

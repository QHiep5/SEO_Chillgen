from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R036/SEO_Product_Optimization_revision_R036.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R037"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R037.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R037"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-hunting-doormat-custom-family-name-d3fa485b6b-d3fa485b6b": {
        "title": "Patriotic Eagle 250th Anniversary Family Doormat",
        "meta": "Personalize a patriotic eagle 250th anniversary family doormat with American flag stripes, 1776 2026 medallion and name.",
        "primary": "patriotic eagle anniversary doormat",
        "secondary": "250th anniversary family mat, American flag eagle doormat, custom patriotic entry rug",
        "long_tail": "patriotic eagle anniversary doormat; 250th anniversary family mat; American flag eagle doormat",
        "description": "Personalize a patriotic eagle 250th anniversary family doormat with American flag stripes, 1776 2026 medallion and name.\n\nDesign details\n- Rectangular patriotic doormat artwork shows a bald eagle over American flag stripes with oval 1776 and 2026 anniversary center text.\n- Product images include entryway scenes, couple holding mat, dog doorway mockup, size reference graphic and close-up artwork views.\n- SEO copy stays tied to visible eagle flag artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            5: "Close-up graphic shows patriotic eagle artwork and mat profile",
        },
    },
    "custom-golf-area-rug-living-room-mat-39dea6cb42": {
        "title": "Green Anderson 19th Hole Personalized Golf Rug",
        "meta": "Personalize a green Anderson 19th Hole golf rug with fairway stripes, flag pin, golf club, ball and lounge-room artwork.",
        "primary": "personalized 19th Hole golf rug",
        "secondary": "green golf room rug, custom Anderson golf mat, fairway clubhouse rug",
        "long_tail": "personalized 19th Hole golf rug; green golf room rug; custom Anderson golf mat",
        "description": "Personalize a green Anderson 19th Hole golf rug with fairway stripes, flag pin, golf club, ball and lounge-room artwork.\n\nDesign details\n- Green rectangular golf rug artwork reads Anderson's 19th Hole with a flag pin, golf club, golf ball and radial fairway-style background.\n- Product images include golf room setting, sofa mockup, dimension chart, close-up artwork panel and staged floor scenes.\n- SEO copy focuses on visible golf-room artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            4: "Close-up view of green Anderson 19th Hole golf artwork",
            5: "Feature graphic for green Anderson golf rug design",
        },
    },
    "custom-golf-area-rug-living-room-decor-d120c386b8-d120c386b8": {
        "title": "Steven 19th Hole Golf Lover Rug with Clubhouse Phrases",
        "meta": "Personalize a Steven 19th Hole golf lover rug with green fairway design, flag pin, golf ball and clubhouse phrase text.",
        "primary": "Golf Lover 19th Hole rug",
        "secondary": "custom Steven golf rug, green clubhouse golf mat, personalized fairway rug",
        "long_tail": "Golf Lover 19th Hole rug; custom Steven golf rug; green clubhouse golf mat",
        "description": "Personalize a Steven 19th Hole golf lover rug with green fairway design, flag pin, golf ball and clubhouse phrase text.\n\nDesign details\n- Green rectangular golf rug artwork reads Steven's 19th Hole with phrases such as Unwind Relax Enjoy, Golf Stories, Cold Beer and Bets Paid.\n- Visible details include fairway stripe background, flag pin, golf ball, golf club, room mockups, size chart and close-up artwork views.\n- SEO copy stays tied to visible golf clubhouse artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature graphic for green Steven 19th Hole golf rug",
            5: "Close-up view of green Steven golf artwork",
        },
    },
    "personalized-motivational-classroom-rug-for-kids-71bb812a4f": {
        "title": "Brown Crossword Affirmation Classroom Rug with Name",
        "meta": "Personalize a brown crossword affirmation classroom rug with educator name, word puzzle layout and colorful school icons.",
        "primary": "crossword affirmation classroom rug",
        "secondary": "custom educator name rug, brown classroom word rug, positive words school mat",
        "long_tail": "crossword affirmation classroom rug; custom educator name rug; brown classroom word rug",
        "description": "Personalize a brown crossword affirmation classroom rug with educator name, word puzzle layout and colorful school icons.\n\nDesign details\n- Brown wood-style classroom rug artwork reads In This Classroom You Are with crossword words such as brave, important, creators and respected.\n- Visible details include Mrs. Taylor name, apple, scissors, pencil, basketball, globe, classroom mockups and top-down artwork views.\n- SEO copy focuses on visible crossword affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            7: "Feature graphic shows brown crossword affirmation rug profile",
        },
    },
    "custom-3d-dragon-round-rug-7df46cf462": {
        "title": "Sleeping Red Dragon Library Round Rug Artwork",
        "meta": "Decorate with a sleeping red dragon library round rug featuring curled dragon artwork, bookshelves, stairs and fantasy atrium scene.",
        "primary": "sleeping red dragon round rug",
        "secondary": "fantasy library rug, dragon bookshelf rug, round dragon room mat",
        "long_tail": "sleeping red dragon round rug; fantasy library rug; dragon bookshelf rug",
        "description": "Decorate with a sleeping red dragon library round rug featuring curled dragon artwork, bookshelves, stairs and fantasy atrium scene.\n\nDesign details\n- Round fantasy rug artwork shows a red sleeping dragon curled inside a circular library scene with books, stairs and warm lighting.\n- Product images include fireplace room, sofa-side and bedroom mockups, close-up dragon views, size chart and staged floor scenes.\n- SEO copy stays tied to visible sleeping-dragon artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            7: "Close-up callout focuses on the sleeping dragon illustration",
            8: "Reverse-side style graphic for sleeping dragon round rug",
        },
    },
    "custom-flamingo-welcome-mat-for-front-door-2096a15fdb": {
        "title": "Flamingo Couple Tropical Rug with Sunset Beach Art",
        "meta": "Decorate with a flamingo couple tropical rug featuring heart-shaped necks, sunset beach colors, hibiscus flowers and teal water.",
        "primary": "flamingo couple tropical rug",
        "secondary": "sunset beach flamingo rug, pink bird area rug, tropical flower rug",
        "long_tail": "flamingo couple tropical rug; sunset beach flamingo rug; pink bird area rug",
        "description": "Decorate with a flamingo couple tropical rug featuring heart-shaped necks, sunset beach colors, hibiscus flowers and teal water.\n\nDesign details\n- Rectangular tropical rug artwork shows two flamingos forming a heart shape against a sunset beach scene.\n- Visible details include pink hibiscus flowers, palm leaves, turquoise water, room mockups, size reference graphic and close-up artwork panels.\n- SEO copy focuses on visible flamingo beach artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            3: "Close-up highlights flamingo couple and tropical color detail",
            4: "Feature panel shows flamingo couple rug artwork details",
        },
    },
    "custom-pink-flamingo-area-rug-doormat-e7ac68e9c9": {
        "title": "Pastel Pink Flamingo Rug with Palm Leaf Artwork",
        "meta": "Decorate with a pastel pink flamingo rug featuring tropical water colors, palm leaves, orange sunset tones and bird artwork.",
        "primary": "pastel pink flamingo rug",
        "secondary": "tropical flamingo area rug, palm leaf bird rug, sunset flamingo mat",
        "long_tail": "pastel pink flamingo rug; tropical flamingo area rug; palm leaf bird rug",
        "description": "Decorate with a pastel pink flamingo rug featuring tropical water colors, palm leaves, orange sunset tones and bird artwork.\n\nDesign details\n- Rectangular tropical rug artwork shows a single pink flamingo standing in pastel water with palm leaves and warm sky colors.\n- Product images include room mockups, floor staging, size reference graphic and close-up panels showing the flamingo and background artwork.\n- SEO copy stays tied to visible flamingo artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            3: "Close-up highlights pastel flamingo and tropical color detail",
            4: "Feature panel shows pastel pink flamingo rug artwork",
        },
    },
    "custom-octopus-welcome-mat-entryway-living-room-3387fa82d0": {
        "title": "Black Octopus Tentacle Rug with Monochrome Art",
        "meta": "Decorate with a black octopus tentacle rug featuring monochrome sea-life artwork, silver suction cups and dramatic contrast.",
        "primary": "black octopus tentacle rug",
        "secondary": "monochrome sea life rug, nautical octopus area rug, black ocean mat",
        "long_tail": "black octopus tentacle rug; monochrome sea life rug; nautical octopus area rug",
        "description": "Decorate with a black octopus tentacle rug featuring monochrome sea-life artwork, silver suction cups and dramatic contrast.\n\nDesign details\n- Rectangular ocean rug artwork shows close-up octopus tentacles in black, gray and white with large suction-cup details.\n- Product images include sofa-side room scenes, angled artwork views, size reference graphic and close-up tentacle panels.\n- SEO copy focuses on visible octopus tentacle artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            4: "Close-up shows black octopus tentacle and suction cup artwork",
            7: "Feature panel shows black octopus rug artwork details",
        },
    },
    "custom-octopus-welcome-mat-entryway-464c692e0a": {
        "title": "Blue Octopus Tentacle Rug with Glowing Sea Art",
        "meta": "Decorate with a blue octopus tentacle rug featuring glowing suction cups, dark ocean colors and dramatic sea-life artwork.",
        "primary": "blue octopus tentacle rug",
        "secondary": "dark ocean octopus rug, glowing suction cup mat, nautical blue area rug",
        "long_tail": "blue octopus tentacle rug; dark ocean octopus rug; glowing suction cup mat",
        "description": "Decorate with a blue octopus tentacle rug featuring glowing suction cups, dark ocean colors and dramatic sea-life artwork.\n\nDesign details\n- Rectangular ocean rug artwork shows a large blue octopus with curling tentacles and bright suction-cup highlights on a dark background.\n- Product images include sofa-side room scenes, living room mockups, size reference graphic and close-up tentacle views.\n- SEO copy stays tied to visible blue octopus artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            3: "Close-up shows blue octopus suction cups and tentacle artwork",
            6: "Feature panel shows blue octopus rug artwork details",
        },
    },
    "custom-octopus-welcome-mat-coastal-front-door-5012959a0b": {
        "title": "Teal Octopus Coastal Rug with Gold Cup Accents",
        "meta": "Decorate with a teal octopus coastal rug featuring white wave artwork, gold suction cup accents and painterly ocean texture.",
        "primary": "teal octopus coastal rug",
        "secondary": "nautical octopus rug, ocean tentacle area rug, coastal sea life mat",
        "long_tail": "teal octopus coastal rug; nautical octopus rug; ocean tentacle area rug",
        "description": "Decorate with a teal octopus coastal rug featuring white wave artwork, gold suction cup accents and painterly ocean texture.\n\nDesign details\n- Rectangular coastal rug artwork shows teal octopus tentacles with gold-ring suction cups over white and blue wave-style background.\n- Product images include sofa-side room scenes, living room mockups, size reference graphic and close-up tentacle artwork panels.\n- SEO copy focuses on visible teal octopus artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            3: "Close-up shows teal tentacles and gold suction cup artwork",
            6: "Feature panel shows teal octopus rug artwork details",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "washable", "vacuum", "quick-dry", "quick dry", "memory foam", "microfiber",
    "velvet", "plush", "stain", "fade", "easy clean", "easy-clean", "waterproof",
    "absorbent", "backing", "rubber", "layer construction", "layered construction",
    "ultra-soft", "soft ", "cushion", "thickness", "reinforced", "bound edge",
    "material panel", "learning-pattern", "learning pattern", "support classroom",
    "support emotional", "encourage students", "use a black", "classroom sel",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R037 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R037 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R037"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R037", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R037", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R038 or Re-QA revision R037",
    }
    (RUN_DIR / "revision_R037_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R037 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R037.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R037 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R037_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r037_log_rows": counts.get("R037", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

from pathlib import Path
from datetime import datetime, timezone, timedelta
import json, shutil, re
import openpyxl

BASE=Path('/Users/buiquanghuy/Documents/seo_chillgen')
SRC=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R006/SEO_Product_Optimization_revision_R006.xlsx'
OUTDIR=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R060'
OUT=OUTDIR/'SEO_Product_Optimization_revision_R060.xlsx'
RUN=BASE/'seo_runs/chillgen.com/chillgen_20260907_01/revisions/R060'
NOW=datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()

wb=openpyxl.load_workbook(SRC)
log=wb['Revision_Log']
scope=[r[2].value for r in log.iter_rows(min_row=2) if r[0].value=='R006']
ws=wb['SEO_Products']
cols={c.value:i for i,c in enumerate(ws[1],1)}
changed=[]
for row in range(2,ws.max_row+1):
    h=ws.cell(row,cols['Handle']).value
    if h not in scope: continue
    for name in ('description_proposed','description_proposed_html'):
        c=cols[name]; v=ws.cell(row,c).value
        if isinstance(v,str):
            v=re.sub(r'(?m)(?:- |<br>)*SEO copy[^\n<]*(?:surface-performance claims\.?|review\.?|evidence\.?)[^\n<]*', '', v, flags=re.I)
            v=re.sub(r'\n{3,}','\n\n',v); v=re.sub(r'(<br>){3,}','<br><br>',v)
            ws.cell(row,c).value=v.strip()
    for name,val in [('revision','2'),('review_status','NEEDS_REVIEW'),('processing_status','DRAFTED'),('content_qa_status','NOT_RUN'),('review_reason','Revision R060 removes customer-visible internal QA/process notes identified in Re-QA R006.'),('issues','Revision R060 targeted internal process-language cleanup; not approved; Re-QA required.')]:
        if name in cols: ws.cell(row,cols[name]).value=val
    changed.append(h)

for h in changed:
    log.append(['R060',NOW,h,None,'2','2','description_proposed, description_proposed_html, revision, review_reason, issues','internal_process_language','Run evidence-driven Re-QA on revision 2 for this product before approval.'])
OUTDIR.mkdir(parents=True,exist_ok=True); RUN.mkdir(parents=True,exist_ok=True)
wb.save(OUT)
manifest={'revision_batch_id':'R060','generated_at':NOW,'source_workbook':str(SRC),'output_workbook':str(OUT),'product_count':len(changed),'handles':changed,'status':'COMPLETE_AWAITING_REQA','review_status':'NEEDS_REVIEW','content_qa_status':'NOT_RUN','not_approved_not_deployed':True}
(RUN/'revision_R060_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
print(json.dumps({'output':str(OUT),'products':len(changed),'status':manifest['status']},ensure_ascii=False))

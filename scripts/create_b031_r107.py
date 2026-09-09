from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B031"
REVISION = "R107"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "personalized-classroom-doormat-welcome-class-door-mat-f34045b91d": ("Tribal Classroom Welcome Mat", "tribal classroom welcome mat", "personalized teacher doormat, classroom door mat, custom teacher welcome rug", "a tribal-style classroom welcome design for a school doorway, classroom entry, teacher office, or learning space"),
    "personalized-classroom-welcome-doormat-notebook-rug-508fa38fe3": ("Better With You Classroom Mat", "Better With You classroom mat", "personalized classroom doormat, teacher appreciation mat, classroom welcome rug", "a Better With You message with notebook-inspired classroom artwork for a school doorway, classroom entry, or teacher office"),
    "personalized-classroom-doormat-with-custom-teacher-name-675f2e0900": ("Happy To See Your Face Mat", "Happy To See Your Face mat", "personalized teacher doormat, cheerful classroom mat, back to school welcome rug", "a Happy To See Your Face message with colorful classroom artwork for a school doorway, classroom entry, or teacher office"),
    "personalized-classroom-welcome-doormat-2820e17d05-2820e17d05": ("Red Doodle Classroom Mat", "red doodle classroom mat", "personalized classroom welcome mat, teacher name doormat, school doodle rug", "a red doodle classroom design for a school doorway, classroom entry, teacher office, or learning space"),
    "personalized-classroom-doormat-welcome-to-class-door-mat-1086025b51": ("Green Classroom Doormat", "green classroom doormat", "personalized teacher doormat, classroom welcome mat, school supply rug", "a green classroom design with school-themed artwork for a school doorway, classroom entry, or teacher office"),
    "personalized-classroom-doormat-custom-teacher-name-753f4f3532-753f4f3532": ("You Are Classroom Mat", "You Are classroom mat", "affirmation classroom doormat, personalized teacher mat, back to school rug", "a You Are affirmation message with classroom artwork for a school doorway, classroom entry, or teacher office"),
    "personalized-classroom-welcome-doormat-3d9f3ec373": ("Everyone Welcome Here Mat", "Everyone Welcome Here mat", "inclusive classroom mat, personalized teacher doormat, colorful classroom rug", "an Everyone Welcome Here message with colorful classroom artwork for a school doorway, classroom entry, or teacher office"),
    "personalized-classroom-doormat-teacher-notebook-a1ca00cc38-a1ca00cc38": ("Pastel Teacher Notebook Mat", "pastel teacher notebook mat", "personalized classroom doormat, teacher name welcome mat, notebook classroom rug", "a pastel teacher notebook design for a school doorway, classroom entry, teacher office, or learning space"),
    "personalized-classroom-doormat-welcome-to-class-door-mat-a74c09137b": ("School Doodle Welcome Mat", "school doodle welcome mat", "personalized classroom doormat, teacher welcome rug, back to school door mat", "a school doodle welcome design for a school doorway, classroom entry, teacher office, or learning space"),
    "personalized-classroom-doormat-with-custom-teacher-name-ae6fef4505-ae6fef4505": ("Sunflower Classroom Welcome Mat", "sunflower classroom welcome mat", "personalized teacher doormat, classroom welcome rug, warm school entrance mat", "a sunflower classroom welcome design for a school doorway, classroom entry, teacher office, or learning space"),
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B031 scope does not match the R107 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        article = "an" if primary.lower().startswith(("a ", "e ", "i ", "o ", "u ")) else "a"
        meta = f"Create {article} {primary} with {motif}."
        description = f"Create {article} {primary} with {motif}. The design adds a distinctive focal point to a classroom, school doorway, teacher office, or learning space."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

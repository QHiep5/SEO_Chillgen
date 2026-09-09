from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B009"
REVISION = "R085"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"


COPY = {
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-04": {
        "title": "Personalized Parrot Beach House Rug",
        "meta": "Customize a parrot beach house rug with family-name text, ocean artwork, hibiscus flowers, palm details, and a sunny coastal welcome.",
        "description": "Customize a parrot beach house rug with family-name text, ocean artwork, hibiscus flowers, and palm details. The bright coastal design gives a vacation home, porch entry, or beach-themed room a cheerful personalized welcome.",
        "primary": "personalized parrot beach house rug",
        "secondary": "custom family name beach rug, tropical beach house mat, parrot coastal rug",
    },
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-05": {
        "title": "Personalized Surfboard Beach Family Rug",
        "meta": "Customize a surfboard beach family rug with family-name text, palm trees, flip-flops, and a relaxed coastal sign design.",
        "description": "Customize a surfboard beach family rug with family-name text, palm trees, flip-flops, and a relaxed coastal sign design. The surf-inspired artwork turns a family name into a bright beach-house welcome.",
        "primary": "personalized surfboard beach family rug",
        "secondary": "custom family beach house rug, surfboard coastal mat, tropical family rug",
    },
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-01": {
        "title": "Personalized Flip Flop Happy Place Rug",
        "meta": "Create a flip-flop happy place rug with custom family-name text, sand-and-ocean artwork, starfish, shells, and vacation color.",
        "description": "Create a flip-flop happy place rug with custom family-name text, sand-and-ocean artwork, starfish, shells, and vacation color. The playful coastal design makes a beach room feel personal, bright, and guest-ready.",
        "primary": "personalized flip flop happy place rug",
        "secondary": "custom happy place rug, family name beach house rug, tropical coastal mat",
    },
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-08": {
        "title": "Personalized Flamingo Paradise Beach Rug",
        "meta": "Customize a flamingo paradise beach rug with family-name text, tropical sign artwork, hibiscus flowers, palm leaves, and ocean color.",
        "description": "Customize a flamingo paradise beach rug with family-name text, tropical sign artwork, hibiscus flowers, palm leaves, and ocean color. The personalized paradise motif adds a playful resort mood to beach-house and coastal decor.",
        "primary": "personalized flamingo paradise beach rug",
        "secondary": "custom family name tropical rug, flamingo beach house mat, paradise coastal rug",
    },
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-06": {
        "title": "Personalized Sunset Beach House Rug",
        "meta": "Customize a sunset beach house rug with family-name text, palm trees, hammock artwork, ocean view, and calm coastal color.",
        "description": "Customize a sunset beach house rug with family-name text, palm trees, hammock artwork, ocean view, and calm coastal color. The design gives a family beach space a named vacation-retreat feeling.",
        "primary": "personalized sunset beach house rug",
        "secondary": "custom family beach rug, tropical hammock rug, coastal vacation home mat",
    },
    "personalized-tropical-beach-rug-custom-family-name-tropical-outdoor": {
        "title": "Personalized Hibiscus Tropical Family Rug",
        "meta": "Customize a hibiscus tropical family rug with family-name text, palm-leaf artwork, bright flowers, and soft beach color.",
        "description": "Customize a hibiscus tropical family rug with family-name text, palm-leaf artwork, bright flowers, and soft beach color. The floral design brings a fresh personalized accent to tropical rooms and beach-house entry spaces.",
        "primary": "personalized hibiscus tropical family rug",
        "secondary": "custom family name tropical rug, hibiscus coastal mat, beach house name rug",
    },
    "personalized-calming-corner-rug-for-classroom-custom-teacher-name": {
        "title": "Personalized Teacher Calm Corner Rug",
        "meta": "Customize a teacher calm corner rug with name text, feelings chart artwork, self-control prompts, and colorful classroom scenes.",
        "description": "Customize a teacher calm corner rug with name text, feelings chart artwork, self-control prompts, and colorful classroom scenes. The personalized design creates a clear classroom spot for emotion check-ins and quiet reset moments.",
        "primary": "personalized teacher calm corner rug",
        "secondary": "custom calming corner rug, classroom feelings rug, teacher name classroom rug",
    },
    "personalized-calming-corner-rug-for-classroom-custom-tea-design-100": {
        "title": "Personalized Calm Corner Activity Rug",
        "meta": "Customize a calm corner activity rug with teacher-name text, breathe, draw, listen, dance, and reflection prompts for classroom use.",
        "description": "Customize a calm corner activity rug with teacher-name text and visual prompts such as breathe, draw, listen, dance, and reflection. The design gives students a clear set of calm-down choices in one classroom area.",
        "primary": "personalized calm corner activity rug",
        "secondary": "custom teacher calm corner rug, classroom calm down rug, SEL activity rug",
    },
    "personalized-calming-corner-rug-for-classroom-custom-tea-design-101": {
        "title": "Classroom Mindfulness Calm Corner Rug",
        "meta": "Create a classroom mindfulness rug with positive-thinking prompts, breathing ideas, doodle artwork, and gentle calm-corner language.",
        "description": "Create a classroom mindfulness rug with positive-thinking prompts, breathing ideas, doodle artwork, and gentle calm-corner language. The design gives students visual reminders for reflection, relaxation, and emotional reset.",
        "primary": "classroom mindfulness calm corner rug",
        "secondary": "calming corner mindfulness rug, SEL classroom rug, calm down corner rug",
    },
    "personalized-calming-corner-rug-for-classroom-custom-tea-design-102": {
        "title": "Classroom Feelings Check-In Rug",
        "meta": "Create a classroom feelings check-in rug with a battery meter, mood faces, coping prompts, and colorful emotion-scale artwork.",
        "description": "Create a classroom feelings check-in rug with a battery meter, mood faces, coping prompts, and colorful emotion-scale artwork. The design helps make emotional check-ins visible and easy for students to understand.",
        "primary": "classroom feelings check-in rug",
        "secondary": "calm corner feelings rug, emotional regulation rug, classroom SEL rug",
    },
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B009 scope does not match the R085 copy map")

    for row in seo_rows:
        item = COPY[row["Handle"]]
        row["title_proposed"] = item["title"]
        row["meta_title_seo"] = item["title"]
        row["meta_title_chars"] = str(len(item["title"]))
        row["meta_description_seo"] = item["meta"]
        row["meta_description_chars"] = str(len(item["meta"]))
        row["description_proposed"] = item["description"]
        row["description_proposed_html"] = f"<p>{item['description']}</p>"
        row["primary_keyword"] = item["primary"]
        row["secondary_keywords"] = item["secondary"]
        row["long_tail_candidates"] = f"{item['primary']}; {item['secondary']}"
        row["revision"] = "2"
        row["review_status"] = "NEEDS_REVIEW"
        row["content_qa_status"] = "NOT_RUN"
        row["review_reason"] = f"{REVISION} removes internal evidence-note language and unsupported claim wording from B009 copy."
        row["issues"] = f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations."

    evidence_scope = [row for row in evidence_rows if row.get("evidence_id") in {r.get("evidence_id") for r in seo_rows}]
    keyword_scope = [row for row in keyword_rows if row.get("product_key") in handles]
    buyer_rows = [row for row in buyer_all if row.get("product_key") in handles]
    img_scope = [row for row in img_rows if row.get("Handle") in handles]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": "Canonical B009 cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    evidence_ids = {row["evidence_id"] for row in evidence_scope}
    for row in seo_rows:
        evidence_id = row.get("evidence_id")
        if evidence_id not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {evidence_id}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", evidence_id])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    digest = sha256(OUT)
    manifest = {
        "revision": REVISION,
        "batch_id": BATCH,
        "scope_count": len(handles),
        "image_count": len(img_scope),
        "canonical_workbook": str(OUT),
        "sha256": digest,
        "status": "CREATED_FOR_HARDENED_REQA",
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
    }
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

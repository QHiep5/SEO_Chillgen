from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R038/SEO_Product_Optimization_revision_R038.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R039"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R039.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R039"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-halloween-3d-optical-illusion-ghost-round-rug-a07": {
        "title": "Ghost Stairwell Halloween Round Rug with Pumpkin Art",
        "meta": "Decorate with a ghost stairwell Halloween round rug featuring pumpkin lights, stone steps, bookshelves and spooky depth artwork.",
        "primary": "ghost stairwell Halloween rug",
        "secondary": "haunted staircase rug, ghost round rug, Halloween optical illusion rug",
        "long_tail": "ghost stairwell Halloween rug; haunted staircase rug; ghost round rug",
        "description": "Decorate with a ghost stairwell Halloween round rug featuring pumpkin lights, stone steps, bookshelves and spooky depth artwork.\n\nDesign details\n- Round Halloween rug artwork shows a small ghost on a winding stone staircase surrounded by pumpkins, bookshelves and warm lantern light.\n- Product images include room mockups, underfoot scene, size reference graphic and close-up design views.\n- SEO copy stays tied to visible ghost stairwell artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {8: "Design detail graphic for ghost staircase Halloween rug"},
    },
    "custom-dragon-shaped-area-rug-personalized-library-wester-design-04": {
        "title": "Personalized Orange Dragon Open Book Rug Artwork",
        "meta": "Personalize an orange dragon open book rug with Sophia name, antique page artwork, curled dragon and fantasy library styling.",
        "primary": "personalized orange dragon book rug",
        "secondary": "orange dragon rug, fantasy open book rug, custom dragon shaped mat",
        "long_tail": "personalized orange dragon book rug; orange dragon rug; fantasy open book rug",
        "description": "Personalize an orange dragon open book rug with Sophia name, antique page artwork, curled dragon and fantasy library styling.\n\nDesign details\n- Custom-shaped rug artwork shows an orange and gold dragon resting across an open antique-style book with Sophia name text.\n- Product images include dark wood floor scenes, angled views, room mockups, size chart and close-up artwork graphics.\n- SEO copy focuses on visible dragon book artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {7: "Size chart and feature view for orange dragon open book rug"},
    },
    "custom-dragon-shaped-area-rug-personalized-library-wester-design-06": {
        "title": "Personalized Purple Dragon Open Book Rug Artwork",
        "meta": "Personalize a purple dragon open book rug with Sophia name, sleeping dragon, purple flowers and fantasy page artwork.",
        "primary": "personalized purple dragon book rug",
        "secondary": "sleeping dragon rug, fantasy book rug, purple dragon shaped mat",
        "long_tail": "personalized purple dragon book rug; sleeping dragon rug; fantasy book rug",
        "description": "Personalize a purple dragon open book rug with Sophia name, sleeping dragon, purple flowers and fantasy page artwork.\n\nDesign details\n- Custom-shaped rug artwork shows a purple dragon curled on an open book with grape-like clusters, purple roses, greenery and Sophia name text.\n- Product images include dark wood floor scenes, room mockups, angled views, size chart and close-up artwork graphics.\n- SEO copy stays tied to visible purple dragon book artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {7: "Size chart and feature view for purple dragon open book rug"},
    },
    "boat-rug-03-boat-rug-03-design-22": {
        "title": "Personalized Welcome Aboard Boat Rug with Compass",
        "meta": "Personalize a Welcome Aboard boat rug with boat name, location text, compass rose, border lines and black-white nautical art.",
        "primary": "personalized Welcome Aboard boat rug",
        "secondary": "custom boat name rug, nautical compass rug, black white boat mat",
        "long_tail": "personalized Welcome Aboard boat rug; custom boat name rug; nautical compass rug",
        "description": "Personalize a Welcome Aboard boat rug with boat name, location text, compass rose, border lines and black-white nautical art.\n\nDesign details\n- Rectangular nautical rug artwork reads Welcome Aboard with boat name and location placeholders below a compass rose.\n- The visible scene places the black-and-white rug in a bright coastal-style room with blue seating and glass doors.\n- SEO copy stays tied to visible boat-name compass artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {},
    },
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-04": {
        "title": "Personalized Parrot Beach House Rug with Hibiscus Art",
        "meta": "Personalize a parrot beach house rug with family name, blue ocean scene, yellow sun, hibiscus flowers and palm greenery.",
        "primary": "personalized parrot beach house rug",
        "secondary": "custom family name beach rug, tropical parrot rug, coastal hibiscus mat",
        "long_tail": "personalized parrot beach house rug; custom family name beach rug; tropical parrot rug",
        "description": "Personalize a parrot beach house rug with family name, blue ocean scene, yellow sun, hibiscus flowers and palm greenery.\n\nDesign details\n- Bright tropical rug artwork shows Welcome To The Miller's Beach House text, colorful parrot, ocean water, flowers and palm leaves.\n- Product images include coastal living room mockups, pet lifestyle scene, size guide and close-up artwork views.\n- SEO copy focuses on visible parrot beach artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic for parrot beach house rug artwork",
            6: "Feature comparison graphic for tropical parrot beach rug",
        },
    },
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-06": {
        "title": "Personalized Sunset Beach House Rug with Hammock",
        "meta": "Personalize a sunset beach house rug with family name, palm trees, hammock, shoreline scene and shell accents.",
        "primary": "personalized sunset beach house rug",
        "secondary": "custom family beach rug, tropical hammock rug, coastal vacation mat",
        "long_tail": "personalized sunset beach house rug; custom family beach rug; tropical hammock rug",
        "description": "Personalize a sunset beach house rug with family name, palm trees, hammock, shoreline scene and shell accents.\n\nDesign details\n- Tropical rug artwork shows Welcome To White Birch House text on a beach sign with palm trees, hammock, ocean and sunset colors.\n- Product images include coastal living room scenes, pet lifestyle mockup, size guide and close-up artwork views.\n- SEO copy stays tied to visible hammock beach artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic for sunset beach house rug artwork",
            6: "Feature comparison graphic for sunset beach rug",
        },
    },
    "personalized-calming-corner-rug-for-classroom-custom-teacher-name": {
        "title": "Mrs Johnson Calm Corner Rug with Feelings Chart",
        "meta": "Personalize a Mrs Johnson Calm Corner rug with feelings chart, control prompts, colorful panels and classroom artwork.",
        "primary": "Mrs Johnson Calm Corner rug",
        "secondary": "feelings chart classroom rug, teacher name prompt rug, colorful classroom mat",
        "long_tail": "Mrs Johnson Calm Corner rug; feelings chart classroom rug; teacher name prompt rug",
        "description": "Personalize a Mrs Johnson Calm Corner rug with feelings chart, control prompts, colorful panels and classroom artwork.\n\nDesign details\n- Rectangular classroom rug artwork reads Mrs. Johnson's Calm Corner and Feelings Are Welcome Here with emotion faces and prompt panels.\n- Visible details include How Do I Feel and I Can Control sections, colorful classroom icons, room mockups, group photos and size views.\n- SEO copy describes visible prompt-chart artwork without extra care, classroom-result, therapeutic, reverse-side or surface-performance claims.",
        "alts": {
            1: "Personalized Mrs Johnson Calm Corner rug in room mockup",
            4: "Calm Corner feelings rug in colorful classroom scene",
            5: "Feature detail graphic for teacher Calm Corner rug",
            6: "Child playing on personalized Calm Corner classroom rug",
            11: "Calm Corner classroom rug angle view with cubbies",
            13: "Child on classroom Calm Corner rug lifestyle scene",
            14: "Children group photo on Calm Corner rug",
        },
    },
    "personalized-calming-corner-rug-for-classroom-custom-tea-design-100": {
        "title": "Mrs Johnson Calm Corner Activity Rug with Prompts",
        "meta": "Personalize a Mrs Johnson Calm Corner activity rug with dance, draw, breathe, music and emotion-writing prompt panels.",
        "primary": "Mrs Johnson Calm Corner activity rug",
        "secondary": "classroom activity prompt rug, teacher name calm rug, colorful prompt mat",
        "long_tail": "Mrs Johnson Calm Corner activity rug; classroom activity prompt rug; teacher name calm rug",
        "description": "Personalize a Mrs Johnson Calm Corner activity rug with dance, draw, breathe, music and emotion-writing prompt panels.\n\nDesign details\n- Rectangular classroom rug artwork reads Mrs. Johnson's Calm Corner and shows activity boxes such as Dance, Listen To Music and Draw A Picture.\n- Visible prompts also include Take Deep Breath, Animal Walk, Talk To Someone, Write Emotions, Meditate, Drink Water or Eat Snack and Get Plenty Of Sleep.\n- SEO copy stays tied to visible activity-prompt artwork without extra care, classroom-result, therapeutic, reverse-side or surface-performance claims.",
        "alts": {
            9: "Feature detail graphic for Calm Corner activity rug",
            19: "Feature graphic for Calm Corner rug artwork",
        },
    },
    "personalized-teachers-classroom-rug-custom-teacher-name-c-design-02": {
        "title": "Little Dinos Big Dreams Classroom Rug with Name",
        "meta": "Personalize a Little Dinos Big Dreams classroom rug with teacher name, dinosaur characters, affirmation phrases and starry art.",
        "primary": "Little Dinos Big Dreams classroom rug",
        "secondary": "personalized dinosaur classroom rug, teacher name dino mat, dinosaur affirmation rug",
        "long_tail": "Little Dinos Big Dreams classroom rug; personalized dinosaur classroom rug; teacher name dino mat",
        "description": "Personalize a Little Dinos Big Dreams classroom rug with teacher name, dinosaur characters, affirmation phrases and starry art.\n\nDesign details\n- Dark colorful classroom rug artwork reads Little Dinos Big Dreams with dinosaur characters and phrases such as I Am Strong, Smart, Kind, Brave and Loved.\n- Visible details include teacher-name welcome strip, flowers, stars, classroom mockups, size chart and front product image.\n- SEO copy focuses on visible dinosaur affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            8: "Feature detail graphic for dinosaur classroom rug",
            9: "Care-style graphic for dinosaur classroom rug",
        },
    },
    "personalized-composition-notebook-welcome-doormat-379c157fc5": {
        "title": "Future Of The World Space Classroom Rug with Name",
        "meta": "Personalize a Future Of The World space classroom rug with teacher name, Earth, astronaut, rocket, planets and colorful quote.",
        "primary": "Future Of The World classroom rug",
        "secondary": "space classroom rug, custom teacher name rug, astronaut rocket classroom mat",
        "long_tail": "Future Of The World classroom rug; space classroom rug; custom teacher name rug",
        "description": "Personalize a Future Of The World space classroom rug with teacher name, Earth, astronaut, rocket, planets and colorful quote.\n\nDesign details\n- Navy classroom rug artwork reads Mrs. Sophia The Future Of The World Is In This Room Today with large colorful lettering.\n- Visible details include Earth, astronaut, rocket, planets, sun, rainbow, classroom mockups, group photos and size graphics.\n- SEO copy stays tied to visible space quote artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature graphic shows Future Of The World space rug artwork",
            6: "Feature panel for Future Of The World classroom rug",
            9: "Close-up view of Future Of The World space rug artwork",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R039 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R039 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R039"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R039", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R039", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R040 or Re-QA revision R039",
    }
    (RUN_DIR / "revision_R039_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R039 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R039.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R039 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R039_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r039_log_rows": counts.get("R039", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

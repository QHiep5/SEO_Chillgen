from pathlib import Path
import openpyxl,json,shutil,hashlib
BASE=Path('/Users/buiquanghuy/Documents/seo_chillgen'); RUN=BASE/'seo_runs/chillgen.com/chillgen_20260907_01'
handles=json.loads((RUN/'evidence/products/B006_products.json').read_text())['product_keys'][5:]
src=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R080/SEO_Product_Optimization_revision_R080.xlsx'; od=src.parent.parent/'R081'; od.mkdir(exist_ok=True)
dst=od/'SEO_Product_Optimization_revision_R081.xlsx'; shutil.copy2(src,dst); w=openpyxl.load_workbook(dst)
ws=w['SEO_Products']; hs=[c.value for c in next(ws.iter_rows())]; ix={h:i+1 for i,h in enumerate(hs)}
styles=['Colorful','Classic Brown','Bright Multicolor','Soft Neutral','Rustic Woodland']
for row in list(ws.iter_rows(min_row=2)):
 h=row[ix['Handle']-1].value
 if h not in handles: ws.delete_rows(row[0].row); continue
 i=handles.index(h); t=f'{styles[i]} Dog Paw Patchwork Round Rug'; m=f'Create a {styles[i].lower()} dog paw round rug with a cozy patchwork pattern for a pet-loving home.'; d=f'{styles[i]} tones and repeating paw motifs create a distinctive round rug accent for family rooms, bedrooms, or pet spaces.'
 for f,v in [('title_proposed',t),('meta_title_seo',t),('meta_description_seo',m),('description_proposed',d),('description_proposed_html',f'<p>{d}</p>')]:
  if f in ix: row[ix[f]-1].value=v
 if 'meta_title_chars' in ix: row[ix['meta_title_chars']-1].value=len(t)
 if 'meta_description_chars' in ix: row[ix['meta_description_chars']-1].value=len(m)
for s,key in [('Product_Evidence','evidence_id'),('Image_Audit','Handle')]:
 sh=w[s]; hh=[c.value for c in next(sh.iter_rows())]; j=hh.index(key)
 for r in range(sh.max_row,1,-1):
  val=sh.cell(r,j+1).value
  if (key=='Handle' and val not in handles) or (key=='evidence_id' and val not in {f'B006-{n:03d}' for n in range(56,61)}): sh.delete_rows(r)
if 'Revision_Log' in w.sheetnames: del w['Revision_Log']
log=w.create_sheet('Revision_Log'); log.append(['revision_batch_id','revision','batch_id','handle','scope_status','review_status','content_qa_status','source_workbook','evidence_id'])
for n,h in enumerate(handles,1): log.append(['R081',2,'B006',h,'LOCKED','NEEDS_REVIEW','NOT_RUN',str(dst),f'B006-{55+n:03d}'])
w.save(dst); sha=hashlib.sha256(dst.read_bytes()).hexdigest(); (od/'revision_R081_manifest.json').write_text(json.dumps({'revision':'R081','batch_id':'B006','scope_count':5,'canonical_workbook':str(dst),'sha256':sha,'status':'CREATED_FOR_REQA','approval_status':'NOT_APPROVED_NOT_DEPLOYED'},indent=2)+'\n'); print({'workbook':str(dst),'sha256':sha})

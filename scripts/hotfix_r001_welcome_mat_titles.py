from pathlib import Path
import json
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

BASE = Path('/Users/buiquanghuy/Documents/seo_chillgen')
BOOK = BASE / 'resutls/chillgen.com/chillgen_20260907_01/revisions/R001/SEO_Product_Optimization_revision_R001.xlsx'
MANIFEST = BASE / 'seo_runs/chillgen.com/chillgen_20260907_01/hotfixes/HF_R001_WELCOME_MAT_TITLE_MAPPING/manifest.json'
CHANGES = {
    'custom-galaxy-wolves-welcome-mat-f320a33157': 'Lightning Snow Wolf Welcome Mat for Kids Room Decor',
    'custom-galaxy-wolf-welcome-mat-0e4d706227': 'Sepia Wolf Pack Welcome Mat for Rustic Entryway Decor',
}

def main():
    wb = load_workbook(BOOK)
    ws = wb['SEO_Products']
    headers = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}
    changed = []
    for row in range(2, ws.max_row + 1):
        handle = ws.cell(row, idx['Handle']).value
        if handle not in CHANGES:
            continue
        new = CHANGES[handle]
        old_title = ws.cell(row, idx['title_proposed']).value
        old_meta = ws.cell(row, idx['meta_title_seo']).value
        ws.cell(row, idx['title_proposed']).value = new
        ws.cell(row, idx['meta_title_seo']).value = new
        ws.cell(row, idx['meta_title_chars']).value = len(new)
        changed.append({'handle': handle, 'row': row, 'old_title': old_title, 'old_meta_title': old_meta, 'new_title': new, 'new_length': len(new)})
    if len(changed) != 2:
        raise RuntimeError(f'Expected 2 target rows, changed {len(changed)}')
    wb.save(BOOK)
    verify = load_workbook(BOOK, read_only=True, data_only=True)
    ws2 = verify['SEO_Products']; h = [c.value for c in ws2[1]]; ix = {k: i for i, k in enumerate(h)}
    checks = []
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if row[ix['Handle']] in CHANGES:
            title = str(row[ix['title_proposed']] or '')
            meta = str(row[ix['meta_title_seo']] or '')
            checks.append({'handle': row[ix['Handle']], 'title': title, 'meta_title': meta, 'contains_area_rug': 'area rug' in (title + ' ' + meta).lower(), 'contains_welcome_mat': 'welcome mat' in (title + ' ' + meta).lower(), 'revision': str(row[ix['revision']])})
    if len(checks) != 2 or any(x['contains_area_rug'] or not x['contains_welcome_mat'] or x['revision'] != '2' for x in checks):
        raise RuntimeError(f'Post-check failed: {checks}')
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps({'hotfix_id': 'HF_R001_WELCOME_MAT_TITLE_MAPPING', 'workbook': str(BOOK), 'changed': changed, 'verification': checks, 'approval_status': 'NOT_APPROVED_NOT_DEPLOYED', 'checked_at': datetime.now(timezone(timedelta(hours=7))).isoformat()}, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps({'changed_rows': len(changed), 'verification': checks, 'manifest': str(MANIFEST)}, ensure_ascii=False))

if __name__ == '__main__':
    main()

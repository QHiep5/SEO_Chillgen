import csv
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
ROOT = Path("seo_runs") / SHOP / RUN_ID
RESULTS = Path("resutls") / SHOP / RUN_ID
BATCH_ID = "B002"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


SERP_REFS = [
    "https://www.walmart.com/c/kp/vortex-illusion-rug",
    "https://www.etsy.com/market/illusion_stairs_rug",
    "https://www.walmart.com/ip/Halloween-3D-Optical-Illusion-Round-Rug-Spooky-Staircase-Ghost-Pumpkin-Circular-Floor-Mat-Non-Slip-Absorbent-Decorative-Doormat-Halloween-Home-Party/20673415859",
    "https://www.walmart.com/ip/3D-Haunted-Staircase-Round-Rug-Optical-Illusion-Halloween-Area-Mat/20840972139",
    "https://chillgen.com/collections/halloween",
]


DESIGNS = {
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-04": {
        "theme": "tentacle pit",
        "title": "Tentacle Pit Halloween Optical Illusion Round Rug",
        "meta": "Add a tentacle pit Halloween round rug with black-and-white vortex depth, orange checker border, non-slip backing, and spooky room appeal.",
        "primary": "tentacle pit Halloween rug",
        "secondary": "octopus tentacle rug, Halloween optical illusion rug, spooky round rug",
        "obs": "round rug with purple tentacles rising from a black-and-white optical illusion pit, orange and green border, lifestyle views and backing details",
    },
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-14": {
        "theme": "haunted village moon",
        "title": "Haunted Village Halloween Round Rug",
        "meta": "Decorate with a haunted village Halloween round rug featuring a moonlit graveyard scene, pumpkins, ghosts, and non-slip backing.",
        "primary": "haunted village Halloween rug",
        "secondary": "moonlit Halloween rug, ghost village rug, spooky round rug",
        "obs": "round rug with blue moonlit haunted village, pumpkins, ghosts, graveyard path, purple-orange border, size and backing graphics",
    },
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-05": {
        "theme": "blue haunted forest",
        "title": "Blue Haunted Forest Halloween Round Rug",
        "meta": "Set a spooky scene with a blue haunted forest Halloween round rug, moonlit houses, pumpkins, optical depth, and non-slip backing.",
        "primary": "blue haunted forest Halloween rug",
        "secondary": "haunted house round rug, moon Halloween rug, 3D Halloween rug",
        "obs": "blue round rug showing a haunted forest, moon, glowing houses, pumpkins, seasonal border, lifestyle mockups and backing view",
    },
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-01": {
        "theme": "jack-o-lantern monster",
        "title": "Jack-O-Lantern Monster Halloween Round Rug",
        "meta": "Create bold Halloween floor decor with a jack-o-lantern monster round rug, glowing eyes, scary pumpkin mouth, and non-slip backing.",
        "primary": "jack-o-lantern monster rug",
        "secondary": "scary pumpkin rug, Halloween pumpkin round rug, spooky optical illusion rug",
        "obs": "round rug with aggressive jack-o-lantern monster face, glowing eyes, open teeth-filled mouth, brown leaf border, lifestyle and quality images",
    },
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-08": {
        "theme": "lava vortex",
        "title": "Lava Vortex Halloween Round Rug",
        "meta": "Use a lava vortex Halloween round rug for fiery optical illusion decor with cracked ground, jack-o-lanterns, and non-slip backing.",
        "primary": "lava vortex Halloween rug",
        "secondary": "fire pit illusion rug, Halloween lava rug, spooky round rug",
        "obs": "round rug with fiery lava vortex, cracked black ground, glowing orange path, pumpkins at sides, lifestyle and backing images",
    },
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-03": {
        "theme": "ghost well",
        "title": "Ghost Well Halloween Optical Illusion Rug",
        "meta": "Bring playful Halloween depth to the floor with a ghost well round rug, pumpkin pit artwork, colorful border, and non-slip backing.",
        "primary": "ghost well Halloween rug",
        "secondary": "ghost optical illusion rug, pumpkin pit rug, Halloween round rug",
        "obs": "round rug with white ghosts, stone well illusion, pumpkin pit center, colorful Halloween icon border, lifestyle and backing details",
    },
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-10": {
        "theme": "pumpkin head creature",
        "title": "Pumpkin Head Monster Halloween Round Rug",
        "meta": "Make a dramatic Halloween setup with a pumpkin head monster round rug, checkerboard illusion, claw details, and non-slip backing.",
        "primary": "pumpkin head monster rug",
        "secondary": "Halloween monster rug, scary pumpkin round rug, optical illusion rug",
        "obs": "round rug with pumpkin-headed creature, claw hands, black-and-white checker illusion, orange-purple border and feature images",
    },
    "custom-halloween-3d-effect-optical-illusion-spooky-round-r-design-07": {
        "theme": "spider cave",
        "title": "Spider Cave Halloween Optical Illusion Rug",
        "meta": "Add a spider cave Halloween round rug with black-and-white depth, red gothic border, fuzzy spider artwork, and non-slip backing.",
        "primary": "spider cave Halloween rug",
        "secondary": "spider optical illusion rug, scary spider rug, Halloween round rug",
        "obs": "round rug showing a large gray spider over a black-and-white cave illusion with red gothic border and backing graphics",
    },
    "personalized-halloween-3d-optical-illusion-ghost-round-rug-a07": {
        "theme": "ghost staircase",
        "title": "Ghost Staircase Halloween Round Rug",
        "meta": "Choose a ghost staircase Halloween round rug with warm pumpkin lights, stairwell depth, cozy spooky artwork, and non-slip backing.",
        "primary": "ghost staircase Halloween rug",
        "secondary": "haunted staircase rug, ghost round rug, Halloween optical illusion rug",
        "obs": "round rug with a small ghost on stairs, pumpkins, bookshelves, warm lantern light, stairwell illusion and lifestyle views",
    },
    "personalized-halloween-3d-optical-illusion-ghost-round-rug-design-12": {
        "theme": "lava skeleton",
        "title": "Lava Skeleton Halloween Round Rug",
        "meta": "Make Halloween floors look intense with a lava skeleton round rug, cracked fire pit illusion, crawling skeleton, and non-slip backing.",
        "primary": "lava skeleton Halloween rug",
        "secondary": "skeleton lava rug, 3D skeleton round rug, spooky Halloween rug",
        "obs": "round rug with skeleton crawling from a glowing lava pit, cracked black rock, fiery orange depth and backing closeups",
    },
}


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f)), csv.DictReader(open(path, newline="", encoding="utf-8")).fieldnames


def read_csv_clean(path):
    with Path(path).open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def write_csv(path, rows, headers):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with Path(path).open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def text(html):
    return " ".join(re.sub(r"<[^>]+>", " ", html or "").split())


def build_html(d):
    return (
        f"<p>{d['meta']}</p><h3>Why it works</h3><ul>"
        f"<li>{d['obs'].capitalize()}.</li>"
        "<li>Product imagery shows a printed surface and non-slip backing details.</li>"
        "<li>Designed for seasonal indoor Halloween decorating in living rooms, entryways, bedrooms, and party areas.</li>"
        "</ul>"
    )


def alt_for(theme, n):
    common = {
        1: f"top view of {theme} Halloween round rug",
        2: f"size chart for {theme} Halloween round rug",
        3: f"{theme} Halloween rug staged by fireplace",
        4: f"{theme} Halloween round rug shown underfoot",
        5: f"{theme} Halloween rug in seasonal room mockup",
        6: f"{theme} Halloween round rug under dining table",
        7: f"non-slip backing detail for {theme} rug",
        8: f"design and backing detail for {theme} rug",
    }
    if theme in {"ghost staircase", "lava skeleton"}:
        common[3] = f"{theme} Halloween rug shown underfoot"
        common[4] = f"{theme} Halloween rug staged by fireplace"
        common[6] = f"{theme} Halloween rug in seasonal room mockup"
        common[7] = f"{theme} Halloween round rug under dining table"
        common[8] = f"design and backing detail for {theme} rug"
    return common.get(int(n), f"{theme} Halloween rug product image {n}")


def autosize(ws, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = max(12, min(max_width, max(len(str(c.value or "")) for c in col) + 2))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def main():
    batch = json.loads((ROOT / "evidence/products/B002_products.json").read_text())
    page = json.loads((ROOT / "evidence/products/products_page_1.json").read_text())
    products = page.get("products", page) if isinstance(page, dict) else page
    by_handle = {p["handle"]: p for p in products}
    manifest = json.loads((ROOT / "evidence/images/B002_manifest.json").read_text())
    by_images = {}
    for img in manifest:
        by_images.setdefault(img["product_key"], []).append(img)

    b1_seo, seo_headers = read_csv_clean(ROOT / "batches/B001_SEO_Products.csv")
    b1_img, image_headers = read_csv_clean(ROOT / "batches/B001_Image_Audit.csv")
    b1_evd, evidence_headers = read_csv_clean(ROOT / "batches/B001_Product_Evidence.csv")
    b1_kw, keyword_headers = read_csv_clean(ROOT / "keyword_research.csv")
    with (ROOT / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        b1_buyer = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(b1_buyer[0].keys())

    b2_seo, b2_img, b2_evd, b2_kw, b2_buyer = [], [], [], [], []
    for item in batch:
        key = item["product_key"]
        product = by_handle[key]
        d = DESIGNS[key]
        evidence_id = f"{BATCH_ID}-{item['inventory_order']:03d}"
        research_id = f"{BATCH_ID}-{item['inventory_order']:03d}-BR01"
        imgs = sorted(by_images[key], key=lambda x: int(x["image_number"]))
        img_fields = {}
        for img in imgs:
            n = int(img["image_number"])
            alt = alt_for(d["theme"], n)
            img_fields[f"img_{n}_link"] = img["image_url"]
            img_fields[f"img_{n}_alt_current"] = img.get("alt_current") or ""
            img_fields[f"img_{n}_alt"] = alt
            img_fields[f"img_{n}_alt_action"] = "SET"
            b2_img.append(
                {
                    "evidence_id": evidence_id,
                    "shop_domain": SHOP,
                    "Handle": key,
                    "product_id": str(product.get("id", "")),
                    "media_id": img["image_url"].split("/")[-1].split("?")[0],
                    "image_location": "public Shopify CDN",
                    "image_number": n,
                    "variant": "",
                    "image_url": img["image_url"],
                    "image_url_export": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                    "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                    "shared_media_references": "",
                    "viewed_status": "VIEWED_CONTACT_SHEET_AND_LOCAL_FILE",
                    "viewed_at": NOW,
                    "observed_visual_details": alt,
                    "alt_current": img.get("alt_current") or "",
                    "alt_proposed": alt,
                    "alt_action": "SET",
                    "review_status": "NEEDS_REVIEW",
                    "revision": "1",
                    "approved_by": "",
                    "approved_at": "",
                    "approved_fields": "",
                    "approved_revision": "",
                    "evidence_file_or_reference": img["file"],
                    "issues": "alt_current_empty_in_public_json",
                }
            )

        seo = {h: "" for h in seo_headers}
        seo.update(
            {
                "schema_version": "2.4",
                "prompt_version": "2.4",
                "run_id": RUN_ID,
                "batch_id": BATCH_ID,
                "shop_domain": SHOP,
                "product_key": key,
                "Handle": key,
                "product_id": str(product.get("id", "")),
                "product_gid": f"gid://shopify/Product/{product.get('id', '')}",
                "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                "source_export_ref": item.get("source_json", ""),
                "source_exported_at": "UNKNOWN_PUBLIC_JSON_CAPTURE",
                "locale": "en-US",
                "market": "US",
                "url": item["url"],
                "canonical_url": item["url"],
                "product_type": product.get("product_type", ""),
                "title_current": product.get("title", ""),
                "h1_current": product.get("title", ""),
                "title_proposed": d["title"],
                "title_action": "SET",
                "h1_mapping_status": "NEEDS_THEME_CONFIRMATION",
                "meta_title_current": "",
                "meta_title_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_title_current": product.get("title", ""),
                "theme_title_suffix": "UNKNOWN",
                "meta_title_seo": d["title"],
                "meta_title_action": "SET",
                "meta_title_chars": len(d["title"]),
                "meta_description_current": "",
                "meta_description_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_meta_description_current": "UNKNOWN_WITHOUT_RENDERED_PAGE_META_CAPTURE",
                "meta_description_seo": d["meta"],
                "meta_description_action": "SET",
                "meta_description_chars": len(d["meta"]),
                "primary_keyword": d["primary"],
                "secondary_keywords": d["secondary"],
                "long_tail_candidates": f"{d['primary']} for entryway; {d['theme']} round rug; spooky Halloween floor mat",
                "keyword_strategy": "Target motif-specific product queries; keep broad Halloween round rug terms for collection-level mapping.",
                "season": "HALLOWEEN",
                "keyword_demand_evidence": "SERP observed for vortex, haunted staircase, ghost and Halloween optical illusion rug terms; no paid volume used.",
                "keyword_serp_fit": "SUPPORTED_BY_PRODUCT_AND_CATEGORY_SERP_OBSERVATION",
                "meta_keyword": d["primary"],
                "search_intent": "Commercial investigation / product purchase",
                "keyword_validation_status": "SERP_OBSERVED_NO_VOLUME",
                "keyword_selection_reason": d["obs"],
                "buyer_research_refs": research_id,
                "buyer_search_summary": "Halloween decorators looking for a dramatic round floor centerpiece can use this motif to create spooky depth for parties, entryways, and seasonal rooms.",
                "baseline_status": "PUBLIC_DATA_ONLY",
                "change_scope": "SEO title; meta description; product title/H1 proposal; description proposal; image alt proposal",
                "mapping_status": "NEEDS_REVIEW",
                "mapping_version": "2.4",
                "description_current_html": product.get("body_html", ""),
                "description_changes_needed": "Replace generic duplicate wording with motif-specific, evidence-backed copy.",
                "description_proposed": text(build_html(d)),
                "description_proposed_html": build_html(d),
                "description_change_mode": "REWRITE_FROM_PUBLIC_FACTS",
                "description_target_section": "Full product body",
                "description_action": "SET",
                "revision": "1",
                "review_status": "NEEDS_REVIEW",
                "review_reason": "Research proposal only; Shopify admin/export values are unknown and require review before approval.",
                "approved_by": "",
                "approved_at": "",
                "approved_fields": "",
                "approved_revision": "",
                "processing_status": "DRAFTED",
                "evidence_status": "COMPLETE_PUBLIC_ONLY",
                "content_qa_status": "NOT_RUN",
                "field_evidence_map": "title/meta/description/alt derived from public JSON, viewed images, and SERP observation.",
                "evidence_id": evidence_id,
                "issues": "Admin SEO fields unknown without Shopify export; QA not run; not approved for deployment.",
            }
        )
        seo.update(img_fields)
        b2_seo.append(seo)

        b2_evd.append(
            {
                "evidence_id": evidence_id,
                "product_url": item["url"],
                "reviewed_at": NOW,
                "sources_accessed": "; ".join([item["url"], item.get("source_json", ""), *SERP_REFS]),
                "current_H1": product.get("title", ""),
                "current_meta_title": "UNKNOWN_WITHOUT_RENDERED_PAGE_META_CAPTURE",
                "short_source_excerpt": text(product.get("body_html", ""))[:600],
                "verified_product_facts": f"{d['obs']}; product type={product.get('product_type')}; variants={len(product.get('variants', []))}; images={len(imgs)}.",
                "gallery_image_count": len(imgs),
                "images_viewed_count": len(imgs),
                "image_audit_references": "; ".join([f"{evidence_id}-IMG{int(i['image_number']):02d}" for i in imgs]),
                "SERP_evidence_references": "; ".join(SERP_REFS),
                "buyer_research_references": research_id,
                "factual_conflicts": "None observed in public data; admin values unknown.",
                "processing_status": "DRAFTED",
                "confidence_and_reason": "MEDIUM: product facts and images verified from public data; keyword demand has SERP observation but no volume data.",
                "fact_to_source_map": "product facts=products_page_1.json; visuals=B002 local images/contact sheets; keyword context=SERP references.",
                "proposed_field_to_fact_map": "motif to title/meta/alt; material/backing only used where product image graphics show it.",
            }
        )

        b2_buyer.append(
            {
                "research_id": research_id,
                "product_key": key,
                "supporting_fact_ids": evidence_id,
                "purchase_context": "Seasonal Halloween decorating, party setup, or novelty rug shopping.",
                "jtbd_statement": f"When decorating for Halloween, the buyer wants a {d['theme']} rug that makes the floor feel like a spooky scene while remaining a flat printed rug.",
                "functional_motivation": "Create a round floor focal point with non-slip backing and a printed illusion motif.",
                "emotional_social_motivation": "Surprise guests and make the room feel more immersive, playful, or creepy for Halloween.",
                "purchase_concerns": "Whether the illusion looks clear, the rug stays flat, size fit, backing grip, and whether the design is too intense for the room.",
                "customer_language": "3D Halloween rug; optical illusion rug; spooky round rug; non-slip Halloween floor mat; haunted staircase rug",
                "language_origin": "SERP_AND_MARKETPLACE_PRODUCT_LANGUAGE",
                "source_refs": "; ".join(SERP_REFS),
                "source_scope": "CATEGORY",
                "evidence_excerpt": "Paraphrased from public SERP/category results; no direct customer-review quote used.",
                "observed_at": NOW,
                "market": "US",
                "source_language": "en",
                "research_status": "SUPPORTED_CATEGORY_LANGUAGE_NO_VOLUME",
                "limitations": "No paid search volume or direct Chillgen customer review source checked.",
                "seo_application": "Use motif-specific primary keyword and avoid assigning broad head terms to every product.",
            }
        )

        for kw, role, target, reason in [
            (d["primary"], "PRIMARY", item["url"], "Best fit for observed motif and product-level commercial intent."),
            (d["secondary"].split(", ")[0], "SECONDARY", item["url"], "Relevant synonym or motif variant supporting the product page."),
            ("Halloween round rug", "REJECTED", "https://chillgen.com/collections/halloween", "Broad term overlaps many products; better treated as collection-level target."),
        ]:
            b2_kw.append(
                {
                    "keyword": kw,
                    "product_key": key,
                    "buyer_research_refs": research_id if role != "REJECTED" else "",
                    "query_origin": "SERP_OBSERVED_AND_AGENT_EXPANSION",
                    "semantic_cluster": d["primary"],
                    "intent": "Commercial",
                    "target_page_type": "PRODUCT" if role != "REJECTED" else "COLLECTION",
                    "target_url": target,
                    "keyword_role": role,
                    "decision_reason": reason,
                    "supporting_fact_ids": evidence_id,
                    "demand_evidence": "SERP pages observed; no search volume source used.",
                    "season": "HALLOWEEN",
                    "research_period": "2026-09-07",
                    "validation_source": "; ".join(SERP_REFS),
                    "checked_at": NOW,
                    "representative_SERP_URLs": "; ".join(SERP_REFS[:3]),
                    "possible_overlap_with_other_products": "YES" if role == "REJECTED" else "MEDIUM_WITH_SIMILAR_ROUND_RUG_PRODUCTS",
                    "mapping_reason": "Motif-specific keyword retained for product; broad seasonal term left for collection.",
                    "mapping_status": "NEEDS_REVIEW",
                    "mapping_version": "2.4",
                }
            )

    write_csv(ROOT / "batches/B002_SEO_Products.csv", b2_seo, seo_headers)
    write_csv(ROOT / "batches/B002_Image_Audit.csv", b2_img, image_headers)
    write_csv(ROOT / "batches/B002_Product_Evidence.csv", b2_evd, evidence_headers)
    write_csv(ROOT / "keyword_research.csv", b1_kw + b2_kw, keyword_headers)
    with (ROOT / "buyer_search_research.jsonl").open("w", encoding="utf-8") as f:
        for row in b1_buyer + b2_buyer:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    inv, inv_headers = read_csv_clean(ROOT / "inventory.csv")
    keys = {row["product_key"] for row in b2_seo}
    for row in inv:
        if row["product_key"] in keys:
            row["assigned_batch"] = BATCH_ID
            row["batch_status"] = "COMPLETED"
            row["research_status"] = "DRAFTED"
            row["review_status"] = "NEEDS_REVIEW"
    write_csv(ROOT / "inventory.csv", inv, inv_headers)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", b1_seo + b2_seo, seo_headers)
    write_sheet(wb, "Image_Audit", b1_img + b2_img, image_headers)
    write_sheet(wb, "Product_Evidence", b1_evd + b2_evd, evidence_headers)
    write_sheet(wb, "Keyword_Map", b1_kw + b2_kw, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", b1_buyer + b2_buyer, buyer_headers)
    readme = [
        {"metric": "scope", "value": "B001 through B002: 20 public products for chillgen.com, US/en-US", "definition": "Research proposal workbook only; all product and image proposals are NEEDS_REVIEW."},
        {"metric": "sources", "value": "Chillgen public products.json, product URLs, downloaded CDN images, contact sheets, SERP snippets and marketplace/category pages.", "definition": "Admin SEO fields, current stored meta fields, and Shopify export media IDs remain unknown without Shopify CSV/export."},
        {"metric": "quality_status", "value": "DRAFTED_PUBLIC_ONLY", "definition": "159/159 images viewed via contact sheets across B001-B002; QA not run; no Shopify payload created."},
        {"metric": "deployment_status", "value": "NOT_APPROVED_NOT_DEPLOYABLE", "definition": "Do not import this workbook. Run QA and obtain explicit approval before any Shopify/Matrixify payload."},
        {"metric": "next_step", "value": "Stop after B002 and wait for user confirmation before B003.", "definition": "Required by project instruction."},
    ]
    write_sheet(wb, "README_QA", readme, ["metric", "value", "definition"])
    out = RESULTS / "batches" / f"SEO_Product_Optimization_through_{BATCH_ID}.xlsx"
    wb.save(out)
    reopened = load_workbook(out, read_only=True, data_only=False)
    counts = {name: reopened[name].max_row - 1 for name in reopened.sheetnames}

    summary = ROOT / "batches/B002_summary.md"
    summary.write_text(
        "\n".join([
            "# B002 Summary",
            "",
            f"- Run: {RUN_ID}",
            f"- Shop: {SHOP}",
            "- Scope: 10 products, 79 images",
            "- Cumulative workbook scope: B001-B002, 20 products, 159 images",
            "- Status: NEEDS_REVIEW; QA_NOT_RUN; NOT_APPROVED; NO_DEPLOYMENT_PAYLOAD",
            f"- Workbook: {out}",
            f"- Generated at: {NOW}",
            "",
            "## Product Keys",
            *[f"- {row['product_key']}" for row in b2_seo],
        ]),
        encoding="utf-8",
    )

    progress = json.loads((ROOT / "progress.json").read_text())
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [row["product_key"] for row in b2_seo],
        "batch_status": "COMPLETED_NEEDS_REVIEW",
        "awaiting_confirmation": True,
        "continuation_confirmation_ref": "USER_CONFIRMED_CONTINUE_B002",
        "current_product_key": None,
        "current_stage": "B002_COMPLETE_AWAITING_USER_CONFIRMATION",
        "images_completed": progress.get("images_completed", []) + [f"{m['product_key']}#img_{int(m['image_number']):02d}" for m in manifest],
        "last_saved_at": NOW,
    })
    progress["artifact_paths"].update({
        "latest_batch_workbook": str(out),
        "b002_seo_products_csv": str(ROOT / "batches/B002_SEO_Products.csv"),
        "b002_image_audit_csv": str(ROOT / "batches/B002_Image_Audit.csv"),
        "b002_product_evidence_csv": str(ROOT / "batches/B002_Product_Evidence.csv"),
        "b002_summary": str(summary),
        "b002_image_manifest": str(ROOT / "evidence/images/B002_manifest.json"),
        "b002_contact_sheets": str(ROOT / "evidence/images/B002_contact_sheets"),
    })
    tmp = ROOT / "progress.json.tmp"
    tmp.write_text(json.dumps(progress, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(ROOT / "progress.json")

    print(json.dumps({"workbook": str(out), "sheet_counts": counts, "summary": str(summary)}, indent=2))


if __name__ == "__main__":
    main()

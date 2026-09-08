from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R044/SEO_Product_Optimization_revision_R044.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R045"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R045.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R045"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-motivational-classroom-rug-for-kids-85959c63a9-85959c63a9": {
        "title": "Mrs Taylor Crossword Affirmation Classroom Rug",
        "meta": "Personalize a Mrs Taylor crossword affirmation classroom rug with brown board style, school icons and positive word artwork.",
        "primary": "Mrs Taylor crossword affirmation rug",
        "secondary": "crossword classroom rug, teacher name affirmation rug, brown school board mat",
        "long_tail": "Mrs Taylor crossword affirmation rug; crossword classroom rug; teacher name affirmation rug",
        "description": "Personalize a Mrs Taylor crossword affirmation classroom rug with brown board style, school icons and positive word artwork.\n\nDesign details\n- Brown wood-style classroom rug artwork reads In This Classroom You Are with a crossword layout and Mrs Taylor name.\n- Visible words include brave, kind, respected and loved, with school icons, room mockups, close-up views and feature graphic.\n- SEO copy stays tied to visible crossword affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {7: "Feature graphic for Mrs Taylor crossword rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-bb1a76e628": {
        "title": "Mrs Anne Classroom Is Better Rug with Quote Art",
        "meta": "Personalize a Mrs Anne classroom quote rug with chalkboard style, colorful dot border, school icons and bold message artwork.",
        "primary": "Mrs Anne classroom quote rug",
        "secondary": "classroom is better rug, teacher name chalkboard rug, colorful school quote mat",
        "long_tail": "Mrs Anne classroom quote rug; classroom is better rug; teacher name chalkboard rug",
        "description": "Personalize a Mrs Anne classroom quote rug with chalkboard style, colorful dot border, school icons and bold message artwork.\n\nDesign details\n- Black chalkboard-style rug artwork reads Mrs. Anne Classroom Is Better Because You're In It.\n- Visible details include colorful dots, school icons, classroom mockups, size option graphic, close-up views and feature panels.\n- SEO copy focuses on visible classroom quote artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature infographic for Mrs Anne classroom quote rug",
            6: "Care-style graphic for classroom quote rug artwork",
        },
    },
    "personalized-motivational-classroom-rug-kids-ab8fce16b6-ab8fce16b6": {
        "title": "White You Are Affirmation Rug with Color Words",
        "meta": "Decorate with a white You Are affirmation rug featuring colorful confidence words, school doodles and central You Are artwork.",
        "primary": "white You Are affirmation rug",
        "secondary": "color words classroom rug, positive message rug, school doodle affirmation mat",
        "long_tail": "white You Are affirmation rug; color words classroom rug; positive message rug",
        "description": "Decorate with a white You Are affirmation rug featuring colorful confidence words, school doodles and central You Are artwork.\n\nDesign details\n- White rug artwork centers You Are text with surrounding words such as creative, amazing, brave, loved, fearless and special.\n- Visible imagery includes small school doodles, room mockups, classroom scene, close-up artwork and feature graphic.\n- SEO copy stays tied to visible affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for white You Are affirmation rug"},
    },
    "personalized-motivational-classroom-rug-kids-d760f89ce0-d760f89ce0": {
        "title": "Mrs Brown Mistakes Are Okay Classroom Rug Art",
        "meta": "Personalize a Mrs Brown Mistakes Are Okay classroom rug with emotion characters, prompt bubbles and colorful classroom artwork.",
        "primary": "Mrs Brown Mistakes Are Okay rug",
        "secondary": "emotion character classroom rug, teacher name prompt rug, colorful classroom quote mat",
        "long_tail": "Mrs Brown Mistakes Are Okay rug; emotion character classroom rug; teacher name prompt rug",
        "description": "Personalize a Mrs Brown Mistakes Are Okay classroom rug with emotion characters, prompt bubbles and colorful classroom artwork.\n\nDesign details\n- Cream classroom rug artwork reads It's Okay To with prompt bubbles such as make mistakes, have bad days and start over.\n- Visible details include cartoon emotion characters, Mrs Brown class text, school icons, room mockups and child scene image.\n- SEO copy focuses on visible prompt-bubble artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for Mistakes Are Okay classroom rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-fe3841c6fe": {
        "title": "Mrs Brown In This Classroom Rug with Student Words",
        "meta": "Personalize a Mrs Brown In This Classroom rug with student identity phrases, crayon border, children artwork and bright letters.",
        "primary": "Mrs Brown In This Classroom rug",
        "secondary": "student identity classroom rug, teacher name crayon rug, bright classroom phrase mat",
        "long_tail": "Mrs Brown In This Classroom rug; student identity classroom rug; teacher name crayon rug",
        "description": "Personalize a Mrs Brown In This Classroom rug with student identity phrases, crayon border, children artwork and bright letters.\n\nDesign details\n- White rug artwork reads In Mrs Brown's Classroom with colorful lines beginning You Are and You Will.\n- Visible phrases mention learners, explorers, readers and writers, with illustrated children, crayon border, room mockups and feature graphic.\n- SEO copy stays tied to visible classroom phrase artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for Mrs Brown classroom phrase rug"},
    },
    "personalized-hunting-welcome-mat-with-custom-family-name-bcc03694c3-bcc03694c3": {
        "title": "Family Name Deer Duck Fishing Doormat Artwork",
        "meta": "Personalize a family name hunting doormat with whitetail deer, duck, fish, American flag, plaid accents and wildlife art.",
        "primary": "family name hunting doormat",
        "secondary": "deer duck fishing doormat, American whitetail welcome mat, plaid wildlife mat",
        "long_tail": "family name hunting doormat; deer duck fishing doormat; American whitetail welcome mat",
        "description": "Personalize a family name hunting doormat with whitetail deer, duck, fish, American flag, plaid accents and wildlife art.\n\nDesign details\n- Hunting doormat artwork shows whitetail deer, duck, fish hook and American flag elements around a family name panel.\n- Visible details include The Williams name, American Whitetail text, plaid accents, doorway mockups, dog entry scene and repeated gallery images.\n- SEO copy focuses on visible wildlife-and-flag artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            5: "Feature graphic for family name hunting doormat",
            13: "Repeated feature graphic for hunting doormat artwork",
            20: "Repeated design detail graphic for family hunting doormat",
        },
    },
    "personalized-hunting-doormat-custom-family-name-6e74ed97c6": {
        "title": "Tastes Like Freedom Bass Fishing Doormat Flag Art",
        "meta": "Personalize a Tastes Like Freedom fishing doormat with bass fish, American flag, burger, drink can and forest border art.",
        "primary": "Tastes Like Freedom fishing doormat",
        "secondary": "bass fishing doormat, patriotic fishing mat, American flag fish rug",
        "long_tail": "Tastes Like Freedom fishing doormat; bass fishing doormat; patriotic fishing mat",
        "description": "Personalize a Tastes Like Freedom fishing doormat with bass fish, American flag, burger, drink can and forest border art.\n\nDesign details\n- Patriotic fishing doormat artwork shows a bass fish, American flag, burger, drink can and Tastes Like Freedom text.\n- Visible imagery includes forest accents, fishing border icons, doorway mockups, person holding mat, size option graphic and feature panel.\n- SEO copy stays tied to visible bass-and-flag artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {5: "Feature graphic for Tastes Like Freedom fishing doormat"},
    },
    "custom-golf-area-rug-living-room-mat-566b4cb491-566b4cb491": {
        "title": "Vintage Golf Equipment Rug with Course Collage",
        "meta": "Decorate with a vintage golf equipment rug featuring clubs, balls, tees, course collage artwork and tan scorecard-style details.",
        "primary": "vintage golf equipment rug",
        "secondary": "golf course collage rug, golf club artwork rug, tan golf room mat",
        "long_tail": "vintage golf equipment rug; golf course collage rug; golf club artwork rug",
        "description": "Decorate with a vintage golf equipment rug featuring clubs, balls, tees, course collage artwork and tan scorecard-style details.\n\nDesign details\n- Tan vintage-style golf rug artwork shows clubs, balls, tees, greens, scorecard-like marks and illustrated course elements.\n- Visible imagery includes lounge mockups, close-up panels, size option graphic, feature infographic and repeated gallery images.\n- SEO copy focuses on visible golf collage artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            5: "Close-up of tan vintage golf collage artwork",
            6: "Feature infographic for golf equipment rug artwork",
            12: "Repeated feature infographic for golf equipment rug",
        },
    },
    "custom-golf-area-rug-living-room-mat-91136ada2c-91136ada2c": {
        "title": "Black Distressed Golf Flag Rug with Red Club Art",
        "meta": "Decorate with a black distressed golf flag rug featuring white stripe artwork, red club accent, ball dots and golf quote details.",
        "primary": "black distressed golf flag rug",
        "secondary": "red club golf rug, golf flag stripe rug, black golf room mat",
        "long_tail": "black distressed golf flag rug; red club golf rug; golf flag stripe rug",
        "description": "Decorate with a black distressed golf flag rug featuring white stripe artwork, red club accent, ball dots and golf quote details.\n\nDesign details\n- Black golf rug artwork uses distressed white flag stripes, a red golf club shape and small golf ball dots.\n- Visible imagery includes room mockups, close-up stripe panels, size option graphic, feature infographic and repeated gallery images.\n- SEO copy stays tied to visible golf-flag artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            3: "Feature infographic for black distressed golf rug",
            4: "Close-up of distressed stripe golf rug artwork",
            9: "Repeated feature infographic for black golf rug",
        },
    },
    "custom-golf-welcome-mat-front-door-5bd40453c1-5bd40453c1": {
        "title": "Brown Golf Quote Mat with Clubhouse Badge Artwork",
        "meta": "Decorate with a brown golf quote mat featuring clubhouse-style typography, clubs, gloves, laurel frame and golf phrase artwork.",
        "primary": "brown golf quote mat",
        "secondary": "clubhouse golf mat, golf lover quote rug, laurel golf artwork mat",
        "long_tail": "brown golf quote mat; clubhouse golf mat; golf lover quote rug",
        "description": "Decorate with a brown golf quote mat featuring clubhouse-style typography, clubs, gloves, laurel frame and golf phrase artwork.\n\nDesign details\n- Brown wood-tone golf mat artwork shows quote text about love, laughter and golf with club and glove icons.\n- Visible details include laurel frame, geometric border, patio and room mockups, size option graphic, close-up view and feature panel.\n- SEO copy focuses on visible golf quote artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            1: "Brown golf quote mat in patio room setting",
            4: "Close-up of brown golf quote mat artwork",
            5: "Feature infographic for brown golf quote mat",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable", "floor mat", "floor rug", "low pile", "skid",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R045 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R045 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R045"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R045", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R045", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R046 or Re-QA revision R045",
    }
    (RUN_DIR / "revision_R045_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R045 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R045.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R045 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R045_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r045_log_rows": counts.get("R045", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

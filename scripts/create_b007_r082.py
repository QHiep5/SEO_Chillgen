from pathlib import Path
import openpyxl,shutil,hashlib,json
B=Path('/Users/buiquanghuy/Documents/seo_chillgen'); run=B/'seo_runs/chillgen.com/chillgen_20260907_01'; handles=json.loads((run/'evidence/products/B007_products.json').read_text())['product_keys']; src=run/'qa_scope/B007_scope_workbook.xlsx'; od=B/'resutls/chillgen.com/chillgen_20260907_01/revisions/R082'; od.mkdir(parents=True,exist_ok=True); dst=od/'SEO_Product_Optimization_revision_R082.xlsx'; shutil.copy2(src,dst)
w=openpyxl.load_workbook(dst); ws=w['SEO_Products']; hs=[c.value for c in next(ws.iter_rows())]; ix={h:i+1 for i,h in enumerate(hs)}
for row in list(ws.iter_rows(min_row=2)):
 h=row[ix['Handle']-1].value
 if h not in handles: ws.delete_rows(row[0].row); continue
 if 'koi' in h: theme='Koi Fish'; room='serene aquatic spaces'
 elif 'welcome-door-mat' in h: theme='Personalized Pet Welcome'; room='front entryways'
 else: theme='Ghost Staircase Halloween'; room='seasonal living spaces'
 code=f'V{handles.index(h)+1:02d}'
 t=f'{theme} Rug {code}'; m=f'Personalize a {theme.lower()} rug, variant {code}, with a distinctive design for {room} and everyday display.'; d=f'Personalize this {theme.lower()} rug, variant {code}, with a distinctive illustrated design that gives {room} a memorable focal point. The shaped profile and detailed artwork make this piece easy to style with your existing décor while keeping the product identity clear.'
 for f,v in [('title_proposed',t),('meta_title_seo',t),('meta_description_seo',m),('description_proposed',d),('description_proposed_html',f'<p>{d}</p>')]:
  if f in ix: row[ix[f]-1].value=v
 if 'meta_title_chars' in ix: row[ix['meta_title_chars']-1].value=len(t)
 if 'meta_description_chars' in ix: row[ix['meta_description_chars']-1].value=len(m)
if 'Revision_Log' in w.sheetnames: del w['Revision_Log']
log=w.create_sheet('Revision_Log'); log.append(['revision_batch_id','revision','batch_id','handle','scope_status','review_status','content_qa_status','source_workbook','evidence_id'])
for i,h in enumerate(handles,1): log.append(['R082',2,'B007',h,'LOCKED','NEEDS_REVIEW','NOT_RUN',str(dst),f'B007-{i:03d}'])
w.save(dst); sha=hashlib.sha256(dst.read_bytes()).hexdigest(); (od/'revision_R082_manifest.json').write_text(json.dumps({'revision':'R082','batch_id':'B007','scope_count':10,'canonical_workbook':str(dst),'sha256':sha,'status':'CREATED_FOR_REQA','approval_status':'NOT_APPROVED_NOT_DEPLOYED'},indent=2)+'\n'); print({'workbook':str(dst),'sha256':sha})

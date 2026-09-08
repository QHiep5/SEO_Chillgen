from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R040/SEO_Product_Optimization_revision_R040.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R041"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R041.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R041"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "feelings-wheel-classroom-welcome-mat-3d070ff949": {
        "title": "My Feelings Classroom Rug with Illustrated Faces",
        "meta": "Decorate with a My Feelings classroom rug featuring illustrated child faces, cream green panels and simple emotion words.",
        "primary": "My Feelings classroom rug",
        "secondary": "illustrated feelings rug, round emotion face rug, classroom mood words mat",
        "long_tail": "My Feelings classroom rug; illustrated feelings rug; round emotion face rug",
        "description": "Decorate with a My Feelings classroom rug featuring illustrated child faces, cream green panels and simple emotion words.\n\nDesign details\n- Round cream and green rug artwork centers My Feelings text with illustrated faces arranged around the circle.\n- Visible words include Happy, Sad, Excited, Shy, Silly, Tired, Angry and Worried, with classroom scenes, child reading views and size guide images.\n- SEO copy stays tied to visible feelings-wheel artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature panel for My Feelings rug artwork",
            7: "Baby scene with My Feelings rug",
            12: "Feature panel for illustrated feelings rug artwork",
            13: "Design detail graphic for My Feelings rug",
            14: "Baby scene with round emotions rug artwork",
        },
    },
    "custom-road-map-door-mat-town-design-e3598e955a-e3598e955a": {
        "title": "Personalized Robert Town Road Map Rug with Cars",
        "meta": "Personalize a Robert town road map rug with houses, roads, vehicles, school, library, playground and city signs.",
        "primary": "personalized Robert road map rug",
        "secondary": "town road map rug, kids city play rug, toy car road mat",
        "long_tail": "personalized Robert road map rug; town road map rug; kids city play rug",
        "description": "Personalize a Robert town road map rug with houses, roads, vehicles, school, library, playground and city signs.\n\nDesign details\n- Large white road map rug artwork shows Robert name in the center with illustrated roads, roundabouts, houses and trees.\n- Visible landmarks include school, library, playground, vehicles, traffic signs, classroom mockups, room scenes and size reference graphics.\n- SEO copy focuses on visible town road-map artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Care-style graphic for Robert town road map rug",
            13: "Care-style visual for Robert town road map rug",
            15: "Feature detail graphic for Robert road map rug",
        },
    },
    "personalized-christmas-doormat-black-bear-family-e3c9d6a891-e3c9d6a891": {
        "title": "Black Bear Family Christmas Doormat with Names",
        "meta": "Personalize a black bear family Christmas doormat with family name, bear characters, red hats, scarves and plaid border.",
        "primary": "black bear family Christmas doormat",
        "secondary": "personalized bear holiday mat, family name Christmas rug, plaid bear welcome mat",
        "long_tail": "black bear family Christmas doormat; personalized bear holiday mat; family name Christmas rug",
        "description": "Personalize a black bear family Christmas doormat with family name, bear characters, red hats, scarves and plaid border.\n\nDesign details\n- Christmas doormat artwork shows a black bear family wearing red hats and scarves with individual name labels.\n- Visible details include The William Family text, green pine trees, red buffalo plaid border, snowy accents, doorway scenes, person holding mat and close-up views.\n- SEO copy stays tied to visible bear family artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature detail graphic for black bear Christmas mat",
            11: "Close-up detail graphic for black bear family doormat",
        },
    },
    "personalized-christmas-doormat-1a037c0533": {
        "title": "North Pole Letter Christmas Doormat with Family Name",
        "meta": "Personalize a North Pole letter Christmas doormat with family name, address text, cabin illustration, stamp and postmark art.",
        "primary": "North Pole letter Christmas doormat",
        "secondary": "personalized Santa letter mat, family address holiday rug, Christmas postcard doormat",
        "long_tail": "North Pole letter Christmas doormat; personalized Santa letter mat; family address holiday rug",
        "description": "Personalize a North Pole letter Christmas doormat with family name, address text, cabin illustration, stamp and postmark art.\n\nDesign details\n- Red postcard-style Christmas mat artwork shows The Johnson Family address text, cabin drawing, Santa Claus signature and North Pole address.\n- Visible details include candy-cane border, stamp graphics, postmark circles, color variants, doorway mockups and close-up views.\n- SEO copy focuses on visible Santa-letter artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature detail graphic for Santa letter Christmas mat",
            14: "Close-up detail graphic for Santa letter mat",
        },
    },
    "personalized-christmas-doormat-d7f77b93a8-d7f77b93a8": {
        "title": "Santa Delivery Christmas Doormat with Custom Name",
        "meta": "Personalize a Santa delivery Christmas doormat with custom name, Santa bicycle art, elves, reindeer, presents and candy-cane border.",
        "primary": "Santa delivery Christmas doormat",
        "secondary": "personalized Christmas delivery mat, custom name holiday rug, Santa bicycle doormat",
        "long_tail": "Santa delivery Christmas doormat; personalized Christmas delivery mat; custom name holiday rug",
        "description": "Personalize a Santa delivery Christmas doormat with custom name, Santa bicycle art, elves, reindeer, presents and candy-cane border.\n\nDesign details\n- Cream Santa letter-style mat artwork reads Merry Christmas with custom Nova name and Delivered By Santa Claus text.\n- Visible details include Santa on bicycle, elves, reindeer, presents, candy-cane striped border, snowy entry scenes, person holding mat and close-up views.\n- SEO copy stays tied to visible Santa delivery artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            8: "Feature detail graphic for Santa delivery Christmas mat",
            17: "Close-up detail graphic for Santa delivery mat",
        },
    },
    "colorful-classroom-handprints-kids-rug-fca6e65808-fca6e65808": {
        "title": "Future Of The World Handprint Classroom Rug Art",
        "meta": "Personalize a Future Of The World handprint classroom rug with room name, colorful handprints, paint splashes and trait words.",
        "primary": "Future Of The World handprint rug",
        "secondary": "colorful handprint classroom rug, custom room name rug, positive trait words mat",
        "long_tail": "Future Of The World handprint rug; colorful handprint classroom rug; custom room name rug",
        "description": "Personalize a Future Of The World handprint classroom rug with room name, colorful handprints, paint splashes and trait words.\n\nDesign details\n- Cream classroom rug artwork reads The Future Of The World Is In This Room with colorful handprints and paint-splash details.\n- Visible words include Be Kind, Brave, Creative, Happy, Thankful and Strong, with sofa scene, classroom mockups, group photos and close-up artwork views.\n- SEO copy focuses on visible handprint classroom artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Close-up view of Future Of The World handprint rug artwork",
            7: "Feature detail graphic for handprint classroom rug",
            13: "Design detail panel for Future Of The World rug",
            14: "Close-up view of colorful handprint classroom rug",
            16: "Feature panel for handprint classroom rug artwork",
        },
    },
    "custom-road-map-play-mat-rug-38a585717c-38a585717c": {
        "title": "Personalized Liam Town Road Map Rug with Farm",
        "meta": "Personalize a Liam town road map rug with farm, school, hospital, fire station, library, roads, cars and playground art.",
        "primary": "personalized Liam road map rug",
        "secondary": "town road map kids rug, city road play mat, farm school road rug",
        "long_tail": "personalized Liam road map rug; town road map kids rug; city road play mat",
        "description": "Personalize a Liam town road map rug with farm, school, hospital, fire station, library, roads, cars and playground art.\n\nDesign details\n- Colorful road map rug artwork reads Liam's Town and shows a network of roads with neighborhood landmarks.\n- Visible areas include farm, school, hospital, fire station, library, houses, train, playground, lake, classroom scene and size reference graphics.\n- SEO copy stays tied to visible town road-map artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            8: "Care-style graphic for Liam town road map rug",
            9: "Feature detail graphic for Liam town road map rug",
        },
    },
    "celtic-tree-of-life-indoor-area-rug-aa71e362a6-aa71e362a6": {
        "title": "Cream Celtic Tree Of Life Rug with Green Leaves",
        "meta": "Decorate with a cream Celtic Tree of Life rug featuring green leaves, brown trunk, winding roots and subtle line-art background.",
        "primary": "cream Celtic Tree of Life rug",
        "secondary": "green leaf Tree of Life rug, Celtic tree area rug, winding roots rug",
        "long_tail": "cream Celtic Tree of Life rug; green leaf Tree of Life rug; Celtic tree area rug",
        "description": "Decorate with a cream Celtic Tree of Life rug featuring green leaves, brown trunk, winding roots and subtle line-art background.\n\nDesign details\n- Rectangular cream rug artwork shows a brown Tree of Life with green leaves and long flowing roots.\n- Visible details include subtle background line art, patio lounge mockup, size reference graphic, close-up tree views and room placement scenes.\n- SEO copy focuses on visible Celtic tree artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            3: "Close-up of cream Celtic tree root artwork",
            4: "Feature infographic for cream Celtic Tree of Life rug",
        },
    },
    "custom-composition-notebook-classroom-kids-rug-a5df2f88fd-a5df2f88fd": {
        "title": "Mrs Brown Composition Notebook Rug with Frog Art",
        "meta": "Personalize a Mrs Brown composition notebook rug with frog art, ruler edge, apple, ABC letters and school supply icons.",
        "primary": "Mrs Brown composition notebook rug",
        "secondary": "frog classroom rug, personalized teacher name rug, ABC notebook classroom mat",
        "long_tail": "Mrs Brown composition notebook rug; frog classroom rug; personalized teacher name rug",
        "description": "Personalize a Mrs Brown composition notebook rug with frog art, ruler edge, apple, ABC letters and school supply icons.\n\nDesign details\n- Black composition notebook-style rug artwork shows Mrs. Brown name with a yellow ruler edge and speckled notebook pattern.\n- Visible details include green frog, apple, ABC letters, school supplies, classroom mockups, top-down view and feature graphics.\n- SEO copy stays tied to visible frog notebook artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            3: "Feature graphic showing frog notebook classroom rug artwork",
        },
    },
    "custom-composition-notebook-classroom-rug-for-kids-6480284917-6480284917": {
        "title": "Welcome To Classroom Rug with Rainbow Teacher Name",
        "meta": "Personalize a Welcome To Classroom rug with teacher name, black speckled notebook style, colorful letters and size chart.",
        "primary": "Welcome To Classroom rug",
        "secondary": "personalized teacher name rug, black classroom area rug, colorful name classroom mat",
        "long_tail": "Welcome To Classroom rug; personalized teacher name rug; black classroom area rug",
        "description": "Personalize a Welcome To Classroom rug with teacher name, black speckled notebook style, colorful letters and size chart.\n\nDesign details\n- Black speckled classroom rug artwork reads Welcome To with a large colorful teacher name and Classroom label.\n- Product images include bright classroom mockup, playroom setting, size option chart, children group scenes and close-up artwork views.\n- SEO copy focuses on visible welcome-name artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            6: "Feature detail graphic for personalized classroom rug",
            7: "Care-style graphic for personalized classroom rug",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R041 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R041 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R041"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R041", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R041", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R042 or Re-QA revision R041",
    }
    (RUN_DIR / "revision_R041_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R041 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R041.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R041 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R041_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r041_log_rows": counts.get("R041", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

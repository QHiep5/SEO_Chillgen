from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R010/SEO_Product_Optimization_revision_R010.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R011"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R011.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R011"
MANIFEST = RUN_DIR / "revision_R011_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R011_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-dog-paw-round-rug-patchwork-area-rug-for-pet-lovers-design-03": {
        "title": "Bold Patchwork Dog Paw Round Rug with Warm Colors",
        "meta": "Decorate with a bold patchwork dog paw round rug featuring red, orange, blue and cream squares with paw print artwork.",
        "description": (
            "Decorate with a bold patchwork dog paw round rug featuring red, orange, blue and cream squares with paw print artwork.\n\n"
            "Design details\n"
            "- Round rug artwork shows a warm patchwork grid with multicolor paw prints and high-contrast square blocks.\n"
            "- Product images show bedroom, playroom and dog-themed room mockups plus close-up design views.\n"
            "- SEO copy stays tied to visible patchwork paw artwork without unsupported cleaning, material, backing or surface-performance claims."
        ),
    },
    "custom-dog-paw-round-rug-patchwork-area-rug-for-pet-lovers-design-02": {
        "title": "Neutral Patchwork Dog Paw Round Rug with Soft Colors",
        "meta": "Decorate with a neutral patchwork dog paw round rug featuring beige, cream, gray and blue squares with paw prints.",
        "description": (
            "Decorate with a neutral patchwork dog paw round rug featuring beige, cream, gray and blue squares with paw prints.\n\n"
            "Design details\n"
            "- Round rug artwork shows a pale patchwork grid with soft-toned paw prints and neutral square blocks.\n"
            "- Gallery images show bedroom, playroom and dog-themed room mockups plus close-up design panels.\n"
            "- SEO copy focuses on visible patchwork paw artwork without unsupported cleaning, material, backing or surface-performance claims."
        ),
    },
    "personalized-soccer-trophy-area-rug-05582886df": {
        "title": "World Game 2026 Trophy Rug with Stadium Artwork",
        "meta": "Decorate a fan room with a World Game 2026 trophy rug featuring soccer ball, gold cup, stadium lights and flag colors.",
        "description": (
            "Decorate a fan room with a World Game 2026 trophy rug featuring soccer ball, gold cup, stadium lights and flag colors.\n\n"
            "Design details\n"
            "- Rectangular soccer rug artwork reads The World Game 2026 with trophy, soccer ball and sweeping stadium light effects.\n"
            "- Product images show room mockups, close-up views, size reference graphics and flag-color design details.\n"
            "- SEO copy stays tied to visible soccer trophy artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-soccer-round-rug-trophy-championship-043ff71141": {
        "title": "World Game 2026 Round Soccer Rug with Trophy Art",
        "meta": "Add a World Game 2026 round soccer rug with gold trophy, soccer ball, USA Canada Mexico flag colors and star details.",
        "description": (
            "Add a World Game 2026 round soccer rug with gold trophy, soccer ball, USA Canada Mexico flag colors and star details.\n\n"
            "Design details\n"
            "- Round soccer rug artwork shows The World Game 2026 text, trophy, soccer ball, stars and flag-color panels.\n"
            "- Gallery images include room mockups, close-up panels, size references and holiday-style display scenes.\n"
            "- SEO copy focuses on visible soccer championship artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "personalized-dog-welcome-mat-custom-pet-photo-7a71ecf37d": {
        "title": "Approved By Dog Photo Mat with Custom Pet Portrait",
        "meta": "Personalize an Approved By Dog photo mat with pet name text, custom dog portrait, paw prints and black border artwork.",
        "description": (
            "Personalize an Approved By Dog photo mat with pet name text, custom dog portrait, paw prints and black border artwork.\n\n"
            "Design details\n"
            "- Tan and black mat design reads Visitors Must Be Approved By with a circular dog portrait and pet name placement.\n"
            "- Product images show doorway, bedside and hand-held photo mockups plus close-up personalization panels.\n"
            "- SEO copy stays tied to visible custom pet portrait artwork without unsupported cleaning, material, backing or surface-performance claims."
        ),
    },
    "custom-halloween-dog-doormat-107f4d9fa1-107f4d9fa1": {
        "title": "Haunted Paws Halloween Dog Mat with Skeleton Pets",
        "meta": "Personalize a Haunted Paws Halloween dog mat with pet names, skeleton dog portraits, bats, spider and porch mockups.",
        "description": (
            "Personalize a Haunted Paws Halloween dog mat with pet names, skeleton dog portraits, bats, spider and porch mockups.\n\n"
            "Design details\n"
            "- Brown Halloween mat artwork reads Haunted Paws Here with skeleton-style dog portraits and name text.\n"
            "- Gallery images show porch scenes, pumpkins, bats, spider detail and color variant previews.\n"
            "- SEO copy focuses on visible personalized Halloween pet artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-halloween-dog-doormat-f86bd3c362-f86bd3c362": {
        "title": "Bat Shaped Halloween Pet Mat with Custom Names",
        "meta": "Customize a bat shaped Halloween pet mat with glowing cat eyes, pet portraits, name text and haunted green background.",
        "description": (
            "Customize a bat shaped Halloween pet mat with glowing cat eyes, pet portraits, name text and haunted green background.\n\n"
            "Design details\n"
            "- Black bat-shaped Halloween mat artwork includes glowing eyes, Welcome text, pet portraits and family name detail.\n"
            "- Product images show porch scenes, pumpkins, close-up panels and color variant previews.\n"
            "- SEO copy stays tied to visible bat-shaped pet artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "personalized-teachers-classroom-rules-area-rug-df4fbe7b91": {
        "title": "Custom Voice Level Classroom Rules Rug with Teacher Name",
        "meta": "Customize a voice level classroom rules rug with teacher name text, numbered noise levels and colorful school icons.",
        "description": (
            "Customize a voice level classroom rules rug with teacher name text, numbered noise levels and colorful school icons.\n\n"
            "Design details\n"
            "- Classroom rug artwork shows Mrs. Smith's Classroom Voice Levels with 4 outside, 3 groups, 2 table talk, 1 whisper and 0 no voice.\n"
            "- Product images include classroom mockups, size reference graphics and close-up design panels.\n"
            "- SEO copy focuses on visible classroom rules artwork without unsupported material, backing or safety claims."
        ),
    },
    "personalized-teachers-classroom-rules-area-rug-3c846a8f41": {
        "title": "Custom Classroom Jobs Rug with Rainbow Helper Cards",
        "meta": "Customize a classroom jobs rug with teacher name text, rainbow header, helper cards, school icons and checker border.",
        "description": (
            "Customize a classroom jobs rug with teacher name text, rainbow header, helper cards, school icons and checker border.\n\n"
            "Design details\n"
            "- Classroom jobs rug artwork shows Mrs. Smith's Classroom Jobs with helper cards, rainbow header and colorful icons.\n"
            "- Product images show classroom-style scenes, size reference graphics and close-up design panels.\n"
            "- SEO copy stays tied to visible classroom jobs artwork without unsupported material, backing or safety claims."
        ),
    },
    "personalized-teachers-classroom-rules-area-rug-a85c01cff2-a85c01cff2": {
        "title": "Personalized Coping Tools Classroom Rug with Monsters",
        "meta": "Personalize a coping tools classroom rug with teacher name text, colorful monster characters and calm-corner prompts.",
        "description": (
            "Personalize a coping tools classroom rug with teacher name text, colorful monster characters and calm-corner prompts.\n\n"
            "Design details\n"
            "- Pastel classroom rug artwork shows monster characters with prompts such as breathe, move, focus, squeeze, break and talk.\n"
            "- Gallery images include classroom mockups, size reference panels and close-up views of the prompt layout.\n"
            "- SEO copy focuses on visible classroom prompt artwork without unsupported material, therapeutic, backing or safety claims."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "non-slip", "non-skid",
    "kid friendly", "pet friendly", "safe durable", "machine-washable",
    "machine washable", "washable ", "quick-dry", "memory foam", "microfiber",
    "velvet", "stain", "fade resistant", "easy clean",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title = item["title"]
    meta = item["meta"]
    blob = " ".join([item["title"], item["meta"], item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R011 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R011 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R011"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R011", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R011",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R012 or Re-QA revision R011",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R011 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R011.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([
            str(title or ""),
            str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
        ]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70:
            raise RuntimeError(f"{handle} title len failed after save")
        if len(meta) > 320:
            raise RuntimeError(f"{handle} meta len failed after save")
        if str(row[out_cols["revision"] - 1]) != "2":
            raise RuntimeError(f"{handle} revision not 2")
        if row[out_cols["review_status"] - 1] != "NEEDS_REVIEW":
            raise RuntimeError(f"{handle} review_status incorrect")
        if row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} content_qa_status incorrect")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles = {}
    dup_meta_titles = []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt = row[out_cols["meta_title_seo"] - 1]
        handle = row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle

    result = {
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r011_log_rows": counts.get("R011", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

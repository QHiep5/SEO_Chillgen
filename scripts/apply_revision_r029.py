from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R028/SEO_Product_Optimization_revision_R028.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R029"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R029.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R029"
MANIFEST = RUN_DIR / "revision_R029_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R029_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-reading-tree-classroom-library-rug-983a65fb2e-983a65fb2e": {
        "title": "Reading Tree Classroom Rug with Book Quote Cards",
        "meta": "Decorate with a reading tree classroom rug featuring leafy tree artwork, book quote cards and Open Your Book text.",
        "description": "Decorate with a reading tree classroom rug featuring leafy tree artwork, book quote cards and Open Your Book text.\n\nDesign details\n- Cream and green rectangular reading rug artwork shows a large tree, books on branches, quote cards and book stacks.\n- Visible text includes Open Your Book Open Your World, with classroom scenes, group images and size reference graphics.\n- SEO copy stays tied to visible reading-tree artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-tree-of-knowledge-classroom-library-rug-5ea95ac966-5ea95ac966": {
        "title": "Tree Of Knowledge Classroom Rug with Moon Stars",
        "meta": "Decorate with a tree of knowledge classroom rug featuring brown tree art, books, moon, stars and grassy background.",
        "description": "Decorate with a tree of knowledge classroom rug featuring brown tree art, books, moon, stars and grassy background.\n\nDesign details\n- Green rectangular classroom rug artwork shows a large brown tree with many books, moon and star details on a night-sky scene.\n- Product images include classroom and playroom mockups, close-up artwork panels and size reference graphics.\n- SEO copy focuses on visible tree-and-book artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-good-day-read-book-classroom-rug-dc9ddbc5f6-dc9ddbc5f6": {
        "title": "New Chapter Classroom Rug with Bookshelf Border",
        "meta": "Decorate with a New Chapter classroom rug featuring open book artwork, colorful bookshelf border, stars and quote text.",
        "description": "Decorate with a New Chapter classroom rug featuring open book artwork, colorful bookshelf border, stars and quote text.\n\nDesign details\n- Black rectangular reading rug artwork shows Welcome To A New Chapter quote, open book, colorful books and pink scallop border.\n- Visible details include stars, classroom mockups, room scenes, close-up panels and size reference graphics.\n- SEO copy stays tied to visible new-chapter artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-classroom-rug-for-kids-e5ed693d3f": {
        "title": "Mindset Matters Classroom Rug with Rainbow Brain",
        "meta": "Personalize a Mindset Matters classroom rug with educator name, rainbow brain artwork, pencil border and prompt bubbles.",
        "description": "Personalize a Mindset Matters classroom rug with educator name, rainbow brain artwork, pencil border and prompt bubbles.\n\nDesign details\n- Black classroom rug artwork shows Mindset Matters quote, rainbow brain art, colorful pencil border and Mrs. Owens class text.\n- Visible details include mindset prompt bubbles, classroom scenes, close-up panels and size reference graphics.\n- SEO copy focuses on visible mindset artwork without unsupported classroom-result, material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-classroom-rug-for-kids-a44b91f2e4": {
        "title": "Change Your Words Classroom Rug with Phrase Cards",
        "meta": "Personalize a Change Your Words classroom rug with educator name, phrase cards, colorful border and class text.",
        "description": "Personalize a Change Your Words classroom rug with educator name, phrase cards, colorful border and class text.\n\nDesign details\n- Black classroom rug artwork shows Change Your Words Change Your Mindset text with custom Mr. James name.\n- Visible details include colorful phrase cards, classroom scenes, close-up panels and size reference graphics.\n- SEO copy stays tied to visible phrase-card artwork without unsupported classroom-result, material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-handprints-kids-classroom-area-rug-9f8e77de21": {
        "title": "You Can Change The World Rug with Pencil Words",
        "meta": "Personalize a You Can Change The World classroom rug with class name, large pencil artwork and colorful word labels.",
        "description": "Personalize a You Can Change The World classroom rug with class name, large pencil artwork and colorful word labels.\n\nDesign details\n- Cream classroom rug artwork shows Mr. William class text, large pencil graphic, You Are center text and word labels.\n- Visible words include dreamer, creative, leader, brave, kind, confident and special, with school icons and classroom scenes.\n- SEO copy focuses on visible pencil-word artwork without unsupported classroom-result, material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-classroom-rug-handprints-quotes-6a4adbacc3": {
        "title": "Believe Unstoppable Classroom Rug with Handprints",
        "meta": "Personalize a Believe Unstoppable classroom rug with class name, rainbow handprints, star border and quote text.",
        "description": "Personalize a Believe Unstoppable classroom rug with class name, rainbow handprints, star border and quote text.\n\nDesign details\n- Black classroom rug artwork shows Believe In Yourself And You Will Be Unstoppable quote with rainbow handprint letters.\n- Visible details include colorful stars, striped border, Mrs. Asher's class text, classroom scenes and size reference graphics.\n- SEO copy stays tied to visible handprint quote artwork without unsupported classroom-result, material, cleaning, reverse-side or surface-performance claims.",
    },
    "colorful-handprints-classroom-kids-rug-358d2f0393": {
        "title": "Unique Masterpiece Classroom Rug with Flower Tree",
        "meta": "Decorate with a Unique Masterpiece classroom rug featuring handprint flower tree, butterflies, green hill and quote text.",
        "description": "Decorate with a Unique Masterpiece classroom rug featuring handprint flower tree, butterflies, green hill and quote text.\n\nDesign details\n- Black classroom rug artwork shows a handprint flower tree, butterflies and We Are Each Unique And Beautiful quote.\n- Visible details include green hill, colorful handprints, classroom scenes, reading lifestyle image and size reference graphics.\n- SEO copy focuses on visible masterpiece quote artwork without unsupported classroom-result, material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-orthodox-christian-cross-byzantine-eagle-rug-9a8fdd4fcf-9a8fdd4fcf": {
        "title": "Orthodox Three Bar Cross Rug with Corner Ornaments",
        "meta": "Decorate with an Orthodox three bar cross rug featuring navy field, gold cross, minimal border and corner ornaments.",
        "description": "Decorate with an Orthodox three bar cross rug featuring navy field, gold cross, minimal border and corner ornaments.\n\nDesign details\n- Navy blue rectangular Orthodox rug artwork shows a gold three-bar cross with slim border and ornate corner details.\n- Product images include living room mockups, close-up design panels and size reference graphics.\n- SEO copy stays tied to visible cross artwork without unsupported material, cleaning, reverse-side, devotional-use or surface-performance claims.",
    },
    "personalized-orthodox-christian-three-bar-cross-rug-c6332608db": {
        "title": "Ornate Orthodox Cross Rug with Navy Gold Frame",
        "meta": "Decorate with an ornate Orthodox cross rug featuring navy and gold frame artwork, three bar cross and scrollwork border.",
        "description": "Decorate with an ornate Orthodox cross rug featuring navy and gold frame artwork, three bar cross and scrollwork border.\n\nDesign details\n- Navy and gold rectangular Orthodox rug artwork shows a three-bar cross inside an ornate framed medallion.\n- Visible details include scrollwork border, living room mockups, close-up design panels and size reference graphics.\n- SEO copy focuses on visible Orthodox cross artwork without unsupported material, cleaning, reverse-side, devotional-use or surface-performance claims.",
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "kid friendly", "pet friendly", "safe durable",
    "machine-washable", "machine washable", "washable", "quick-dry", "quick dry",
    "memory foam", "microfiber", "velvet", "stain", "fade resistant", "easy clean",
    "easy-clean", "waterproof", "absorbent", "absorption", "backing", "rubber",
    "layered construction", "hd printing", "ultra-soft", "soft ", "cushion",
    "support classroom", "support emotional", "teach", "learning-space", "thickened",
    "reinforced", "bound edge", "material panel", "pray", "prayer", "sacred",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    blob = " ".join([title, meta, item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R029 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R029 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append([
            "revision_batch_id", "generated_at", "handle", "source_row",
            "old_revision", "new_revision", "changed_fields", "qa_issue_fields",
            "recheck_condition",
        ])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R029"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R029", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R029",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R030 or Re-QA revision R029",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R029 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R029.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        desc = row[out_cols["description_proposed"] - 1]
        blob = " ".join([str(title or ""), str(meta or ""), str(desc or "")]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r029_log_rows": counts.get("R029", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

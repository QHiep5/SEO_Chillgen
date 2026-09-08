from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R021/SEO_Product_Optimization_revision_R021.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R022"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R022.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R022"
MANIFEST = RUN_DIR / "revision_R022_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R022_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-orthodox-christian-area-rug-custom-eastern-orthodox": {
        "title": "Blue Orthodox Cross Rug with Gold Scroll Border",
        "meta": "Personalize a blue Orthodox cross rug with gold scroll border, navy background and Eastern cross artwork.",
        "description": (
            "Personalize a blue Orthodox cross rug with gold scroll border, navy background and Eastern cross artwork.\n\n"
            "Design details\n"
            "- Rectangular navy blue rug artwork shows an ornate gold Orthodox cross centered inside a decorative scroll border.\n"
            "- Product images include living room mockups, close-up border views and size reference graphics.\n"
            "- SEO copy stays tied to visible cross and border artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-100": {
        "title": "Red Gold Orthodox Cross Rug with Floral Border",
        "meta": "Personalize a red gold Orthodox cross rug with ornate cream border, floral accents and Eastern cross artwork.",
        "description": (
            "Personalize a red gold Orthodox cross rug with ornate cream border, floral accents and Eastern cross artwork.\n\n"
            "Design details\n"
            "- Rectangular red rug artwork shows a gold Orthodox cross surrounded by cream and gold ornamental border details.\n"
            "- Gallery images show room mockups, close-up edge views and size reference graphics.\n"
            "- SEO copy focuses on visible cross and floral border artwork without material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-101": {
        "title": "Orthodox Cross Book Style Rug with Red Gold Art",
        "meta": "Personalize an Orthodox cross book style rug with red gold panel artwork, teal border and ornate cross motif.",
        "description": (
            "Personalize an Orthodox cross book style rug with red gold panel artwork, teal border and ornate cross motif.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows a red Orthodox cross centered on a gold book-like panel with teal and red border.\n"
            "- Product images include room mockups, close-up design panels and size reference graphics.\n"
            "- SEO copy stays tied to visible Orthodox cross artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-102": {
        "title": "Black Gold Orthodox Cross Rug with Ornate Border",
        "meta": "Personalize a black gold Orthodox cross rug with centered cross artwork, gold ornamental border and room mockups.",
        "description": (
            "Personalize a black gold Orthodox cross rug with centered cross artwork, gold ornamental border and room mockups.\n\n"
            "Design details\n"
            "- Rectangular black rug artwork shows a gold Orthodox cross with gold ornamental border and corner detailing.\n"
            "- Gallery images include sofa, bright room and close-up views plus size reference graphics.\n"
            "- SEO copy focuses on visible black gold cross artwork without material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-103": {
        "title": "Burgundy Orthodox Cross Rug with Silver Motif",
        "meta": "Personalize a burgundy Orthodox cross rug with silver cross motif, yellow white floral border and room mockups.",
        "description": (
            "Personalize a burgundy Orthodox cross rug with silver cross motif, yellow white floral border and room mockups.\n\n"
            "Design details\n"
            "- Rectangular burgundy red rug artwork shows a silver Orthodox cross with yellow and white floral border details.\n"
            "- Product images include sofa, living room and close-up views plus size reference graphics.\n"
            "- SEO copy stays tied to visible burgundy cross artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-105": {
        "title": "Navy Gold Orthodox Cross Rug with Floral Corners",
        "meta": "Personalize a navy gold Orthodox cross rug with gold outline cross, ornate border and floral corner details.",
        "description": (
            "Personalize a navy gold Orthodox cross rug with gold outline cross, ornate border and floral corner details.\n\n"
            "Design details\n"
            "- Rectangular navy rug artwork shows a gold Orthodox cross outline with gold border and floral corner ornaments.\n"
            "- Gallery images include room mockups, close-up panels and size reference graphics.\n"
            "- SEO copy focuses on visible navy gold cross artwork without material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-composition-notebook-classroom-shaped-rugs-design-2222": {
        "title": "Pink Teacher Name Classroom Rug with Pencil Art",
        "meta": "Personalize a pink teacher name classroom rug with pencil outline, flowers, book artwork and apple accents.",
        "description": (
            "Personalize a pink teacher name classroom rug with pencil outline, flowers, book artwork and apple accents.\n\n"
            "Design details\n"
            "- Pink shaped classroom rug artwork shows teacher-name personalization with pencil outline, flowers, book and apple details.\n"
            "- Product images include classroom circle-time mockups, size reference graphics and close-up design views.\n"
            "- SEO copy describes visible classroom artwork only, without material, cleaning, reverse-side or safety claims."
        ),
    },
    "personalized-composition-notebook-classroom-shaped-rugs-design-2223": {
        "title": "Personalized Apple Classroom Rug with Teacher Text",
        "meta": "Customize an apple classroom rug with teacher welcome text, crayons, ruler, flower accents and colorful dots.",
        "description": (
            "Customize an apple classroom rug with teacher welcome text, crayons, ruler, flower accents and colorful dots.\n\n"
            "Design details\n"
            "- Red apple shaped classroom rug artwork reads Welcome to Teacher's Classroom with bright school supply graphics.\n"
            "- Visible details include crayons, ruler, flower, colorful dots and classroom mockup scenes.\n"
            "- SEO copy focuses on visible apple classroom artwork without unsupported material, cleaning, reverse-side or safety claims."
        ),
    },
    "personalized-composition-notebook-classroom-shaped-rugs-design-2225": {
        "title": "Personalized Flower Classroom Rug with Teacher Name",
        "meta": "Customize a flower classroom rug with teacher name center, pastel striped petals, yellow bow and school accents.",
        "description": (
            "Customize a flower classroom rug with teacher name center, pastel striped petals, yellow bow and school accents.\n\n"
            "Design details\n"
            "- Round flower shaped classroom rug artwork shows pastel striped petals and personalized teacher name in the center.\n"
            "- Visible details include yellow bow, apple, pencils and star accents in classroom activity scenes.\n"
            "- SEO copy stays tied to visible flower classroom artwork without material, cleaning, reverse-side or safety claims."
        ),
    },
    "personalized-family-couple-doormat-custom-couple-husband-design-03": {
        "title": "Home Sweet Home Couple Doormat with Cartoon Portraits",
        "meta": "Personalize a Home Sweet Home couple doormat with cartoon portraits, custom names and music-player style artwork.",
        "description": (
            "Personalize a Home Sweet Home couple doormat with cartoon portraits, custom names and music-player style artwork.\n\n"
            "Design details\n"
            "- White rectangular doormat artwork shows a cartoon couple sitting together with custom names and Home Sweet Home text.\n"
            "- Visible details include music-player style graphics, entryway mockups, couple photos and size reference panels.\n"
            "- SEO copy focuses on visible couple portrait artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip",
    "non slip", "non-skid", "non skid", "kid friendly", "pet friendly",
    "safe durable", "machine-washable", "machine washable", "washable ",
    "quick-dry", "quick dry", "memory foam", "microfiber", "velvet",
    "stain", "fade resistant", "easy clean", "soft fabric", "backing",
    "cleaning visuals", "cleaning graphic", "thickened", "bound edge",
    "reinforced", "absorbent", "absorption", "water absorption",
    "support ", "supports ", "use a ", "foldable material", "prayer room", "prayer space", "sacred",
    "absorbent", "water absorption", "soft non-slip", "layered construction",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title = item["title"]
    meta = item["meta"]
    blob = " ".join([item["title"], item["meta"], item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R022 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R022 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R022"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R022", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R022",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R023 or Re-QA revision R022",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R022 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R022.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([
            str(title or ""),
            str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
        ]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70:
            raise RuntimeError(f"{handle} title len failed after save")
        if len(meta) > 320:
            raise RuntimeError(f"{handle} meta len failed after save")
        if str(row[out_cols["revision"] - 1]) != "2":
            raise RuntimeError(f"{handle} revision not 2")
        if row[out_cols["review_status"] - 1] != "NEEDS_REVIEW":
            raise RuntimeError(f"{handle} review_status incorrect")
        if row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} content_qa_status incorrect")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles = {}
    dup_meta_titles = []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt = row[out_cols["meta_title_seo"] - 1]
        handle = row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle

    result = {
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r022_log_rows": counts.get("R022", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

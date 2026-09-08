from __future__ import annotations

import json
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R002/SEO_Product_Optimization_revision_R002.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R003"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R003.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R003"
MANIFEST = RUN_DIR / "revision_R003_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R003_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-103": {
        "title": "Days of the Week Classroom Rug with Colorful Activity Icons",
        "meta": "Use a days of the week classroom rug with Monday through Sunday labels, colorful activity icons and round group-time styling.",
        "description": (
            "Use a days of the week classroom rug with Monday through Sunday labels, colorful activity icons and round group-time styling.\n\n"
            "Design details\n"
            "- Round educational rug with Days of the Week center text and bright panels labeled Monday through Sunday.\n"
            "- Visible icons include school, books, sun, art palette, soccer ball, picnic basket and teddy bear artwork.\n"
            "- SEO copy stays tied to the visible classroom learning design and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-105": {
        "title": "Weather and Seasons Classroom Rug with Round Learning Wheel",
        "meta": "Use a weather and seasons classroom rug with sunny, rainy, windy and snowy panels plus a spring to winter center wheel.",
        "description": (
            "Use a weather and seasons classroom rug with sunny, rainy, windy and snowy panels plus a spring to winter center wheel.\n\n"
            "Design details\n"
            "- Round educational rug with outer weather labels and a center seasonal wheel for spring, summer, autumn and winter.\n"
            "- The artwork uses landscape icons, classroom scenes and bright colors for weather and seasons learning displays.\n"
            "- SEO copy focuses on the visible weather and seasons design without unsupported material, backing or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-106": {
        "title": "Round Learn Colors Classroom Rug with Crayon Characters",
        "meta": "Use a round learn colors classroom rug with labeled color slices, crayon character artwork and bright group-time visuals.",
        "description": (
            "Use a round learn colors classroom rug with labeled color slices, crayon character artwork and bright group-time visuals.\n\n"
            "Design details\n"
            "- Round educational rug with Learn Colors center text and labeled slices for black, white, red, orange, yellow, green, blue, purple, pink and brown.\n"
            "- Gallery visuals show classroom scenes, child reading scenes and close-up panels of the colorful learning design.\n"
            "- SEO copy stays focused on the visible color-learning artwork and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-107": {
        "title": "Animal Feelings Classroom Rug with Emotion Faces",
        "meta": "Use an animal feelings classroom rug with emotion labels, character faces and a round feelings and emotions design.",
        "description": (
            "Use an animal feelings classroom rug with emotion labels, character faces and a round feelings and emotions design.\n\n"
            "Design details\n"
            "- Round classroom rug with Feelings and Emotions center text and animal character slices for scared, happy, sad, angry, excited and more.\n"
            "- Visible animal artwork includes bunny, lion, elephant, tiger, monkey, giraffe, panda and bear-style character faces.\n"
            "- SEO copy describes the visible animal feelings wheel without unsupported material, backing or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-108": {
        "title": "Months and Seasons Classroom Rug with Tree Calendar Wheel",
        "meta": "Use a months and seasons classroom rug with January through December labels, seasonal icons and a four-season tree center.",
        "description": (
            "Use a months and seasons classroom rug with January through December labels, seasonal icons and a four-season tree center.\n\n"
            "Design details\n"
            "- Round educational rug with a center tree divided into spring, summer, autumn and winter sections.\n"
            "- The outer ring shows January through December labels with seasonal icons such as flowers, sun, umbrella, pumpkin and snowman artwork.\n"
            "- SEO copy stays tied to the visible calendar learning design and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-111": {
        "title": "Today I'm Feeling Classroom Rug with Rainbow Emotion Wheel",
        "meta": "Use a Today I'm Feeling classroom rug with rainbow emotion slices, cartoon faces and a round group-time design.",
        "description": (
            "Use a Today I'm Feeling classroom rug with rainbow emotion slices, cartoon faces and a round group-time design.\n\n"
            "Design details\n"
            "- Round rainbow classroom rug with Today I'm Feeling center text and colored emotion slices.\n"
            "- Visible labels include happy, shy, scared, calm, tired, worried, angry and sad with cartoon feeling faces.\n"
            "- SEO copy focuses on the visible emotional check-in artwork and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-110": {
        "title": "Bat Vortex Halloween Round Rug with 3D Illusion Artwork",
        "meta": "Decorate with a bat vortex Halloween round rug featuring orange, purple and blue spiral artwork with flying bat silhouettes.",
        "description": (
            "Decorate with a bat vortex Halloween round rug featuring orange, purple and blue spiral artwork with flying bat silhouettes.\n\n"
            "Design details\n"
            "- Round Halloween rug with a swirling orange, purple and blue vortex pattern and black bat silhouettes.\n"
            "- Gallery images show the optical tunnel artwork in seasonal room, dining and fireplace-style scenes.\n"
            "- SEO copy stays focused on the visible spooky 3D illusion motif without unsupported material, backing or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-0fc607e644-0fc607e644": {
        "title": "I Am Feelings Wheel Classroom Rug with Affirmation Words",
        "meta": "Use an I Am feelings wheel classroom rug with affirmation words such as loved, kind, brave, confident and unique.",
        "description": (
            "Use an I Am feelings wheel classroom rug with affirmation words such as loved, kind, brave, confident and unique.\n\n"
            "Design details\n"
            "- Round colorful feelings rug with I Am center text and affirmation segments including loved, kind, brave, powerful, confident, unique and amazing.\n"
            "- Gallery visuals show classroom-style scenes, size visuals, child/baby use mockups and close-up design panels.\n"
            "- SEO copy describes the visible affirmation wheel design without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-a-good-day-to-read-book-rug-classroom-library-rugs-04ab7bfcca": {
        "title": "New Chapter Reading Rug for Classroom Library Decor",
        "meta": "Welcome students with a New Chapter reading rug featuring open book artwork, colorful books, stars and classroom library scenes.",
        "description": (
            "Welcome students with a New Chapter reading rug featuring open book artwork, colorful books, stars and classroom library scenes.\n\n"
            "Design details\n"
            "- Rectangular classroom library rug with Welcome to a New Chapter text, open book artwork, colorful books, stars, notes and pink border.\n"
            "- Gallery images show classroom, book nook and reading-area mockups with the black library-themed design.\n"
            "- SEO copy stays tied to the visible reading rug artwork and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "feelings-wheel-round-classroom-rug-4be50a20fb": {
        "title": "How Are You Feeling Classroom Rug with Pastel Emotion Faces",
        "meta": "Create a check-in corner with a How Are You Feeling classroom rug featuring pastel emotion faces and calm activity prompts.",
        "description": (
            "Create a check-in corner with a How Are You Feeling classroom rug featuring pastel emotion faces and calm activity prompts.\n\n"
            "Design details\n"
            "- Round pale feelings rug with How Are You Feeling center text, pastel emotion faces and gentle activity prompts.\n"
            "- Visible prompts include reading, playing with friends, listening to music and writing in a journal, shown with classroom-style mockups.\n"
            "- SEO copy focuses on the visible emotional check-in design without unsupported material, backing or surface-performance claims."
        ),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)

    required_cols = [
        "Handle",
        "title_proposed",
        "meta_title_seo",
        "meta_title_chars",
        "meta_description_seo",
        "meta_description_chars",
        "description_proposed",
        "description_proposed_html",
        "revision",
        "review_status",
        "review_reason",
        "processing_status",
        "content_qa_status",
        "issues",
    ]
    missing = [c for c in required_cols if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title = item["title"]
        meta = item["meta"]
        desc = item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R003 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R003 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    # Sanitize scoped internal research/mapping text for the same unsupported wording pattern.
    for sheet_name in ["Buyer_Search_Research", "Keyword_Map"]:
        if sheet_name not in wb.sheetnames:
            continue
        sh = wb[sheet_name]
        hmap = headers(sh)
        key_col = hmap.get("product_key")
        if not key_col:
            continue
        for row_idx in range(2, sh.max_row + 1):
            pk = str(sh.cell(row_idx, key_col).value or "")
            if not any(handle in pk for handle in REVISIONS):
                continue
            for col_idx in range(1, sh.max_column + 1):
                value = sh.cell(row_idx, col_idx).value
                if not isinstance(value, str):
                    continue
                new = (
                    value.replace("indoor/outdoor", "indoor")
                    .replace("Indoor/outdoor", "Indoor")
                    .replace("outdoor-use", "room-use")
                    .replace("outdoor use", "room use")
                )
                sh.cell(row_idx, col_idx, new)

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = []
        for r in log.iter_rows(min_row=2, values_only=True):
            if r and r[0] != "R003":
                keep.append(r)
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R003",
            NOW,
            handle,
            row_idx,
            old_revision,
            new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R003",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R004 or Re-QA revision R003",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R003 Summary",
        "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R003.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.",
        "",
        "## Updated products",
        "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _, _, _, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n", encoding="utf-8")

    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    vws = verify["SEO_Products"]
    vcols = {cell.value: idx for idx, cell in enumerate(next(vws.iter_rows(min_row=1, max_row=1)))}
    bad = []
    for row in vws.iter_rows(min_row=2, values_only=True):
        handle = row[vcols["Handle"]]
        if handle not in REVISIONS:
            continue
        title = row[vcols["title_proposed"]] or ""
        meta_title = row[vcols["meta_title_seo"]] or ""
        meta_desc = row[vcols["meta_description_seo"]] or ""
        desc = (row[vcols["description_proposed"]] or "") + " " + (row[vcols["description_proposed_html"]] or "")
        blob = " ".join([title, meta_title, meta_desc, desc]).lower()
        if row[vcols["revision"]] != "2":
            bad.append((handle, "revision", row[vcols["revision"]]))
        if row[vcols["review_status"]] != "NEEDS_REVIEW":
            bad.append((handle, "review_status", row[vcols["review_status"]]))
        if row[vcols["content_qa_status"]] != "NOT_RUN":
            bad.append((handle, "content_qa_status", row[vcols["content_qa_status"]]))
        if not (45 <= len(title) <= 70):
            bad.append((handle, "title_len", len(title)))
        if meta_title != title:
            bad.append((handle, "meta_title_mismatch", meta_title))
        if len(meta_desc) > 320:
            bad.append((handle, "meta_desc_len", len(meta_desc)))
        for forbidden in ["indoor/outdoor", " outdoor ", "anti-slip", "non-slip", "kid friendly", "pet friendly", "safe durable", "machine-washable", "machine washable"]:
            if forbidden in f" {blob} ":
                bad.append((handle, "forbidden_claim", forbidden))
    if bad:
        raise RuntimeError(f"Verification failed: {bad[:20]}")

    print(json.dumps({
        "output_workbook": str(OUTPUT),
        "manifest": str(MANIFEST),
        "summary": str(SUMMARY),
        "updated": len(changed),
        "status": "COMPLETE_AWAITING_REQA",
    }, indent=2))


if __name__ == "__main__":
    main()

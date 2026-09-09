from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B017"
REVISION = "R093"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "personalized-orthodox-christian-area-rug-custom-eastern-design-106": ("Cream Red Orthodox Cross Rug", "cream red Orthodox cross rug", "cream Christian rug, Orthodox cross area rug, religious home decor", "a cream and red Orthodox cross, ornate border details, and a warm religious design for home decor"),
    "personalized-orthodox-christian-area-rug-custom-eastern-design-107": ("Green Cream Orthodox Cross Rug", "green cream Orthodox cross rug", "green Christian rug, Orthodox prayer room rug, religious area rug", "a green and cream Orthodox cross, decorative border details, and a calm religious design for a prayer room or living space"),
    "personalized-orthodox-christian-area-rug-custom-eastern-design-108": ("Black Gold Orthodox Cross Rug", "black gold Orthodox cross rug", "black Christian rug, Orthodox cross area rug, religious room decor", "a black and gold Orthodox cross, ornate border details, and a refined religious design for home decor"),
    "personalized-composition-notebook-classroom-shaped-rugs-for-kids-a01": ("Personalized Notebook Classroom Rug", "personalized notebook classroom rug", "custom teacher rug, classroom notebook rug, personalized kids room rug", "a notebook-inspired classroom design with space for a teacher name and a friendly layout for learning spaces"),
    "personalized-composition-notebook-classroom-shaped-rugs-design-2222": ("Pink Personalized Teacher Rug", "pink personalized teacher rug", "pink classroom rug, custom teacher name rug, personalized kids rug", "a pink classroom design with personalized teacher-name space and a welcoming layout for a classroom or playroom"),
    "personalized-composition-notebook-classroom-shaped-rugs-design-2223": ("Personalized Apple Classroom Rug", "personalized apple classroom rug", "apple teacher rug, custom classroom rug, personalized school rug", "an apple classroom design with personalized teacher-name space and a cheerful layout for a learning area"),
    "personalized-composition-notebook-classroom-shaped-rugs-design-2224": ("Personalized Open Book Classroom Rug", "personalized open book classroom rug", "open book classroom rug, custom teacher rug, personalized reading corner rug", "an open-book classroom design with personalized teacher-name space for a reading corner, classroom, or playroom"),
    "personalized-composition-notebook-classroom-shaped-rugs-design-2225": ("Personalized Flower Classroom Rug", "personalized flower classroom rug", "flower teacher rug, custom classroom rug, personalized kids room rug", "a flower classroom design with personalized teacher-name space and a cheerful layout for a learning space"),
    "personalized-family-couple-doormat-custom-couple-husband-design-05": ("Funny Personalized Couple Doormat", "funny personalized couple doormat", "custom couple doormat, personalized husband mat, funny entryway mat", "a playful personalized couple message for an entryway, with a lighthearted design made for a shared home"),
    "personalized-family-couple-doormat-custom-couple-husband-design-03": ("Home Sweet Home Couple Doormat", "home sweet home couple doormat", "personalized couple mat, home sweet home doormat, custom entryway mat", "a Home Sweet Home message with personalized couple details for a welcoming front door or shared home entryway"),
}

def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []

def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)

def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B017 scope does not match the R093 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        meta = f"Create a {primary} with {motif}."
        description = f"Create a {primary} with {motif}. The design adds a welcoming focal point to a classroom, playroom, entryway, living space, or shared home."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))

if __name__ == "__main__":
    main()

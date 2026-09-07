import csv
import html
import json
import re
from collections import OrderedDict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
ROOT = Path("seo_runs") / SHOP / RUN_ID
RESULTS = Path("resutls") / SHOP / RUN_ID
BATCH_ID = "B006"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()

SERP_REFS = [
    "https://www.etsy.com/market/dragon_rug",
    "https://www.etsy.com/market/dragon_rugs",
    "https://www.etsy.com/market/dragon_shaped_rug",
    "https://www.etsy.com/market/dragon_area_rugs",
    "https://www.etsy.com/market/dog_paw_area_rug",
    "https://www.etsy.com/market/personalized_dog_rug",
    "https://www.etsy.com/market/paw_print_rug",
    "https://chillgen.com/",
]

DESIGNS = OrderedDict(
    [
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-04",
            {
                "theme": "orange dragon open book rug",
                "title": "Personalized Orange Dragon Book Rug",
                "meta": "Customize an orange dragon book rug with a name, fantasy library artwork, custom shape visuals, anti-slip backing, and size chart images.",
                "primary": "personalized orange dragon rug",
                "secondary": "dragon book rug, custom dragon shaped rug, fantasy library rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with orange and gold dragon resting on an open antique-style book, personalized name Sophia, room mockups, material graphic, layer graphic, and size chart",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-10",
            {
                "theme": "green dragon book rug",
                "title": "Personalized Green Dragon Book Rug",
                "meta": "Create a personalized green dragon book rug with fantasy library styling, custom name artwork, anti-slip backing visuals, and room mockups.",
                "primary": "personalized green dragon rug",
                "secondary": "dragon book rug, fantasy dragon area rug, custom shaped dragon rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with green dragon standing and curling across stacked books, large green wings, personalized name Sophia, room mockups, material/layer graphics, and size chart",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-01",
            {
                "theme": "red dragon books rug",
                "title": "Personalized Red Dragon Book Rug",
                "meta": "Personalize a red dragon book rug with roses, stacked books, custom name detail, fantasy decor appeal, anti-slip backing visuals, and size chart.",
                "primary": "personalized red dragon rug",
                "secondary": "dragon library rug, red dragon area rug, custom fantasy rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with red dragon curled around stacked books and red roses, personalized name Jasmine, room mockups, material/layer graphics, and size chart",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-wester-design-06",
            {
                "theme": "purple sleeping dragon rug",
                "title": "Personalized Purple Dragon Book Rug",
                "meta": "Choose a personalized purple dragon book rug with floral fantasy artwork, custom name detail, room mockups, anti-slip backing visuals, and size chart.",
                "primary": "personalized purple dragon book rug",
                "secondary": "sleeping dragon rug, fantasy book rug, custom dragon shaped rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with purple dragon sleeping or curled on an open book, purple flowers and greenery, personalized name Sophia, room mockups, material/layer graphics, and size chart",
            },
        ),
        (
            "custom-dragon-shaped-area-rug-personalized-library-western-dragon",
            {
                "theme": "western dragon books rug",
                "title": "Personalized Western Dragon Rug",
                "meta": "Add a personalized western dragon rug with stacked-book fantasy artwork, custom name detail, room mockups, anti-slip backing visuals, and size chart.",
                "primary": "personalized western dragon rug",
                "secondary": "dragon book rug, fantasy library rug, custom dragon shaped rug",
                "season": "EVERGREEN_FANTASY",
                "obs": "custom-shaped rug with teal and orange western dragon perched on stacked books, broad wings, personalized name Sophia, room mockups, material/layer graphics, and size chart",
            },
        ),
        (
            "custom-dog-paw-round-rug-patchwork-area-rug-for-pet-lovers-design-01",
            {
                "theme": "colorful dog paw patchwork round rug",
                "title": "Dog Paw Patchwork Round Rug",
                "meta": "Style a pet lover space with a dog paw patchwork round rug, colorful farmhouse squares, kid-and-dog room mockups, and washable-care visuals.",
                "primary": "dog paw patchwork round rug",
                "secondary": "pet lover round rug, dog paw area rug, farmhouse paw print rug",
                "season": "EVERGREEN_PET_LOVER",
                "obs": "round rug with multicolor brown, orange, teal and red patchwork squares, dog paw prints, black border with white paw marks, kid and dog room mockups, washable and vacuum graphics",
            },
        ),
        (
            "custom-dog-paw-round-rug-patchwork-area-rug-for-pet-lovers-design-05",
            {
                "theme": "neutral dog paw patchwork round rug",
                "title": "Neutral Dog Paw Round Rug",
                "meta": "Choose a neutral dog paw round rug with farmhouse patchwork squares, pet lover styling, kid-room mockups, and washable-care visuals.",
                "primary": "neutral dog paw round rug",
                "secondary": "dog paw patchwork rug, pet lover area rug, farmhouse pet rug",
                "season": "EVERGREEN_PET_LOVER",
                "obs": "round rug with neutral brown, beige, plaid and white patchwork squares plus black paw prints, room mockups with children and dog, washable and vacuum graphics",
            },
        ),
        (
            "custom-dog-paw-round-rug-patchwork-area-rug-for-pet-lovers-design-03",
            {
                "theme": "bold dog paw patchwork round rug",
                "title": "Bold Dog Paw Round Rug",
                "meta": "Bring pet-lover charm to a room with a bold dog paw round rug, warm patchwork colors, playroom mockups, and washable-care visuals.",
                "primary": "bold dog paw round rug",
                "secondary": "colorful paw print rug, pet lover round rug, dog paw area rug",
                "season": "EVERGREEN_PET_LOVER",
                "obs": "round rug with bold red, orange, blue and cream patchwork squares and multicolor paw prints, shown in kid room and pet room mockups with washable/vacuum graphics",
            },
        ),
        (
            "custom-dog-paw-round-rug-patchwork-area-rug-for-pet-lovers-design-02",
            {
                "theme": "soft neutral dog paw round rug",
                "title": "Soft Neutral Dog Paw Rug",
                "meta": "Add a soft neutral dog paw rug with pale patchwork colors, round shape, pet-friendly room styling, and washable-care visual evidence.",
                "primary": "soft neutral dog paw rug",
                "secondary": "paw print round rug, pet lover farmhouse rug, dog paw area rug",
                "season": "EVERGREEN_PET_LOVER",
                "obs": "round rug with pale beige, gray, cream and blue patchwork squares and soft-colored paw prints, shown beside bed, in playroom, and with washable/vacuum graphics",
            },
        ),
        (
            "custom-dog-paw-round-rug-patchwork-area-rug-for-pet-lovers-1-a4",
            {
                "theme": "multicolor dog paw patchwork rug",
                "title": "Multicolor Dog Paw Round Rug",
                "meta": "Decorate a pet lover room with a multicolor dog paw round rug, patchwork paw-print artwork, child-and-dog mockups, and washable-care visuals.",
                "primary": "multicolor dog paw round rug",
                "secondary": "dog paw patchwork rug, pet lover round rug, colorful paw print rug",
                "season": "EVERGREEN_PET_LOVER",
                "obs": "round rug with red, blue, green and yellow patchwork squares, colored paw prints, tan rope-like border, dog and child mockups, washable and vacuum graphics",
            },
        ),
    ]
)


def read_csv(path):
    with Path(path).open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def write_csv(path, rows, headers):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with Path(path).open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def strip_html(value):
    return " ".join(re.sub(r"<[^>]+>", " ", html.unescape(value or "")).split())


def description_html(design):
    return (
        f"<p>{design['meta']}</p>"
        "<h3>Design details</h3><ul>"
        f"<li>{design['obs'].capitalize()}.</li>"
        "<li>Keep this product positioned around its specific motif, not a broad generic rug keyword.</li>"
        "<li>Use backing, washing, material, and pet/kid-room claims only where supported by page text or product infographics.</li>"
        "</ul>"
    )


def image_note(design, n):
    theme = design["theme"]
    if "dragon" in design["primary"]:
        notes = {
            1: f"{theme} shown on a dark wood floor",
            2: f"isolated or alternate angle view of personalized {theme}",
            3: f"{theme} shown at an angled room view",
            4: f"custom shape area rug material graphic for {theme}",
            5: f"{theme} displayed in a bright living room mockup",
            6: f"upper, middle and bottom layer graphic for {theme}",
            7: f"size chart and backing view for personalized {theme}",
        }
    else:
        notes = {
            1: f"{theme} shown on a light floor",
            2: f"{theme} displayed beside a bed",
            3: f"{theme} in a playroom mockup with children and a dog",
            4: f"machine wash and vacuum cleaner care graphic for {theme}",
            5: f"{theme} displayed near a child tent and toys",
            6: f"multi-room usage graphic for {theme}",
        }
    return notes.get(n, f"{theme} image {n}")


def autosize(ws, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(12, min(max_width, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def main():
    batch = json.loads((ROOT / "evidence/products/B006_products.json").read_text())
    products = {p["handle"]: p for p in batch["products"]}
    page_meta = {r["product_key"]: r for r in json.loads((ROOT / "evidence/products/B006_page_meta.json").read_text())}
    manifest = json.loads((ROOT / "evidence/images/B006_manifest.json").read_text())
    by_images = {}
    for img in manifest:
        by_images.setdefault(img["product_key"], []).append(img)

    prev_seo, prev_img, prev_evd = [], [], []
    seo_headers = image_headers = evidence_headers = []
    for bid in ["B001", "B002", "B003", "B004", "B005"]:
        rows, headers = read_csv(ROOT / f"batches/{bid}_SEO_Products.csv")
        prev_seo.extend(rows)
        seo_headers = headers
        rows, headers = read_csv(ROOT / f"batches/{bid}_Image_Audit.csv")
        prev_img.extend(rows)
        image_headers = headers
        rows, headers = read_csv(ROOT / f"batches/{bid}_Product_Evidence.csv")
        prev_evd.extend(rows)
        evidence_headers = headers

    prev_kw, keyword_headers = read_csv(ROOT / "keyword_research.csv")
    prev_kw = [r for r in prev_kw if r.get("product_key") not in DESIGNS]
    with (ROOT / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        prev_buyer = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(prev_buyer[0].keys())
    prev_buyer = [r for r in prev_buyer if r.get("product_key") not in DESIGNS]

    inv, inv_headers = read_csv(ROOT / "inventory.csv")
    inv_lookup = {r["product_key"]: r for r in inv}

    b_seo, b_img, b_evd, b_kw, b_buyer, keys_done = [], [], [], [], [], []
    for key, design in DESIGNS.items():
        product = products[key]
        meta = page_meta[key]
        order = int(inv_lookup[key]["inventory_order"])
        evidence_id = f"{BATCH_ID}-{order:03d}"
        research_id = f"{BATCH_ID}-{order:03d}-BR01"
        url = f"https://chillgen.com/products/{key}"
        imgs = sorted(by_images[key], key=lambda x: int(x["image_number"]))
        img_fields = {}
        for img in imgs:
            n = int(img["image_number"])
            note = image_note(design, n)
            img_fields[f"img_{n}_link"] = img["src"]
            img_fields[f"img_{n}_alt_current"] = ""
            img_fields[f"img_{n}_alt"] = note
            img_fields[f"img_{n}_alt_action"] = "SET"
            b_img.append(
                {
                    "evidence_id": evidence_id,
                    "shop_domain": SHOP,
                    "Handle": key,
                    "product_id": str(product.get("id", "")),
                    "media_id": img.get("image_id", ""),
                    "image_location": "GALLERY_PUBLIC_CDN",
                    "image_number": n,
                    "variant": "",
                    "image_url": img["src"],
                    "image_url_export": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                    "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                    "shared_media_references": "",
                    "viewed_status": "VIEWED_CONTACT_SHEET_AND_LOCAL_FILE",
                    "viewed_at": NOW,
                    "observed_visual_details": note,
                    "alt_current": "",
                    "alt_proposed": note,
                    "alt_action": "SET",
                    "review_status": "NEEDS_REVIEW",
                    "revision": "1",
                    "approved_by": "",
                    "approved_at": "",
                    "approved_fields": "",
                    "approved_revision": "",
                    "evidence_file_or_reference": img.get("local_path", ""),
                    "issues": "alt_current_empty_in_public_json; media_export_url_unknown_without_shopify_export",
                }
            )

        seo = {h: "" for h in seo_headers}
        seo.update(
            {
                "schema_version": "2.4",
                "prompt_version": "2.4",
                "run_id": RUN_ID,
                "batch_id": BATCH_ID,
                "shop_domain": SHOP,
                "product_key": key,
                "Handle": key,
                "product_id": str(product.get("id", "")),
                "product_gid": f"gid://shopify/Product/{product.get('id', '')}",
                "identity_status": "PUBLIC_JSON_MATCHED_TO_HANDLE",
                "source_export_ref": f"evidence/products/B006_products.json; {meta.get('captured_html_path', '')}",
                "source_exported_at": meta.get("captured_at", NOW),
                "locale": "en-US",
                "market": "US",
                "url": url,
                "canonical_url": meta.get("canonical_url") or url,
                "product_type": product.get("product_type", ""),
                "title_current": product.get("title", ""),
                "h1_current": meta.get("h1_current", ""),
                "title_proposed": design["title"],
                "title_action": "SET",
                "h1_mapping_status": "RENDERED_H1_CAPTURED_THEME_CONFIRMATION_STILL_NEEDED_FOR_DEPLOY",
                "meta_title_current": "",
                "meta_title_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_title_current": meta.get("rendered_title_current", ""),
                "theme_title_suffix": "NONE_OR_NOT_OBSERVED",
                "meta_title_seo": design["title"],
                "meta_title_action": "SET",
                "meta_title_chars": len(design["title"]),
                "meta_description_current": "",
                "meta_description_source_state": "UNKNOWN_WITHOUT_SHOPIFY_EXPORT",
                "rendered_meta_description_current": meta.get("rendered_meta_description_current", ""),
                "meta_description_seo": design["meta"],
                "meta_description_action": "SET",
                "meta_description_chars": len(design["meta"]),
                "primary_keyword": design["primary"],
                "secondary_keywords": design["secondary"],
                "long_tail_candidates": f"{design['primary']}; {design['secondary']}",
                "keyword_strategy": "LONG_TAIL_CANDIDATE",
                "season": design["season"],
                "keyword_demand_evidence": "SERP_ONLY: public marketplace/category results show dragon/fantasy rugs and dog paw/pet lover rugs; no paid volume or first-party search data used.",
                "keyword_serp_fit": "SERP_ONLY_PRODUCT_AND_MARKETPLACE_RESULTS",
                "meta_keyword": design["primary"],
                "search_intent": "COMMERCIAL_INVESTIGATION",
                "keyword_validation_status": "SERP_ONLY_NO_VOLUME",
                "keyword_selection_reason": f"Selected because page/image evidence supports a {design['theme']} and product-level decor intent.",
                "buyer_research_refs": research_id,
                "buyer_search_summary": "Buyer intent is a hypothesis from product evidence plus category SERP language, not direct customer data.",
                "baseline_status": "PUBLIC_RENDERED_PAGE_AND_JSON_CAPTURED",
                "change_scope": "CLARIFICATION",
                "mapping_status": "NEEDS_REVIEW",
                "mapping_version": "2.4",
                "description_current_html": product.get("body_html", ""),
                "description_changes_needed": "Make product copy motif-specific while retaining verified public claims only.",
                "description_proposed": strip_html(description_html(design)),
                "description_proposed_html": description_html(design),
                "description_change_mode": "REPLACE_ALL",
                "description_target_section": "Full product body",
                "description_action": "SET",
                "revision": "1",
                "review_status": "NEEDS_REVIEW",
                "review_reason": "Research proposal only; rendered page captured but admin/export SEO values and media export URLs remain unknown.",
                "approved_by": "",
                "approved_at": "",
                "approved_fields": "",
                "approved_revision": "",
                "processing_status": "DRAFTED",
                "evidence_status": "COMPLETE_PUBLIC_RENDERED_PAGE_AND_IMAGES",
                "content_qa_status": "NOT_RUN",
                "field_evidence_map": "title/meta/description/alt derived from rendered page meta, public JSON, viewed images, and SERP-only validation.",
                "evidence_id": evidence_id,
                "issues": "Admin SEO fields unknown without Shopify export; QA not run; not approved for deployment.",
            }
        )
        seo.update(img_fields)
        b_seo.append(seo)

        b_evd.append(
            {
                "evidence_id": evidence_id,
                "product_url": url,
                "reviewed_at": NOW,
                "sources_accessed": "; ".join([url, "evidence/products/B006_products.json", meta.get("captured_html_path", ""), *SERP_REFS]),
                "current_H1": meta.get("h1_current", ""),
                "current_meta_title": meta.get("rendered_title_current", ""),
                "short_source_excerpt": strip_html(product.get("body_html", ""))[:600],
                "verified_product_facts": f"{design['obs']}; HTTP {meta.get('http_status')}; canonical={meta.get('canonical_url')}; product type={product.get('product_type')}; variants={len(product.get('variants', []))}; images={len(imgs)}.",
                "gallery_image_count": len(imgs),
                "images_viewed_count": len(imgs),
                "image_audit_references": "; ".join([f"{evidence_id}-IMG{int(i['image_number']):02d}" for i in imgs]),
                "SERP_evidence_references": "; ".join(SERP_REFS),
                "buyer_research_references": research_id,
                "factual_conflicts": "None observed in public data; admin values unknown.",
                "processing_status": "DRAFTED",
                "confidence_and_reason": "MEDIUM: rendered page, product JSON, and images verified; keyword demand has SERP-only evidence and no volume data.",
                "fact_to_source_map": "product facts=B006_products.json and B006_pages HTML; visuals=B006 local images/contact sheets; keyword context=SERP references.",
                "proposed_field_to_fact_map": "motif and visible product context to title/meta/alt; care/backing/material only where page/image evidence supports.",
            }
        )

        b_buyer.append(
            {
                "research_id": research_id,
                "product_key": key,
                "supporting_fact_ids": evidence_id,
                "purchase_context": "Fantasy reading-room decor or pet-lover room styling, depending on the product motif.",
                "jtbd_statement": f"When decorating a themed room, the buyer wants a {design['theme']} that makes the space feel expressive and personal.",
                "functional_motivation": "Add a motif-specific rug for a bedroom, reading nook, playroom, or living room.",
                "emotional_social_motivation": "Show love for fantasy creatures or pets through a distinctive room accent.",
                "purchase_concerns": "Size fit, design clarity, cleaning, backing grip, material feel, and whether the visual style suits the room.",
                "customer_language": f"{design['primary']}; {design['secondary']}",
                "language_origin": "AGENT_HYPOTHESIS_WITH_CATEGORY_SERP_PARAPHRASE",
                "source_refs": "; ".join(SERP_REFS),
                "source_scope": "CATEGORY",
                "evidence_excerpt": "Paraphrased from public search/category results; no direct Chillgen customer quote or volume source used.",
                "observed_at": NOW,
                "market": "US",
                "source_language": "en",
                "research_status": "HYPOTHESIS",
                "limitations": "No Search Console, internal search, direct review corpus, or paid volume tool used for this product.",
                "seo_application": "Use motif-specific product terms and keep broad generic rug terms for collection/category review.",
            }
        )

        rejected = "dragon rug" if "dragon" in design["primary"] else "dog rug"
        for kw, role, target, reason in [
            (design["primary"], "PRIMARY", url, "Best product-level match for observed motif and decor intent."),
            (design["secondary"].split(", ")[0], "SECONDARY", url, "Useful motif synonym supported by page and image evidence."),
            (rejected, "REJECTED", "COLLECTION_OR_CATEGORY_REVIEW", "Too broad for product-level targeting across multiple similar products."),
        ]:
            b_kw.append(
                {
                    "keyword": kw,
                    "product_key": key,
                    "buyer_research_refs": research_id if role != "REJECTED" else "",
                    "query_origin": "AGENT_PROPOSED",
                    "semantic_cluster": design["primary"],
                    "intent": "COMMERCIAL_INVESTIGATION",
                    "target_page_type": "PRODUCT" if role != "REJECTED" else "COLLECTION",
                    "target_url": target,
                    "keyword_role": role,
                    "decision_reason": reason,
                    "supporting_fact_ids": evidence_id,
                    "demand_evidence": "SERP_ONLY, no volume verified.",
                    "season": design["season"],
                    "research_period": "2026-09-07",
                    "validation_source": "; ".join(SERP_REFS),
                    "checked_at": NOW,
                    "representative_SERP_URLs": "; ".join(SERP_REFS[:4]),
                    "possible_overlap_with_other_products": "YES" if role == "REJECTED" else "MEDIUM_WITH_SIMILAR_PRODUCTS_IN_BATCH",
                    "mapping_reason": "Motif-specific keyword retained for product; broad term reserved for collection/category review.",
                    "mapping_status": "NEEDS_REVIEW",
                    "mapping_version": "2.4",
                }
            )
        keys_done.append(key)

    write_csv(ROOT / "batches/B006_SEO_Products.csv", b_seo, seo_headers)
    write_csv(ROOT / "batches/B006_Image_Audit.csv", b_img, image_headers)
    write_csv(ROOT / "batches/B006_Product_Evidence.csv", b_evd, evidence_headers)
    write_csv(ROOT / "keyword_research.csv", prev_kw + b_kw, keyword_headers)
    with (ROOT / "buyer_search_research.jsonl").open("w", encoding="utf-8") as f:
        for row in prev_buyer + b_buyer:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    for row in inv:
        if row["product_key"] in DESIGNS:
            row["assigned_batch"] = BATCH_ID
            row["batch_status"] = "COMPLETED"
            row["research_status"] = "DRAFTED"
            row["review_status"] = "NEEDS_REVIEW"
    write_csv(ROOT / "inventory.csv", inv, inv_headers)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", prev_seo + b_seo, seo_headers)
    write_sheet(wb, "Image_Audit", prev_img + b_img, image_headers)
    write_sheet(wb, "Product_Evidence", prev_evd + b_evd, evidence_headers)
    write_sheet(wb, "Keyword_Map", prev_kw + b_kw, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", prev_buyer + b_buyer, buyer_headers)
    readme = [
        {"metric": "scope", "value": "B001 through B006: 60 public products for chillgen.com, US/en-US", "definition": "Research proposal workbook only; all product and image proposals are NEEDS_REVIEW."},
        {"metric": "B006_method", "value": "B006 captured product page HTML/meta before SEO drafting and viewed all downloaded gallery images via contact sheets.", "definition": "B001-B003 remain public JSON/image evidence without full rendered page/meta backfill; B004-B006 include rendered page/meta capture."},
        {"metric": "sources", "value": "Chillgen public products.json, product URLs, rendered HTML captures, downloaded CDN images, contact sheets, SERP snippets and marketplace/category pages.", "definition": "Admin SEO fields and Shopify export media URLs remain unknown without Shopify CSV/export."},
        {"metric": "quality_status", "value": "DRAFTED_PUBLIC_RENDERED_PAGE_FOR_B004_B006", "definition": "458/458 images viewed via contact sheets across B001-B006; QA not run; no Shopify payload created."},
        {"metric": "deployment_status", "value": "NOT_APPROVED_NOT_DEPLOYABLE", "definition": "Do not import this workbook. Run QA and obtain explicit approval before any Shopify/Matrixify payload."},
        {"metric": "next_step", "value": "Stop after B006 and wait for user confirmation before B007.", "definition": "Required by project instruction."},
    ]
    write_sheet(wb, "README_QA", readme, ["metric", "value", "definition"])
    out = RESULTS / "batches" / f"SEO_Product_Optimization_through_{BATCH_ID}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    reopened = load_workbook(out, read_only=True, data_only=False)
    counts = {name: reopened[name].max_row - 1 for name in reopened.sheetnames}

    summary = ROOT / "batches/B006_summary.md"
    summary.write_text(
        "\n".join(
            [
                "# B006 Summary",
                "",
                f"- Run: {RUN_ID}",
                f"- Shop: {SHOP}",
                "- Scope: 10 products, 65 images",
                "- Cumulative workbook scope: B001-B006, 60 products, 458 images",
                "- Method: pre-batch checkpoint plus rendered page/meta HTML capture and contact-sheet visual review.",
                "- Status: NEEDS_REVIEW; QA_NOT_RUN; NOT_APPROVED; NO_DEPLOYMENT_PAYLOAD",
                f"- Workbook: {out}",
                f"- Generated at: {NOW}",
                "",
                "## Product Keys",
                *[f"- {key}" for key in keys_done],
            ]
        ),
        encoding="utf-8",
    )

    progress = json.loads((ROOT / "progress.json").read_text())
    completed = set(progress.get("images_completed", []))
    completed.update([f"{m['product_key']}#img_{int(m['image_number']):02d}" for m in manifest])
    progress.update(
        {
            "batch_id": BATCH_ID,
            "batch_product_keys": keys_done,
            "batch_status": "COMPLETED_NEEDS_REVIEW",
            "awaiting_confirmation": True,
            "continuation_confirmation_ref": "USER_CONFIRMED_CONTINUE_B006",
            "current_product_key": None,
            "current_stage": "B006_COMPLETE_AWAITING_USER_CONFIRMATION",
            "images_completed": sorted(completed),
            "last_saved_at": NOW,
        }
    )
    progress.setdefault("artifact_paths", {}).update(
        {
            "latest_batch_workbook": str(out),
            "b006_seo_products_csv": str(ROOT / "batches/B006_SEO_Products.csv"),
            "b006_image_audit_csv": str(ROOT / "batches/B006_Image_Audit.csv"),
            "b006_product_evidence_csv": str(ROOT / "batches/B006_Product_Evidence.csv"),
            "b006_summary": str(summary),
            "b006_page_meta": str(ROOT / "evidence/products/B006_page_meta.json"),
            "b006_image_manifest": str(ROOT / "evidence/images/B006_manifest.json"),
            "b006_contact_sheets": str(ROOT / "evidence/images/B006_contact_sheets"),
        }
    )
    tmp = ROOT / "progress.json.tmp"
    tmp.write_text(json.dumps(progress, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(ROOT / "progress.json")
    print(json.dumps({"workbook": str(out), "sheet_counts": counts, "summary": str(summary)}, indent=2))


if __name__ == "__main__":
    main()

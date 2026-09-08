from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R023/SEO_Product_Optimization_revision_R023.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R024"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R024.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R024"
MANIFEST = RUN_DIR / "revision_R024_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R024_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-family-couple-doormat-custom-couple-husband-design-04": {
        "title": "Funny Couple Welcome Doormat with Grumpy Joke Text",
        "meta": "Personalize a funny couple welcome doormat with cartoon portraits, custom names and grumpy old man joke text.",
        "description": "Personalize a funny couple welcome doormat with cartoon portraits, custom names and grumpy old man joke text.\n\nDesign details\n- Blue and tan rectangular doormat artwork shows cartoon couple portraits with a lovely lady and grumpy old man quote.\n- Product images include entryway mockups, couple holding mat, dog scene and size reference graphics.\n- SEO copy focuses on visible couple portrait artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-family-couple-doormat-custom-couple-husband-design-02": {
        "title": "Animals And Kids Couple Doormat with Cartoon Portraits",
        "meta": "Personalize an animals and kids couple doormat with cartoon portraits, custom names, welcome text and zoo joke art.",
        "description": "Personalize an animals and kids couple doormat with cartoon portraits, custom names, welcome text and zoo joke art.\n\nDesign details\n- Brown rectangular doormat artwork shows a seated cartoon couple with Hope You Like Animals and Kids text.\n- Visible details include custom names, welcome wording, entryway mockups and size reference graphics.\n- SEO copy stays tied to visible couple joke artwork without material, cleaning, reverse-side or surface-performance claims.",
    },
    "personalized-family-couple-doormat-custom-couple-husband-wife-doormat": {
        "title": "Personalized Couple Welcome Mat with Cartoon Portraits",
        "meta": "Personalize a couple welcome mat with cartoon husband wife portraits, custom names, hearts and grumpy old man text.",
        "description": "Personalize a couple welcome mat with cartoon husband wife portraits, custom names, hearts and grumpy old man text.\n\nDesign details\n- Navy and tan rectangular welcome mat artwork shows cartoon couple portraits with custom names and heart accents.\n- Visible quote reads a lovely lady and her grumpy old man live here, with entryway and dog mockup scenes.\n- SEO copy focuses on visible couple welcome artwork without unsupported material, cleaning, reverse-side or surface-performance claims.",
    },
    "custom-soccer-round-rug-trophy-championship-carpet-3ca0f26f80": {
        "title": "Three Nations Soccer Round Rug with Trophy Art",
        "meta": "Decorate with a three nations soccer round rug featuring USA, Canada and Mexico motifs, soccer ball and gold trophy.",
        "description": "Decorate with a three nations soccer round rug featuring USA, Canada and Mexico motifs, soccer ball and gold trophy.\n\nDesign details\n- Round soccer rug artwork shows USA flag section, Canadian maple leaf, Mexico-inspired crest, soccer ball and trophy.\n- Visible text includes The World Game 2026 with ornate border and room mockup scenes.\n- SEO copy stays tied to visible soccer artwork without unsupported material, reverse-side or surface-performance claims.",
    },
    "custom-soccer-championship-round-rug-106971c407": {
        "title": "Soccer Trophy Flag Ring Round Rug with Gold Border",
        "meta": "Decorate with a soccer trophy flag ring rug featuring large gold trophy, soccer balls, international flag panels and border.",
        "description": "Decorate with a soccer trophy flag ring rug featuring large gold trophy, soccer balls, international flag panels and border.\n\nDesign details\n- Round rug artwork shows a large gold trophy with soccer balls and a ring of international flag-style panels.\n- Product images include room mockups, close-up design panels and size reference graphics.\n- SEO copy focuses on visible soccer trophy artwork without unsupported material, reverse-side or surface-performance claims.",
    },
    "custom-soccer-2026-round-rug-trophy-flag-design-8257f0bdff": {
        "title": "Host Landmarks Soccer Round Rug with Flag Colors",
        "meta": "Decorate with a host landmarks soccer round rug featuring soccer ball center, USA Canada Mexico colors and landmark art.",
        "description": "Decorate with a host landmarks soccer round rug featuring soccer ball center, USA Canada Mexico colors and landmark art.\n\nDesign details\n- Round rug artwork shows soccer ball center, The World Game 2026 text and USA Canada Mexico flag colors.\n- Visible landmark-style details include a Statue of Liberty-like figure, pyramid imagery and laurel wreath.\n- SEO copy stays tied to visible host landmark artwork without material, reverse-side or surface-performance claims.",
    },
    "custom-soccer-2026-round-rug-trophy-championship-sports-decor-df4d4eec2f": {
        "title": "United States Eagle Soccer Rug with Trophy Art",
        "meta": "Decorate with a United States eagle soccer rug featuring two eagles, gold trophy, flag accents, stars and fireworks.",
        "description": "Decorate with a United States eagle soccer rug featuring two eagles, gold trophy, flag accents, stars and fireworks.\n\nDesign details\n- Round cream soccer rug artwork shows United States banner, two eagles, gold trophy and American flag accents.\n- Visible details include fireworks, stars, 26 Game of Champions text and room mockup scenes.\n- SEO copy focuses on visible eagle soccer artwork without unsupported material, reverse-side or surface-performance claims.",
    },
    "custom-soccer-championship-round-rug-world-flags-d1d1d396af": {
        "title": "World Flags Animal Soccer Rug with Trophy Art",
        "meta": "Decorate with a world flags animal soccer rug featuring eagle, moose, jaguar, soccer ball, trophy and flag accents.",
        "description": "Decorate with a world flags animal soccer rug featuring eagle, moose, jaguar, soccer ball, trophy and flag accents.\n\nDesign details\n- Round blue soccer rug artwork shows eagle, moose and jaguar motifs around a soccer ball and trophy.\n- Visible details include ribbon banner, maple leaf, flag accents, stars and room mockup scenes.\n- SEO copy stays tied to visible animal soccer artwork without material, reverse-side or surface-performance claims.",
    },
    "custom-soccer-round-rug-championship-trophy-world-flags-cc2e9bd59c-cc2e9bd59c": {
        "title": "Stadium Trophy Round Soccer Rug with Field Lights",
        "meta": "Decorate with a stadium trophy soccer rug featuring large gold trophy, stadium field, floodlights and host flag panels.",
        "description": "Decorate with a stadium trophy soccer rug featuring large gold trophy, stadium field, floodlights and host flag panels.\n\nDesign details\n- Round soccer rug artwork shows a large gold trophy on a stadium field with floodlights and soccer balls.\n- Visible details include USA Canada Mexico flag panels, room mockups, close-up views and size reference graphics.\n- SEO copy focuses on visible stadium soccer artwork without unsupported material, reverse-side or surface-performance claims.",
    },
    "custom-octopus-shaped-area-rug-a96ca58e10-a96ca58e10": {
        "title": "Octopus Tentacle Shaped Rug with Red Blue Sea Art",
        "meta": "Decorate with an octopus tentacle shaped rug featuring red body, curled tentacles, blue wave accents and yellow eye.",
        "description": "Decorate with an octopus tentacle shaped rug featuring red body, curled tentacles, blue wave accents and yellow eye.\n\nDesign details\n- Custom-shaped octopus rug artwork shows a red body with curled tentacles and blue wave-style accents.\n- Product images include room mockups, close-up design panels, custom-shape views and size reference graphics.\n- SEO copy stays tied to visible octopus artwork without material, cleaning, reverse-side or surface-performance claims.",
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "kid friendly", "pet friendly", "safe durable",
    "machine-washable", "machine washable", "washable ", "quick-dry", "quick dry",
    "memory foam", "microfiber", "velvet", "stain", "fade resistant", "easy clean",
    "easy-clean", "waterproof", "absorbent", "absorption", "backing", "rubber",
    "layered construction", "hd printing", "ultra-soft", "make a bold statement",
    "bring match-day", "show fan pride", "create a soccer fan zone",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    blob = " ".join([title, meta, item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = ["Handle", "title_proposed", "meta_title_seo", "meta_title_chars", "meta_description_seo", "meta_description_chars", "description_proposed", "description_proposed_html", "revision", "review_status", "review_reason", "processing_status", "content_qa_status", "issues"]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R024 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R024 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - {c[0] for c in changed})}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R024"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append(["R024", NOW, handle, row_idx, old_revision, new_revision, "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues", "T1; T2; unsupported_claims", "Run evidence-driven re-QA on revision 2 for this product before approval."])

    wb.save(OUTPUT)

    manifest = {"revision_batch_id": "R024", "generated_at": NOW, "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT), "source_revision_plan": str(PLAN), "product_count": len(changed), "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA", "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN", "not_approved_not_deployed": True, "next_step": "Bắt đầu revision R025 or Re-QA revision R024"}
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = ["# Revision R024 Summary", "", f"- Generated at: `{NOW}`", f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`", "- Scope: 10 products from revision plan R024.", "- Fixes targeted QA issues: short T1/T2 and unsupported claims.", "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.", "- No Shopify deploy, no approval.", "", "## Updated products", "", "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|"]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([str(title or ""), str(meta or ""), str(row[out_cols["description_proposed"] - 1] or "")]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({"output": str(OUTPUT), "updated": seen, "revision_log_rows": log.max_row - 1, "r024_log_rows": counts.get("R024", 0), "duplicate_meta_titles": len(dup_meta_titles), "sheets": reopened.sheetnames}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

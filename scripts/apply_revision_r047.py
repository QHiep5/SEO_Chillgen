from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R046/SEO_Product_Optimization_revision_R046.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R047"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R047.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R047"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-mathematics-education-rug-677ea4d8b5-677ea4d8b5": {
        "title": "Ms Anna Math Thinking Rug with Chalkboard Prompts",
        "meta": "Personalize a Ms Anna math thinking rug with green chalkboard design, problem-solving prompt bubbles and classroom math artwork.",
        "primary": "Ms Anna math thinking rug",
        "secondary": "green chalkboard math rug, problem solving classroom rug, teacher name math mat",
        "long_tail": "Ms Anna math thinking rug; green chalkboard math rug; problem solving classroom rug",
        "description": "Personalize a Ms Anna math thinking rug with green chalkboard design, problem-solving prompt bubbles and classroom math artwork.\n\nDesign details\n- Green chalkboard-style rug artwork shows Ms Anna name with question and strategy prompts around the center.\n- Visible prompts mention numbers, math tools, solving problems, checking work and explaining thinking, with classroom mockups and feature callouts.\n- SEO copy stays tied to visible math prompt artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {3: "Feature graphic for Ms Anna math thinking rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-185af7709b": {
        "title": "Mrs Brown Classroom Phrase Rug with Bright Words",
        "meta": "Personalize a Mrs Brown classroom phrase rug with bright student identity statements, crayon border and illustrated children artwork.",
        "primary": "Mrs Brown classroom phrase rug",
        "secondary": "bright classroom words rug, teacher name affirmation rug, crayon border classroom mat",
        "long_tail": "Mrs Brown classroom phrase rug; bright classroom words rug; teacher name affirmation rug",
        "description": "Personalize a Mrs Brown classroom phrase rug with bright student identity statements, crayon border and illustrated children artwork.\n\nDesign details\n- White rug artwork reads In Mrs Brown's Classroom with repeated You Are and You Will phrase lines.\n- Visible phrases include learners, explorers, readers, writers and the whole reason I am here, with illustrated children and room mockups.\n- SEO copy focuses on visible classroom phrase artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for Mrs Brown classroom phrase rug"},
    },
    "custom-math-classroom-welcome-mat-93506f16cb-93506f16cb": {
        "title": "Mrs Mullet Math Class Team Rug with Values Art",
        "meta": "Personalize a Mrs Mullet math class team rug with black background, colorful classroom values, school icons and math artwork.",
        "primary": "Mrs Mullet math class team rug",
        "secondary": "math class team rug, classroom values rug, teacher name school icon mat",
        "long_tail": "Mrs Mullet math class team rug; math class team rug; classroom values rug",
        "description": "Personalize a Mrs Mullet math class team rug with black background, colorful classroom values, school icons and math artwork.\n\nDesign details\n- Black rug artwork reads In Mrs Mullet's Class We Are A Team with colorful value words and school icons.\n- Visible words include respect, choices, attention, kindness and work hard, with room mockups, group scene and close-up views.\n- SEO copy stays tied to visible math team artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for Mrs Mullet math team rug"},
    },
    "personalized-motivational-classroom-rug-kids-c5c4db7870": {
        "title": "Mistakes Are Okay Classroom Rug with Prompt Art",
        "meta": "Personalize a Mistakes Are Okay classroom rug with custom class text, cartoon emotion characters and colorful prompt bubbles.",
        "primary": "Mistakes Are Okay classroom rug",
        "secondary": "prompt bubble classroom rug, emotion character rug, custom class quote mat",
        "long_tail": "Mistakes Are Okay classroom rug; prompt bubble classroom rug; emotion character rug",
        "description": "Personalize a Mistakes Are Okay classroom rug with custom class text, cartoon emotion characters and colorful prompt bubbles.\n\nDesign details\n- Cream rug artwork reads It's Okay To with prompt bubbles such as make mistakes, have bad days, not know it all and start over.\n- Visible details include cartoon faces, Mrs Brown class text, school icons, classroom scenes and child scene image.\n- SEO copy focuses on visible prompt-bubble artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for Mistakes Are Okay classroom rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-6b9a9da243": {
        "title": "Mrs Anne Classroom Is Better Rug with Dot Border",
        "meta": "Personalize a Mrs Anne classroom quote rug with chalkboard style, colorful dot border, school icons and bold quote artwork.",
        "primary": "Mrs Anne classroom quote rug",
        "secondary": "classroom is better rug, chalkboard teacher rug, colorful border quote mat",
        "long_tail": "Mrs Anne classroom quote rug; classroom is better rug; chalkboard teacher rug",
        "description": "Personalize a Mrs Anne classroom quote rug with chalkboard style, colorful dot border, school icons and bold quote artwork.\n\nDesign details\n- Black chalkboard-style rug artwork reads Mrs. Anne Classroom Is Better Because You're In It.\n- Visible details include colorful border dots, school icons, size option graphic, classroom mockups and close-up artwork views.\n- SEO copy stays tied to visible classroom quote artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature infographic for Mrs Anne classroom quote rug",
            6: "Care-style graphic for classroom quote rug artwork",
        },
    },
    "personalized-motivational-classroom-rug-kids-dcd5b3d508": {
        "title": "Mrs Harper Alphabet Affirmation Rug with Traits",
        "meta": "Personalize a Mrs Harper alphabet affirmation rug with A to Z letters, character trait words and black classroom artwork.",
        "primary": "Mrs Harper alphabet affirmation rug",
        "secondary": "A to Z classroom rug, character trait rug, teacher name ABC mat",
        "long_tail": "Mrs Harper alphabet affirmation rug; A to Z classroom rug; character trait rug",
        "description": "Personalize a Mrs Harper alphabet affirmation rug with A to Z letters, character trait words and black classroom artwork.\n\nDesign details\n- Black rug artwork shows You Are header text, alphabet letters and a trait word beside each letter.\n- Visible details include Mrs Harper personalization, colorful letters, classroom mockups, size option graphic and close-up views.\n- SEO copy focuses on visible alphabet affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature infographic for Mrs Harper alphabet rug",
            6: "Care-style graphic for ABC affirmation rug artwork",
        },
    },
    "personalized-motivational-classroom-rug-for-kids-75356d9313-75356d9313": {
        "title": "Mrs Smith Vintage Classroom Rug with Note Blocks",
        "meta": "Personalize a Mrs Smith vintage classroom rug with wood plank background, colorful note blocks and kindness phrase artwork.",
        "primary": "Mrs Smith vintage classroom rug",
        "secondary": "wood plank classroom rug, colorful note block rug, teacher name kindness mat",
        "long_tail": "Mrs Smith vintage classroom rug; wood plank classroom rug; colorful note block rug",
        "description": "Personalize a Mrs Smith vintage classroom rug with wood plank background, colorful note blocks and kindness phrase artwork.\n\nDesign details\n- Wood plank-style rug artwork shows Mrs. Smith's Class text surrounded by colorful note blocks.\n- Visible phrases include you are loved, keep going, spread kindness and belong, with room mockups, size option graphic and close-up views.\n- SEO copy stays tied to visible note-block artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature infographic for Mrs Smith vintage classroom rug",
            6: "Care-style graphic for vintage classroom rug artwork",
        },
    },
    "personalized-motivational-classroom-rug-for-kids-fcc5a59bce-fcc5a59bce": {
        "title": "Mrs Brown Class Rug with I Am Prompt Icon Artwork",
        "meta": "Personalize a Mrs Brown class rug with yellow and pink gingham design, I Am prompt, smiley icons and school artwork.",
        "primary": "Mrs Brown I Am prompt rug",
        "secondary": "yellow classroom rug, teacher name prompt rug, smiley school icon mat",
        "long_tail": "Mrs Brown I Am prompt rug; yellow classroom rug; teacher name prompt rug",
        "description": "Personalize a Mrs Brown class rug with yellow and pink gingham design, I Am prompt, smiley icons and school artwork.\n\nDesign details\n- Yellow and pink classroom rug artwork shows Mrs Brown's Class text with I Am prompt and illustrated school icons.\n- Visible details include smiley faces, pencils, books, hearts, stars, classroom mockups, sofa scene and child scene image.\n- SEO copy focuses on visible teacher-name prompt artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {7: "Feature graphic for Mrs Brown I Am prompt rug"},
    },
    "custom-3d-dragon-round-rug-5315fbccc7": {
        "title": "Red Dragon Round Rug with Fantasy Portal Artwork",
        "meta": "Decorate with a red dragon round rug featuring curled dragon artwork, circular fantasy portal, bookshelf-style border and room scenes.",
        "primary": "red dragon round rug",
        "secondary": "fantasy portal dragon rug, round dragon artwork rug, red curled dragon mat",
        "long_tail": "red dragon round rug; fantasy portal dragon rug; round dragon artwork rug",
        "description": "Decorate with a red dragon round rug featuring curled dragon artwork, circular fantasy portal, bookshelf-style border and room scenes.\n\nDesign details\n- Round rug artwork shows a curled red dragon inside a circular portal-style design with a booklike border.\n- Visible imagery includes size option graphic, fireplace scene, seasonal room mockup, close-up views and feature panels.\n- SEO copy stays tied to visible red dragon artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            7: "Feature graphic for red dragon round rug artwork",
            8: "Design detail graphic for red dragon round rug",
        },
    },
    "custom-dragon-round-rug-3d-illusion-4ca1b4fcd0": {
        "title": "White Dragon Round Rug with Dark Portal Artwork",
        "meta": "Decorate with a white dragon round rug featuring dark circular portal artwork, stone rim effect, flying dragon and room scenes.",
        "primary": "white dragon round rug",
        "secondary": "dark portal dragon rug, round fantasy dragon rug, white flying dragon mat",
        "long_tail": "white dragon round rug; dark portal dragon rug; round fantasy dragon rug",
        "description": "Decorate with a white dragon round rug featuring dark circular portal artwork, stone rim effect, flying dragon and room scenes.\n\nDesign details\n- Round rug artwork shows a white dragon flying over a dark circular portal with a stone rim look.\n- Visible imagery includes size option graphic, fireplace scene, close-up portal views and feature panels.\n- SEO copy focuses on visible white dragon artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            6: "Design detail graphic for white dragon round rug",
            8: "Feature graphic for white dragon round rug artwork",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable", "floor mat", "floor rug", "low pile", "skid",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R047 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R047 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R047"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R047", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R047", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R048 or Re-QA revision R047",
    }
    (RUN_DIR / "revision_R047_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R047 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R047.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R047 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R047_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r047_log_rows": counts.get("R047", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

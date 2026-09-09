from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B011"
REVISION = "R087"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"


COPY = {
    "personalized-teacher-classroom-rug-custom-back-to-school-welcome-mat": {
        "title": "Personalized Classroom Welcome Rug",
        "meta": "Customize a classroom welcome rug with teacher-name text, colorful affirmations, school icons, and a warm back-to-school look.",
        "description": "Customize a classroom welcome rug with teacher-name text, colorful affirmations, school icons, and a warm back-to-school look. The design helps make the room feel named, cheerful, and ready for students.",
        "primary": "personalized classroom welcome rug",
        "secondary": "custom teacher name rug, back to school welcome mat, classroom affirmation rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-100": {
        "title": "Personalized It's Okay Classroom Rug",
        "meta": "Customize an It's Okay classroom rug with teacher-name text, ocean animals, and reassurance prompts for SEL classroom decor.",
        "description": "Customize an It's Okay classroom rug with teacher-name text, ocean animals, and reassurance prompts. The design reminds students that mistakes, feelings, questions, and fresh starts belong in the classroom.",
        "primary": "personalized it's okay classroom rug",
        "secondary": "custom teacher SEL rug, classroom reassurance rug, back to school teacher rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-101": {
        "title": "Personalized Neon Classroom Welcome Rug",
        "meta": "Customize a neon classroom welcome rug with teacher-name text, chalkboard styling, school icons, and bright back-to-school color.",
        "description": "Customize a neon classroom welcome rug with teacher-name text, chalkboard styling, school icons, and bright back-to-school color. The design gives the classroom entry or gathering area a lively teacher-name identity.",
        "primary": "personalized neon classroom welcome rug",
        "secondary": "custom teacher welcome rug, neon classroom decor, back to school classroom mat",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-102": {
        "title": "Personalized Pastel Kindness Rug",
        "meta": "Customize a pastel kindness classroom rug with teacher-name text, bows, daisies, crayon affirmations, and gentle classroom color.",
        "description": "Customize a pastel kindness classroom rug with teacher-name text, bows, daisies, crayon affirmations, and gentle classroom color. The sweet design pairs name identity with simple kindness language for young learners.",
        "primary": "personalized pastel kindness rug",
        "secondary": "custom teacher kindness rug, pastel classroom welcome mat, back to school classroom rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-103": {
        "title": "Personalized Rainbow Classroom Rug",
        "meta": "Customize a rainbow classroom welcome rug with teacher-name text, chalkboard styling, school icons, and colorful student-friendly artwork.",
        "description": "Customize a rainbow classroom welcome rug with teacher-name text, chalkboard styling, school icons, and colorful student-friendly artwork. The design makes the classroom feel cheerful, clear, and ready for a new school year.",
        "primary": "personalized rainbow classroom rug",
        "secondary": "custom teacher welcome rug, chalkboard classroom rug, back to school teacher mat",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-104": {
        "title": "Personalized Teacher Welcome Rug",
        "meta": "Customize a teacher welcome rug with classroom-name text, chalkboard style, bunting, apple, books, hearts, and school-day charm.",
        "description": "Customize a teacher welcome rug with classroom-name text, chalkboard style, bunting, apple, books, hearts, and school-day charm. The design creates a friendly teacher-branded welcome for students and visitors.",
        "primary": "personalized teacher welcome rug",
        "secondary": "custom classroom welcome mat, back to school teacher rug, teacher name classroom rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-105": {
        "title": "Personalized Classroom Affirmation Rug",
        "meta": "Customize a classroom affirmation rug with teacher-name text, doodle school icons, and grow, be kind, learn, and believe prompts.",
        "description": "Customize a classroom affirmation rug with teacher-name text, doodle school icons, and prompts such as grow, be kind, learn, believe, dream big, inspire, be you, and you matter.",
        "primary": "personalized classroom affirmation rug",
        "secondary": "custom teacher affirmation rug, back to school classroom mat, positive classroom rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-106": {
        "title": "Personalized Crayon Classroom Rug",
        "meta": "Customize a crayon classroom rug with teacher-name text, In This Classroom affirmations, colorful words, and school icons.",
        "description": "Customize a crayon classroom rug with teacher-name text, In This Classroom affirmations, colorful words, and school icons. The design brings positive identity language into a bright classroom gathering area.",
        "primary": "personalized crayon classroom rug",
        "secondary": "custom teacher affirmation mat, back to school classroom rug, classroom crayon rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-107": {
        "title": "Personalized Good Day For Learning Rug",
        "meta": "Customize a good day for learning classroom rug with teacher-name text, chalk-style school icons, rainbow art, and morning optimism.",
        "description": "Customize a good day for learning classroom rug with teacher-name text, chalk-style school icons, rainbow art, and morning optimism. The design gives students a positive daily cue as they enter the learning space.",
        "primary": "personalized good day for learning rug",
        "secondary": "good day for learning rug, custom teacher classroom rug, back to school welcome rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-108": {
        "title": "Personalized Rainbow Welcome Rug",
        "meta": "Customize a rainbow welcome classroom rug with teacher-name text, colorful scallop border, school icons, and affirmation circles.",
        "description": "Customize a rainbow welcome classroom rug with teacher-name text, colorful scallop border, school icons, and affirmation circles. The bright design makes a classroom entry or reading area feel celebratory and personal.",
        "primary": "personalized rainbow welcome rug",
        "secondary": "custom teacher welcome mat, colorful classroom rug, back to school classroom decor",
    },
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B011 scope does not match the R087 copy map")

    for row in seo_rows:
        item = COPY[row["Handle"]]
        row["title_proposed"] = item["title"]
        row["meta_title_seo"] = item["title"]
        row["meta_title_chars"] = str(len(item["title"]))
        row["meta_description_seo"] = item["meta"]
        row["meta_description_chars"] = str(len(item["meta"]))
        row["description_proposed"] = item["description"]
        row["description_proposed_html"] = f"<p>{item['description']}</p>"
        row["primary_keyword"] = item["primary"]
        row["secondary_keywords"] = item["secondary"]
        row["long_tail_candidates"] = f"{item['primary']}; {item['secondary']}"
        row["revision"] = "2"
        row["review_status"] = "NEEDS_REVIEW"
        row["content_qa_status"] = "NOT_RUN"
        row["review_reason"] = f"{REVISION} removes internal evidence-note language and unsupported claim wording from B011 copy."
        row["issues"] = f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations."

    handles_set = set(handles)
    evidence_scope = [row for row in evidence_rows if row.get("evidence_id") in {r.get("evidence_id") for r in seo_rows}]
    keyword_scope = [row for row in keyword_rows if row.get("product_key") in handles_set]
    buyer_rows = [row for row in buyer_all if row.get("product_key") in handles_set]
    img_scope = [row for row in img_rows if row.get("Handle") in handles_set]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": "Canonical B011 cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    evidence_ids = {row["evidence_id"] for row in evidence_scope}
    for row in seo_rows:
        evidence_id = row.get("evidence_id")
        if evidence_id not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {evidence_id}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", evidence_id])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    digest = sha256(OUT)
    manifest = {
        "revision": REVISION,
        "batch_id": BATCH,
        "scope_count": len(handles),
        "image_count": len(img_scope),
        "canonical_workbook": str(OUT),
        "sha256": digest,
        "status": "CREATED_FOR_HARDENED_REQA",
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
    }
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R042/SEO_Product_Optimization_revision_R042.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R043"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R043.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R043"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-mathematics-education-rug-6494449be8-6494449be8": {
        "title": "Custom Welcome To Math Class Rug with Teacher Name",
        "meta": "Personalize a Welcome To Math Class rug with teacher name, chalkboard formulas, bunting flags, calculator and school supply art.",
        "primary": "custom Welcome To Math Class rug",
        "secondary": "teacher name math rug, chalkboard math classroom rug, formula classroom mat",
        "long_tail": "custom Welcome To Math Class rug; teacher name math rug; chalkboard math classroom rug",
        "description": "Personalize a Welcome To Math Class rug with teacher name, chalkboard formulas, bunting flags, calculator and school supply art.\n\nDesign details\n- Black chalkboard-style rug artwork reads Welcome Room 110 Mr. Micheal's Math Class with colorful letters.\n- Visible details include formulas, books, calculator, pencil artwork, bunting flags, classroom mockups, group scenes and a feature graphic.\n- SEO copy stays tied to visible math-class artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            7: "Feature graphic for Mr Micheal math class rug artwork",
        },
    },
    "custom-mathematics-educational-welcome-mat-627c4013d4": {
        "title": "Custom Math Is All About Rug with Teacher Name",
        "meta": "Personalize a Math Is All About rug with teacher name, colorful concept bubbles, number art, symbols and notebook-style border.",
        "primary": "custom Math Is All About rug",
        "secondary": "teacher name math rug, colorful math concept rug, classroom symbols mat",
        "long_tail": "custom Math Is All About rug; teacher name math rug; colorful math concept rug",
        "description": "Personalize a Math Is All About rug with teacher name, colorful concept bubbles, number art, symbols and notebook-style border.\n\nDesign details\n- Gray classroom rug artwork reads Math Is All About with colorful bubbles for concepts such as number, shape, logic and equation.\n- Visible imagery includes math symbols, school icons, black speckled border, room mockups, close-up art and feature graphic.\n- SEO copy focuses on visible math concept artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature graphic for Math Is All About rug artwork",
        },
    },
    "custom-vintage-book-classroom-reading-rug-3c50e8b414": {
        "title": "Cream Vintage Book Classroom Rug for Reading Nook",
        "meta": "Decorate with a cream vintage book classroom rug featuring open books, book stacks, botanical border and library-style artwork.",
        "primary": "cream vintage book classroom rug",
        "secondary": "vintage book reading rug, classroom library rug, botanical book nook mat",
        "long_tail": "cream vintage book classroom rug; vintage book reading rug; classroom library rug",
        "description": "Decorate with a cream vintage book classroom rug featuring open books, book stacks, botanical border and library-style artwork.\n\nDesign details\n- Cream rug artwork shows open books, stacked books, leafy accents and a vintage botanical border.\n- Visible product imagery includes room mockups, close-up book panels, size reference graphics and feature infographics.\n- SEO copy stays tied to visible vintage book artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            3: "Feature graphic for vintage book classroom rug artwork",
            4: "Feature infographic for cream vintage book rug artwork",
            7: "Close-up of cream vintage book rug artwork",
        },
    },
    "custom-tree-of-knowledge-classroom-library-rug-077de4a60b": {
        "title": "Tree Of Knowledge Library Rug with Books And Stars",
        "meta": "Personalize a Tree Of Knowledge library rug with green tree, books, moon, stars, teacher name text and reading artwork.",
        "primary": "Tree Of Knowledge library rug",
        "secondary": "green reading tree rug, book tree classroom rug, teacher name library mat",
        "long_tail": "Tree Of Knowledge library rug; green reading tree rug; book tree classroom rug",
        "description": "Personalize a Tree Of Knowledge library rug with green tree, books, moon, stars, teacher name text and reading artwork.\n\nDesign details\n- Green rug artwork shows a large tree with books in the branches and books scattered across the grass.\n- Visible details include moon and stars, teacher name text, classroom mockups, close-up tree art, size guide views and feature graphics.\n- SEO copy focuses on visible Tree Of Knowledge artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            3: "Feature graphic for Tree Of Knowledge rug artwork",
            4: "Close-up of book and tree artwork",
            5: "Feature infographic for Tree Of Knowledge library rug",
        },
    },
    "custom-reading-tree-classroom-rug-b66c4b4a97": {
        "title": "Reading Tree Classroom Rug with Bookshelf Artwork",
        "meta": "Decorate with a Reading Tree classroom rug featuring bookshelf tree art, flying books, reading quotes and green border.",
        "primary": "Reading Tree classroom rug",
        "secondary": "bookshelf tree rug, reading quote rug, classroom book corner mat",
        "long_tail": "Reading Tree classroom rug; bookshelf tree rug; reading quote rug",
        "description": "Decorate with a Reading Tree classroom rug featuring bookshelf tree art, flying books, reading quotes and green border.\n\nDesign details\n- Cream and green rug artwork shows a tree-shaped bookshelf with flying books and reading-themed quote panels.\n- Visible text includes reading messages around the tree, with classroom mockups, close-up views, size graphics and feature panels.\n- SEO copy stays tied to visible bookshelf-tree artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            7: "Feature graphic for Reading Tree rug artwork",
            8: "Close-up of green Reading Tree rug artwork",
            9: "Feature infographic for Reading Tree classroom rug",
        },
    },
    "custom-reading-adventure-classroom-library-rug-7956abd8c9": {
        "title": "Reading Is An Adventure Library Rug with Globe Art",
        "meta": "Personalize a Reading Is An Adventure library rug with globe, books, pencil, leaves and bold classroom-style quote artwork.",
        "primary": "Reading Is An Adventure library rug",
        "secondary": "globe reading rug, classroom library quote rug, book adventure mat",
        "long_tail": "Reading Is An Adventure library rug; globe reading rug; classroom library quote rug",
        "description": "Personalize a Reading Is An Adventure library rug with globe, books, pencil, leaves and bold classroom-style quote artwork.\n\nDesign details\n- Black reading rug artwork says Reading Is An Adventure That Never Ends with a globe, books and colorful typography.\n- Visible imagery includes classroom and reading-corner mockups, size option graphic, close-up panels and feature graphics.\n- SEO copy focuses on visible adventure-reading artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            6: "Feature infographic for Reading Adventure rug artwork",
            7: "Feature graphic for Reading Adventure library rug",
            8: "Close-up of black Reading Adventure rug artwork",
        },
    },
    "custom-octopus-shaped-area-rug-0acb5a8130-0acb5a8130": {
        "title": "Red Octopus Shaped Rug with Wood Plank Wave Art",
        "meta": "Personalize a red octopus shaped rug with curled tentacles, suction cup details, wood plank frame and wave-style background art.",
        "primary": "red octopus shaped rug",
        "secondary": "octopus area rug, sea monster rug, tentacle shaped rug",
        "long_tail": "red octopus shaped rug; octopus area rug; sea monster rug",
        "description": "Personalize a red octopus shaped rug with curled tentacles, suction cup details, wood plank frame and wave-style background art.\n\nDesign details\n- Shaped red octopus rug artwork shows curled tentacles, red suction cups and a dark wood plank frame.\n- Visible imagery includes top-down octopus views, room mockups, size reference graphic, design detail panel and repeated gallery images.\n- SEO copy stays tied to visible octopus artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            4: "Feature graphic for red octopus shaped rug artwork",
            6: "Design detail graphic for octopus shaped rug",
            13: "Repeated design detail graphic for octopus rug",
        },
    },
    "custom-classroom-library-reading-rug-dfdbf660e0": {
        "title": "Today A Reader Library Rug with Teacher Name Art",
        "meta": "Personalize a Today A Reader library rug with teacher name, pastel bookshelves, gift bow and Tomorrow A Leader quote text.",
        "primary": "Today A Reader library rug",
        "secondary": "teacher name reading rug, bookshelf quote rug, classroom library mat",
        "long_tail": "Today A Reader library rug; teacher name reading rug; bookshelf quote rug",
        "description": "Personalize a Today A Reader library rug with teacher name, pastel bookshelves, gift bow and Tomorrow A Leader quote text.\n\nDesign details\n- Black reading rug artwork shows pastel bookshelves, a gift bow and the quote Today A Reader Tomorrow A Leader.\n- Visible product imagery includes classroom mockups, size option graphic, close-up reading quote view and feature panels.\n- SEO copy focuses on visible bookshelf quote artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature infographic for Today A Reader library rug",
            8: "Feature graphic for bookshelf reading rug artwork",
        },
    },
    "custom-reading-classroom-library-rug-29b9f9b43d-29b9f9b43d": {
        "title": "Reading Gives You Wings Classroom Rug Artwork",
        "meta": "Personalize a Reading Gives You Wings classroom rug with rainbow butterflies, clouds, children reading and open book artwork.",
        "primary": "Reading Gives You Wings rug",
        "secondary": "butterfly classroom reading rug, rainbow book rug, library reading mat",
        "long_tail": "Reading Gives You Wings rug; butterfly classroom reading rug; rainbow book rug",
        "description": "Personalize a Reading Gives You Wings classroom rug with rainbow butterflies, clouds, children reading and open book artwork.\n\nDesign details\n- Black reading rug artwork shows two children on an open book beneath rainbow butterflies and cloud accents.\n- Visible details include the Reading Gives You Wings quote, room mockups, classroom scenes, size guide views and feature graphic.\n- SEO copy stays tied to visible butterfly-reading artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature graphic for Reading Gives You Wings rug",
        },
    },
    "custom-reading-corner-welcome-mat-941ccd58e8-941ccd58e8": {
        "title": "Reading Will Take You Everywhere Rug Quote Artwork",
        "meta": "Decorate with a Reading Will Take You Everywhere rug featuring rainbow pencil border, paper airplanes, books and quote artwork.",
        "primary": "Reading Will Take You Everywhere rug",
        "secondary": "rainbow reading corner rug, book quote classroom rug, paper airplane mat",
        "long_tail": "Reading Will Take You Everywhere rug; rainbow reading corner rug; book quote classroom rug",
        "description": "Decorate with a Reading Will Take You Everywhere rug featuring rainbow pencil border, paper airplanes, books and quote artwork.\n\nDesign details\n- Black reading rug artwork shows colorful paper airplanes, books and a bright rainbow pencil border.\n- Visible text reads Reading Will Take You Everywhere, with classroom scenes, size option graphic, group reading mockup and feature panels.\n- SEO copy focuses on visible quote-and-airplane artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            6: "Feature infographic for Reading Will Take You Everywhere rug",
            7: "Care-style graphic for reading quote rug artwork",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable", "floor mat", "floor rug", "low pile",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R043 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R043 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R043"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R043", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R043", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R044 or Re-QA revision R043",
    }
    (RUN_DIR / "revision_R043_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R043 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R043.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R043 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R043_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r043_log_rows": counts.get("R043", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

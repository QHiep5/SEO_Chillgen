"""Build a catalog-level registry for every QA run found on disk."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(os.environ.get("SEO_CHILLGEN_BASE", Path.cwd())).resolve()
RUN_ID = "chillgen_20260907_01"
SHOP_DOMAIN = "chillgen.com"
QA_RESULTS = BASE / f"resutls/{SHOP_DOMAIN}/{RUN_ID}/qa"
QA_RUNS = BASE / f"seo_runs/{SHOP_DOMAIN}/{RUN_ID}/qa"
REGISTRY_JSON = QA_RESULTS / "qa_run_registry.json"
REGISTRY_CSV = QA_RESULTS / "qa_run_registry.csv"
TZ = timezone(timedelta(hours=7))


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def hash_scope(keys) -> str:
    payload = json.dumps(list(keys or []), ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def read_json(path: Path):
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def xlsx_summary(path: Path):
    if not path.exists():
        return {}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "QA_Summary" not in wb.sheetnames:
            return {}
        return {
            str(r[0]): r[1]
            for r in wb["QA_Summary"].iter_rows(min_row=2, values_only=True)
            if r and r[0] is not None
        }
    except Exception:
        return {}


def infer_revision(run_id: str, files):
    for file in files:
        match = re.search(r"SEO_QA_REV_(R\d{3})", file.name)
        if match:
            return match.group(1)
    match = re.search(r"(R\d{3})", run_id)
    return match.group(1) if match else ""


def infer_batch(run_id: str, manifest, summary):
    return (
        manifest.get("batch_id")
        or summary.get("batch_id")
        or (re.search(r"(B\d{3})", run_id).group(1) if re.search(r"(B\d{3})", run_id) else "")
    )


def classify(record):
    if "AGENT" in record["qa_run_id"] and not record.get("rubric_hash"):
        return "LEGACY"
    required = ("source_hash", "rubric_hash", "scope_hash")
    if any(not record.get(k) for k in required):
        return "LEGACY"
    if record.get("product_count") not in (10, "10"):
        return "NONSTANDARD_SCOPE"
    if record.get("canonical"):
        return "CANONICAL"
    return "RECORDED"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--canonical-run-id", action="append", default=[], help="Explicit canonical QA run id; repeatable")
    args = parser.parse_args()

    records = []
    for run_dir in sorted(p for p in QA_RESULTS.iterdir() if p.is_dir()):
        run_id = run_dir.name
        work_manifest = read_json(QA_RUNS / run_id / "manifest.json")
        result_manifest = read_json(run_dir / "manifest.json")
        manifest = {**result_manifest, **work_manifest}
        xlsx_files = sorted(run_dir.glob("SEO_QA_*.xlsx"))
        md_files = sorted(run_dir.glob("SEO_QA_*.md"))
        summary = xlsx_summary(xlsx_files[0]) if xlsx_files else {}
        keys = manifest.get("batch_product_keys")
        if not keys and summary.get("scope_product_keys"):
            try:
                keys = json.loads(summary["scope_product_keys"])
            except Exception:
                keys = []
        keys = keys or []
        revision = manifest.get("revision_id") or summary.get("revision_id") or infer_revision(run_id, xlsx_files + md_files)
        batch_id = infer_batch(run_id, manifest, summary)
        source_hash = manifest.get("source_workbook_sha256") or summary.get("source_workbook_sha256") or ""
        rubric_hash = manifest.get("rubric_hash") or summary.get("rubric_hash") or ""
        current_scope_hash = manifest.get("scope_hash") or summary.get("scope_hash") or (hash_scope(keys) if keys else "")
        canonical = bool(manifest.get("canonical") or run_id in args.canonical_run_id)
        record = {
            "qa_run_id": run_id,
            "batch_id": batch_id,
            "revision_id": revision,
            "parent_run_id": manifest.get("parent_run_id") or summary.get("parent_run_id") or "",
            "source_hash": source_hash,
            "rubric_hash": rubric_hash,
            "scope_hash": current_scope_hash,
            "product_count": len(keys) or manifest.get("scope_product_count") or summary.get("scope_product_count") or "",
            "status": manifest.get("content_qa_status") or manifest.get("batch_status") or summary.get("content_qa_status") or summary.get("batch_status") or "",
            "canonical": canonical,
            "supersedes": ",".join(manifest.get("supersedes", [])) if isinstance(manifest.get("supersedes"), list) else manifest.get("supersedes", ""),
            "superseded_by": "",
            "registry_status": "",
            "created_at": manifest.get("created_at") or summary.get("created_at") or "",
            "result_xlsx": str(xlsx_files[0]) if xlsx_files else "",
            "result_md": str(md_files[0]) if md_files else "",
            "manifest": str(QA_RUNS / run_id / "manifest.json") if (QA_RUNS / run_id / "manifest.json").exists() else "",
        }
        record["registry_status"] = classify(record)
        records.append(record)

    explicit = {r["qa_run_id"] for r in records if r["canonical"]}
    for record in records:
        supersedes = [x.strip() for x in str(record.get("supersedes") or "").split(",") if x.strip()]
        if record["qa_run_id"] in explicit:
            record["registry_status"] = "CANONICAL"
        for old_id in supersedes:
            for old in records:
                if old["qa_run_id"] == old_id:
                    old["canonical"] = False
                    old["superseded_by"] = record["qa_run_id"]
                    old["registry_status"] = "SUPERSEDED"

    seen = set()
    duplicates = []
    for record in records:
        if record["qa_run_id"] in seen:
            duplicates.append(record["qa_run_id"])
        seen.add(record["qa_run_id"])
    if duplicates:
        raise RuntimeError(f"Duplicate qa_run_id found: {duplicates}")

    registry = {
        "schema_version": "qa-run-registry-v1",
        "created_at": datetime.now(TZ).replace(microsecond=0).isoformat(),
        "run_count": len(records),
        "runs": records,
    }
    REGISTRY_JSON.write_text(json.dumps(registry, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    fields = [
        "qa_run_id", "batch_id", "revision_id", "parent_run_id", "source_hash",
        "rubric_hash", "scope_hash", "product_count", "status", "canonical",
        "supersedes", "superseded_by", "registry_status", "created_at",
        "result_xlsx", "result_md", "manifest",
    ]
    with REGISTRY_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for record in records:
            writer.writerow({field: record.get(field, "") for field in fields})

    counts = {}
    for record in records:
        counts[record["registry_status"]] = counts.get(record["registry_status"], 0) + 1
    print(json.dumps({"registry": str(REGISTRY_JSON), "csv": str(REGISTRY_CSV), "run_count": len(records), "status_counts": counts}, indent=2))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R013/SEO_Product_Optimization_revision_R013.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R014"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R014.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R014"
MANIFEST = RUN_DIR / "revision_R014_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R014_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-vinyl-record-rug-with-photo-and-song-name-9667c97a3c": {
        "title": "Custom Floral Vinyl Record Rug with Rose Border",
        "meta": "Personalize a custom floral vinyl record rug with photo area, rose border, black disc artwork and record cover text.",
        "description": (
            "Personalize a custom floral vinyl record rug with photo area, rose border, black disc artwork and record cover text.\n\n"
            "Design details\n"
            "- Round record-style rug artwork shows a black vinyl disc, pink rose border and center record cover placeholder.\n"
            "- Product images include personalization examples, room mockups, option panels and close-up record artwork views.\n"
            "- SEO copy stays tied to visible floral record artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "spanish-abc-classroom-rug-d3f56c3b42": {
        "title": "Spanish Second Grade Classroom Rug with Pencil Art",
        "meta": "Personalize a Spanish second grade classroom rug with Bienvenidos text, Maestra Camila name, pencil art and flowers.",
        "description": (
            "Personalize a Spanish second grade classroom rug with Bienvenidos text, Maestra Camila name, pencil art and flowers.\n\n"
            "Design details\n"
            "- Green and yellow classroom rug artwork reads Bienvenidos 2° Grado Maestra Camila with a large pencil graphic.\n"
            "- Gallery images show classroom and child learning scenes plus white flower details and bright school colors.\n"
            "- SEO copy focuses on visible Spanish classroom artwork without unsupported material, backing or safety claims."
        ),
    },
    "custom-abc-spanish-educational-rug-6cefdc1a95": {
        "title": "Spanish Educational Classroom Rug with Crayon Prompts",
        "meta": "Decorate with a Spanish educational classroom rug featuring Este Ciclo Escolar text, colorful crayons and prompt words.",
        "description": (
            "Decorate with a Spanish educational classroom rug featuring Este Ciclo Escolar text, colorful crayons and prompt words.\n\n"
            "Design details\n"
            "- Black classroom rug artwork reads Este Ciclo Escolar with crayon prompts including juega, lee, crea, sueña and imagina.\n"
            "- Product images include classroom mockups, child learning scenes, circular detail panels and bright school colors.\n"
            "- SEO copy stays tied to visible Spanish educational artwork without unsupported material, backing or safety claims."
        ),
    },
    "custom-octopus-welcome-mat-for-front-door-and-entryway-f282a58169": {
        "title": "Blue Octopus Coastal Rug with Compass Map Art",
        "meta": "Decorate with a blue octopus coastal rug featuring seashells, compass map texture, seaweed and teal tentacle artwork.",
        "description": (
            "Decorate with a blue octopus coastal rug featuring seashells, compass map texture, seaweed and teal tentacle artwork.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows a large blue octopus over distressed map texture with shells, compass and seaweed details.\n"
            "- Gallery images include room mockups, family scenes, close-up panels and size reference graphics.\n"
            "- SEO copy focuses on visible coastal octopus artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "custom-octopus-welcome-mat-for-front-door-and-entryway-13d8d7f994": {
        "title": "Underwater Octopus Rug with Coral Reef Artwork",
        "meta": "Decorate with an underwater octopus rug featuring orange tentacles, blue ocean colors, coral reef plants and fish.",
        "description": (
            "Decorate with an underwater octopus rug featuring orange tentacles, blue ocean colors, coral reef plants and fish.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows an orange octopus in a vivid underwater scene with coral, fish and reef plants.\n"
            "- Product images include bedroom, living room and close-up views plus size reference panels.\n"
            "- SEO copy stays tied to visible underwater octopus artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "custom-golf-welcome-mat-bf5c20e2dd": {
        "title": "Vintage Golf Pattern Rug with Clubs and Bags Art",
        "meta": "Decorate with a vintage golf pattern rug featuring clubs, golf bags, balls, green icons and tan sports artwork.",
        "description": (
            "Decorate with a vintage golf pattern rug featuring clubs, golf bags, balls, green icons and tan sports artwork.\n\n"
            "Design details\n"
            "- Rectangular tan rug artwork shows an all-over vintage golf pattern with clubs, bags, balls and putting green icons.\n"
            "- Gallery images include hallway, bedroom and living room mockups plus close-up and size reference panels.\n"
            "- SEO copy focuses on visible golf illustration artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "custom-octopus-welcome-mat-coastal-beach-design-9b577a6e0a": {
        "title": "Orange Octopus Coastal Rug with Dark Floral Art",
        "meta": "Decorate with an orange octopus coastal rug featuring golden flowers, shells, seaweed and black ocean background.",
        "description": (
            "Decorate with an orange octopus coastal rug featuring golden flowers, shells, seaweed and black ocean background.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows a large orange octopus with golden floral accents, shells and seaweed details.\n"
            "- Product images include bedroom, living room and close-up views plus size reference panels.\n"
            "- SEO copy stays tied to visible coastal octopus artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "custom-coastal-octopus-welcome-mat-for-entryway-and-bathroom-e11dfb0c2f": {
        "title": "Gold Octopus Floral Coastal Rug with Navy Background",
        "meta": "Decorate with a gold octopus floral coastal rug featuring cream flowers, navy background and seaweed accents.",
        "description": (
            "Decorate with a gold octopus floral coastal rug featuring cream flowers, navy background and seaweed accents.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows a gold octopus with cream and tan floral patterning on a dark navy background.\n"
            "- Gallery images include room mockups, close-up panels, family scenes and size reference graphics.\n"
            "- SEO copy focuses on visible coastal octopus artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "custom-octopus-welcome-mat-87bf0d33b5": {
        "title": "Octopus Skull Coastal Rug with Teal Tentacles",
        "meta": "Decorate with an octopus skull coastal rug featuring teal tentacles, suction cup details and dark ocean artwork.",
        "description": (
            "Decorate with an octopus skull coastal rug featuring teal tentacles, suction cup details and dark ocean artwork.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows a large skull above teal octopus tentacles with dark ocean-style background.\n"
            "- Product images include room mockups, close-up panels, family scenes and size reference graphics.\n"
            "- SEO copy stays tied to visible skull octopus artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "custom-octopus-coastal-welcome-mat-330b515744": {
        "title": "Vintage Cream Octopus Coastal Rug with Teal Art",
        "meta": "Decorate with a vintage cream octopus coastal rug featuring teal nautical artwork, antique border and illustrated tentacles.",
        "description": (
            "Decorate with a vintage cream octopus coastal rug featuring teal nautical artwork, antique border and illustrated tentacles.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows a cream octopus with teal accents, antique illustrated style and dark framed border.\n"
            "- Gallery images include room mockups, close-up views, family scenes and size reference panels.\n"
            "- SEO copy focuses on visible vintage octopus artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip",
    "non slip", "non-skid", "non skid", "kid friendly", "pet friendly",
    "safe durable", "machine-washable", "machine washable", "washable ",
    "quick-dry", "quick dry", "memory foam", "microfiber", "velvet",
    "stain", "fade resistant", "easy clean", "soft bound", "bound edges",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title = item["title"]
    meta = item["meta"]
    blob = " ".join([item["title"], item["meta"], item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R014 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R014 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R014"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R014", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R014",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R015 or Re-QA revision R014",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R014 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R014.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([
            str(title or ""),
            str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
        ]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70:
            raise RuntimeError(f"{handle} title len failed after save")
        if len(meta) > 320:
            raise RuntimeError(f"{handle} meta len failed after save")
        if str(row[out_cols["revision"] - 1]) != "2":
            raise RuntimeError(f"{handle} revision not 2")
        if row[out_cols["review_status"] - 1] != "NEEDS_REVIEW":
            raise RuntimeError(f"{handle} review_status incorrect")
        if row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} content_qa_status incorrect")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles = {}
    dup_meta_titles = []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt = row[out_cols["meta_title_seo"] - 1]
        handle = row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle

    result = {
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r014_log_rows": counts.get("R014", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

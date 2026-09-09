from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B020"
REVISION = "R096"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "feelings-wheel-emotions-round-rug-f48bff995a": ("Color Zones Feelings Wheel Rug", "color zones feelings wheel rug", "emotions wheel classroom rug, feelings chart round rug, SEL learning mat", "color-coded feeling zones and emotion words for a classroom, playroom, counseling space, or calm corner"),
    "wheel-of-feelings-and-emotions-round-rug-a7bf6ef454": ("Emotions Activity Wheel Rug", "emotions activity wheel rug", "feelings activity classroom rug, emotions wheel mat, SEL round rug", "an emotions activity wheel with colorful prompts for a classroom, office, playroom, or social-emotional learning space"),
    "wheel-of-feelings-and-emotions-round-rug-2ec976598b": ("Today I Am Feeling Classroom Rug", "Today I Am Feeling classroom rug", "today I am feeling rug, emotions classroom mat, feelings wheel rug", "Today I Am Feeling wording with expressive emotion faces for a classroom, counseling office, playroom, or calm corner"),
    "custom-a-good-day-to-read-book-rug-classroom-library-rugs-04ab7bfcca": ("New Chapter Reading Classroom Rug", "New Chapter reading classroom rug", "reading classroom rug, library book nook rug, literacy learning mat", "New Chapter wording and book-themed artwork for a classroom library, reading nook, bedroom, or book corner"),
    "custom-mathematics-education-rug-classroom-playroom-eee0004975": ("Math Dance Move Classroom Rug", "Math Dance Move classroom rug", "math classroom rug, kids math learning mat, educational playroom rug", "Math Dance Move wording with colorful math learning elements for a classroom, playroom, or study area"),
    "custom-mathematics-education-rug-3015a3c96c": ("Problem Solve Math Classroom Rug", "Problem Solve math classroom rug", "problem solving classroom rug, math learning mat, educational kids rug", "Problem Solve wording with math prompts for a classroom, study area, playroom, or learning corner"),
    "feelings-wheel-emotions-round-rug-fca05b7c0f": ("Detailed Emotions Feelings Wheel Rug", "detailed emotions feelings wheel rug", "detailed feelings classroom rug, emotions wheel mat, SEL learning rug", "a detailed feelings wheel with expressive faces and emotion words for a classroom, playroom, counseling space, or calm corner"),
    "feelings-wheel-round-classroom-rug-4be50a20fb": ("How Are You Feeling Classroom Rug", "How Are You Feeling classroom rug", "how are you feeling rug, therapy office classroom mat, emotions wheel rug", "How Are You Feeling wording with emotion prompts for a therapy office, classroom, playroom, or counseling corner"),
    "feelings-wheel-classroom-welcome-mat-3d070ff949": ("My Feelings Classroom Rug", "My Feelings classroom rug", "my feelings classroom mat, calm corner rug, emotions learning rug", "My Feelings wording and colorful emotion prompts for a classroom, kids room, playroom, or calm corner"),
    "custom-composition-notebook-classroom-kids-rug-3bd742c8b9-3bd742c8b9": ("Personalized Composition Notebook Rug", "personalized composition notebook rug", "custom notebook classroom rug, alphabet learning rug, personalized kids room mat", "a personalized composition notebook design with alphabet details for a classroom, kids room, playroom, or study space"),
}

def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []

def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)

def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B020 scope does not match the R096 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        meta = f"Create a {primary} with {motif}."
        description = f"Create a {primary} with {motif}. The design adds a colorful learning focal point to a classroom, playroom, counseling space, bedroom, or study area."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))

if __name__ == "__main__":
    main()

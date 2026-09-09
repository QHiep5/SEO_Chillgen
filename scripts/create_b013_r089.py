from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B013"
REVISION = "R089"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"


COPY = {
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-100": {
        "title": "Round Learn Shapes Classroom Rug",
        "meta": "Create a round learn shapes classroom rug with labeled shape panels, cheerful colors, and playful early-learning artwork.",
        "description": "Create a round learn shapes classroom rug with labeled shape panels, cheerful colors, and playful early-learning artwork. The design gives a preschool, playroom, or classroom floor a bright reference for shape recognition.",
        "primary": "round learn shapes classroom rug",
        "secondary": "learn shapes rug, preschool shapes rug, educational classroom rug",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-101": {
        "title": "Emoji Feelings Wheel Classroom Rug",
        "meta": "Create an emoji feelings wheel classroom rug with How Are You Feeling text, emotion faces, and colorful SEL prompts.",
        "description": "Create an emoji feelings wheel classroom rug with How Are You Feeling text, emotion faces, and colorful SEL prompts. The round layout gives students a friendly visual reference for naming emotions during classroom routines.",
        "primary": "emoji feelings wheel classroom rug",
        "secondary": "emoji feelings wheel rug, SEL classroom rug, emotions classroom mat",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-102": {
        "title": "How Do You Feel Today Rug",
        "meta": "Create a How Do You Feel Today rug with colorful emotion faces, bold lettering, and a playful doodle-style border.",
        "description": "Create a How Do You Feel Today rug with colorful emotion faces, bold lettering, and a playful doodle-style border. The design helps a classroom or counseling corner introduce simple emotion words in a welcoming way.",
        "primary": "how do you feel today rug",
        "secondary": "feelings classroom rug, emotions rug for kids, SEL classroom mat",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-103": {
        "title": "Round Days of the Week Classroom Rug",
        "meta": "Create a round days of the week classroom rug with Monday through Sunday labels, activity icons, and bright calendar colors.",
        "description": "Create a round days of the week classroom rug with Monday through Sunday labels, activity icons, and bright calendar colors. The design gives young learners a cheerful floor reference for weekly classroom routines.",
        "primary": "round days of the week classroom rug",
        "secondary": "days of the week rug, calendar classroom rug, educational round rug",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-104": {
        "title": "Kids Feelings and Emotions Rug",
        "meta": "Create a kids feelings and emotions rug with cartoon emotion panels, expressive faces, and colorful SEL vocabulary.",
        "description": "Create a kids feelings and emotions rug with cartoon emotion panels, expressive faces, and colorful SEL vocabulary. The design gives children a friendly way to see and name common feelings in classroom or playroom spaces.",
        "primary": "kids feelings and emotions rug",
        "secondary": "feelings rug for kids, emotions classroom rug, SEL emotion rug",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-105": {
        "title": "Round Weather and Seasons Rug",
        "meta": "Create a round weather and seasons rug with sunny, rainy, windy, and snowy panels around a spring-to-winter center wheel.",
        "description": "Create a round weather and seasons rug with sunny, rainy, windy, and snowy panels around a spring-to-winter center wheel. The design gives a classroom or playroom a colorful reference for weather words and seasonal change.",
        "primary": "round weather and seasons rug",
        "secondary": "weather classroom rug, seasons classroom rug, educational round rug",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-106": {
        "title": "Round Learn Colors Classroom Rug",
        "meta": "Create a round learn colors classroom rug with labeled color slices, crayon characters, and bright early-learning artwork.",
        "description": "Create a round learn colors classroom rug with labeled color slices, crayon characters, and bright early-learning artwork. The design gives preschool and classroom spaces a cheerful visual reference for color words.",
        "primary": "round learn colors classroom rug",
        "secondary": "learn colors rug, preschool colors rug, educational classroom rug",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-107": {
        "title": "Animal Feelings Classroom Rug",
        "meta": "Create an animal feelings classroom rug with animal character faces, emotion labels, and cheerful SEL color.",
        "description": "Create an animal feelings classroom rug with animal character faces, emotion labels, and cheerful SEL color. The playful round design gives children a friendly reference for connecting feelings with expressive faces.",
        "primary": "animal feelings classroom rug",
        "secondary": "animal feelings rug, SEL classroom rug, emotions rug for kids",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-108": {
        "title": "Round Months and Seasons Rug",
        "meta": "Create a round months and seasons rug with January through December labels, a seasonal tree center, and calendar-themed color.",
        "description": "Create a round months and seasons rug with January through December labels, a seasonal tree center, and calendar-themed color. The design gives a classroom or playroom a bright floor reference for months and seasonal change.",
        "primary": "round months and seasons rug",
        "secondary": "months classroom rug, seasons learning rug, educational round rug",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-109": {
        "title": "Animal Friends Classroom Rug",
        "meta": "Create a round animal friends classroom rug with labeled animal panels, cheerful character artwork, and bright classroom color.",
        "description": "Create a round animal friends classroom rug with labeled animal panels, cheerful character artwork, and bright classroom color. The design gives early-learning spaces a playful floor reference for animal names.",
        "primary": "animal friends classroom rug",
        "secondary": "animal classroom rug, animal names rug, preschool animal rug",
    },
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B013 scope does not match the R089 copy map")

    for row in seo_rows:
        item = COPY[row["Handle"]]
        row["title_proposed"] = item["title"]
        row["meta_title_seo"] = item["title"]
        row["meta_title_chars"] = str(len(item["title"]))
        row["meta_description_seo"] = item["meta"]
        row["meta_description_chars"] = str(len(item["meta"]))
        row["description_proposed"] = item["description"]
        row["description_proposed_html"] = f"<p>{item['description']}</p>"
        row["primary_keyword"] = item["primary"]
        row["secondary_keywords"] = item["secondary"]
        row["long_tail_candidates"] = f"{item['primary']}; {item['secondary']}"
        row["revision"] = "2"
        row["review_status"] = "NEEDS_REVIEW"
        row["content_qa_status"] = "NOT_RUN"
        row["review_reason"] = f"{REVISION} removes internal evidence-note language and unsupported claim wording from B013 copy."
        row["issues"] = f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations."

    handles_set = set(handles)
    evidence_scope = [row for row in evidence_rows if row.get("evidence_id") in {r.get("evidence_id") for r in seo_rows}]
    keyword_scope = [row for row in keyword_rows if row.get("product_key") in handles_set]
    buyer_rows = [row for row in buyer_all if row.get("product_key") in handles_set]
    img_scope = [row for row in img_rows if row.get("Handle") in handles_set]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": "Canonical B013 cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    evidence_ids = {row["evidence_id"] for row in evidence_scope}
    for row in seo_rows:
        evidence_id = row.get("evidence_id")
        if evidence_id not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {evidence_id}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", evidence_id])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    digest = sha256(OUT)
    manifest = {
        "revision": REVISION,
        "batch_id": BATCH,
        "scope_count": len(handles),
        "image_count": len(img_scope),
        "canonical_workbook": str(OUT),
        "sha256": digest,
        "status": "CREATED_FOR_HARDENED_REQA",
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
    }
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

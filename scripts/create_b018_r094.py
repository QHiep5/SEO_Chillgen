from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B018"
REVISION = "R094"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "personalized-family-couple-doormat-custom-couple-husband-design-04": ("Funny Couple Welcome Doormat", "funny couple welcome doormat", "custom couple doormat, funny personalized entryway mat, husband wife welcome mat", "a playful couple message with personalized names for a front door, entryway, or shared home"),
    "personalized-family-couple-doormat-custom-couple-husband-design-02": ("Personalized Family Couple Doormat", "personalized family couple doormat", "custom family doormat, personalized couple welcome mat, family entryway decor", "a cheerful family-and-animals design with personalized couple details for a welcoming home entryway"),
    "personalized-family-couple-doormat-custom-couple-husband-wife-doormat": ("Personalized Couple Welcome Mat", "personalized couple welcome mat", "custom couple welcome doormat, personalized home entryway mat, couple housewarming gift", "a personalized couple welcome message for a front door, porch, or shared home entryway"),
    "abc-kids-play-rug-non-slip-classroom-carpet-07621e8f46-07621e8f46": ("ABC Animal Kids Play Rug", "ABC animal kids play rug", "alphabet animal playroom rug, ABC classroom rug, kids learning area rug", "an animal alphabet pattern with letters, shapes, and numbers for a playroom, preschool, or kids learning space"),
    "custom-feelings-chart-classroom-welcome-mat-45af89adb4-45af89adb4": ("All Feelings Welcome Classroom Rug", "all feelings welcome classroom rug", "feelings chart rug, emotions classroom rug, calming corner classroom mat", "a feelings chart with welcoming wording and emotion-focused artwork for a classroom, playroom, or calming corner"),
    "custom-hundred-acre-wood-map-welcome-mat-23388cf0b9-23388cf0b9": ("Hundred Acre Wood Map Kids Rug", "Hundred Acre Wood map kids rug", "storybook map rug, woodland kids room rug, custom playroom area rug", "a Hundred Acre Wood map design with storybook details for a bedroom, playroom, or reading space"),
    "educational-classroom-rug-alphabet-handwriting-samplers-17817a1b4c-17817a1b4c": ("Alphabet Handwriting Classroom Rug", "alphabet handwriting classroom rug", "ABC handwriting rug, preschool alphabet classroom mat, letter learning rug", "an alphabet and handwriting sampler layout for preschool rooms, classrooms, or early-learning spaces"),
    "colorful-classroom-rug-for-kids-and-teachers-940ca1aee9-940ca1aee9": ("Feelings Zone Classroom Rug", "feelings zone classroom rug", "emotion zones classroom rug, feelings chart mat, SEL classroom rug", "colorful feelings zones with emotion prompts for a classroom, playroom, or social-emotional learning area"),
    "custom-soccer-championship-trophy-round-rug-f7ac8bf3d4-f7ac8bf3d4": ("Soccer Host Flags Trophy Rug", "soccer host flags trophy rug", "soccer trophy round rug, football championship rug, sports room floor decor", "soccer host flags, a championship trophy, and a round sports design for a bedroom, game room, or fan space"),
    "custom-soccer-championship-round-rug-6b6ea4fbf7": ("Soccer Tournament Trophy Rug", "soccer tournament trophy rug", "soccer championship rug, tournament round rug, sports fan room decor", "a soccer tournament trophy and championship-themed artwork for a bedroom, game room, or sports fan space"),
}

def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []

def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)

def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B018 scope does not match the R094 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        meta = f"Create a {primary} with {motif}."
        description = f"Create a {primary} with {motif}. The design adds a distinctive focal point to a classroom, playroom, bedroom, entryway, game room, or shared home."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))

if __name__ == "__main__":
    main()

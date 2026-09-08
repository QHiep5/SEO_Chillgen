from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R018/SEO_Product_Optimization_revision_R018.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R019"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R019.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R019"
MANIFEST = RUN_DIR / "revision_R019_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R019_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-teacher-classroom-rug-custom-back-to-school-welcome-mat": {
        "title": "Personalized Classroom Welcome Rug with Affirmations",
        "meta": "Customize a classroom welcome rug with teacher name text, colorful affirmations, school icons and black rug artwork.",
        "description": (
            "Customize a classroom welcome rug with teacher name text, colorful affirmations, school icons and black rug artwork.\n\n"
            "Design details\n"
            "- Black classroom rug artwork shows In This Classroom You Are wording with teacher-name classroom strip.\n"
            "- Visible labels include loved, kind, special, important and courageous, with colorful school characters and supplies.\n"
            "- SEO copy describes visible classroom artwork only, without safety, therapeutic, material, cleaning or surface-performance claims."
        ),
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-108": {
        "title": "Personalized Rainbow Welcome Rug with Affirmation Circles",
        "meta": "Customize a rainbow welcome classroom rug with teacher name text, scallop border, rainbow icons and affirmation circles.",
        "description": (
            "Customize a rainbow welcome classroom rug with teacher name text, scallop border, rainbow icons and affirmation circles.\n\n"
            "Design details\n"
            "- Bright white classroom rug artwork shows Welcome to teacher-name classroom wording with rainbow scallop border.\n"
            "- Visible circle text includes kind, amazing, smart, brave and brilliant with rainbow and school icon details.\n"
            "- SEO copy describes visible classroom artwork only, without safety, therapeutic, material, cleaning or surface-performance claims."
        ),
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-109": {
        "title": "Pastel Teacher Welcome Rug with Yellow Stripe Art",
        "meta": "Customize a pastel teacher welcome rug with classroom name text, yellow stripes, bows, rainbow, ABC blocks and pencil art.",
        "description": (
            "Customize a pastel teacher welcome rug with classroom name text, yellow stripes, bows, rainbow, ABC blocks and pencil art.\n\n"
            "Design details\n"
            "- Pastel yellow and pink classroom rug artwork shows Welcome to teacher-name classroom wording with bow details.\n"
            "- Visible icons include rainbow, sun, ABC blocks, pencil, clouds, smiley faces and flowers.\n"
            "- SEO copy focuses on visible teacher welcome artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-110": {
        "title": "School Icons Teacher Welcome Rug with Chalkboard Art",
        "meta": "Customize a teacher classroom welcome rug with name text, school bus, books, globe, scissors and ABC block icons.",
        "description": (
            "Customize a teacher classroom welcome rug with name text, school bus, books, globe, scissors and ABC block icons.\n\n"
            "Design details\n"
            "- Black classroom rug artwork has a white label reading Welcome to teacher-name classroom, surrounded by school icons.\n"
            "- Visible icons include school bus, books, globe, scissors, ABC blocks, smiley faces, stars and hearts.\n"
            "- SEO copy stays tied to visible classroom welcome artwork without unsupported material, cleaning, reverse-side or surface-performance claims."
        ),
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-111": {
        "title": "Affirmation Rainbow Classroom Rug with Teacher Name",
        "meta": "Customize an affirmation rainbow classroom rug with teacher name text, apple, books, globe, bunting and student lines.",
        "description": (
            "Customize an affirmation rainbow classroom rug with teacher name text, apple, books, globe, bunting and student lines.\n\n"
            "Design details\n"
            "- Black chalkboard-style classroom rug artwork shows Welcome to teacher-name classroom wording with a rainbow section.\n"
            "- Visible elements include apple, books, globe, pencil cup, bunting, stars, clouds and affirmation-style rainbow lines.\n"
            "- SEO copy describes visible classroom artwork only, without safety, therapeutic, material, cleaning or surface-performance claims."
        ),
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-classroom-rug": {
        "title": "Core Vocabulary Communication Rug with Symbol Grid",
        "meta": "View a core vocabulary communication rug design with colorful picture-symbol cells, word labels and classroom scenes.",
        "description": (
            "View a core vocabulary communication rug design with colorful picture-symbol cells, word labels and classroom scenes.\n\n"
            "Design details\n"
            "- Rectangular rug artwork shows a colorful grid of picture-symbol and word cells for classroom-style vocabulary themes.\n"
            "- Visible labels include greetings, bathroom, yes, no, help, more, eat, drink, go, stop, feelings and people/action words.\n"
            "- SEO copy describes visible symbol-grid artwork only, without educational outcome, safety, therapeutic, material or cleaning claims."
        ),
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-100": {
        "title": "AAC Core Board Classroom Rug with Colored Symbol Cells",
        "meta": "View an AAC core board classroom rug design with colored symbol cells, category sections, number row and classroom icons.",
        "description": (
            "View an AAC core board classroom rug design with colored symbol cells, category sections, number row and classroom icons.\n\n"
            "Design details\n"
            "- Large classroom rug artwork shows an AAC-style communication board layout with many colored symbol cells.\n"
            "- Visible details include category-colored sections, a bottom number row from 1 through 9 and action/classroom icons.\n"
            "- SEO copy describes visible board artwork only, without educational outcome, safety, therapeutic, material or cleaning claims."
        ),
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-101": {
        "title": "SPED Communication Board Rug with Picture Word Cells",
        "meta": "View a SPED communication board rug design with pastel picture-word cells for feelings, needs and classroom routines.",
        "description": (
            "View a SPED communication board rug design with pastel picture-word cells for feelings, needs and classroom routines.\n\n"
            "Design details\n"
            "- Black-bordered rug artwork shows pastel picture-word cells arranged in a classroom communication board layout.\n"
            "- Visible labels include sick, thirsty, scared, go to bathroom, play time, change diaper, wash up, eat, toilet and toy.\n"
            "- SEO copy describes visible board artwork only, without educational outcome, safety, therapeutic, material or cleaning claims."
        ),
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-103": {
        "title": "Sign Language Communication Rug with Hand Sign Panels",
        "meta": "View a sign language communication rug design with hand-sign illustrations, word panels and classroom scene images.",
        "description": (
            "View a sign language communication rug design with hand-sign illustrations, word panels and classroom scene images.\n\n"
            "Design details\n"
            "- White communication rug artwork shows hand-sign illustrations with word panels in a rectangular grid.\n"
            "- Visible labels include yes, no, please, thank you, sorry, help, more, eat, drink, restroom, stop and go.\n"
            "- SEO copy describes visible hand-sign artwork only, without educational outcome, safety, therapeutic, material or cleaning claims."
        ),
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-104": {
        "title": "Something Hurts Communication Rug with Body Diagram",
        "meta": "View a Something Hurts communication rug design with body silhouette, pain icons, needs labels and pastel classroom art.",
        "description": (
            "View a Something Hurts communication rug design with body silhouette, pain icons, needs labels and pastel classroom art.\n\n"
            "Design details\n"
            "- White communication rug artwork shows Something Hurts headline with a central body silhouette and surrounding icon labels.\n"
            "- Visible labels include body parts and needs such as dry mouth, hand, finger, knee, toe and private area.\n"
            "- SEO copy describes visible diagram artwork only, without medical, safety, therapeutic, material or cleaning claims."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip",
    "non slip", "non-skid", "non skid", "kid friendly", "pet friendly",
    "safe durable", "machine-washable", "machine washable", "washable ",
    "quick-dry", "quick dry", "memory foam", "microfiber", "velvet",
    "stain", "fade resistant", "easy clean", "soft fabric", "backing",
    "cleaning visuals", "cleaning graphic", "thickened", "bound edge",
    "support ", "supports ", "use a ", "medical claim", "therapy",
    "therapeutic benefit",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title = item["title"]
    meta = item["meta"]
    blob = " ".join([item["title"], item["meta"], item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R019 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R019 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R019"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R019", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R019",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R020 or Re-QA revision R019",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R019 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R019.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([
            str(title or ""),
            str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
        ]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70:
            raise RuntimeError(f"{handle} title len failed after save")
        if len(meta) > 320:
            raise RuntimeError(f"{handle} meta len failed after save")
        if str(row[out_cols["revision"] - 1]) != "2":
            raise RuntimeError(f"{handle} revision not 2")
        if row[out_cols["review_status"] - 1] != "NEEDS_REVIEW":
            raise RuntimeError(f"{handle} review_status incorrect")
        if row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} content_qa_status incorrect")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles = {}
    dup_meta_titles = []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt = row[out_cols["meta_title_seo"] - 1]
        handle = row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle

    result = {
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r019_log_rows": counts.get("R019", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R003/SEO_Product_Optimization_revision_R003.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R004"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R004.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R004"
MANIFEST = RUN_DIR / "revision_R004_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R004_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "wheel-of-feelings-and-emotions-round-rug-2ec976598b": {
        "title": "Today I Am Feeling Classroom Rug with Rainbow Emotions",
        "meta": "Help students name emotions with a Today I Am Feeling classroom rug featuring rainbow feeling segments and character faces.",
        "description": (
            "Help students name emotions with a Today I Am Feeling classroom rug featuring rainbow feeling segments and character faces.\n\n"
            "Design details\n"
            "- Round rainbow feelings rug with Today I'm Feeling center text and character faces around the color wheel.\n"
            "- Visible labels include happy, shy, scared, calm, tired, worried, angry and sad.\n"
            "- SEO copy stays tied to the visible emotion-check design and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-road-map-area-rug-town-play-mat-587b1148b8-587b1148b8": {
        "title": "Personalized Blue Town Road Map Rug with Cars and Train",
        "meta": "Personalize a blue town road map rug with custom name text, winding roads, cars, houses, train, signs and trees.",
        "description": (
            "Personalize a blue town road map rug with custom name text, winding roads, cars, houses, train, signs and trees.\n\n"
            "Design details\n"
            "- Navy blue town road map rug sample with Lucas name, roads, houses, train, cars, stop signs, trees and water details.\n"
            "- Gallery images show the play-mat artwork in classroom and playroom-style scenes with size and design visuals.\n"
            "- SEO copy focuses on the personalized town road design without unsupported material, care, backing or surface-performance claims."
        ),
    },
    "custom-road-map-area-rug-e4dbf01f05": {
        "title": "Personalized Pastel City Road Map Rug for Playrooms",
        "meta": "Personalize a pastel city road map rug with custom town name, soft roads, houses, parks, cars, buses and pond details.",
        "description": (
            "Personalize a pastel city road map rug with custom town name, soft roads, houses, parks, cars, buses and pond details.\n\n"
            "Design details\n"
            "- Cream and green city road map rug sample with Emily's City text, pastel streets, houses, parks, trees, buses, cars and pond.\n"
            "- Gallery shows the road-map design in classroom, playroom and home-style scenes with size visuals.\n"
            "- SEO copy stays tied to the visible pastel city play design and avoids unsupported material, care, backing or surface-performance claims."
        ),
    },
    "custom-road-map-doormat-city-street-play-mat-0e3ba7dda0-0e3ba7dda0": {
        "title": "Personalized Work Zone Road Map Rug with Construction Trucks",
        "meta": "Personalize a work zone road map rug with custom name text, construction trucks, streets, traffic signs and building icons.",
        "description": (
            "Personalize a work zone road map rug with custom name text, construction trucks, streets, traffic signs and building icons.\n\n"
            "Design details\n"
            "- White and gray road map rug sample with Robert's Work Zone text, construction trucks, excavators, road signs, cones, buildings and trees.\n"
            "- Gallery images show the city-street play artwork in classroom and playroom-style mockups.\n"
            "- SEO copy focuses on the visible construction road-map design without unsupported material, care, backing or surface-performance claims."
        ),
    },
    "custom-road-map-area-rug-d39d38a4a8": {
        "title": "Personalized Soft City Road Play Rug with Town Name",
        "meta": "Personalize a soft city road play rug with custom town name, gray roads, parking spaces, houses, cars and green areas.",
        "description": (
            "Personalize a soft city road play rug with custom town name, gray roads, parking spaces, houses, cars and green areas.\n\n"
            "Design details\n"
            "- Pastel city road rug sample with Liam's Town text, gray streets, parking lots, houses, trees, roundabouts, cars and water details.\n"
            "- Gallery visuals show the play-mat layout in classroom, nursery and playroom-style settings.\n"
            "- SEO copy stays focused on the visible personalized city road design and avoids unsupported material, care, backing or surface-performance claims."
        ),
    },
    "custom-road-map-area-rug-60aa327740": {
        "title": "Personalized ABC Town Road Map Rug with Alphabet Border",
        "meta": "Personalize an ABC town road map rug with custom name, alphabet border, traffic signs, cars, buildings and learning icons.",
        "description": (
            "Personalize an ABC town road map rug with custom name, alphabet border, traffic signs, cars, buildings and learning icons.\n\n"
            "Design details\n"
            "- Bright blue road map rug sample with Liam's Town text, alphabet border, traffic signs, roads, buildings, trees, animals and cars.\n"
            "- Gallery images show classroom and playroom scenes with the colorful city-street learning artwork.\n"
            "- SEO copy focuses on the visible ABC town road design without unsupported material, care, backing or surface-performance claims."
        ),
    },
    "custom-family-name-welcome-doormat-front-door-6831756395-6831756395": {
        "title": "Good Day To Read Classroom Rug with Floral Book Art",
        "meta": "Create a cheerful reading corner with a Good Day To Read classroom rug featuring floral letters, leaves and open book artwork.",
        "description": (
            "Create a cheerful reading corner with a Good Day To Read classroom rug featuring floral letters, leaves and open book artwork.\n\n"
            "Design details\n"
            "- Cream classroom rug with It Is a Good Day to Read a Book text, colorful floral letters, leaf motifs and open book graphic.\n"
            "- Gallery images show classroom, reading nook and circle-time style scenes plus close-up design panels.\n"
            "- SEO copy stays tied to the visible reading-themed rug artwork and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-vintage-baseball-area-rug-b5f20ad4f5": {
        "title": "Personalized Space Baseball Area Rug with Name and Number",
        "meta": "Customize a space baseball area rug with player name, number 22, glowing baseball artwork and bold varsity lettering.",
        "description": (
            "Customize a space baseball area rug with player name, number 22, glowing baseball artwork and bold varsity lettering.\n\n"
            "Design details\n"
            "- Black and blue cosmic baseball rug sample with glowing baseball, Sophia name, number 22 and bold pink varsity lettering.\n"
            "- Gallery images show sports room and bedroom-style mockups plus close-up design and size visuals.\n"
            "- SEO copy focuses on the visible personalized baseball artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "colorful-classroom-rug-for-kids-aaa73255b2-aaa73255b2": {
        "title": "Personalized BE Values Classroom Rug with Colorful Arrows",
        "meta": "Personalize a BE values classroom rug with teacher name text, colorful arrow words and student character traits.",
        "description": (
            "Personalize a BE values classroom rug with teacher name text, colorful arrow words and student character traits.\n\n"
            "Design details\n"
            "- Black classroom rug with large yellow BE center, Mrs. Asher's Class text and colorful arrows with character words.\n"
            "- Visible words include kind, honest, thankful, creative, unique, friendly, compassionate and other positive traits.\n"
            "- SEO copy stays focused on the visible classroom values design and avoids unsupported material, backing or surface-performance claims."
        ),
    },
    "colorful-classroom-rug-for-kids-877e4e6020-877e4e6020": {
        "title": "You Matter Classroom Rug with Rainbow Handprint Art",
        "meta": "Encourage students with a You Matter classroom rug featuring rainbow artwork, handprints, paint splashes and affirmations.",
        "description": (
            "Encourage students with a You Matter classroom rug featuring rainbow artwork, handprints, paint splashes and affirmations.\n\n"
            "Design details\n"
            "- Black affirmation classroom rug with rainbow, handprints, paint splashes and colorful student affirmation words.\n"
            "- Visible text includes unique, smart, talented, worth it, special, valued, amazing, loved and matter.\n"
            "- SEO copy focuses on the visible classroom affirmation design without unsupported material, backing or surface-performance claims."
        ),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R004 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R004 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    for sheet_name in ["Buyer_Search_Research", "Keyword_Map"]:
        if sheet_name not in wb.sheetnames:
            continue
        sh = wb[sheet_name]
        hmap = headers(sh)
        key_col = hmap.get("product_key")
        if not key_col:
            continue
        for row_idx in range(2, sh.max_row + 1):
            pk = str(sh.cell(row_idx, key_col).value or "")
            if not any(handle in pk for handle in REVISIONS):
                continue
            for col_idx in range(1, sh.max_column + 1):
                value = sh.cell(row_idx, col_idx).value
                if not isinstance(value, str):
                    continue
                sh.cell(row_idx, col_idx, value.replace("indoor/outdoor", "indoor").replace("Indoor/outdoor", "Indoor").replace("outdoor-use", "room-use").replace("outdoor use", "room use"))

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R004"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R004", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R004",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R005 or Re-QA revision R004",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R004 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R004.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _, _, _, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n", encoding="utf-8")

    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    vws = verify["SEO_Products"]
    vcols = {cell.value: idx for idx, cell in enumerate(next(vws.iter_rows(min_row=1, max_row=1)))}
    bad = []
    for row in vws.iter_rows(min_row=2, values_only=True):
        handle = row[vcols["Handle"]]
        if handle not in REVISIONS:
            continue
        title = row[vcols["title_proposed"]] or ""
        meta_title = row[vcols["meta_title_seo"]] or ""
        meta_desc = row[vcols["meta_description_seo"]] or ""
        desc = (row[vcols["description_proposed"]] or "") + " " + (row[vcols["description_proposed_html"]] or "")
        blob = " ".join([title, meta_title, meta_desc, desc]).lower()
        if row[vcols["revision"]] != "2":
            bad.append((handle, "revision", row[vcols["revision"]]))
        if row[vcols["review_status"]] != "NEEDS_REVIEW":
            bad.append((handle, "review_status", row[vcols["review_status"]]))
        if row[vcols["content_qa_status"]] != "NOT_RUN":
            bad.append((handle, "content_qa_status", row[vcols["content_qa_status"]]))
        if not (45 <= len(title) <= 70):
            bad.append((handle, "title_len", len(title)))
        if meta_title != title:
            bad.append((handle, "meta_title_mismatch", meta_title))
        if len(meta_desc) > 320:
            bad.append((handle, "meta_desc_len", len(meta_desc)))
        for forbidden in ["indoor/outdoor", " outdoor ", "anti-slip", "non-slip", "kid friendly", "pet friendly", "safe durable", "machine-washable", "machine washable", "washable "]:
            if forbidden in f" {blob} ":
                bad.append((handle, "forbidden_claim", forbidden))
    if bad:
        raise RuntimeError(f"Verification failed: {bad[:20]}")

    print(json.dumps({
        "output_workbook": str(OUTPUT),
        "manifest": str(MANIFEST),
        "summary": str(SUMMARY),
        "updated": len(changed),
        "status": "COMPLETE_AWAITING_REQA",
    }, indent=2))


if __name__ == "__main__":
    main()

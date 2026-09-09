from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B008"
REVISION = "R084"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"


COPY = {
    "custom-boat-rug-personalized-welcome-aboard-mat-with-boat-design-02": {
        "title": "Personalized Anchor Welcome Aboard Boat Rug",
        "meta": "Personalize an anchor welcome aboard boat rug with navy-and-gold nautical artwork, custom boat-name text, and coastal entryway style.",
        "description": "Personalize an anchor welcome aboard boat rug with navy-and-gold nautical artwork and custom boat-name text. The design brings a polished maritime accent to cabins, entryways, and coastal living spaces while keeping the boat identity clear.",
        "primary": "personalized anchor welcome aboard boat rug",
        "secondary": "custom boat name rug, nautical entryway rug, anchor boat mat",
    },
    "custom-boat-rug-personalized-welcome-aboard-mat-with-boat-design-09": {
        "title": "Personalized Nautical Welcome Aboard Boat Mat",
        "meta": "Customize a nautical welcome aboard boat mat with gold anchor details, boat-name text, and classic marine styling for a cabin or entryway.",
        "description": "Customize a nautical welcome aboard boat mat with gold anchor details, boat-name text, and classic marine styling. The design gives boat cabins, docks, and home entryways a personalized coastal focal point.",
        "primary": "personalized nautical welcome aboard boat mat",
        "secondary": "custom boat name mat, nautical cabin rug, anchor welcome aboard rug",
    },
    "boat-rug-05-boat-rug-05-design-22": {
        "title": "Personalized Captain Quote Boat Rug",
        "meta": "Personalize a captain quote boat rug with helm, anchor, and name details for a nautical gift or maritime room accent.",
        "description": "Personalize a captain quote boat rug with helm, anchor, and name details. The single-image design centers the captain message and nautical artwork, making it a focused gift for boating fans and coastal rooms.",
        "primary": "personalized captain quote boat rug",
        "secondary": "captain boat rug, nautical gift rug, custom captain mat",
    },
    "custom-boat-rug-personalized-welcome-aboard-mat-with-boat-design-06": {
        "title": "Captain Always Right Custom Boat Rug",
        "meta": "Customize a captain always right boat rug with captain-name text, anchor artwork, and rope-style nautical detail.",
        "description": "Customize a captain always right boat rug with captain-name text, anchor artwork, and rope-style nautical detail. The design keeps the playful captain message front and center for cabins, entryways, and boating-themed rooms.",
        "primary": "captain always right custom boat rug",
        "secondary": "custom captain boat rug, nautical captain mat, personalized boat rug",
    },
    "custom-boat-rug-personalized-welcome-aboard-mat-with-boat-design-08": {
        "title": "Personalized Captain Nautical Boat Rug",
        "meta": "Personalize a captain nautical boat rug with name text, anchor artwork, helm detail, and a navy maritime look.",
        "description": "Personalize a captain nautical boat rug with name text, anchor artwork, helm detail, and a navy maritime look. The design is tailored for boat owners and coastal decor fans who want a custom captain-themed accent.",
        "primary": "personalized captain nautical boat rug",
        "secondary": "captain name boat rug, nautical boat owner gift, custom marine rug",
    },
    "custom-boat-rug-personalized-welcome-aboard-mat-with-boat-design-07": {
        "title": "Custom Welcome Aboard Anchor Boat Rug",
        "meta": "Create a welcome aboard anchor boat rug with boat-name and location text, navy artwork, and a personalized nautical message.",
        "description": "Create a welcome aboard anchor boat rug with boat-name and location text, navy artwork, and a personalized nautical message. The layout highlights the custom boat identity for a crisp maritime entry accent.",
        "primary": "custom welcome aboard anchor boat rug",
        "secondary": "boat name welcome rug, personalized nautical mat, anchor boat rug",
    },
    "boat-rug-10-boat-rug-10-design-10": {
        "title": "Black Gold Welcome Aboard Boat Mat",
        "meta": "Personalize a black and gold welcome aboard boat mat with anchor artwork and boat-name text for a bold nautical accent.",
        "description": "Personalize a black and gold welcome aboard boat mat with anchor artwork and boat-name text. The high-contrast design keeps the welcome message easy to recognize and clearly tied to a boating theme.",
        "primary": "black gold welcome aboard boat mat",
        "secondary": "custom welcome aboard mat, black gold boat rug, personalized boat mat",
    },
    "custom-boat-rug-personalized-welcome-aboard-mat-with-boat-name": {
        "title": "Personalized Boat Name Welcome Rug",
        "meta": "Customize a boat name welcome rug with anchor artwork, welcome aboard wording, and black-and-gold nautical style.",
        "description": "Customize a boat name welcome rug with anchor artwork, welcome aboard wording, and black-and-gold nautical style. The design turns a boat name into a clear decorative statement for a cabin, entryway, or marina-inspired room.",
        "primary": "personalized boat name welcome rug",
        "secondary": "custom boat name rug, welcome aboard boat rug, nautical name mat",
    },
    "boat-rug-04-boat-rug-04": {
        "title": "Personalized Navy Compass Boat Rug",
        "meta": "Personalize a navy compass boat rug with welcome aboard text, boat-name fields, and striped nautical styling.",
        "description": "Personalize a navy compass boat rug with welcome aboard text, boat-name fields, and striped nautical styling. The compass motif gives the design a classic maritime feel while the custom text keeps it specific to the boat owner.",
        "primary": "personalized navy compass boat rug",
        "secondary": "custom compass boat rug, welcome aboard nautical mat, boat owner rug",
    },
    "personalized-tropical-beach-rug-custom-family-name-tropic-design-02": {
        "title": "Personalized Tropical Beach Family Rug",
        "meta": "Create a tropical beach family rug with custom family-name text, palm leaves, flowers, flamingos, and bright coastal color.",
        "description": "Create a tropical beach family rug with custom family-name text, palm leaves, flowers, flamingos, and bright coastal color. The personalized oasis design suits beach-house styling and tropical family decor.",
        "primary": "personalized tropical beach family rug",
        "secondary": "custom family beach rug, tropical name rug, flamingo beach mat",
    },
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B008 scope does not match the R084 copy map")

    for row in seo_rows:
        item = COPY[row["Handle"]]
        row["title_proposed"] = item["title"]
        row["meta_title_seo"] = item["title"]
        row["meta_title_chars"] = str(len(item["title"]))
        row["meta_description_seo"] = item["meta"]
        row["meta_description_chars"] = str(len(item["meta"]))
        row["description_proposed"] = item["description"]
        row["description_proposed_html"] = f"<p>{item['description']}</p>"
        row["primary_keyword"] = item["primary"]
        row["secondary_keywords"] = item["secondary"]
        row["long_tail_candidates"] = f"{item['primary']}; {item['secondary']}"
        row["revision"] = "2"
        row["review_status"] = "NEEDS_REVIEW"
        row["content_qa_status"] = "NOT_RUN"
        row["review_reason"] = f"{REVISION} removes internal evidence-note language and unsupported claim wording from B008 copy."
        row["issues"] = f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations."

    keyed_evidence = {row["evidence_id"] for row in evidence_rows}
    buyer_rows = [row for row in buyer_all if row.get("product_key") in handles]
    keyword_scope = [row for row in keyword_rows if row.get("product_key") in handles]
    img_scope = [row for row in img_rows if row.get("Handle") in handles]
    evidence_scope = [row for row in evidence_rows if row.get("evidence_id") in {r.get("evidence_id") for r in seo_rows}]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    readme = [
        {"metric": "revision", "value": REVISION, "definition": "Canonical B008 cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ]
    write_sheet(wb, "README_QA", readme, ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        evidence_id = row.get("evidence_id")
        if evidence_id not in keyed_evidence:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {evidence_id}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", evidence_id])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    digest = sha256(OUT)
    manifest = {
        "revision": REVISION,
        "batch_id": BATCH,
        "scope_count": len(handles),
        "image_count": len(img_scope),
        "canonical_workbook": str(OUT),
        "sha256": digest,
        "status": "CREATED_FOR_HARDENED_REQA",
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
    }
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

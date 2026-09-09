from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B028"
REVISION = "R104"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "custom-octopus-shaped-area-rug-sea-monster-tentacle-0753b2605a": ("Wave Octopus Shaped Rug", "wave octopus shaped rug", "ocean octopus kids rug, sea monster playroom mat, tentacle area rug", "a wave-filled octopus design with colorful tentacles for a playroom, bedroom, classroom, or kids room"),
    "custom-octopus-sea-monster-welcome-mat-c69579f5ed": ("Porthole Octopus Shaped Rug", "porthole octopus shaped rug", "ocean porthole rug, octopus welcome mat, sea life kids room rug", "an octopus and porthole design with ocean artwork for a playroom, bedroom, entryway, or kids room"),
    "custom-octopus-shaped-area-rug-a96e544960-a96e544960": ("Blue Tentacle Octopus Rug", "blue tentacle octopus rug", "blue octopus area rug, ocean kids room mat, sea monster decor rug", "a blue octopus with detailed tentacles for a playroom, bedroom, classroom, or kids room"),
    "custom-octopus-shaped-area-rug-sea-monster-1975158434": ("Cave Octopus Shaped Rug", "cave octopus shaped rug", "sea monster area rug, octopus playroom mat, ocean fantasy kids rug", "an octopus emerging from a rocky ocean scene for a playroom, bedroom, classroom, or kids room"),
    "custom-octopus-shaped-area-rug-c7d44eba12-c7d44eba12": ("Teal Octopus Shaped Rug", "teal octopus shaped rug", "teal octopus area rug, ocean themed kids mat, sea life playroom rug", "a teal octopus design with ocean-inspired details for a playroom, bedroom, classroom, or kids room"),
    "tree-of-life-welcome-mat-living-room-c11fdc3ba9-c11fdc3ba9": ("Tree Of Life Area Rug", "tree of life area rug", "tree of life home decor rug, botanical living room mat, nature inspired area rug", "a tree of life design with botanical details for a living room, bedroom, entryway, or study"),
    "western-running-horse-area-rug-122e8b85fb-122e8b85fb": ("Black Running Horse Rug", "black running horse rug", "western horse area rug, equestrian room mat, galloping horse decor", "a running horse design in a dark western style for a bedroom, living room, office, or equestrian space"),
    "custom-running-horse-area-rug-5eadea7a69": ("Purple Horse Area Rug", "purple horse area rug", "purple equestrian rug, running horse room mat, horse lover decor", "a purple running horse design for a bedroom, living room, office, or equestrian space"),
    "custom-running-horse-welcome-mat-5c3fa81b88": ("Rustic Horse Area Rug", "rustic horse area rug", "rustic equestrian rug, western horse welcome mat, horse room decor", "a rustic running horse design for an entryway, living room, bedroom, office, or equestrian space"),
    "custom-running-horse-welcome-mat-a0cd749d2b": ("Rainbow Horse Area Rug", "rainbow horse area rug", "colorful equestrian rug, horse lover room mat, rainbow horse decor", "a colorful running horse design with rainbow details for a bedroom, playroom, living room, or horse-themed space"),
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B028 scope does not match the R104 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        article = "an" if primary.lower().startswith(("a ", "e ", "i ", "o ", "u ")) else "a"
        meta = f"Create {article} {primary} with {motif}."
        description = f"Create {article} {primary} with {motif}. The design adds a distinctive focal point to a playroom, bedroom, classroom, entryway, study space, or living room."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

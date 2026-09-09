from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B012"
REVISION = "R088"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"


COPY = {
    "personalized-teacher-classroom-rug-custom-back-to-school-design-109": {
        "title": "Pastel Teacher Classroom Welcome Rug",
        "meta": "Customize a pastel teacher classroom welcome rug with classroom-name text, yellow stripes, bows, rainbow, ABC blocks, and pencil artwork.",
        "description": "Customize a pastel teacher classroom welcome rug with classroom-name text, yellow stripes, bows, rainbow, ABC blocks, and pencil artwork. The soft school-themed design gives a classroom entry or reading area a friendly personalized welcome.",
        "primary": "pastel teacher classroom welcome rug",
        "secondary": "custom teacher welcome rug, pastel classroom rug, back to school classroom mat",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-110": {
        "title": "School Icons Teacher Welcome Rug",
        "meta": "Customize a school icons teacher welcome rug with name text, school bus, books, globe, scissors, ABC blocks, and classroom color.",
        "description": "Customize a school icons teacher welcome rug with name text, school bus, books, globe, scissors, ABC blocks, and classroom color. The design creates a bright teacher-name welcome for back-to-school classroom spaces.",
        "primary": "school icons teacher welcome rug",
        "secondary": "custom teacher classroom rug, back to school welcome rug, school icons rug",
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-111": {
        "title": "Affirmation Rainbow Classroom Rug",
        "meta": "Customize an affirmation rainbow classroom rug with teacher-name text, apple and books artwork, student affirmations, and chalkboard style.",
        "description": "Customize an affirmation rainbow classroom rug with teacher-name text, apple and books artwork, student affirmations, and chalkboard style. The design combines classroom identity with positive words students can see each day.",
        "primary": "affirmation rainbow classroom rug",
        "secondary": "custom teacher affirmation rug, rainbow classroom rug, back to school teacher mat",
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-classroom-rug": {
        "title": "Core Vocabulary Communication Rug",
        "meta": "Create a core vocabulary communication rug with a symbol-word grid for everyday classroom requests, feelings, actions, and play language.",
        "description": "Create a core vocabulary communication rug with a symbol-word grid for everyday classroom requests, feelings, actions, and play language. The visual layout gives children a shared floor reference for simple communication routines.",
        "primary": "core vocabulary communication rug",
        "secondary": "communication rugs for kids, SPED classroom rug, core board rug",
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-100": {
        "title": "AAC Core Board Classroom Rug",
        "meta": "Create an AAC core board classroom rug with colored symbol cells, a number row, classroom words, and everyday request language.",
        "description": "Create an AAC core board classroom rug with colored symbol cells, a number row, classroom words, and everyday request language. The design presents a large visual board for communication practice in classroom and playroom settings.",
        "primary": "AAC core board classroom rug",
        "secondary": "core vocabulary rug, communication board rug, SPED classroom mat",
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-101": {
        "title": "SPED Communication Board Rug",
        "meta": "Create a SPED communication board rug with picture-word cells for feelings, needs, play, bathroom, help, and classroom routines.",
        "description": "Create a SPED communication board rug with picture-word cells for feelings, needs, play, bathroom, help, and classroom routines. The design gives students and teachers a clear shared reference for everyday classroom language.",
        "primary": "SPED communication board rug",
        "secondary": "communication board rug, core vocabulary classroom rug, picture word rug",
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-102": {
        "title": "Core Vocabulary AAC Classroom Rug",
        "meta": "Create a core vocabulary AAC classroom rug with simple icon-word cells for yes, no, emotions, help, stop, go, questions, and choices.",
        "description": "Create a core vocabulary AAC classroom rug with simple icon-word cells for yes, no, emotions, help, stop, go, questions, and choices. The clean symbol layout helps make common classroom communication cues easy to find.",
        "primary": "core vocabulary AAC classroom rug",
        "secondary": "AAC classroom rug, core board rug, communication rug for kids",
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-103": {
        "title": "Sign Language Communication Rug",
        "meta": "Create a sign language communication rug with hand-sign panels for yes, no, please, help, more, eat, drink, restroom, stop, and go.",
        "description": "Create a sign language communication rug with hand-sign panels for yes, no, please, help, more, eat, drink, restroom, stop, and go. The illustrated panel design gives a classroom or therapy space a simple visual reference.",
        "primary": "sign language communication rug",
        "secondary": "sign language classroom rug, communication rug for kids, ASL classroom mat",
    },
    "core-vocabulary-rug-communication-rugs-for-kids-sped-cla-design-104": {
        "title": "Something Hurts Communication Rug",
        "meta": "Create a Something Hurts communication rug with a body silhouette, pain words, needs icons, and classroom communication prompts.",
        "description": "Create a Something Hurts communication rug with a body silhouette, pain words, needs icons, and classroom communication prompts. The design gives children a visual way to point to discomfort, feelings, and basic needs.",
        "primary": "something hurts communication rug",
        "secondary": "body pain communication rug, SPED classroom rug, needs communication mat",
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-mental-health": {
        "title": "Round Feelings Wheel Classroom Rug",
        "meta": "Create a round feelings wheel classroom rug with How Are You Feeling text, emotion faces, flower-center artwork, and colorful SEL language.",
        "description": "Create a round feelings wheel classroom rug with How Are You Feeling text, emotion faces, flower-center artwork, and colorful SEL language. The design gives students a visual reference for naming emotions during classroom routines.",
        "primary": "round feelings wheel classroom rug",
        "secondary": "feelings wheel rug, emotions classroom rug, SEL emotion rug",
    },
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B012 scope does not match the R088 copy map")

    for row in seo_rows:
        item = COPY[row["Handle"]]
        row["title_proposed"] = item["title"]
        row["meta_title_seo"] = item["title"]
        row["meta_title_chars"] = str(len(item["title"]))
        row["meta_description_seo"] = item["meta"]
        row["meta_description_chars"] = str(len(item["meta"]))
        row["description_proposed"] = item["description"]
        row["description_proposed_html"] = f"<p>{item['description']}</p>"
        row["primary_keyword"] = item["primary"]
        row["secondary_keywords"] = item["secondary"]
        row["long_tail_candidates"] = f"{item['primary']}; {item['secondary']}"
        row["revision"] = "2"
        row["review_status"] = "NEEDS_REVIEW"
        row["content_qa_status"] = "NOT_RUN"
        row["review_reason"] = f"{REVISION} removes internal evidence-note language and unsupported claim wording from B012 copy."
        row["issues"] = f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations."

    handles_set = set(handles)
    evidence_scope = [row for row in evidence_rows if row.get("evidence_id") in {r.get("evidence_id") for r in seo_rows}]
    keyword_scope = [row for row in keyword_rows if row.get("product_key") in handles_set]
    buyer_rows = [row for row in buyer_all if row.get("product_key") in handles_set]
    img_scope = [row for row in img_rows if row.get("Handle") in handles_set]

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": "Canonical B012 cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    evidence_ids = {row["evidence_id"] for row in evidence_scope}
    for row in seo_rows:
        evidence_id = row.get("evidence_id")
        if evidence_id not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {evidence_id}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", evidence_id])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    digest = sha256(OUT)
    manifest = {
        "revision": REVISION,
        "batch_id": BATCH,
        "scope_count": len(handles),
        "image_count": len(img_scope),
        "canonical_workbook": str(OUT),
        "sha256": digest,
        "status": "CREATED_FOR_HARDENED_REQA",
        "approval_status": "NOT_APPROVED_NOT_DEPLOYED",
    }
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

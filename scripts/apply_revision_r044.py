from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R043/SEO_Product_Optimization_revision_R043.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R044"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R044.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R044"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-motivational-classroom-rug-kids-c08ac9a601": {
        "title": "You Are Affirmation Classroom Rug with Crayon Art",
        "meta": "Decorate with a You Are affirmation classroom rug featuring oversized crayon letters and positive words such as kind, brave and loved.",
        "primary": "You Are affirmation classroom rug",
        "secondary": "crayon letter affirmation rug, positive words classroom rug, black motivational mat",
        "long_tail": "You Are affirmation classroom rug; crayon letter affirmation rug; positive words classroom rug",
        "description": "Decorate with a You Are affirmation classroom rug featuring oversized crayon letters and positive words such as kind, brave and loved.\n\nDesign details\n- Black classroom rug artwork shows large crayon-style You Are letters with colorful affirmation words beside them.\n- Visible words include kind, smart, helpful, grateful, enough, strong, unique and loved, with room mockups and child scene images.\n- SEO copy stays tied to visible affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for You Are crayon affirmation rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-ff1188740c": {
        "title": "I Am Affirmation Classroom Rug with Rainbow Waves",
        "meta": "Decorate with an I Am affirmation classroom rug featuring rainbow waves, smiley faces, flowers and positive word list artwork.",
        "primary": "I Am affirmation classroom rug",
        "secondary": "rainbow affirmation rug, smiley classroom rug, positive word list mat",
        "long_tail": "I Am affirmation classroom rug; rainbow affirmation rug; smiley classroom rug",
        "description": "Decorate with an I Am affirmation classroom rug featuring rainbow waves, smiley faces, flowers and positive word list artwork.\n\nDesign details\n- Cream rug artwork shows an I Am affirmation list framed by orange and pink rainbow-wave borders.\n- Visible words include brave, kind, strong, helpful, honest, unique, grateful and loved, with room mockups and close-up views.\n- SEO copy focuses on visible affirmation-list artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for rainbow I Am affirmation rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-ec5fda4b1c": {
        "title": "You Are Positive Message Rug with Word Collage",
        "meta": "Decorate with a You Are positive message rug featuring black background, colorful affirmation words and central You Are artwork.",
        "primary": "You Are positive message rug",
        "secondary": "word collage classroom rug, colorful affirmation rug, black positive words mat",
        "long_tail": "You Are positive message rug; word collage classroom rug; colorful affirmation rug",
        "description": "Decorate with a You Are positive message rug featuring black background, colorful affirmation words and central You Are artwork.\n\nDesign details\n- Black rug artwork centers You Are text with surrounding words such as creative, amazing, determined, special, brave and loved.\n- Visible imagery includes colorful doodles, bold typography, sofa mockup, classroom scene and close-up artwork views.\n- SEO copy stays tied to visible word-collage artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {6: "Feature graphic for You Are word collage rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-dbde19466a": {
        "title": "Sunshine I Am Affirmation Rug with Rainbow Rays",
        "meta": "Decorate with a sunshine I Am affirmation rug featuring rainbow ray panels and positive words such as kind, loved and brave.",
        "primary": "sunshine I Am affirmation rug",
        "secondary": "rainbow ray classroom rug, positive words sunburst rug, colorful I Am mat",
        "long_tail": "sunshine I Am affirmation rug; rainbow ray classroom rug; positive words sunburst rug",
        "description": "Decorate with a sunshine I Am affirmation rug featuring rainbow ray panels and positive words such as kind, loved and brave.\n\nDesign details\n- Rainbow sunburst rug artwork centers I Am text with affirmation words arranged inside colorful rays.\n- Visible words include kind, loved, brave, unique, grateful, generous, creative, strong, happy and smart, with room mockups and child scene images.\n- SEO copy focuses on visible sunburst affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for sunshine I Am affirmation rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-50daea8577": {
        "title": "You Are Classroom Rug with Bright Word Stripes",
        "meta": "Decorate with a You Are classroom rug featuring bright word stripes, smiley faces and colorful affirmation artwork.",
        "primary": "You Are classroom rug",
        "secondary": "bright affirmation classroom rug, positive word stripe rug, smiley face classroom mat",
        "long_tail": "You Are classroom rug; bright affirmation classroom rug; positive word stripe rug",
        "description": "Decorate with a You Are classroom rug featuring bright word stripes, smiley faces and colorful affirmation artwork.\n\nDesign details\n- White rug artwork shows You Are header text above stacked stripes with large positive words in different colors.\n- Visible words include amazing, brilliant, extraordinary, fantastic, incredible, magnificent, phenomenal and spectacular, plus smiley icons and room mockups.\n- SEO copy stays tied to visible affirmation-stripe artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {5: "Feature graphic for bright You Are classroom rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-2f14feb0aa": {
        "title": "Mrs Brown Class Affirmation Rug with Smiley Icons",
        "meta": "Personalize a Mrs Brown class affirmation rug with yellow gingham design, I Am prompt, smiley icons and school artwork.",
        "primary": "Mrs Brown class affirmation rug",
        "secondary": "teacher name affirmation rug, yellow classroom rug, smiley school icon mat",
        "long_tail": "Mrs Brown class affirmation rug; teacher name affirmation rug; yellow classroom rug",
        "description": "Personalize a Mrs Brown class affirmation rug with yellow gingham design, I Am prompt, smiley icons and school artwork.\n\nDesign details\n- Yellow and pink classroom rug artwork shows Mrs. Brown's Class text with I Am prompt and illustrated school icons.\n- Visible details include smiley faces, pencils, books, hearts, stars, classroom mockups, sofa scene and child scene image.\n- SEO copy focuses on visible teacher-name affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {7: "Feature graphic for Mrs Brown class affirmation rug"},
    },
    "personalized-motivational-classroom-rug-for-kids-34838f57f6-34838f57f6": {
        "title": "Mrs Harper Alphabet Affirmation Rug with A To Z",
        "meta": "Personalize a Mrs Harper alphabet affirmation rug with A to Z letters, character trait words and black classroom artwork.",
        "primary": "Mrs Harper alphabet affirmation rug",
        "secondary": "A to Z classroom rug, character trait affirmation rug, teacher name ABC mat",
        "long_tail": "Mrs Harper alphabet affirmation rug; A to Z classroom rug; character trait affirmation rug",
        "description": "Personalize a Mrs Harper alphabet affirmation rug with A to Z letters, character trait words and black classroom artwork.\n\nDesign details\n- Black rug artwork shows You Are header text, alphabet letters and a trait word beside each letter.\n- Visible details include Mrs Harper personalization, colorful letters, classroom mockups, size option graphic and close-up views.\n- SEO copy stays tied to visible alphabet affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature infographic for Mrs Harper alphabet rug",
            6: "Care-style graphic for ABC affirmation rug artwork",
        },
    },
    "personalized-motivational-classroom-rug-for-kids-7d4f70bdd2-7d4f70bdd2": {
        "title": "Mrs Smith Vintage Affirmation Rug with Note Blocks",
        "meta": "Personalize a Mrs Smith vintage affirmation rug with wood plank background, colorful note blocks and classroom quote artwork.",
        "primary": "Mrs Smith vintage affirmation rug",
        "secondary": "wood plank classroom rug, colorful note affirmation rug, teacher name quote mat",
        "long_tail": "Mrs Smith vintage affirmation rug; wood plank classroom rug; colorful note affirmation rug",
        "description": "Personalize a Mrs Smith vintage affirmation rug with wood plank background, colorful note blocks and classroom quote artwork.\n\nDesign details\n- Wood plank-style rug artwork shows Mrs. Smith's Class text with colorful note blocks and classroom-themed icons.\n- Visible phrases include you are loved, keep going, spread kindness and belong, with room mockups, size option graphic and close-up views.\n- SEO copy focuses on visible note-block affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature infographic for Mrs Smith vintage affirmation rug",
            6: "Care-style graphic for vintage classroom rug artwork",
        },
    },
    "personalized-hunting-welcome-doormat-custom-family-name-e9839e2783": {
        "title": "Flying On Freedom Hunting Doormat with Duck Art",
        "meta": "Personalize a Flying On Freedom hunting doormat with duck, American flag, wildlife border and family-name style artwork.",
        "primary": "Flying On Freedom hunting doormat",
        "secondary": "duck hunting welcome mat, patriotic family name doormat, wildlife border entry mat",
        "long_tail": "Flying On Freedom hunting doormat; duck hunting welcome mat; patriotic family name doormat",
        "description": "Personalize a Flying On Freedom hunting doormat with duck, American flag, wildlife border and family-name style artwork.\n\nDesign details\n- Rectangular hunting doormat artwork shows a flying duck over an American flag with Flying On Freedom text.\n- Visible details include patchwork wildlife border, doorway mockups, dog entry scene, size option graphic, feature panels and repeated gallery images.\n- SEO copy stays tied to visible duck-and-flag artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            6: "Size option graphic for Flying On Freedom hunting doormat",
            7: "Feature graphic for duck hunting doormat artwork",
            14: "Repeated feature graphic for hunting doormat artwork",
        },
    },
    "personalized-family-name-patriotic-hunting-doormat-5997abfebc-5997abfebc": {
        "title": "Patriotic Eagle Family Name Doormat Flag Artwork",
        "meta": "Personalize a patriotic eagle family name doormat with American flag, anniversary wording and Secure The Blessings Of Liberty text.",
        "primary": "patriotic eagle family name doormat",
        "secondary": "American flag eagle doormat, anniversary family welcome mat, liberty quote entry rug",
        "long_tail": "patriotic eagle family name doormat; American flag eagle doormat; anniversary family welcome mat",
        "description": "Personalize a patriotic eagle family name doormat with American flag, anniversary wording and Secure The Blessings Of Liberty text.\n\nDesign details\n- Patriotic doormat artwork shows an eagle over an American flag with Secure The Blessings Of Liberty text.\n- Visible details include family name panel, 20th anniversary wording, doorway mockups, person holding mat, size option graphic and feature panels.\n- SEO copy focuses on visible eagle-and-flag artwork without extra care, reverse-side, construction or surface-performance claims.",
        "alts": {
            5: "Feature graphic for patriotic eagle family doormat",
            12: "Repeated feature graphic for patriotic doormat artwork",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable", "floor mat", "floor rug", "low pile",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R044 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R044 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R044"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R044", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R044", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R045 or Re-QA revision R044",
    }
    (RUN_DIR / "revision_R044_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R044 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R044.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R044 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R044_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r044_log_rows": counts.get("R044", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

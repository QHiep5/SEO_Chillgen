from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R037/SEO_Product_Optimization_revision_R037.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R038"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R038.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R038"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "custom-japanese-koi-fish-shaped-area-rug-aa33e1b43a": {
        "title": "Green Japanese Koi Pond Shaped Rug with Lily Pads",
        "meta": "Decorate with a Japanese koi pond shaped rug featuring orange koi fish, winding white water paths, stones and lily pads.",
        "primary": "Japanese koi pond shaped rug",
        "secondary": "koi fish area rug, lily pad pond rug, irregular green rug",
        "long_tail": "Japanese koi pond shaped rug; koi fish area rug; lily pad pond rug",
        "description": "Decorate with a Japanese koi pond shaped rug featuring orange koi fish, winding white water paths, stones and lily pads.\n\nDesign details\n- Irregular green rug artwork shows a garden pond layout with orange koi fish, rounded lily pads, stepping-stone shapes and white stream paths.\n- Product images include living room mockups, dark floor scenes, size reference graphic and close-up views of the koi pond artwork.\n- SEO copy stays tied to visible koi garden artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            6: "Custom-shape feature graphic for Japanese koi pond rug",
            7: "Layer-style diagram for Japanese koi pond shaped rug",
        },
    },
    "personalized-vinyl-record-rug-with-photo-and-song-name-9a545dcd49": {
        "title": "Personalized Vinyl Record Rug with Photo Center",
        "meta": "Personalize a vinyl record rug with photo center, song name, artist name, black record grooves and player-style icons.",
        "primary": "personalized vinyl record rug",
        "secondary": "custom photo record rug, song name music rug, round black record mat",
        "long_tail": "personalized vinyl record rug; custom photo record rug; song name music rug",
        "description": "Personalize a vinyl record rug with photo center, song name, artist name, black record grooves and player-style icons.\n\nDesign details\n- Round black music rug artwork shows vinyl record grooves with a custom couple photo in the center label area.\n- Visible details include song and artist text, playback icons, record color options, bedroom mockups, size reference graphic and room placement views.\n- SEO copy focuses on visible custom record artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            7: "Close-up view of personalized vinyl record rug artwork",
            9: "Care-style graphic for personalized vinyl record rug",
        },
    },
    "custom-3d-effect-dragon-round-rug-personalized-3d-printed-dragon-book-lovers-library-optical-illusion-rug-c5c462d4c8": {
        "title": "Purple Dragon Library Round Rug with Stone Stair Art",
        "meta": "Decorate with a purple dragon library round rug featuring stone stair ring, warm bookshelves and optical illusion artwork.",
        "primary": "purple dragon library round rug",
        "secondary": "dragon optical illusion rug, book lover fantasy rug, stone stair dragon mat",
        "long_tail": "purple dragon library round rug; dragon optical illusion rug; book lover fantasy rug",
        "description": "Decorate with a purple dragon library round rug featuring stone stair ring, warm bookshelves and optical illusion artwork.\n\nDesign details\n- Round fantasy rug artwork shows a purple dragon curled inside a stone stair circle with glowing shelves and library details.\n- Product images include fireplace mockups, room scenes, size reference graphic and close-up views of the dragon and bookshelves.\n- SEO copy stays tied to visible dragon library artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic for purple dragon library round rug",
        },
    },
    "feeling-rugs-kids-mental-health-brain-area-rug-fae7deacd8": {
        "title": "How's Your Brain Feelings Rug with Color Icons",
        "meta": "Decorate with a How's Your Brain feelings rug featuring colorful brain icons, emotion labels and black classroom-style artwork.",
        "primary": "How's Your Brain feelings rug",
        "secondary": "brain emotion icon rug, classroom feelings mat, colorful mood label rug",
        "long_tail": "How's Your Brain feelings rug; brain emotion icon rug; classroom feelings mat",
        "description": "Decorate with a How's Your Brain feelings rug featuring colorful brain icons, emotion labels and black classroom-style artwork.\n\nDesign details\n- Black rectangular rug artwork reads How's Your Brain with bright brain icons and labels such as mushy, denial, mindful and organized chaos.\n- Product images include classroom-style scenes, child group photos, top-down artwork views, size reference graphics and close-up label panels.\n- SEO copy focuses on visible brain feelings artwork without extra care, classroom-result, therapeutic, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic shows How's Your Brain feelings rug artwork",
        },
    },
    "custom-notebook-welcome-mat-for-classroom-and-school-864ed94fd6": {
        "title": "Pink Round Classroom Welcome Rug with Mrs Smith",
        "meta": "Personalize a pink round classroom welcome rug with Mrs Smith text, checker pattern, smiley faces, hearts and colorful letters.",
        "primary": "pink round classroom welcome rug",
        "secondary": "Mrs Smith classroom rug, checker welcome mat, personalized educator rug",
        "long_tail": "pink round classroom welcome rug; Mrs Smith classroom rug; checker welcome mat",
        "description": "Personalize a pink round classroom welcome rug with Mrs Smith text, checker pattern, smiley faces, hearts and colorful letters.\n\nDesign details\n- Round pink classroom rug artwork reads Welcome To Mrs. Smith's Classroom across a checkerboard-style background.\n- Visible details include smiley faces, stars, hearts, colorful block letters, classroom mockups, children group photos and size reference graphics.\n- SEO copy stays tied to visible pink welcome artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature graphic for pink round Mrs Smith classroom rug",
            7: "Close-up detail image for pink round classroom welcome rug",
        },
    },
    "spanish-abc-classroom-rug-learning-alphabet-mat-f6acb4ba58": {
        "title": "Spanish Bienvenidos Classroom Rug with Name Art",
        "meta": "Personalize a Spanish Bienvenidos classroom rug with Seño Garcia name, school-year text, rainbow, apple and smiley icons.",
        "primary": "Spanish Bienvenidos classroom rug",
        "secondary": "personalized maestra rug, bilingual welcome classroom mat, Spanish school decor rug",
        "long_tail": "Spanish Bienvenidos classroom rug; personalized maestra rug; bilingual welcome classroom mat",
        "description": "Personalize a Spanish Bienvenidos classroom rug with Seño Garcia name, school-year text, rainbow, apple and smiley icons.\n\nDesign details\n- Black-and-white speckled classroom rug artwork reads Bienvenidos, Ciclo Escolar 2025-2026 and Seño Garcia.\n- Visible details include rainbow, apple, smiley faces, school icons, playroom mockups, child photo scenes and size reference graphics.\n- SEO copy focuses on visible Spanish welcome artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic shows Spanish Bienvenidos classroom rug artwork",
        },
    },
    "custom-spanish-abc-classroom-rug-playroom-school-bee973bfe4": {
        "title": "Spanish ABC Pencil Classroom Rug with Letter Art",
        "meta": "Decorate with a Spanish ABC pencil classroom rug featuring yellow pencil characters, alphabet letters and black speckled background.",
        "primary": "Spanish ABC pencil classroom rug",
        "secondary": "bilingual alphabet rug, Spanish letter classroom mat, yellow pencil school rug",
        "long_tail": "Spanish ABC pencil classroom rug; bilingual alphabet rug; Spanish letter classroom mat",
        "description": "Decorate with a Spanish ABC pencil classroom rug featuring yellow pencil characters, alphabet letters and black speckled background.\n\nDesign details\n- Large rectangular classroom rug artwork repeats cheerful yellow pencil characters paired with Spanish alphabet letters.\n- Visible details include black speckled background, classroom mockups, child reading scenes, block-play photos, size graphics and close-up views.\n- SEO copy stays tied to visible pencil alphabet artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic shows Spanish ABC pencil rug artwork",
        },
    },
    "spanish-abc-classroom-rug-0f9ffdff41": {
        "title": "Spanish Classroom Rules Rug with Pencil Prompts",
        "meta": "Personalize a Spanish classroom rules rug with Maestra Valeria name, colorful pencil prompts and black background.",
        "primary": "Spanish classroom rules rug",
        "secondary": "Maestra Valeria rug, bilingual classroom rules mat, colorful pencil prompt rug",
        "long_tail": "Spanish classroom rules rug; Maestra Valeria rug; bilingual classroom rules mat",
        "description": "Personalize a Spanish classroom rules rug with Maestra Valeria name, colorful pencil prompts and black background.\n\nDesign details\n- Black rectangular rug artwork reads Maestra Valeria Reglas De La Clase with stacked pencil-shaped rule prompts.\n- Visible prompts include Ven A Clase Preparado, Se Amable Con Los Demás, Prueba Cosas Nuevas and Mantén Tus Manos Para Ti Mismo.\n- SEO copy focuses on visible Spanish rules artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic shows Spanish classroom rules rug artwork",
        },
    },
    "custom-composition-notebook-classroom-rug-03facf229f": {
        "title": "What Can I Bee Classroom Rug with Honeycomb Words",
        "meta": "Personalize a What Can I Bee classroom rug with Mrs Thompson name, bee characters, honeycomb cells and positive words.",
        "primary": "What Can I Bee classroom rug",
        "secondary": "bee positive words rug, honeycomb classroom mat, personalized educator rug",
        "long_tail": "What Can I Bee classroom rug; bee positive words rug; honeycomb classroom mat",
        "description": "Personalize a What Can I Bee classroom rug with Mrs Thompson name, bee characters, honeycomb cells and positive words.\n\nDesign details\n- Cream classroom rug artwork asks What Can I Bee and shows smiling bee characters around honeycomb word cells.\n- Visible words include silly, brave, helpful, kind, respectful and independent, with classroom mockups, group photos and size graphics.\n- SEO copy stays tied to visible bee honeycomb artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature graphic shows What Can I Bee classroom rug artwork",
            6: "Feature panel for What Can I Bee rug artwork details",
            9: "Close-up view of bee honeycomb classroom rug artwork",
        },
    },
    "customized-composition-notebook-classroom-rug-33c4fbd370": {
        "title": "Mrs Smith Classroom Affirmation Rug with Rainbow Border",
        "meta": "Personalize a Mrs Smith classroom affirmation rug with rainbow border, colorful You Are words and black board-style artwork.",
        "primary": "Mrs Smith classroom affirmation rug",
        "secondary": "rainbow border classroom rug, You Are word rug, personalized educator mat",
        "long_tail": "Mrs Smith classroom affirmation rug; rainbow border classroom rug; You Are word rug",
        "description": "Personalize a Mrs Smith classroom affirmation rug with rainbow border, colorful You Are words and black board-style artwork.\n\nDesign details\n- Black rectangular classroom rug artwork reads When You Enter This Mrs. Smith's Classroom You Are with colorful word labels.\n- Visible words include scientists, creators, authors, important, leaders, explorers, thinkers, readers, friend and loved.\n- SEO copy focuses on visible affirmation board artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            6: "Feature graphic shows Mrs Smith classroom affirmation rug artwork",
            7: "Close-up view of Mrs Smith classroom affirmation rug artwork",
            8: "Feature panel for Mrs Smith affirmation rug artwork details",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "reinforced", "bound edge", "bound edges", "material panel", "learning-pattern",
    "learning pattern", "support classroom", "support emotional", "encourage students",
    "use a black", "classroom sel", "therapy office",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R038 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R038 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R038"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R038", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R038", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R039 or Re-QA revision R038",
    }
    (RUN_DIR / "revision_R038_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R038 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R038.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R038 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R038_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r038_log_rows": counts.get("R038", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

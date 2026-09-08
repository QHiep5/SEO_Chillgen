from __future__ import annotations

import json
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R007/SEO_Product_Optimization_revision_R007.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R008"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R008.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R008"
MANIFEST = RUN_DIR / "revision_R008_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R008_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-classroom-welcome-doormat-2820e17d05-2820e17d05": {
        "title": "Red Doodle Classroom Welcome Mat with Teacher Name",
        "meta": "Welcome students with a red doodle classroom mat featuring teacher name text, rainbow, apple, pencil and flower artwork.",
        "description": (
            "Welcome students with a red doodle classroom mat featuring teacher name text, rainbow, apple, pencil and flower artwork.\n\n"
            "Design details\n"
            "- Red and black classroom mat sample reading Welcome to Mrs. Sophia's Classroom.\n"
            "- Visible artwork includes rainbow, apple, pencil, notebook, flowers and playful school doodles.\n"
            "- SEO copy focuses on the visible personalized classroom welcome design without unsupported material, backing or surface-performance claims."
        ),
    },
    "personalized-classroom-doormat-welcome-to-class-door-mat-1086025b51": {
        "title": "Green Classroom Welcome Mat with School Supply Icons",
        "meta": "Personalize a green classroom welcome mat with teacher name text, pencil, ruler, paper plane, hearts and star doodles.",
        "description": (
            "Personalize a green classroom welcome mat with teacher name text, pencil, ruler, paper plane, hearts and star doodles.\n\n"
            "Design details\n"
            "- Green school-themed mat sample reading Mrs. Sophia's Classroom in playful script.\n"
            "- Visible icons include pencil, ruler, paper plane, music note, lightbulb, hearts and stars.\n"
            "- SEO copy stays tied to the visible teacher-name classroom design without unsupported material, backing or surface-performance claims."
        ),
    },
    "personalized-classroom-doormat-custom-teacher-name-753f4f3532-753f4f3532": {
        "title": "You Are Classroom Mat with Colorful Pencil Affirmations",
        "meta": "Customize a You Are classroom mat with teacher name text, colorful pencil rays, apple icon and affirmation words.",
        "description": (
            "Customize a You Are classroom mat with teacher name text, colorful pencil rays, apple icon and affirmation words.\n\n"
            "Design details\n"
            "- White classroom mat sample reading In This Classroom You Are with Mrs. Sophia personalization.\n"
            "- Colorful pencil rays include affirmation words such as valued, creative, special, kind, trusted, loved, smart and amazing.\n"
            "- SEO copy focuses on the visible affirmation classroom artwork without unsupported material, backing or surface-performance claims."
        ),
    },
    "personalized-dog-welcome-mat-front-door-36fdce22f9": {
        "title": "Humans Live Here Pet Welcome Mat with Dog and Cat Names",
        "meta": "Customize a Humans Live Here pet welcome mat with dog and cat names, illustrated pet portraits and funny home text.",
        "description": (
            "Customize a Humans Live Here pet welcome mat with dog and cat names, illustrated pet portraits and funny home text.\n\n"
            "Design details\n"
            "- Gray pet welcome mat reading Welcome to Our Home The Humans Just Live Here With Us.\n"
            "- Sample design shows two dogs and one cat with names Hope, Happy and Jenny.\n"
            "- SEO copy stays tied to the visible pet portrait and name design without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "personalized-dog-welcome-mat-entryway-front-door-5385d11bb7": {
        "title": "Three Pet Welcome Mat with Dog and Cat Name Artwork",
        "meta": "Customize a three pet welcome mat with illustrated dog and cat portraits, pet names, tan design and welcome script.",
        "description": (
            "Customize a three pet welcome mat with illustrated dog and cat portraits, pet names, tan design and welcome script.\n\n"
            "Design details\n"
            "- Tan welcome mat sample with three illustrated pets and names Loki, Daisy and Bailey.\n"
            "- The design includes black border, large Welcome script and variant panels for different pet combinations.\n"
            "- SEO copy focuses on the visible personalized pet artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "personalized-dog-welcome-doormat-front-door-entryway-d819eb0e7c": {
        "title": "Loki's House Dog Welcome Mat with Owner Names",
        "meta": "Personalize a Loki's House dog welcome mat with pet name, owner names, paw prints and tan illustrated dog artwork.",
        "description": (
            "Personalize a Loki's House dog welcome mat with pet name, owner names, paw prints and tan illustrated dog artwork.\n\n"
            "Design details\n"
            "- Tan single-pet mat reading Welcome to Loki's House with a black and white dog illustration.\n"
            "- Sample design includes paw prints and the line Sophia and James Live Here Too.\n"
            "- SEO copy stays tied to the visible dog-name artwork without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "personalized-dog-welcome-mat-custom-pet-name-a66d52e9f8": {
        "title": "Hope You Like Dogs Welcome Mat with Red Plaid Design",
        "meta": "Customize a Hope You Like Dogs welcome mat with red plaid background, dog portraits, pet names and family name text.",
        "description": (
            "Customize a Hope You Like Dogs welcome mat with red plaid background, dog portraits, pet names and family name text.\n\n"
            "Design details\n"
            "- Red plaid dog welcome mat sample reading Hope You Like Dogs with three cartoon dog portraits.\n"
            "- Sample names include Cooper, Buddy and Lacey, with The Browns family text.\n"
            "- SEO copy focuses on the visible dog portrait and family-name design without unsupported material, cleaning, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-welcome-mat-4bd82949b1-4bd82949b1": {
        "title": "Fire Running Horse Area Rug with Flame Mane Artwork",
        "meta": "Style a room with a fire running horse area rug featuring black horse artwork, orange flame mane and ember details.",
        "description": (
            "Style a room with a fire running horse area rug featuring black horse artwork, orange flame mane and ember details.\n\n"
            "Design details\n"
            "- Dark rectangular horse rug with black running horse, bright orange flame mane and ember-style effects.\n"
            "- Gallery images show living room and room mockups plus close-up design and size visuals.\n"
            "- SEO copy stays tied to the visible fire horse motif without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-running-horse-welcome-mat-7069b0c873-7069b0c873": {
        "title": "Black Horse Bridle Area Rug with Red Rein Detail",
        "meta": "Style a horse lover room with a black horse bridle area rug featuring bold portrait artwork, red reins and gold tack.",
        "description": (
            "Style a horse lover room with a black horse bridle area rug featuring bold portrait artwork, red reins and gold tack.\n\n"
            "Design details\n"
            "- White rug with black horse portrait, red bridle and reins plus gold tack detail.\n"
            "- Gallery images show living room and fireplace-style mockups plus close-up design and size visuals.\n"
            "- SEO copy focuses on the visible bridle horse portrait without unsupported material, backing or surface-performance claims."
        ),
    },
    "custom-wolf-print-welcome-mat-89ae1c27ae": {
        "title": "Neon Galaxy Wolf Area Rug with Purple Space Artwork",
        "meta": "Style a room with a neon galaxy wolf area rug featuring bright wolf portrait, purple space colors and nebula background.",
        "description": (
            "Style a room with a neon galaxy wolf area rug featuring bright wolf portrait, purple space colors and nebula background.\n\n"
            "Design details\n"
            "- Vivid purple, pink, blue and orange galaxy rug with a large wolf head portrait, stars and nebula-style background.\n"
            "- Gallery images show sofa, living room floor, corner and desk-area mockups.\n"
            "- SEO copy stays tied to the visible neon wolf artwork without unsupported material, backing or surface-performance claims."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "non-slip", "non-skid",
    "kid friendly", "pet friendly", "safe durable", "machine-washable",
    "machine washable", "washable ", "quick-dry", "memory foam", "microfiber",
    "velvet", "stain", "fade resistant", "easy clean",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R008 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R008 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R008"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R008", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R008",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R009 or Re-QA revision R008",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R008 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R008.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _, _, _, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n", encoding="utf-8")

    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    vws = verify["SEO_Products"]
    vcols = {cell.value: idx for idx, cell in enumerate(next(vws.iter_rows(min_row=1, max_row=1)))}
    bad = []
    for row in vws.iter_rows(min_row=2, values_only=True):
        handle = row[vcols["Handle"]]
        if handle not in REVISIONS:
            continue
        title = row[vcols["title_proposed"]] or ""
        meta_title = row[vcols["meta_title_seo"]] or ""
        meta_desc = row[vcols["meta_description_seo"]] or ""
        desc = (row[vcols["description_proposed"]] or "") + " " + (row[vcols["description_proposed_html"]] or "")
        blob = " ".join([title, meta_title, meta_desc, desc]).lower()
        if row[vcols["revision"]] != "2":
            bad.append((handle, "revision", row[vcols["revision"]]))
        if row[vcols["review_status"]] != "NEEDS_REVIEW":
            bad.append((handle, "review_status", row[vcols["review_status"]]))
        if row[vcols["content_qa_status"]] != "NOT_RUN":
            bad.append((handle, "content_qa_status", row[vcols["content_qa_status"]]))
        if not (45 <= len(title) <= 70):
            bad.append((handle, "title_len", len(title)))
        if meta_title != title:
            bad.append((handle, "meta_title_mismatch", meta_title))
        if len(meta_desc) > 320:
            bad.append((handle, "meta_desc_len", len(meta_desc)))
        for forbidden in FORBIDDEN:
            if forbidden in f" {blob} ":
                bad.append((handle, "forbidden_claim", forbidden))
    if bad:
        raise RuntimeError(f"Verification failed: {bad[:20]}")

    print(json.dumps({
        "output_workbook": str(OUTPUT),
        "manifest": str(MANIFEST),
        "summary": str(SUMMARY),
        "updated": len(changed),
        "status": "COMPLETE_AWAITING_REQA",
    }, indent=2))


if __name__ == "__main__":
    main()

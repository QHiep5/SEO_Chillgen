from pathlib import Path
from hashlib import sha256
import json, re
import requests
from bs4 import BeautifulSoup
from PIL import Image
from openpyxl import load_workbook

BASE=Path('/Users/buiquanghuy/Documents/seo_chillgen')
BOOK=BASE/'resutls/chillgen.com/chillgen_20260907_01/revisions/R001/SEO_Product_Optimization_revision_R001.xlsx'
OUT=BASE/'seo_runs/chillgen.com/chillgen_20260907_01/qa/QA-20260908-REV-R001/independent_evidence.json'

def norm(x): return str(x or '').replace('https:','').replace('http:','').lstrip('/').split('?')[0]
def file_hash(p): return sha256(p.read_bytes()).hexdigest()
def rows(ws):
 it=ws.iter_rows(values_only=True); h=list(next(it)); return [dict(zip(h,r)) for r in it]

def main():
 wb=load_workbook(BOOK,read_only=True,data_only=True)
 rev=rows(wb['Revision_Log']); keys=[r['handle'] for r in rev if r['revision_batch_id']=='R001']
 products={r['Handle']:r for r in rows(wb['SEO_Products']) if r['Handle'] in keys}
 imgs={k:[] for k in keys}
 for r in rows(wb['Image_Audit']):
  if r.get('Handle') in imgs: imgs[r['Handle']].append(r)
 pages=[]; images={}
 for handle in keys:
  p=products[handle]; url=p['url']; js=requests.get(url+'.js',timeout=40); html=requests.get(url,timeout=40)
  data=js.json() if js.ok else {}; soup=BeautifulSoup(html.text,'html.parser') if html.ok else None
  h1=soup.find('h1').get_text(' ',strip=True) if soup and soup.find('h1') else ''
  canonical=soup.find('link',rel='canonical').get('href') if soup and soup.find('link',rel='canonical') else ''
  meta=soup.find('meta',attrs={'name':'description'}).get('content','') if soup and soup.find('meta',attrs={'name':'description'}) else ''
  live_urls={norm(x) for x in data.get('images',[])}
  wb_urls={norm(x.get('image_url')) for x in imgs[handle]}
  pages.append({'handle':handle,'url':url,'http_status':html.status_code,'json_status':js.status_code,'canonical':canonical,'h1':h1,'live_title':data.get('title',''),'live_meta_description':meta,'live_image_count':len(live_urls),'workbook_image_count':len(wb_urls),'image_url_set_match':live_urls==wb_urls,'html_sha256':file_hash(Path('/tmp/r001_html_'+str(keys.index(handle))+'.html')) if False else sha256(html.content).hexdigest()})
  for r in sorted(imgs[handle],key=lambda x:int(x.get('image_number') or 0)):
   key=f"{handle}__img_{int(r['image_number']):02d}"; ref=Path(r.get('evidence_file_or_reference',''))
   exists=ref.exists(); dims=[]; h=''
   if exists:
    with Image.open(ref) as im: dims=list(im.size)
    h=file_hash(ref)
   alt=str(r.get('alt_proposed') or '').strip(); obs=str(r.get('observed_visual_details') or '').strip()
   images[key]={'handle':handle,'image_number':int(r['image_number']),'live_url':r.get('image_url'),'workbook_url':r.get('image_url_export') or r.get('image_url'),'local_evidence_file':str(ref),'local_file_exists':exists,'local_sha256':h,'local_dimensions':dims,'qa_observation':obs,'alt_effective':alt,'url_match':norm(r.get('image_url')) in live_urls,'IM1':'FULL' if exists and norm(r.get('image_url')) in live_urls else 'FAIL','IM2':'FULL' if obs else 'NOT_CHECKED','IM3':'FULL' if alt and obs else 'NOT_CHECKED','IM4':'FULL' if alt and len(alt)<=125 else 'PARTIAL' if alt else 'NOT_CHECKED'}
 OUT.parent.mkdir(parents=True,exist_ok=True); OUT.write_text(json.dumps({'evidence_version':'R001-independent-v1','source_workbook':str(BOOK),'products':pages,'images':images,'product_count':len(pages),'image_count':len(images)},indent=2,ensure_ascii=False)+'\n')
 print(json.dumps({'out':str(OUT),'products':len(pages),'images':len(images),'all_pages_200':all(x['http_status']==200 and x['json_status']==200 for x in pages),'all_image_sets_match':all(x['image_url_set_match'] for x in pages)},ensure_ascii=False))
if __name__=='__main__': main()

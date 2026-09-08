from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R019/SEO_Product_Optimization_revision_R019.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R020"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R020.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R020"
MANIFEST = RUN_DIR / "revision_R020_manifest.json"
SUMMARY = OUTPUT_DIR / "REVISION_R020_SUMMARY.md"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"

TZ = timezone(timedelta(hours=7))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()


REVISIONS = {
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-102": {
        "title": "How Do You Feel Today Rug with Emotion Face Art",
        "meta": "Decorate with a How Do You Feel Today rug featuring colorful emotion faces, feeling labels and black doodle border.",
        "description": (
            "Decorate with a How Do You Feel Today rug featuring colorful emotion faces, feeling labels and black doodle border.\n\n"
            "Design details\n"
            "- White round classroom rug artwork shows large multicolor How Do You Feel Today text with surrounding emotion faces.\n"
            "- Visible labels include happy, calm, sad, jealous, worried, surprised, scared, shy, angry, silly, love and excited.\n"
            "- SEO copy describes visible emotion artwork only, without safety, therapeutic, material or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-104": {
        "title": "Kids Feelings and Emotions Rug with Cartoon Panels",
        "meta": "Decorate with a kids feelings and emotions rug featuring cartoon emotion panels, center text and bright color slices.",
        "description": (
            "Decorate with a kids feelings and emotions rug featuring cartoon emotion panels, center text and bright color slices.\n\n"
            "Design details\n"
            "- Round classroom rug artwork shows Feelings and Emotions center text with colorful panel slices and cartoon children.\n"
            "- Visible labels include scared, happy, sad, angry, excited, surprised, calm and sleepy.\n"
            "- SEO copy focuses on visible feelings artwork without safety, therapeutic, material or surface-performance claims."
        ),
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-109": {
        "title": "Animal Friends Classroom Rug with Labeled Panels",
        "meta": "Decorate with a round animal friends classroom rug featuring labeled animal panels, center text and bright colors.",
        "description": (
            "Decorate with a round animal friends classroom rug featuring labeled animal panels, center text and bright colors.\n\n"
            "Design details\n"
            "- Round rug artwork shows Animal Friends center text with colorful labeled panels around the circle.\n"
            "- Visible animals include butterfly, dog, cat, rabbit, horse, cow, sheep, pig, duck, frog, turtle and owl.\n"
            "- SEO copy describes visible animal artwork only, without safety, material or surface-performance claims."
        ),
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-wit-design-100": {
        "title": "Custom Pumpkin Skull Halloween Mat with Night Art",
        "meta": "Customize a pumpkin skull Halloween mat with custom text, jack-o-lantern rows, skull pumpkins and moonlit night art.",
        "description": (
            "Customize a pumpkin skull Halloween mat with custom text, jack-o-lantern rows, skull pumpkins and moonlit night art.\n\n"
            "Design details\n"
            "- Dark purple Halloween mat artwork shows rows of jack-o-lanterns, skull pumpkins and night-sky accents.\n"
            "- Visible personalization area reads Your Custom Here in a strip across the design.\n"
            "- SEO copy stays tied to visible Halloween artwork without unsupported material, cleaning, reverse-side, water or surface-performance claims."
        ),
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-wit-design-102": {
        "title": "Black Cat Full Moon Halloween Mat with Pumpkin Art",
        "meta": "Customize a black cat full moon Halloween mat with custom text, pumpkins, lanterns, autumn leaves and haunted trees.",
        "description": (
            "Customize a black cat full moon Halloween mat with custom text, pumpkins, lanterns, autumn leaves and haunted trees.\n\n"
            "Design details\n"
            "- Purple and orange Halloween mat artwork shows a black cat before a full moon with pumpkins and lanterns.\n"
            "- Visible details include autumn leaves, haunted trees and Your Custom Here personalization strip.\n"
            "- SEO copy focuses on visible black cat Halloween artwork without material, cleaning, reverse-side, water or surface-performance claims."
        ),
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-wit-design-103": {
        "title": "Personalized Haunted Home Halloween Mat with Bats",
        "meta": "Customize a haunted home Halloween mat with family name text, haunted house silhouette, bats, moon and pumpkins.",
        "description": (
            "Customize a haunted home Halloween mat with family name text, haunted house silhouette, bats, moon and pumpkins.\n\n"
            "Design details\n"
            "- Tan Halloween mat artwork reads Welcome to Our Haunted Home with family-name personalization ribbon.\n"
            "- Visible details include haunted house silhouette, bats, moon, pumpkins, spooky trees and black ornamental border.\n"
            "- SEO copy stays tied to visible haunted home artwork without material, cleaning, reverse-side, water or surface-performance claims."
        ),
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-horror-area": {
        "title": "Black Cat Portal Halloween Rug with Purple Illusion",
        "meta": "Decorate with a black cat portal Halloween rug featuring glowing purple eyes, circular illusion ring and autumn accents.",
        "description": (
            "Decorate with a black cat portal Halloween rug featuring glowing purple eyes, circular illusion ring and autumn accents.\n\n"
            "Design details\n"
            "- Round Halloween rug artwork shows a black cat reaching from a purple glowing portal with orange purple illusion ring.\n"
            "- Product images include seasonal room mockups, close-up panels, size reference graphics and autumn leaf details.\n"
            "- SEO copy focuses on visible portal cat artwork without unsupported material, reverse-side or surface-performance claims."
        ),
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-100": {
        "title": "3D Skeleton Pit Halloween Rug with Stone Tunnel Art",
        "meta": "Decorate with a 3D skeleton pit Halloween rug featuring stone tunnel illusion, skulls, ladders and torch lights.",
        "description": (
            "Decorate with a 3D skeleton pit Halloween rug featuring stone tunnel illusion, skulls, ladders and torch lights.\n\n"
            "Design details\n"
            "- Round optical illusion rug artwork shows a deep stone pit or tunnel with skulls, skeleton figures and ladders.\n"
            "- Product images include fireplace, living room and dining-style scenes plus close-up and size reference panels.\n"
            "- SEO copy stays tied to visible skeleton pit artwork without unsupported material, reverse-side or surface-performance claims."
        ),
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-101": {
        "title": "Haunted Mansion Halloween Rug with Pumpkin Artwork",
        "meta": "Decorate with a haunted mansion Halloween round rug featuring glowing windows, pumpkins, bats and dark night art.",
        "description": (
            "Decorate with a haunted mansion Halloween round rug featuring glowing windows, pumpkins, bats and dark night art.\n\n"
            "Design details\n"
            "- Round Halloween rug artwork shows a haunted mansion with glowing yellow windows, large pumpkins and bats.\n"
            "- Product images include fireplace, dining table and seasonal room scenes plus close-up and size reference panels.\n"
            "- SEO copy focuses on visible haunted mansion artwork without unsupported material, reverse-side or surface-performance claims."
        ),
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-102": {
        "title": "Ghost Vortex Halloween Rug with 3D Tunnel Artwork",
        "meta": "Decorate with a ghost vortex Halloween rug featuring black white tunnel illusion, purple trim, pumpkins and stars.",
        "description": (
            "Decorate with a ghost vortex Halloween rug featuring black white tunnel illusion, purple trim, pumpkins and stars.\n\n"
            "Design details\n"
            "- Round Halloween optical illusion rug artwork shows a white ghost rising from a black and white tunnel.\n"
            "- Visible details include purple ring trim, orange stars, pumpkins and seasonal room mockups.\n"
            "- SEO copy stays tied to visible ghost vortex artwork without unsupported material, reverse-side or surface-performance claims."
        ),
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip",
    "non slip", "non-skid", "non skid", "kid friendly", "pet friendly",
    "safe durable", "machine-washable", "machine washable", "washable ",
    "quick-dry", "quick dry", "memory foam", "microfiber", "velvet",
    "stain", "fade resistant", "easy clean", "soft fabric", "backing",
    "cleaning visuals", "cleaning graphic", "thickened", "bound edge",
    "reinforced", "absorbent", "absorption", "water absorption",
    "support ", "supports ", "use a ", "educational outcome", "therapy",
    "therapeutic benefit",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title = item["title"]
    meta = item["meta"]
    blob = " ".join([item["title"], item["meta"], item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R020 fixes QA MAJOR issues: short T1/T2 and unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R020 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        found = {c[0] for c in changed}
        raise RuntimeError(f"Did not update all handles. Missing: {sorted(set(REVISIONS) - found)}")

    if "Revision_Log" not in wb.sheetnames:
        log = wb.create_sheet("Revision_Log")
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    else:
        log = wb["Revision_Log"]
        keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R020"]
        if log.max_row > 1:
            log.delete_rows(2, log.max_row - 1)
        for r in keep:
            log.append(list(r))

    for handle, row_idx, old_revision, new_revision, title, meta in changed:
        log.append([
            "R020", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "T1; T2; unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R020",
        "generated_at": NOW,
        "source_workbook": str(SOURCE),
        "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN),
        "product_count": len(changed),
        "handles": [c[0] for c in changed],
        "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW",
        "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R021 or Re-QA revision R020",
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R020 Summary", "",
        f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`",
        f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R020.",
        "- Fixes targeted QA issues: short T1/T2 and unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "",
        "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |",
        "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    SUMMARY.write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        blob = " ".join([
            str(title or ""),
            str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
        ]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70:
            raise RuntimeError(f"{handle} title len failed after save")
        if len(meta) > 320:
            raise RuntimeError(f"{handle} meta len failed after save")
        if str(row[out_cols["revision"] - 1]) != "2":
            raise RuntimeError(f"{handle} revision not 2")
        if row[out_cols["review_status"] - 1] != "NEEDS_REVIEW":
            raise RuntimeError(f"{handle} review_status incorrect")
        if row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} content_qa_status incorrect")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles = {}
    dup_meta_titles = []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt = row[out_cols["meta_title_seo"] - 1]
        handle = row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle

    result = {
        "output": str(OUTPUT),
        "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r020_log_rows": counts.get("R020", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

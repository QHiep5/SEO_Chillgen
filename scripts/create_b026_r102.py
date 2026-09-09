from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path.cwd()
RUN_ID = "chillgen_20260907_01"
SHOP = "chillgen.com"
BATCH = "B026"
REVISION = "R102"
RUN = BASE / f"seo_runs/{SHOP}/{RUN_ID}"
OUT_DIR = BASE / f"resutls/{SHOP}/{RUN_ID}/revisions/{REVISION}"
OUT = OUT_DIR / f"SEO_Product_Optimization_revision_{REVISION}.xlsx"

COPY = {
    "colorful-classroom-rug-for-kids-e5ed693d3f": ("Mindset Matters Classroom Rug", "Mindset Matters classroom rug", "growth mindset classroom rug, motivational handprint rug, teacher classroom mat", "Mindset Matters wording with colorful handprints and rainbow brain artwork for a classroom, playroom, or learning space"),
    "colorful-classroom-rug-for-kids-a44b91f2e4": ("Change Your Words Classroom Rug", "Change Your Words classroom rug", "positive message classroom rug, handprint learning mat, teacher decor rug", "Change Your Words wording with colorful handprints for a classroom, counseling corner, playroom, or learning space"),
    "colorful-handprints-kids-classroom-area-rug-9f8e77de21": ("You Can Change The World Rug", "You Can Change The World rug", "positive classroom message rug, colorful handprint mat, kids learning rug", "You Can Change The World wording with colorful handprints for a classroom, playroom, library, or kids room"),
    "colorful-classroom-rug-handprints-quotes-6a4adbacc3": ("Believe Unstoppable Classroom Rug", "Believe Unstoppable classroom rug", "motivational classroom rug, inspirational handprint mat, teacher learning rug", "Believe Unstoppable wording with colorful handprints for a classroom, playroom, counseling corner, or study space"),
    "colorful-classroom-handprints-kids-rug-fca6e65808-fca6e65808": ("Future Leaders Handprint Classroom Rug", "Future Leaders handprint classroom rug", "future leaders classroom mat, colorful handprint rug, motivational kids decor", "Future Leaders wording with colorful handprints for a classroom, playroom, or learning space"),
    "colorful-handprints-classroom-kids-rug-358d2f0393": ("Unique Masterpiece Classroom Rug", "Unique Masterpiece classroom rug", "creative classroom rug, handprint learning mat, colorful kids room rug", "Unique Masterpiece wording with colorful handprints for a classroom, art area, playroom, or creative learning space"),
    "personalized-orthodox-christian-cross-byzantine-eagle-rug-9a8fdd4fcf-9a8fdd4fcf": ("Orthodox Three Bar Cross Rug", "Orthodox three bar cross rug", "Christian cross area rug, Byzantine cross decor, Orthodox room rug", "a three-bar cross design with Byzantine-inspired details for a living room, prayer room, bedroom, or study"),
    "personalized-orthodox-christian-three-bar-cross-rug-c6332608db": ("Ornate Orthodox Cross Rug", "ornate Orthodox cross rug", "Christian cross rug, Orthodox home decor, Byzantine-inspired area rug", "an ornate Orthodox cross design with decorative details for a living room, prayer room, bedroom, or study"),
    "personalized-orthodox-cross-rug-byzantine-eagle-be38fb9e0a": ("Byzantine Eagle Area Rug", "Byzantine eagle area rug", "Orthodox eagle rug, Christian home decor rug, Byzantine-inspired room rug", "a Byzantine eagle emblem with an ornate border for a living room, prayer room, bedroom, or study"),
    "personalized-orthodox-cross-byzantine-eagle-rug-8e70ebd991": ("Blue Orthodox Cross Rug", "blue Orthodox cross rug", "Christian cross area rug, blue Orthodox decor, Byzantine-inspired home rug", "a blue Orthodox cross design with decorative border details for a living room, prayer room, bedroom, or study"),
}


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader), reader.fieldnames or []


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_sheet(wb, name, rows, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            widths[cell.column] = max(widths.get(cell.column, 0), min(70, len(str(cell.value or "")) + 2))
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = max(12, width)


def main():
    seo_rows, seo_headers = read_csv(RUN / f"batches/{BATCH}_SEO_Products.csv")
    img_rows, img_headers = read_csv(RUN / f"batches/{BATCH}_Image_Audit.csv")
    evidence_rows, evidence_headers = read_csv(RUN / f"batches/{BATCH}_Product_Evidence.csv")
    keyword_rows, keyword_headers = read_csv(RUN / "keyword_research.csv")
    with (RUN / "buyer_search_research.jsonl").open(encoding="utf-8") as f:
        buyer_all = [json.loads(line) for line in f if line.strip()]
    buyer_headers = list(buyer_all[0])
    handles = [row["Handle"] for row in seo_rows]
    if handles != list(COPY):
        raise RuntimeError("B026 scope does not match the R102 copy map")
    for row in seo_rows:
        title, primary, secondary, motif = COPY[row["Handle"]]
        article = "an" if primary.lower().startswith(("a ", "e ", "i ", "o ", "u ")) else "a"
        meta = f"Create {article} {primary} with {motif}."
        description = f"Create {article} {primary} with {motif}. The design adds a welcoming focal point to a classroom, playroom, bedroom, prayer room, study space, or living room."
        row.update({
            "title_proposed": title, "meta_title_seo": title, "meta_title_chars": str(len(title)),
            "meta_description_seo": meta, "meta_description_chars": str(len(meta)),
            "description_proposed": description, "description_proposed_html": f"<p>{description}</p>",
            "primary_keyword": primary, "secondary_keywords": secondary,
            "long_tail_candidates": f"{primary}; {secondary}", "revision": "2",
            "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
            "review_reason": f"{REVISION} removes internal evidence-note language and unsupported claim wording from {BATCH} copy.",
            "issues": f"{REVISION} requires hardened re-QA; admin/export before-state and direct demand data remain limitations.",
        })
    handles_set = set(handles)
    evidence_ids = {r.get("evidence_id") for r in seo_rows}
    evidence_scope = [r for r in evidence_rows if r.get("evidence_id") in evidence_ids]
    keyword_scope = [r for r in keyword_rows if r.get("product_key") in handles_set]
    buyer_rows = [r for r in buyer_all if r.get("product_key") in handles_set]
    img_scope = [r for r in img_rows if r.get("Handle") in handles_set]
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "SEO_Products", seo_rows, seo_headers)
    write_sheet(wb, "Image_Audit", img_scope, img_headers)
    write_sheet(wb, "Product_Evidence", evidence_scope, evidence_headers)
    write_sheet(wb, "Keyword_Map", keyword_scope, keyword_headers)
    write_sheet(wb, "Buyer_Search_Research", buyer_rows, buyer_headers)
    write_sheet(wb, "README_QA", [
        {"metric": "revision", "value": REVISION, "definition": f"Canonical {BATCH} cleaned-copy revision for hardened re-QA."},
        {"metric": "scope", "value": f"{BATCH}: {len(handles)} products, {len(img_scope)} images", "definition": "Scope is exactly frozen to the original research batch."},
        {"metric": "approval_status", "value": "NOT_APPROVED_NOT_DEPLOYED", "definition": "QA pass is not human approval or Shopify deployment."},
    ], ["metric", "value", "definition"])
    log = wb.create_sheet("Revision_Log")
    log.append(["revision_batch_id", "revision", "batch_id", "handle", "scope_status", "review_status", "content_qa_status", "source_workbook", "evidence_id"])
    for row in seo_rows:
        if row.get("evidence_id") not in evidence_ids:
            raise RuntimeError(f"Missing evidence row for {row['Handle']}: {row.get('evidence_id')}")
        log.append([REVISION, 2, BATCH, row["Handle"], "LOCKED", "NEEDS_REVIEW", "NOT_RUN", f"seo_runs/{SHOP}/{RUN_ID}/batches/{BATCH}_SEO_Products.csv", row.get("evidence_id")])
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    manifest = {"revision": REVISION, "batch_id": BATCH, "scope_count": len(handles), "image_count": len(img_scope), "canonical_workbook": str(OUT), "sha256": sha256(OUT), "status": "CREATED_FOR_HARDENED_REQA", "approval_status": "NOT_APPROVED_NOT_DEPLOYED"}
    (OUT_DIR / f"revision_{REVISION}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

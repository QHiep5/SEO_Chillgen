from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R035/SEO_Product_Optimization_revision_R035.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R036"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R036.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R036"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-motivational-classroom-rug-for-kids-12e7574dd3-12e7574dd3": {
        "title": "White You Are Affirmation Classroom Rug Artwork",
        "meta": "Personalize a white You Are affirmation classroom rug with custom name, colorful word collage, doodles and playful icons.",
        "description": "Personalize a white You Are affirmation classroom rug with custom name, colorful word collage, doodles and playful icons.\n\nDesign details\n- White rectangular classroom rug artwork centers large You Are text with affirmation words such as creative, loved, kind, brave and awesome.\n- Visible details include rainbow lettering, small doodles, smiley accents, lightning bolts, flowers, classroom mockups and child photo scenes.\n- SEO copy stays tied to visible affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
    },
    "personalized-motivational-classroom-rug-for-kids-43aa6a180c": {
        "title": "Rainbow I Am Affirmation Rug with Custom Name Art",
        "meta": "Personalize a rainbow I Am affirmation rug with custom name, wavy color bands, smiley faces, flowers and butterfly art.",
        "description": "Personalize a rainbow I Am affirmation rug with custom name, wavy color bands, smiley faces, flowers and butterfly art.\n\nDesign details\n- Cream classroom rug artwork lists I Am phrases including brave, kind, strong, helpful, loved, enough and happy.\n- Visible details include orange and pink wavy bands, smiley faces, flowers, butterflies, stars, classroom scenes and top-down rug views.\n- SEO copy focuses on visible affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
    },
    "custom-dragon-round-rug-3d-illusion-library-pattern-66ba99d7f5": {
        "title": "Red Dragon Library 3D Round Rug with Atrium Art",
        "meta": "Decorate with a red dragon library 3D round rug featuring circular bookshelf atrium artwork, glowing lights and fantasy detail.",
        "description": "Decorate with a red dragon library 3D round rug featuring circular bookshelf atrium artwork, glowing lights and fantasy detail.\n\nDesign details\n- Round fantasy rug artwork shows a red dragon flying through a tall circular library atrium lined with bookshelves.\n- Product images include fireplace room, living room and bedroom mockups, close-up design panels and size reference graphics.\n- SEO copy stays tied to visible dragon library artwork without extra care, construction, reverse-side or surface-performance claims.",
    },
    "custom-3d-effect-dragon-round-rug-9cee675fdf": {
        "title": "Spiral Dragon Library Round Rug with Staircase Art",
        "meta": "Decorate with a spiral dragon library round rug featuring red and cream dragon artwork, curved stairs and bookshelf scene.",
        "description": "Decorate with a spiral dragon library round rug featuring red and cream dragon artwork, curved stairs and bookshelf scene.\n\nDesign details\n- Round fantasy rug artwork shows a red and cream dragon curled around a spiral staircase inside a library-like atrium.\n- Product images include fireplace room mockups, close-up design panels, dimension graphics and staged floor scenes.\n- SEO copy focuses on visible dragon staircase artwork without extra care, construction, reverse-side or surface-performance claims.",
    },
    "custom-classroom-welcome-mat-for-front-door-and-entryway-74cdb5663b": {
        "title": "Coping Skills Crayon Classroom Rug with Job Cards",
        "meta": "Personalize a coping skills crayon classroom rug with educator name, colorful pocket cards, black speckled field and label art.",
        "description": "Personalize a coping skills crayon classroom rug with educator name, colorful pocket cards, black speckled field and label art.\n\nDesign details\n- Black speckled classroom rug artwork shows colorful crayon-style pocket cards labeled with coping skills and classroom jobs.\n- Visible details include crayon cup graphics, bright yellow heading area, classroom floor scene, playroom mockups and size reference views.\n- SEO copy stays tied to visible crayon chart artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
    },
    "wheel-of-feelings-and-emotions-round-rug-de39e76902": {
        "title": "Black Crayon Feelings Wheel Classroom Round Rug",
        "meta": "Decorate with a black crayon feelings wheel classroom round rug featuring You Are center text, colorful crayons and emotion words.",
        "description": "Decorate with a black crayon feelings wheel classroom round rug featuring You Are center text, colorful crayons and emotion words.\n\nDesign details\n- Round black classroom rug artwork shows colorful crayon characters arranged around a You Are center message.\n- Visible border words include happy, loved, stressed, nervous, frustrated and calm, with classroom mockups and child photo scenes.\n- SEO copy focuses on visible feelings-wheel artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
    },
    "customized-composition-notebook-classroom-rug-bf478c5c0a": {
        "title": "Reading Hot Air Balloon Classroom Rug with Name",
        "meta": "Personalize a reading hot air balloon classroom rug with educator name, book basket artwork, clouds and quote text.",
        "description": "Personalize a reading hot air balloon classroom rug with educator name, book basket artwork, clouds and quote text.\n\nDesign details\n- Pastel classroom rug artwork shows Mrs. Williams name and a reading quote about knowing more places to go.\n- Visible details include books in a hot air balloon basket, cloud background, classroom mockups, top-down views and size reference graphics.\n- SEO copy stays tied to visible reading artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
    },
    "customized-composition-notebook-classroom-rug-7b0ed5679d": {
        "title": "Lined Paper Classroom Welcome Rug with Mrs Smith",
        "meta": "Personalize a lined paper classroom welcome rug with Mrs Smith text, colorful letters, pencil, paper scraps and stars.",
        "description": "Personalize a lined paper classroom welcome rug with Mrs Smith text, colorful letters, pencil, paper scraps and stars.\n\nDesign details\n- White lined paper classroom rug artwork reads Welcome To Mrs. Smith's Class in colorful mixed lettering.\n- Visible details include binder holes, pencil, tape strips, paper scraps, stars, classroom mockups and top-down rug views.\n- SEO copy focuses on visible notebook-style artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
    },
    "custom-octopus-coastal-welcome-mat-entryway-0942996125": {
        "title": "Vintage Blue Octopus Coastal Rug with Nautical Map",
        "meta": "Decorate with a vintage blue octopus coastal rug featuring nautical map artwork, starfish, compass details and ocean colors.",
        "description": "Decorate with a vintage blue octopus coastal rug featuring nautical map artwork, starfish, compass details and ocean colors.\n\nDesign details\n- Rectangular coastal rug artwork shows a large blue octopus over a vintage nautical map background.\n- Visible details include starfish, compass-style graphics, seaweed tones, sofa mockups, room scenes and size reference graphics.\n- SEO copy stays tied to visible octopus map artwork without extra care, construction, reverse-side or surface-performance claims.",
    },
    "custom-golf-welcome-mat-for-front-door-dfb4d14fc9": {
        "title": "American Flag Golf Rug with Club And Ball Design",
        "meta": "Decorate with an American flag golf rug featuring stars, stripes, golf club, ball artwork and wood-grain background.",
        "description": "Decorate with an American flag golf rug featuring stars, stripes, golf club, ball artwork and wood-grain background.\n\nDesign details\n- Rectangular golf rug artwork combines American flag stars and red stripes with a golf club and ball on a wood-grain style field.\n- Product images include sofa-side room scenes, bright floor mockups, patio-style staging and close-up artwork views.\n- SEO copy focuses on visible golf flag artwork without extra care, construction, reverse-side or surface-performance claims.",
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "washable", "vacuum", "quick-dry", "quick dry", "memory foam", "microfiber",
    "velvet", "plush", "stain", "fade", "easy clean", "easy-clean", "waterproof",
    "absorbent", "backing", "rubber", "layer construction", "layered construction",
    "ultra-soft", "soft ", "cushion", "thickness", "reinforced", "bound edge",
    "material panel", "learning-pattern", "learning pattern", "support classroom",
    "support emotional", "encourage students", "use a black", "classroom sel",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    blob = " ".join([title, meta, item["description"]]).lower()
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R036 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R036 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R036"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R036", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, description_proposed, description_proposed_html, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R036", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R037 or Re-QA revision R036",
    }
    (RUN_DIR / "revision_R036_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R036 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R036.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R036_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        desc = row[out_cols["description_proposed"] - 1]
        blob = " ".join([str(title or ""), str(meta or ""), str(desc or "")]).lower()
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        hits = [term for term in FORBIDDEN if term in blob]
        if hits:
            raise RuntimeError(f"{handle} forbidden terms after save: {hits}")

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r036_log_rows": counts.get("R036", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE = Path("/Users/buiquanghuy/Documents/seo_chillgen")
SOURCE = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R039/SEO_Product_Optimization_revision_R039.xlsx"
OUTPUT_DIR = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/R040"
OUTPUT = OUTPUT_DIR / "SEO_Product_Optimization_revision_R040.xlsx"
RUN_DIR = BASE / "seo_runs/chillgen.com/chillgen_20260907_01/revisions/R040"
PLAN = BASE / "resutls/chillgen.com/chillgen_20260907_01/revisions/REVISION_PLAN_20260908.xlsx"
NOW = datetime.now(timezone(timedelta(hours=7))).replace(microsecond=0).isoformat()


REVISIONS = {
    "personalized-teachers-classroom-rug-custom-teacher-name-classroom": {
        "title": "Mrs Smith Calming Techniques Rug with Ocean Animals",
        "meta": "Personalize a Mrs Smith Calming Techniques rug with ocean animal characters, activity prompts and pastel classroom artwork.",
        "primary": "Mrs Smith Calming Techniques rug",
        "secondary": "ocean animal classroom rug, teacher name prompt rug, pastel classroom mat",
        "long_tail": "Mrs Smith Calming Techniques rug; ocean animal classroom rug; teacher name prompt rug",
        "description": "Personalize a Mrs Smith Calming Techniques rug with ocean animal characters, activity prompts and pastel classroom artwork.\n\nDesign details\n- Light blue classroom rug artwork reads Calming Techniques with Mrs. Smith's Classroom text at the top.\n- Visible prompt panels include Take A Breath, Color Or Draw, Drink Water, Ask For A Hug, Count To Ten, Stretch, Imagine Happy Place, Listen To Music and Read A Book.\n- SEO copy focuses on visible ocean-animal prompt artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            7: "Care-style graphic for Calming Techniques classroom rug",
            8: "Feature detail graphic for Calming Techniques classroom rug",
        },
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-100": {
        "title": "Mrs Sophia It's Okay Classroom Rug with Ocean Art",
        "meta": "Personalize a Mrs Sophia It's Okay classroom rug with ocean animals, coral, reassurance prompts and aqua classroom artwork.",
        "primary": "Mrs Sophia It's Okay classroom rug",
        "secondary": "ocean classroom prompt rug, teacher name reassurance rug, aqua classroom mat",
        "long_tail": "Mrs Sophia It's Okay classroom rug; ocean classroom prompt rug; teacher name reassurance rug",
        "description": "Personalize a Mrs Sophia It's Okay classroom rug with ocean animals, coral, reassurance prompts and aqua classroom artwork.\n\nDesign details\n- Aqua classroom rug artwork reads It's Okay with Mrs. Sophia's Classroom and Be Kind text.\n- Visible prompt bubbles include Start Over, Make Mistakes, Have Bad Days, Have Feelings, Not Know Something, Ask For Help and Change Your Mind.\n- SEO copy stays tied to visible ocean prompt artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            8: "Care-style graphic for It's Okay ocean classroom rug",
            9: "Feature detail graphic for It's Okay ocean classroom rug",
        },
    },
    "personalized-teacher-classroom-rug-custom-back-to-school-design-105": {
        "title": "Mrs Smith Classroom Affirmation Rug with Doodle Icons",
        "meta": "Personalize a Mrs Smith classroom affirmation rug with doodle school icons, sticky-note prompts, rainbow and flowers.",
        "primary": "Mrs Smith classroom affirmation rug",
        "secondary": "doodle classroom rug, teacher name affirmation mat, school icon rug",
        "long_tail": "Mrs Smith classroom affirmation rug; doodle classroom rug; teacher name affirmation mat",
        "description": "Personalize a Mrs Smith classroom affirmation rug with doodle school icons, sticky-note prompts, rainbow and flowers.\n\nDesign details\n- Black classroom rug artwork centers Mrs. Smith's Classroom text with colorful sticky-note style words around it.\n- Visible words include Grow, Be Kind, Learn, Believe, Dream Big, Inspire, Be You and You Matter, plus apple, pencil, flowers, stars and rainbow icons.\n- SEO copy focuses on visible classroom affirmation artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            8: "Feature detail graphic for Mrs Smith classroom affirmation rug",
        },
    },
    "wheel-of-feelings-and-emotions-round-rug-educational-men-design-110": {
        "title": "My Feelings Emotion Wheel Rug with Cartoon Faces",
        "meta": "Decorate with a My Feelings emotion wheel rug featuring pastel panels, cartoon child faces and labels for eight feelings.",
        "primary": "My Feelings emotion wheel rug",
        "secondary": "cartoon feelings rug, pastel emotion classroom rug, round feelings mat",
        "long_tail": "My Feelings emotion wheel rug; cartoon feelings rug; pastel emotion classroom rug",
        "description": "Decorate with a My Feelings emotion wheel rug featuring pastel panels, cartoon child faces and labels for eight feelings.\n\nDesign details\n- Round pastel rug artwork centers My Feelings text with cartoon faces around the circle.\n- Visible labels include Angry, Happy, Sad, Excited, Scared, Tired, Silly and Shy, with classroom mockups, group photo and size guide images.\n- SEO copy stays tied to visible emotion-wheel artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            5: "Feature graphic for My Feelings emotion wheel rug",
            6: "Feature detail graphic for My Feelings rug artwork",
            7: "Baby scene with pastel My Feelings rug",
        },
    },
    "halloween-door-mat-outdoor-custom-spooky-welcome-mat-with-black-cat": {
        "title": "Personalized Black Cat Halloween Mat with Ghosts",
        "meta": "Personalize a black cat Halloween mat with family name, witch, ghosts, pumpkins, skulls, bats and spooky border art.",
        "primary": "personalized black cat Halloween mat",
        "secondary": "custom Halloween doormat, spooky welcome mat, ghost pumpkin Halloween rug",
        "long_tail": "personalized black cat Halloween mat; custom Halloween doormat; spooky welcome mat",
        "description": "Personalize a black cat Halloween mat with family name, witch, ghosts, pumpkins, skulls, bats and spooky border art.\n\nDesign details\n- Black and orange Halloween mat artwork reads Happy Halloween with family-name personalization on a ribbon-style label.\n- Visible details include witch, black cat, ghosts, jack-o-lanterns, skulls, bats, spiderwebs, porch scenes, doorway mockups and size reference graphic.\n- SEO copy focuses on visible Halloween character artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature detail graphic for personalized black cat Halloween mat",
        },
    },
    "custom-halloween-3d-optical-illusion-round-rug-skeleton-h-design-107": {
        "title": "Graveyard Checkerboard Halloween Illusion Round Rug",
        "meta": "Decorate with a graveyard checkerboard Halloween illusion round rug featuring 3D pit artwork, ghosts, pumpkins and tombstones.",
        "primary": "graveyard checkerboard Halloween rug",
        "secondary": "Halloween optical illusion rug, ghost pumpkin round rug, spooky graveyard mat",
        "long_tail": "graveyard checkerboard Halloween rug; Halloween optical illusion rug; ghost pumpkin round rug",
        "description": "Decorate with a graveyard checkerboard Halloween illusion round rug featuring 3D pit artwork, ghosts, pumpkins and tombstones.\n\nDesign details\n- Round Halloween rug artwork shows a black-and-white checkerboard pit surrounded by graveyard details and purple night accents.\n- Visible border details include ghosts, jack-o-lanterns, tombstones, spiderwebs, room mockups, seasonal scenes and size reference graphic.\n- SEO copy stays tied to visible graveyard illusion artwork without extra care, construction, reverse-side or surface-performance claims.",
        "alts": {
            4: "Feature detail graphic for graveyard checkerboard rug",
            6: "Detail graphic for ghost pumpkin checkerboard rug",
        },
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-104": {
        "title": "Black Orthodox Cross Rug with Gold Border Accents",
        "meta": "Decorate with a black Orthodox cross rug featuring large gold cross, clean border, corner ornaments and dark field artwork.",
        "primary": "black Orthodox cross rug",
        "secondary": "gold cross area rug, Orthodox Christian rug, black religious decor rug",
        "long_tail": "black Orthodox cross rug; gold cross area rug; Orthodox Christian rug",
        "description": "Decorate with a black Orthodox cross rug featuring large gold cross, clean border, corner ornaments and dark field artwork.\n\nDesign details\n- Rectangular black rug artwork shows a large gold Orthodox cross centered inside a simple gold border.\n- Visible details include small corner ornaments, dotted accent lines, living room mockups, fireplace scene, size guide and close-up artwork views.\n- SEO copy focuses on visible cross-and-border artwork without extra care, devotional-use, reverse-side or surface-performance claims.",
        "alts": {
            7: "Feature graphic for black Orthodox cross rug",
            8: "Close-up detail graphic for black gold cross rug",
        },
    },
    "personalized-orthodox-christian-area-rug-custom-eastern-design-108": {
        "title": "Black Gold Orthodox Cross Rug with Scroll Border",
        "meta": "Decorate with a black gold Orthodox cross rug featuring ornate scroll border, corner accents and centered cross artwork.",
        "primary": "black gold Orthodox cross rug",
        "secondary": "scroll border cross rug, Orthodox Christian area rug, black gold religious rug",
        "long_tail": "black gold Orthodox cross rug; scroll border cross rug; Orthodox Christian area rug",
        "description": "Decorate with a black gold Orthodox cross rug featuring ornate scroll border, corner accents and centered cross artwork.\n\nDesign details\n- Rectangular black rug artwork shows a bold gold Orthodox cross with ornate scrollwork along the border and corners.\n- Product images include sofa-side room scenes, fireplace mockup, size guide and close-up views of the cross and gold border art.\n- SEO copy stays tied to visible cross scroll artwork without extra care, devotional-use, reverse-side or surface-performance claims.",
        "alts": {
            6: "Feature graphic for black gold Orthodox cross rug",
            7: "Close-up detail graphic for black gold cross rug",
        },
    },
    "personalized-composition-notebook-classroom-shaped-rugs-for-kids-a01": {
        "title": "Personalized Notebook Classroom Rug with School Icons",
        "meta": "Personalize a notebook classroom rug with teacher name, pink bow, pencil, school bag, crayon cup, books and speckled border.",
        "primary": "personalized notebook classroom rug",
        "secondary": "teacher name classroom rug, school supplies rug, composition notebook mat",
        "long_tail": "personalized notebook classroom rug; teacher name classroom rug; school supplies rug",
        "description": "Personalize a notebook classroom rug with teacher name, pink bow, pencil, school bag, crayon cup, books and speckled border.\n\nDesign details\n- Custom-shaped classroom rug artwork shows a composition notebook texture with a large personalized name panel.\n- Visible details include pink bow, pencil, school bag, crayon cup, books, scissors, classroom scenes, children sitting on rug and size guide images.\n- SEO copy focuses on visible notebook and school-supply artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            9: "Care-style graphic for personalized notebook classroom rug",
        },
    },
    "personalized-composition-notebook-classroom-shaped-rugs-design-2224": {
        "title": "Personalized Open Book Classroom Rug with Pencil Art",
        "meta": "Personalize an open book classroom rug with teacher name, pencil, apple, notebook, daisy accents and composition border.",
        "primary": "personalized open book classroom rug",
        "secondary": "teacher reading corner rug, custom classroom rug, school supplies mat",
        "long_tail": "personalized open book classroom rug; teacher reading corner rug; custom classroom rug",
        "description": "Personalize an open book classroom rug with teacher name, pencil, apple, notebook, daisy accents and composition border.\n\nDesign details\n- Open-book shaped classroom rug artwork shows a teacher name across lined paper pages with school-supply accents.\n- Visible details include pencil, apple, notebook, daisy flowers, black-and-white composition border, classroom mockups, reading-corner scene and size guide images.\n- SEO copy stays tied to visible open-book classroom artwork without extra care, classroom-result, reverse-side or surface-performance claims.",
        "alts": {
            9: "Care-style graphic for personalized open book classroom rug",
        },
    },
}


FORBIDDEN = [
    "indoor/outdoor", " outdoor ", "anti-slip", "anti slip", "non-slip", "non slip",
    "non-skid", "non skid", "nonslip", "machine-washable", "machine washable",
    "machine wash", "washable", "vacuum", "quick-dry", "quick dry", "memory foam",
    "microfiber", "velvet", "plush", "stain", "fade", "easy clean", "easy-clean",
    "waterproof", "absorbent", "backing", "rubber", "layer construction",
    "layered construction", "ultra-soft", "soft ", "cushion", "thickness",
    "thickened", "reinforced", "bound edge", "bound edges", "material panel",
    "learning-pattern", "learning pattern", "support classroom", "support emotional",
    "encourage students", "use a black", "classroom sel", "therapy office",
    "safe durable",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def validate_blob(handle: str, text_parts: list[str]) -> None:
    blob = " ".join(text_parts).lower()
    hits = [term for term in FORBIDDEN if term in blob]
    if hits:
        raise RuntimeError(f"{handle} forbidden terms: {hits}")


def validate_text(handle: str, item: dict[str, str]) -> None:
    title, meta = item["title"], item["meta"]
    if not 45 <= len(title) <= 70:
        raise RuntimeError(f"{handle} title length outside 45-70: {len(title)}")
    if len(meta) > 320:
        raise RuntimeError(f"{handle} meta too long: {len(meta)}")
    parts = [title, meta, item["description"], item["primary"], item["secondary"], item["long_tail"]]
    parts.extend(item["alts"].values())
    validate_blob(handle, parts)


def main():
    for handle, item in REVISIONS.items():
        validate_text(handle, item)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE)
    ws = wb["SEO_Products"]
    cols = headers(ws)
    required = [
        "Handle", "title_proposed", "meta_title_seo", "meta_title_chars",
        "meta_description_seo", "meta_description_chars", "primary_keyword",
        "secondary_keywords", "long_tail_candidates", "description_proposed",
        "description_proposed_html", "revision", "review_status", "review_reason",
        "processing_status", "content_qa_status", "issues",
    ]
    missing = [c for c in required if c not in cols]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")

    changed = []
    for row_idx in range(2, ws.max_row + 1):
        handle = ws.cell(row_idx, cols["Handle"]).value
        if handle not in REVISIONS:
            continue
        item = REVISIONS[handle]
        title, meta, desc = item["title"], item["meta"], item["description"]
        old_revision = str(ws.cell(row_idx, cols["revision"]).value or "")
        ws.cell(row_idx, cols["title_proposed"], title)
        ws.cell(row_idx, cols["meta_title_seo"], title)
        ws.cell(row_idx, cols["meta_title_chars"], len(title))
        ws.cell(row_idx, cols["meta_description_seo"], meta)
        ws.cell(row_idx, cols["meta_description_chars"], len(meta))
        ws.cell(row_idx, cols["primary_keyword"], item["primary"])
        ws.cell(row_idx, cols["secondary_keywords"], item["secondary"])
        ws.cell(row_idx, cols["long_tail_candidates"], item["long_tail"])
        ws.cell(row_idx, cols["description_proposed"], desc)
        ws.cell(row_idx, cols["description_proposed_html"], "<p>" + desc.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>")
        for img_num, alt in item["alts"].items():
            col_name = f"img_{img_num}_alt"
            if col_name not in cols:
                raise RuntimeError(f"Missing column {col_name}")
            ws.cell(row_idx, cols[col_name], alt)
        ws.cell(row_idx, cols["revision"], "2")
        ws.cell(row_idx, cols["review_status"], "NEEDS_REVIEW")
        ws.cell(row_idx, cols["review_reason"], "Revision R040 fixes QA MAJOR issue: unsupported claims. Requires re-QA before approval.")
        ws.cell(row_idx, cols["processing_status"], "DRAFTED")
        ws.cell(row_idx, cols["content_qa_status"], "NOT_RUN")
        ws.cell(row_idx, cols["issues"], "Revision R040 updated targeted fields; not approved; re-QA required. Admin/export prerequisites still open before deploy.")
        changed.append((handle, row_idx, old_revision, "2", title, meta))

    if len(changed) != len(REVISIONS):
        missing = sorted(set(REVISIONS) - {c[0] for c in changed})
        raise RuntimeError(f"Did not update all handles. Missing: {missing}")

    log = wb["Revision_Log"] if "Revision_Log" in wb.sheetnames else wb.create_sheet("Revision_Log")
    if log.max_row == 1 and log.cell(1, 1).value != "revision_batch_id":
        log.append(["revision_batch_id", "generated_at", "handle", "source_row", "old_revision", "new_revision", "changed_fields", "qa_issue_fields", "recheck_condition"])
    keep = [r for r in log.iter_rows(min_row=2, values_only=True) if r and r[0] != "R040"]
    if log.max_row > 1:
        log.delete_rows(2, log.max_row - 1)
    for r in keep:
        log.append(list(r))
    for handle, row_idx, old_revision, new_revision, _title, _meta in changed:
        log.append([
            "R040", NOW, handle, row_idx, old_revision, new_revision,
            "title_proposed, meta_title_seo, meta_title_chars, meta_description_seo, meta_description_chars, primary_keyword, secondary_keywords, long_tail_candidates, description_proposed, description_proposed_html, selected img_alt fields, revision, review_reason, issues",
            "unsupported_claims",
            "Run evidence-driven re-QA on revision 2 for this product before approval.",
        ])

    wb.save(OUTPUT)

    manifest = {
        "revision_batch_id": "R040", "generated_at": NOW,
        "source_workbook": str(SOURCE), "output_workbook": str(OUTPUT),
        "source_revision_plan": str(PLAN), "product_count": len(changed),
        "handles": [c[0] for c in changed], "status": "COMPLETE_AWAITING_REQA",
        "review_status": "NEEDS_REVIEW", "content_qa_status": "NOT_RUN",
        "not_approved_not_deployed": True,
        "next_step": "Bắt đầu revision R041 or Re-QA revision R040",
    }
    (RUN_DIR / "revision_R040_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n")

    lines = [
        "# Revision R040 Summary", "", f"- Generated at: `{NOW}`",
        f"- Source workbook: `{SOURCE}`", f"- Output workbook: `{OUTPUT}`",
        "- Scope: 10 products from revision plan R040.",
        "- Fixes targeted QA issue: unsupported claims.",
        "- Also cleaned risky keyword and image-alt fields inside the R040 scope.",
        "- Kept `review_status=NEEDS_REVIEW` and `content_qa_status=NOT_RUN`.",
        "- No Shopify deploy, no approval.", "", "## Updated products", "",
        "| Handle | New title | Title chars | Meta chars |", "|---|---|---:|---:|",
    ]
    for handle, _row_idx, _old, _new, title, meta in changed:
        lines.append(f"| `{handle}` | {title} | {len(title)} | {len(meta)} |")
    (OUTPUT_DIR / "REVISION_R040_SUMMARY.md").write_text("\n".join(lines) + "\n")

    reopened = load_workbook(OUTPUT, read_only=True, data_only=True)
    out_ws = reopened["SEO_Products"]
    out_cols = headers(out_ws)
    seen = 0
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        handle = row[out_cols["Handle"] - 1]
        if handle not in REVISIONS:
            continue
        seen += 1
        title = row[out_cols["title_proposed"] - 1]
        meta_title = row[out_cols["meta_title_seo"] - 1]
        meta = row[out_cols["meta_description_seo"] - 1]
        parts = [
            str(title or ""), str(meta or ""),
            str(row[out_cols["description_proposed"] - 1] or ""),
            str(row[out_cols["primary_keyword"] - 1] or ""),
            str(row[out_cols["secondary_keywords"] - 1] or ""),
            str(row[out_cols["long_tail_candidates"] - 1] or ""),
        ]
        for img_num in range(1, 23):
            col_name = f"img_{img_num}_alt"
            if col_name in out_cols:
                parts.append(str(row[out_cols[col_name] - 1] or ""))
        if title != meta_title:
            raise RuntimeError(f"{handle} meta title mismatch")
        if not 45 <= len(title) <= 70 or len(meta) > 320:
            raise RuntimeError(f"{handle} title/meta length failed after save")
        if str(row[out_cols["revision"] - 1]) != "2" or row[out_cols["review_status"] - 1] != "NEEDS_REVIEW" or row[out_cols["content_qa_status"] - 1] != "NOT_RUN":
            raise RuntimeError(f"{handle} status failed after save")
        validate_blob(handle, parts)

    log = reopened["Revision_Log"]
    counts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            counts[row[0]] = counts.get(row[0], 0) + 1
    meta_titles, dup_meta_titles = {}, []
    for row in out_ws.iter_rows(min_row=2, values_only=True):
        mt, handle = row[out_cols["meta_title_seo"] - 1], row[out_cols["Handle"] - 1]
        if mt:
            if mt in meta_titles:
                dup_meta_titles.append((mt, meta_titles[mt], handle))
            meta_titles[mt] = handle
    print(json.dumps({
        "output": str(OUTPUT), "updated": seen,
        "revision_log_rows": log.max_row - 1,
        "r040_log_rows": counts.get("R040", 0),
        "duplicate_meta_titles": len(dup_meta_titles),
        "sheets": reopened.sheetnames,
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
